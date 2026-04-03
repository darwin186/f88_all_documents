from unittest.mock import patch
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import RequestFactory, TestCase
from django.urls import reverse
from django.utils import timezone

from app_documents.models import GapoScheduledMessage, UserProfile
from app_notification.services import NotificationSendError
from openpyxl import Workbook

from .forms import AdmIncomingDispatchForm, AdmParcelReceiptForm
from .models import (
    AdmAdministrativeDocument,
    AdmCompany,
    AdmContentType,
    AdmDepartment,
    AdmDocumentStatus,
    AdmDocumentType,
    AdmIncomingDispatchType,
    AdmIncomingDispatchStatus,
    AdmIncomingDispatchStatusLog,
    AdmIncomingGapoGroup,
    AdmIncomingDispatch,
    AdmParcelNotificationBatch,
    AdmParcelDynamicTemplate,
    AdmParcelRecipientCatalog,
    AdmParcelRecipientImportBatch,
    AdmParcelReceipt,
    AdmParcelSenderSuggestion,
    AdmSignerRole,
)
from .views import (
    _allowed_status_codes_for_dispatch,
    _build_incoming_workflow_steps,
    _send_parcel_group_notification_now,
)


class AdmAdministrativeDocumentNumberInvariantTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(username="tester", password="secret")
        cls.doc_type = AdmDocumentType.objects.create(code="QD", name="Quyet dinh")
        cls.content_type = AdmContentType.objects.create(code="NEW", name="New")
        cls.signer_role = AdmSignerRole.objects.create(code="GD", title="Giam doc")
        cls.status = AdmDocumentStatus.objects.create(
            code=AdmDocumentStatus.CODE_DRAFT,
            name="Draft",
        )
        cls.company = AdmCompany.objects.create(code="F88", name="F88")
        cls.department = AdmDepartment.objects.create(
            code="HR",
            name="Nhan su",
            company=cls.company,
        )

    def _create_document(self, title="Van ban goc"):
        return AdmAdministrativeDocument.objects.create(
            doc_type=self.doc_type,
            content_type=self.content_type,
            title=title,
            signer_role=self.signer_role,
            issuing_company=self.company,
            issuing_department=self.department,
            status=self.status,
            created_by=self.user,
        )

    def test_number_and_running_number_stay_immutable_after_normal_update(self):
        doc = self._create_document()
        original_number = doc.document_number_full
        original_running = doc.running_number

        doc.title = "Cap nhat tieu de"
        doc.save()
        doc.refresh_from_db()

        self.assertEqual(doc.document_number_full, original_number)
        self.assertEqual(doc.running_number, original_running)

    def test_direct_tampering_is_reverted_on_save(self):
        doc = self._create_document()
        original_number = doc.document_number_full
        original_running = doc.running_number

        doc.running_number = original_running + 99
        doc.document_number_full = "999/2099/FAKE-FAKE/FAKE"
        doc.save()
        doc.refresh_from_db()

        self.assertEqual(doc.document_number_full, original_number)
        self.assertEqual(doc.running_number, original_running)

    def test_void_transition_can_append_void_suffix(self):
        doc = self._create_document()
        original_number = doc.document_number_full
        original_running = doc.running_number

        doc.is_void = True
        doc.document_number_full = f"{original_number}-VOID"
        doc.save(update_fields=["is_void", "document_number_full", "updated_at"])
        doc.refresh_from_db()

        self.assertTrue(doc.is_void)
        self.assertEqual(doc.document_number_full, f"{original_number}-VOID")
        self.assertEqual(doc.running_number, original_running)

    def test_void_suffix_without_void_flag_is_reverted(self):
        doc = self._create_document()
        original_number = doc.document_number_full

        doc.document_number_full = f"{original_number}-VOID"
        doc.save()
        doc.refresh_from_db()

        self.assertFalse(doc.is_void)
        self.assertEqual(doc.document_number_full, original_number)


class AdmIncomingDispatchTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.company_f88 = AdmCompany.objects.create(code="F88", name="F88")
        cls.company_nnx = AdmCompany.objects.create(code="NNX", name="NNX")
        cls.company_dtf88 = AdmCompany.objects.create(
            code="DTF88",
            name="Dau Tu F88",
        )
        cls.dispatch_type, _ = AdmIncomingDispatchType.objects.update_or_create(
            code="buu_pham_buu_kien",
            defaults={
                "name": "Bưu phẩm/bưu kiện",
                "sort_order": 2,
                "is_active": True,
            },
        )
        cls.dispatch_type_doc, _ = AdmIncomingDispatchType.objects.update_or_create(
            code="cong_van",
            defaults={
                "name": "Công văn đến",
                "sort_order": 1,
                "is_active": True,
            },
        )
        cls.dispatch_status, _ = AdmIncomingDispatchStatus.objects.update_or_create(
            code="pkg_received",
            defaults={
                "name": "Đã tiếp nhận",
                "sort_order": 1,
                "is_active": True,
            },
        )
        cls.dispatch_status_2, _ = AdmIncomingDispatchStatus.objects.update_or_create(
            code="pkg_at_clerical",
            defaults={
                "name": "Đang ở Văn thư",
                "sort_order": 2,
                "is_active": True,
            },
        )
        cls.dispatch_status_3, _ = AdmIncomingDispatchStatus.objects.update_or_create(
            code="pkg_processing",
            defaults={
                "name": "Đã xác nhận",
                "sort_order": 3,
                "is_active": True,
            },
        )
        cls.dispatch_status_4, _ = AdmIncomingDispatchStatus.objects.update_or_create(
            code="pkg_done",
            defaults={
                "name": "Đã bàn giao toàn bộ",
                "sort_order": 4,
                "is_active": True,
            },
        )
        cls.dispatch_status_archived, _ = AdmIncomingDispatchStatus.objects.update_or_create(
            code="pkg_archived",
            defaults={
                "name": "Lưu trữ",
                "sort_order": 5,
                "is_active": True,
            },
        )
        cls.dispatch_status_5, _ = AdmIncomingDispatchStatus.objects.update_or_create(
            code="doc_received",
            defaults={
                "name": "Đã tiếp nhận",
                "sort_order": 6,
                "is_active": True,
            },
        )
        cls.dispatch_status_6, _ = AdmIncomingDispatchStatus.objects.update_or_create(
            code="doc_at_clerical",
            defaults={
                "name": "Đang ở Văn thư",
                "sort_order": 7,
                "is_active": True,
            },
        )
        cls.dispatch_status_7, _ = AdmIncomingDispatchStatus.objects.update_or_create(
            code="doc_at_assistant",
            defaults={
                "name": "Đang ở Ban trợ lý",
                "sort_order": 8,
                "is_active": True,
            },
        )
        cls.dispatch_status_8, _ = AdmIncomingDispatchStatus.objects.update_or_create(
            code="doc_archived",
            defaults={
                "name": "Lưu trữ",
                "sort_order": 9,
                "is_active": True,
            },
        )
        cls.gapo_group = AdmIncomingGapoGroup.objects.create(
            code="HCVT_GROUP",
            name="Nhóm HCVT",
            gapo_group_id="gapo-group-001",
            sort_order=1,
            is_active=True,
        )
        cls.department = AdmDepartment.objects.create(
            code="HCNS",
            name="Hanh chinh nhan su",
            company=cls.company_f88,
        )
        cls.admin_user = User.objects.create_user(username="admin", password="secret")
        cls.vanthu_user = User.objects.create_user(username="vanthu", password="secret")
        cls.viewer_user = User.objects.create_user(username="viewer", password="secret")
        cls.recipient_user = User.objects.create_user(
            username="recipient",
            password="secret",
            first_name="Nhan",
            last_name="Vien",
        )
        vanthu_group = Group.objects.create(name="cv van thu")
        cls.vanthu_user.groups.add(vanthu_group)
        UserProfile.objects.create(
            user=cls.recipient_user,
            department="Hanh chinh nhan su",
            employee_code="E001",
            gapo_user_id="10001",
        )
        UserProfile.objects.create(
            user=cls.vanthu_user,
            department="Van thu",
            employee_code="VT01",
        )
        cls.recipient_batch = AdmParcelRecipientImportBatch.objects.create(
            original_name="recipient_test.xlsx",
            file=SimpleUploadedFile(
                "recipient_test.xlsx",
                b"test-recipient-batch",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            is_current=True,
            imported_by=cls.admin_user,
            imported_rows=3,
            active_rows=3,
        )
        cls.recipient_entry = AdmParcelRecipientCatalog.objects.create(
            import_batch=cls.recipient_batch,
            row_number=2,
            gapo_user_id="10001",
            employee_code="E001",
            full_name="Nhan Vien",
            email="recipient@f88.vn",
            phone_number="84900000000",
            phone_number_normalized="84900000000",
            birth_date="01/01/1990",
            department_full="Tap doan F88 || Hanh chinh nhan su",
            department_name="Hanh chinh nhan su",
            is_active_member=True,
        )
        cls.recipient_entry_2 = AdmParcelRecipientCatalog.objects.create(
            import_batch=cls.recipient_batch,
            row_number=3,
            gapo_user_id="10002",
            employee_code="E002",
            full_name="Nhan Vien Moi",
            email="recipient2@f88.vn",
            phone_number="84900000001",
            phone_number_normalized="84900000001",
            birth_date="02/02/1992",
            department_full="Tap doan F88 || Tai chinh ke toan",
            department_name="Tai chinh ke toan",
            is_active_member=True,
        )

    def _parcel_payload(self, **overrides):
        payload = {
            "recipient_department": "Hanh chinh nhan su",
            "recipient_directory": str(self.recipient_entry.id),
            "parcel_type": "hoso",
            "sender_unit": "Viettel post",
            "content": "Buu pham test",
            "tracking_code": "VT123",
            "receiving_company": self.company_f88.code,
        }
        payload.update(overrides)
        return payload

    def test_incoming_dispatch_form_allows_blank_processing_department(self):
        form = AdmIncomingDispatchForm(
            data={
                "document_number": "CV-001",
                "sending_unit": "UBND",
                "signer_name": "Nguyen Van A",
                "summary": "Noi dung test",
                "incoming_item_type": self.dispatch_type_doc.code,
                "receiving_company": self.company_f88.code,
            }
        )
        self.assertTrue(form.is_valid(), form.errors)

    def test_parcel_form_requires_processing_departments(self):
        form = AdmParcelReceiptForm(
            data=self._parcel_payload(recipient_department="")
        )
        self.assertFalse(form.is_valid())
        self.assertIn("recipient_department", form.errors)

    def test_parcel_form_allows_unknown_recipient_and_department(self):
        form = AdmParcelReceiptForm(
            data=self._parcel_payload(
                recipient_department=AdmParcelReceiptForm.UNKNOWN_DEPARTMENT_VALUE,
                recipient_directory="",
                recipient_unknown="1",
            )
        )
        self.assertTrue(form.is_valid(), form.errors)
        parcel = form.save(commit=False)
        self.assertEqual(parcel.recipient_name, AdmParcelReceiptForm.UNKNOWN_RECIPIENT_LABEL)
        self.assertEqual(parcel.recipient_department, AdmParcelReceiptForm.UNKNOWN_RECIPIENT_LABEL)

    def test_responsible_user_and_received_date_are_immutable(self):
        dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-002",
            responsible_user=self.vanthu_user,
            sending_unit="So Tu Phap",
            signer_name="Tran Van B",
            summary="Noi dung test",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_nnx,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_2,
            created_by=self.vanthu_user,
        )
        original_date = dispatch.received_date

        dispatch.responsible_user = self.admin_user
        dispatch.received_date = original_date.replace(day=max(1, original_date.day - 1))
        dispatch.save()
        dispatch.refresh_from_db()

        self.assertEqual(dispatch.responsible_user_id, self.vanthu_user.id)
        self.assertEqual(dispatch.received_date, original_date)

    def test_viewer_can_view_and_create_dispatch(self):
        self.client.login(username="viewer", password="secret")
        list_url = reverse("admindocuments:parcel_receipt_list")

        response = self.client.get(list_url)
        self.assertEqual(response.status_code, 200)

        post_response = self.client.post(
            list_url,
            data=self._parcel_payload(),
        )
        self.assertEqual(post_response.status_code, 302)
        self.assertEqual(AdmParcelReceipt.objects.count(), 1)
        created = AdmParcelReceipt.objects.first()
        self.assertEqual(created.status_id, self.dispatch_status.code)
        self.assertEqual(created.received_by_id, self.viewer_user.id)
        self.assertEqual(GapoScheduledMessage.objects.count(), 0)

    def test_vanthu_can_create_dispatch(self):
        self.client.login(username="vanthu", password="secret")
        list_url = reverse("admindocuments:parcel_receipt_list")
        response = self.client.post(
            list_url,
            data=self._parcel_payload(receiving_company=self.company_dtf88.code),
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(AdmParcelReceipt.objects.count(), 1)
        created = AdmParcelReceipt.objects.first()
        self.assertIsNotNone(created)
        self.assertIn(f"open={created.id}", response.url)
        self.assertIn(f"highlight={created.id}", response.url)
        self.assertIn("pane=list", response.url)
        self.assertEqual(created.received_by_id, self.vanthu_user.id)
        self.assertEqual(created.sender_unit, "Viettel post")
        self.assertEqual(created.status_id, self.dispatch_status.code)
        self.assertTrue(created.confirmation_token)

    def test_can_create_dispatch_without_document_number(self):
        self.client.login(username="vanthu", password="secret")
        list_url = reverse("admindocuments:parcel_receipt_list")
        response = self.client.post(
            list_url,
            data=self._parcel_payload(),
        )
        self.assertEqual(response.status_code, 302)
        created = AdmParcelReceipt.objects.latest("id")
        self.assertIsNone(created.document_number)

    def test_create_parcel_adds_new_sender_suggestion(self):
        self.client.login(username="vanthu", password="secret")
        list_url = reverse("admindocuments:parcel_receipt_list")
        sender_name = "J&T Cargo"
        self.assertFalse(
            AdmParcelSenderSuggestion.objects.filter(name__iexact=sender_name).exists()
        )

        response = self.client.post(
            list_url,
            data=self._parcel_payload(sender_unit=sender_name),
        )

        self.assertEqual(response.status_code, 302)
        self.assertTrue(
            AdmParcelSenderSuggestion.objects.filter(name__iexact=sender_name).exists()
        )

    @patch("app_admindocuments.views.send_gapo_scheduled_message.apply_async")
    @patch("app_admindocuments.views.send_via_gapo")
    def test_send_notification_action_updates_status_immediately(self, mocked_send_via_gapo, mocked_apply_async):
        mocked_send_via_gapo.return_value = {"ok": True}
        self.client.login(username="vanthu", password="secret")
        parcel = AdmParcelReceipt.objects.create(
            document_number="PK-009A",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hoso",
            sender_unit="Viettel Post",
            receiving_company=self.company_f88,
            status=self.dispatch_status,
            created_by=self.vanthu_user,
            confirmation_token="notify-now-token",
        )
        response = self.client.post(
            reverse("admindocuments:parcel_receipt_send_notification", args=[parcel.id]),
            data={"next": f"{reverse('admindocuments:parcel_receipt_list')}?open={parcel.id}"},
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        parcel.refresh_from_db()
        self.assertEqual(parcel.status_id, self.dispatch_status_2.code)
        self.assertIsNotNone(parcel.notified_at)
        self.assertIsNotNone(parcel.notification_schedule)
        self.assertEqual(parcel.notification_schedule.status, GapoScheduledMessage.Status.SENT)
        self.assertIsNotNone(parcel.reminder_schedule)
        self.assertEqual(mocked_apply_async.call_count, 1)
        messages = list(response.context["messages"])
        self.assertTrue(
            any("Đã gửi thông báo nhận thành công." in str(message) for message in messages)
        )

    @patch("app_admindocuments.views.send_gapo_scheduled_message.apply_async")
    @patch("app_admindocuments.views.send_via_gapo")
    def test_send_notification_action_reports_error_immediately(self, mocked_send_via_gapo, mocked_apply_async):
        mocked_send_via_gapo.side_effect = NotificationSendError("Gapo error 500")
        self.client.login(username="vanthu", password="secret")
        parcel = AdmParcelReceipt.objects.create(
            document_number="PK-009B",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hoso",
            sender_unit="Viettel Post",
            receiving_company=self.company_f88,
            status=self.dispatch_status,
            created_by=self.vanthu_user,
            confirmation_token="notify-fail-token",
        )
        response = self.client.post(
            reverse("admindocuments:parcel_receipt_send_notification", args=[parcel.id]),
            data={"next": f"{reverse('admindocuments:parcel_receipt_list')}?open={parcel.id}"},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        parcel.refresh_from_db()
        self.assertEqual(parcel.status_id, self.dispatch_status.code)
        self.assertEqual(mocked_apply_async.call_count, 0)
        messages = list(response.context["messages"])
        self.assertTrue(
            any("Gửi thông báo nhận thất bại" in str(message) for message in messages)
        )

    @patch("app_admindocuments.views.send_via_gapo")
    @patch("app_admindocuments.views.send_gapo_scheduled_message.apply_async")
    def test_group_notification_action_sends_single_message(self, mocked_apply_async, mocked_send_via_gapo):
        mocked_send_via_gapo.return_value = {"ok": True}
        AdmParcelDynamicTemplate.objects.update_or_create(
            template_type=AdmParcelDynamicTemplate.TEMPLATE_PARCEL_NOTIFY_CONFIRM,
            defaults={
                "name": "Notify parcel",
                "title_template": "Xin chao {{recipient_name}}",
                "body_template": "Ban co {{parcel_count}} kien tu {{primary_sender}} cho {{company_name}}.",
                "button_text": "Xac nhan tai day",
                "hero_image_url": "https://cdn.example.com/parcel.png",
                "button_bg_color": "#14532D",
                "button_text_color": "#F8FAFC",
                "card_border_color": "#22C55E",
                "is_active": True,
            },
        )
        parcel_one = AdmParcelReceipt.objects.create(
            document_number="PK-G-01",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hoso",
            sender_unit="Vnpost",
            receiving_company=self.company_f88,
            status=self.dispatch_status,
            created_by=self.vanthu_user,
        )
        parcel_two = AdmParcelReceipt.objects.create(
            document_number="PK-G-02",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hanghoa",
            sender_unit="Shopee",
            receiving_company=self.company_f88,
            status=self.dispatch_status,
            created_by=self.vanthu_user,
        )
        request = RequestFactory().post("/admindocuments/parcel-receipts/notify-selected/")
        request.user = self.vanthu_user
        request.build_absolute_uri = lambda path="": f"http://testserver{path}"

        batch, _, reminder_error = _send_parcel_group_notification_now(
            [parcel_one, parcel_two],
            self.vanthu_user,
            request,
        )

        parcel_one.refresh_from_db()
        parcel_two.refresh_from_db()
        self.assertEqual(parcel_one.status_id, self.dispatch_status_2.code)
        self.assertEqual(parcel_two.status_id, self.dispatch_status_2.code)
        self.assertEqual(batch.parcel_count, 2)
        self.assertEqual(batch.status, AdmParcelNotificationBatch.Status.SENT)
        self.assertEqual(reminder_error, "")
        self.assertIn("Bạn có 2 kiện hàng", batch.message_text)
        mocked_send_via_gapo.assert_called_once()
        _, kwargs = mocked_send_via_gapo.call_args
        self.assertEqual(kwargs["body_type"], "dynamic")
        layout = kwargs["body_metadata"]["metadata"]["layout"]
        self.assertEqual(layout["type"], "container")
        self.assertEqual(layout["children"][0]["type"], "photo")
        self.assertEqual(layout["children"][2]["children"][0]["type"], "button")
        self.assertEqual(layout["children"][0]["photo_url"], "https://cdn.example.com/parcel.png")
        self.assertEqual(layout["children"][1]["children"][0]["text_object"]["text"], "Xin chao Nhan Vien - E001")
        self.assertEqual(
            layout["children"][1]["children"][1]["text_object"]["text"],
            "Ban co 2 kien tu Vnpost cho F88.",
        )
        self.assertEqual(layout["children"][2]["children"][0]["text_object"]["text"], "Xac nhan tai day")
        self.assertEqual(layout["children"][2]["children"][0]["background"], "14532D")
        self.assertEqual(layout["children"][2]["children"][0]["text_object"]["color"], "F8FAFC")
        self.assertEqual(layout["border"]["color"], "#22C55E")
        self.assertIn("/admindocuments/pb/", layout["children"][2]["children"][0]["deep_link"])
        self.assertEqual(layout["deep_link"], layout["children"][2]["children"][0]["deep_link"])
        self.assertIn("Link nhanh:", batch.message_text)
        mocked_apply_async.assert_called_once()

    @patch("app_admindocuments.views.send_via_gapo")
    @patch("app_admindocuments.views.send_gapo_scheduled_message.apply_async")
    def test_group_notification_action_splits_selected_parcels_by_recipient(self, mocked_apply_async, mocked_send_via_gapo):
        mocked_send_via_gapo.return_value = {"ok": True}
        other_recipient = AdmParcelRecipientCatalog.objects.create(
            import_batch=self.recipient_batch,
            row_number=4,
            gapo_user_id="10002",
            employee_code="E002",
            full_name="Nguoi Khac",
            email="other@f88.vn",
            phone_number="84911111111",
            phone_number_normalized="84911111111",
            department_full="Tap doan F88 || Hanh chinh nhan su",
            department_name="Hanh chinh nhan su",
            is_active_member=True,
        )
        parcel_one = AdmParcelReceipt.objects.create(
            document_number="PK-GD-01",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hoso",
            sender_unit="Vnpost",
            receiving_company=self.company_f88,
            status=self.dispatch_status,
            created_by=self.vanthu_user,
        )
        parcel_two = AdmParcelReceipt.objects.create(
            document_number="PK-GD-02",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=other_recipient,
            recipient_name="Nguoi Khac",
            recipient_employee_code="E002",
            recipient_gapo_user_id="10002",
            parcel_type="hanghoa",
            sender_unit="Shopee",
            receiving_company=self.company_f88,
            status=self.dispatch_status,
            created_by=self.vanthu_user,
        )

        self.client.login(username="vanthu", password="secret")
        response = self.client.post(
            reverse("admindocuments:parcel_receipt_send_group_notification"),
            data={
                "selected_parcels": f"{parcel_one.id},{parcel_two.id}",
                "next": reverse("admindocuments:parcel_receipt_list"),
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(mocked_send_via_gapo.call_count, 2)

    def test_batch_confirmation_link_confirms_all_selected_parcels(self):
        parcel_one = AdmParcelReceipt.objects.create(
            document_number="PK-C-01",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hoso",
            sender_unit="Vnpost",
            receiving_company=self.company_f88,
            status=self.dispatch_status_2,
            created_by=self.vanthu_user,
        )
        parcel_two = AdmParcelReceipt.objects.create(
            document_number="PK-C-02",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hanghoa",
            sender_unit="Shopee",
            receiving_company=self.company_f88,
            status=self.dispatch_status_2,
            created_by=self.vanthu_user,
        )
        batch = AdmParcelNotificationBatch.objects.create(
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            recipient_department="Hanh chinh nhan su",
            parcel_count=2,
            status=AdmParcelNotificationBatch.Status.SENT,
        )
        batch.parcels.set([parcel_one, parcel_two])

        response = self.client.post(
            reverse("admindocuments:parcel_batch_confirm", args=[batch.token]),
            data={"confirmed": "1", "actual_receiver_name": "Nhan Vien", "actual_receiver_employee_code": "E001"},
        )

        self.assertEqual(response.status_code, 302)
        parcel_one.refresh_from_db()
        parcel_two.refresh_from_db()
        batch.refresh_from_db()
        self.assertEqual(parcel_one.status_id, self.dispatch_status_4.code)
        self.assertEqual(parcel_two.status_id, self.dispatch_status_4.code)
        self.assertEqual(batch.status, AdmParcelNotificationBatch.Status.CONFIRMED)
        self.assertIsNotNone(parcel_one.confirmed_at)
        self.assertIsNotNone(parcel_one.completed_at)
        self.assertIsNotNone(batch.confirmed_at)

    def test_parcel_batch_confirm_qr_returns_svg(self):
        self.client.login(username="vanthu", password="secret")
        batch = AdmParcelNotificationBatch.objects.create(
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            recipient_department="Hanh chinh nhan su",
            parcel_count=1,
            status=AdmParcelNotificationBatch.Status.SENT,
        )

        response = self.client.get(
            reverse("admindocuments:parcel_batch_confirm_qr", args=[batch.token])
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "image/svg+xml")
        self.assertIn(b"<svg", response.content)

    def test_mark_handed_over_route_is_deprecated_after_three_step_flow(self):
        self.client.login(username="vanthu", password="secret")
        parcel_one = AdmParcelReceipt.objects.create(
            document_number="PK-H-01",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hoso",
            sender_unit="Vnpost",
            receiving_company=self.company_f88,
            status=self.dispatch_status_3,
            created_by=self.vanthu_user,
        )
        parcel_two = AdmParcelReceipt.objects.create(
            document_number="PK-H-02",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hanghoa",
            sender_unit="Shopee",
            receiving_company=self.company_f88,
            status=self.dispatch_status_3,
            created_by=self.vanthu_user,
        )

        response = self.client.post(
            reverse("admindocuments:parcel_receipt_mark_handed_over"),
            data={
                "selected_parcels": f"{parcel_one.id},{parcel_two.id}",
                "next": reverse("admindocuments:parcel_receipt_list"),
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        parcel_one.refresh_from_db()
        parcel_two.refresh_from_db()
        self.assertEqual(parcel_one.status_id, self.dispatch_status_3.code)
        self.assertEqual(parcel_two.status_id, self.dispatch_status_3.code)

    def test_parcel_form_rejects_recipient_outside_department(self):
        other_recipient = AdmParcelRecipientCatalog.objects.create(
            import_batch=self.recipient_batch,
            row_number=4,
            gapo_user_id="10002",
            employee_code="E002",
            full_name="Other User",
            department_full="Tap doan F88 || Khac",
            department_name="Khac",
            is_active_member=True,
        )
        form = AdmParcelReceiptForm(
            data=self._parcel_payload(recipient_directory=str(other_recipient.id))
        )
        self.assertFalse(form.is_valid())
        self.assertIn("recipient_directory", form.errors)

    def test_change_status_by_select_step_uses_type_specific_flow(self):
        superuser = User.objects.create_superuser(
            username="super_dispatch",
            password="secret",
            email="super_dispatch@example.com",
        )
        dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-005",
            responsible_user=self.vanthu_user,
            sending_unit="So Cong Thuong",
            signer_name="Test Signer",
            summary="Noi dung test doi trang thai",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_5,
            created_by=self.vanthu_user,
        )
        dispatch.processing_departments.set([self.department])

        self.client.force_login(superuser)
        url = reverse("admindocuments:incoming_dispatch_change_status", args=[dispatch.id])
        response = self.client.post(url, data={"status": self.dispatch_status_6.code})

        self.assertEqual(response.status_code, 302)
        dispatch.refresh_from_db()
        self.assertEqual(dispatch.status_id, self.dispatch_status_6.code)
        log = AdmIncomingDispatchStatusLog.objects.get(dispatch=dispatch)
        self.assertEqual(log.from_status_id, self.dispatch_status_5.code)
        self.assertEqual(log.to_status_id, self.dispatch_status_6.code)
        self.assertEqual(log.changed_by_id, superuser.id)

    def test_incoming_document_status_follows_document_workflow_options(self):
        superuser = User.objects.create_superuser(
            username="super_dispatch_2",
            password="secret",
            email="super_dispatch_2@example.com",
        )
        dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-005A",
            responsible_user=self.vanthu_user,
            sending_unit="So Cong Thuong",
            signer_name="Test Signer",
            summary="Noi dung test doi trang thai tu do",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_5,
            created_by=self.vanthu_user,
        )
        dispatch.processing_departments.set([self.department])

        self.client.force_login(superuser)
        url = reverse("admindocuments:incoming_dispatch_change_status", args=[dispatch.id])
        response = self.client.post(url, data={"status": self.dispatch_status_7.code})

        self.assertEqual(response.status_code, 302)
        dispatch.refresh_from_db()
        self.assertEqual(dispatch.status_id, self.dispatch_status_7.code)

    def test_incoming_dispatch_update_processing_department_can_set_and_clear(self):
        superuser = User.objects.create_superuser(
            username="super_dispatch_dept",
            password="secret",
            email="super_dispatch_dept@example.com",
        )
        second_department = AdmDepartment.objects.create(
            code="KT",
            name="Ke toan",
            company=self.company_f88,
        )
        dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-005B",
            responsible_user=self.vanthu_user,
            sending_unit="So Cong Thuong",
            signer_name="Test Signer",
            summary="Noi dung test doi phong ban",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_5,
            created_by=self.vanthu_user,
        )
        dispatch.processing_departments.set([self.department])

        self.client.force_login(superuser)
        url = reverse("admindocuments:incoming_dispatch_update_department", args=[dispatch.id])
        response = self.client.post(url, data={"processing_department": second_department.id})

        self.assertEqual(response.status_code, 302)
        dispatch.refresh_from_db()
        self.assertEqual(
            list(dispatch.processing_departments.values_list("id", flat=True)),
            [second_department.id],
        )

        response = self.client.post(url, data={"processing_department": ""})

        self.assertEqual(response.status_code, 302)
        dispatch.refresh_from_db()
        self.assertFalse(dispatch.processing_departments.exists())

    def test_incoming_document_workflow_no_longer_shows_pending_signer_step(self):
        dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-006",
            responsible_user=self.vanthu_user,
            sending_unit="So Cong Thuong",
            signer_name="Test Signer",
            summary="Noi dung test workflow",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_6,
            created_by=self.vanthu_user,
        )

        labels = [step["label"] for step in _build_incoming_workflow_steps(dispatch)]
        self.assertNotIn("Trình ký", labels)

    def test_incoming_document_available_statuses_excludes_initial_received(self):
        dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-006B",
            responsible_user=self.vanthu_user,
            sending_unit="So Cong Thuong",
            signer_name="Test Signer",
            summary="Noi dung test available statuses",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_5,
            created_by=self.vanthu_user,
        )

        available_codes = _allowed_status_codes_for_dispatch(dispatch)

        self.assertNotIn(self.dispatch_status_5.code, available_codes)
        self.assertIn(self.dispatch_status_6.code, available_codes)
        self.assertIn(self.dispatch_status_7.code, available_codes)
        self.assertIn(self.dispatch_status_8.code, available_codes)

    def test_incoming_document_workflow_steps_include_status_change_timestamps(self):
        dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-006C",
            responsible_user=self.vanthu_user,
            sending_unit="So Cong Thuong",
            signer_name="Test Signer",
            summary="Noi dung test workflow timestamps",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_7,
            created_by=self.vanthu_user,
        )
        clerical_changed_at = timezone.now() - timezone.timedelta(hours=3)
        assistant_changed_at = timezone.now() - timezone.timedelta(hours=1)
        first_log = AdmIncomingDispatchStatusLog.objects.create(
            dispatch=dispatch,
            from_status=self.dispatch_status_5,
            to_status=self.dispatch_status_6,
            changed_by=self.admin_user,
        )
        AdmIncomingDispatchStatusLog.objects.filter(pk=first_log.pk).update(changed_at=clerical_changed_at)
        second_log = AdmIncomingDispatchStatusLog.objects.create(
            dispatch=dispatch,
            from_status=self.dispatch_status_6,
            to_status=self.dispatch_status_7,
            changed_by=self.admin_user,
        )
        AdmIncomingDispatchStatusLog.objects.filter(pk=second_log.pk).update(changed_at=assistant_changed_at)

        steps = _build_incoming_workflow_steps(dispatch)
        timestamp_map = {step["code"]: step["timestamp"] for step in steps}

        self.assertEqual(timestamp_map["doc_at_clerical"], clerical_changed_at)
        self.assertEqual(timestamp_map["doc_at_assistant"], assistant_changed_at)

    def test_legacy_incoming_document_status_is_mapped_to_correct_workflow_step(self):
        legacy_status, _ = AdmIncomingDispatchStatus.objects.update_or_create(
            code="to_btl",
            defaults={
                "name": "Chuyển BTL xử lý",
                "sort_order": 98,
                "is_active": True,
            },
        )
        dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-006A",
            responsible_user=self.vanthu_user,
            sending_unit="So Cong Thuong",
            signer_name="Test Signer",
            summary="Noi dung test workflow cu",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=legacy_status,
            created_by=self.vanthu_user,
        )

        steps = _build_incoming_workflow_steps(dispatch)
        current_steps = [step["label"] for step in steps if step["is_current"]]
        self.assertEqual(current_steps, ["Đang ở Ban trợ lý"])

    def test_legacy_to_pc_status_is_mapped_to_assistant_workflow_step(self):
        legacy_status, _ = AdmIncomingDispatchStatus.objects.update_or_create(
            code="to_pc",
            defaults={
                "name": "Chuyển PC xử lý",
                "sort_order": 99,
                "is_active": True,
            },
        )
        dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-006B1",
            responsible_user=self.vanthu_user,
            sending_unit="So Cong Thuong",
            signer_name="Test Signer",
            summary="Noi dung test workflow cu to_pc",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=legacy_status,
            created_by=self.vanthu_user,
        )

        steps = _build_incoming_workflow_steps(dispatch)
        current_steps = [step["label"] for step in steps if step["is_current"]]
        self.assertEqual(current_steps, ["Đang ở Ban trợ lý"])

    def test_incoming_dispatch_upload_accepts_pdf_files(self):
        dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-007",
            responsible_user=self.vanthu_user,
            sending_unit="So Cong Thuong",
            signer_name="Test Signer",
            summary="Noi dung test upload pdf",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_5,
            created_by=self.vanthu_user,
        )
        self.client.force_login(self.admin_user)
        upload = SimpleUploadedFile(
            "test.pdf",
            b"%PDF-1.4 test content",
            content_type="application/pdf",
        )

        response = self.client.post(
            reverse("admindocuments:incoming_dispatch_upload_images", args=[dispatch.id]),
            data={"images": upload, "next": reverse("admindocuments:incoming_document_list")},
        )

        self.assertEqual(response.status_code, 302)
        dispatch.refresh_from_db()
        self.assertEqual(dispatch.images.count(), 1)
        self.assertTrue(dispatch.images.first().image.name.endswith(".pdf"))

    def test_incoming_dispatch_upload_allows_more_than_five_files(self):
        dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-007A",
            responsible_user=self.vanthu_user,
            sending_unit="So Cong Thuong",
            signer_name="Test Signer",
            summary="Noi dung test nhieu tep",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_5,
            created_by=self.vanthu_user,
        )
        self.client.force_login(self.admin_user)
        uploads = [
            SimpleUploadedFile(
                f"test-{idx}.pdf",
                b"%PDF-1.4 test content",
                content_type="application/pdf",
            )
            for idx in range(6)
        ]

        response = self.client.post(
            reverse("admindocuments:incoming_dispatch_upload_images", args=[dispatch.id]),
            data={"images": uploads, "next": reverse("admindocuments:incoming_document_list")},
        )

        self.assertEqual(response.status_code, 302)
        dispatch.refresh_from_db()
        self.assertEqual(dispatch.images.count(), 6)

    def test_legacy_route_redirects_to_incoming_document_screen(self):
        self.client.login(username="viewer", password="secret")
        response = self.client.get(reverse("admindocuments:incoming_dispatch_list"))

        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            response.url,
            reverse("admindocuments:incoming_document_list"),
        )

    def test_each_screen_only_lists_its_own_type(self):
        parcel_dispatch = AdmParcelReceipt.objects.create(
            document_number="PK-001",
            received_by=self.vanthu_user,
            received_at=None,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            recipient_user=self.recipient_user,
            parcel_type="hoso",
            sender_unit="Viettel Post",
            content="Buu pham test",
            receiving_company=self.company_f88,
            status=self.dispatch_status,
            created_by=self.vanthu_user,
            confirmation_token="tokentest1",
        )
        doc_dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-006",
            responsible_user=self.vanthu_user,
            sending_unit="UBND",
            signer_name="Signer",
            summary="Cong van test",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_5,
            created_by=self.vanthu_user,
        )
        doc_dispatch.processing_departments.set([self.department])

        self.client.login(username="viewer", password="secret")
        doc_response = self.client.get(reverse("admindocuments:incoming_document_list"))
        parcel_response = self.client.get(reverse("admindocuments:parcel_receipt_list"))

        self.assertContains(doc_response, "CV-006")
        self.assertNotContains(doc_response, "PK-001")
        self.assertContains(parcel_response, "Viettel Post")
        self.assertNotContains(parcel_response, "CV-006")

    def test_incoming_document_list_hides_archived_by_default(self):
        archived_dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-ARCH-01",
            responsible_user=self.vanthu_user,
            sending_unit="UBND",
            signer_name="Signer",
            summary="Cong van da luu tru",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_8,
            created_by=self.vanthu_user,
        )
        archived_dispatch.processing_departments.set([self.department])

        self.client.login(username="viewer", password="secret")
        response = self.client.get(reverse("admindocuments:incoming_document_list"))

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "CV-ARCH-01")

    def test_incoming_document_list_can_filter_archived_records(self):
        archived_dispatch = AdmIncomingDispatch.objects.create(
            document_number="CV-ARCH-02",
            responsible_user=self.vanthu_user,
            sending_unit="UBND",
            signer_name="Signer",
            summary="Cong van da luu tru 2",
            incoming_item_type=self.dispatch_type_doc,
            receiving_company=self.company_f88,
            gapo_group=self.gapo_group,
            status=self.dispatch_status_8,
            created_by=self.vanthu_user,
        )
        archived_dispatch.processing_departments.set([self.department])

        self.client.login(username="viewer", password="secret")
        response = self.client.get(
            reverse("admindocuments:incoming_document_list"),
            {"status": self.dispatch_status_8.code},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "CV-ARCH-02")

    def test_parcel_list_hides_completed_items(self):
        AdmParcelReceipt.objects.create(
            document_number="PK-010",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hoso",
            sender_unit="Da Hoan Tat",
            tracking_code="DONE-001",
            receiving_company=self.company_f88,
            status=self.dispatch_status_4,
            created_by=self.vanthu_user,
            confirmation_token="done-token",
        )
        AdmParcelReceipt.objects.create(
            document_number="PK-011",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hoso",
            sender_unit="Dang Mo",
            tracking_code="OPEN-001",
            receiving_company=self.company_f88,
            status=self.dispatch_status_2,
            created_by=self.vanthu_user,
            confirmation_token="open-token",
        )

        self.client.login(username="viewer", password="secret")
        response = self.client.get(reverse("admindocuments:parcel_receipt_list"))

        rendered_names = [item.sender_unit for item in response.context["dispatches"].object_list]
        self.assertIn("Dang Mo", rendered_names)
        self.assertNotIn("Da Hoan Tat", rendered_names)

    @patch("app_admindocuments.views.send_gapo_scheduled_message.apply_async")
    def test_confirmation_link_marks_parcel_confirmed(self, mocked_apply_async):
        parcel = AdmParcelReceipt.objects.create(
            document_number="PK-009",
            received_by=self.vanthu_user,
            received_at=None,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            recipient_user=self.recipient_user,
            parcel_type="hoso",
            sender_unit="Viettel Post",
            content="Buu pham test",
            receiving_company=self.company_f88,
            status=self.dispatch_status_2,
            created_by=self.vanthu_user,
            confirmation_token="confirmtoken",
        )
        reminder = GapoScheduledMessage.objects.create(
            receiver_id="10001",
            message="reminder",
            schedule_at=timezone.now(),
            created_by=self.vanthu_user,
        )
        parcel.reminder_schedule = reminder
        parcel.save(update_fields=["reminder_schedule"])

        response = self.client.post(
            reverse("admindocuments:parcel_receipt_confirm", args=["confirmtoken"]),
            data={"confirmed": "1", "actual_receiver_name": "Nhan Vien", "actual_receiver_employee_code": "E001"},
        )

        self.assertEqual(response.status_code, 302)
        parcel.refresh_from_db()
        reminder.refresh_from_db()
        self.assertEqual(parcel.status_id, self.dispatch_status_4.code)
        self.assertIsNotNone(parcel.confirmed_at)
        self.assertIsNotNone(parcel.completed_at)
        self.assertEqual(parcel.actual_receiver_name, "Nhan Vien")
        self.assertEqual(reminder.status, GapoScheduledMessage.Status.CANCELLED)

    def test_assign_recipient_after_unknown_receipt(self):
        self.client.login(username="vanthu", password="secret")
        parcel = AdmParcelReceipt.objects.create(
            document_number="PK-UNKNOWN-01",
            received_by=self.vanthu_user,
            recipient_department=AdmParcelReceiptForm.UNKNOWN_RECIPIENT_LABEL,
            recipient_name=AdmParcelReceiptForm.UNKNOWN_RECIPIENT_LABEL,
            parcel_type="hoso",
            sender_unit="Viettel Post",
            receiving_company=self.company_f88,
            status=self.dispatch_status,
            created_by=self.vanthu_user,
        )

        response = self.client.post(
            reverse("admindocuments:parcel_receipt_assign_recipient", args=[parcel.id]),
            data={
                "recipient_directory_id": str(self.recipient_entry.id),
                "recipient_department": self.recipient_entry.department_name,
                "next": reverse("admindocuments:parcel_receipt_list"),
            },
        )

        self.assertEqual(response.status_code, 302)
        parcel.refresh_from_db()
        self.assertEqual(parcel.recipient_directory_id, self.recipient_entry.id)
        self.assertEqual(parcel.recipient_name, self.recipient_entry.full_name)
        self.assertEqual(parcel.recipient_department, self.recipient_entry.department_name)

    def test_completed_parcel_with_legacy_status_is_hidden_from_list(self):
        self.client.login(username="vanthu", password="secret")
        completed_parcel = AdmParcelReceipt.objects.create(
            document_number="PK-LEGACY-DONE",
            received_by=self.vanthu_user,
            recipient_department="Phong CNTT",
            recipient_name="Nhan Vien Cu",
            recipient_employee_code="E777",
            parcel_type="hoso",
            sender_unit="Viettel Post",
            receiving_company=self.company_f88,
            status=self.dispatch_status_2,
            created_by=self.vanthu_user,
            confirmed_at=timezone.now(),
            completed_at=timezone.now(),
        )

        response = self.client.get(reverse("admindocuments:parcel_receipt_list"))

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, completed_parcel.document_number)

    def test_parcel_list_filters_recipient_by_text_query(self):
        self.client.login(username="vanthu", password="secret")
        matched = AdmParcelReceipt.objects.create(
            document_number="PK-FILTER-RECIPIENT-1",
            received_by=self.vanthu_user,
            recipient_directory=self.recipient_entry,
            recipient_department=self.recipient_entry.department_name,
            recipient_name=self.recipient_entry.full_name,
            recipient_employee_code=self.recipient_entry.employee_code,
            recipient_gapo_user_id=self.recipient_entry.gapo_user_id,
            parcel_type="hoso",
            sender_unit="Viettel Post",
            receiving_company=self.company_f88,
            status=self.dispatch_status,
            created_by=self.vanthu_user,
        )
        other = AdmParcelReceipt.objects.create(
            document_number="PK-FILTER-RECIPIENT-2",
            received_by=self.vanthu_user,
            recipient_directory=self.recipient_entry_2,
            recipient_department=self.recipient_entry_2.department_name,
            recipient_name=self.recipient_entry_2.full_name,
            recipient_employee_code=self.recipient_entry_2.employee_code,
            recipient_gapo_user_id=self.recipient_entry_2.gapo_user_id,
            parcel_type="hanghoa",
            sender_unit="Shopee",
            receiving_company=self.company_f88,
            status=self.dispatch_status,
            created_by=self.vanthu_user,
        )

        response = self.client.get(
            reverse("admindocuments:parcel_receipt_list"),
            {"recipient": self.recipient_entry.employee_code},
        )

        self.assertEqual(response.status_code, 200)
        rendered_names = [item.sender_unit for item in response.context["dispatches"].object_list]
        self.assertIn(matched.sender_unit, rendered_names)
        self.assertNotIn(other.sender_unit, rendered_names)

    def test_reassign_recipient_after_initial_assignment(self):
        self.client.login(username="vanthu", password="secret")
        parcel = AdmParcelReceipt.objects.create(
            document_number="PK-REASSIGN-01",
            received_by=self.vanthu_user,
            recipient_directory=self.recipient_entry,
            recipient_department=self.recipient_entry.department_name,
            recipient_name=self.recipient_entry.full_name,
            recipient_employee_code=self.recipient_entry.employee_code,
            recipient_gapo_user_id=self.recipient_entry.gapo_user_id,
            parcel_type="hoso",
            sender_unit="Viettel Post",
            receiving_company=self.company_f88,
            status=self.dispatch_status,
            created_by=self.vanthu_user,
        )

        response = self.client.post(
            reverse("admindocuments:parcel_receipt_assign_recipient", args=[parcel.id]),
            data={
                "recipient_directory_id": str(self.recipient_entry_2.id),
                "recipient_department": self.recipient_entry_2.department_name,
                "next": reverse("admindocuments:parcel_receipt_list"),
            },
        )

        self.assertEqual(response.status_code, 302)
        parcel.refresh_from_db()
        self.assertEqual(parcel.recipient_directory_id, self.recipient_entry_2.id)
        self.assertEqual(parcel.recipient_name, self.recipient_entry_2.full_name)
        self.assertEqual(parcel.recipient_department, self.recipient_entry_2.department_name)

    def test_register_proxy_saves_name_and_employee_code(self):
        self.client.login(username="vanthu", password="secret")
        parcel = AdmParcelReceipt.objects.create(
            document_number="PK-PROXY-01",
            received_by=self.vanthu_user,
            recipient_department="Hanh chinh nhan su",
            recipient_directory=self.recipient_entry,
            recipient_name="Nhan Vien",
            recipient_employee_code="E001",
            recipient_gapo_user_id="10001",
            parcel_type="hoso",
            sender_unit="Viettel Post",
            receiving_company=self.company_f88,
            status=self.dispatch_status_2,
            created_by=self.vanthu_user,
        )

        response = self.client.post(
            reverse("admindocuments:parcel_receipt_register_proxy", args=[parcel.id]),
            data={
                "claim_mode": "proxy",
                "proxy_receiver_name": "Nguoi Nhan Ho",
            },
        )

        self.assertEqual(response.status_code, 302)
        parcel.refresh_from_db()
        self.assertEqual(parcel.proxy_receiver_name, "Nguoi Nhan Ho")
        self.assertEqual(parcel.proxy_receiver_employee_code, "")

    def test_recipient_directory_upload_imports_merged_cells(self):
        superuser = User.objects.create_superuser(
            username="super_import",
            password="secret",
            email="super_import@example.com",
        )
        workbook = Workbook()
        ws = workbook.active
        ws.title = "Danh sach user"
        headers = [
            "STT",
            "GAPO User ID",
            "Mã nhân viên",
            "Tên thành viên",
            "Email",
            "Số điện thoại",
            "Trạng thái",
            "Quyền",
            "Sơ đồ tổ chức",
            "Chức vụ",
            "Phòng ban đầy đủ",
            "Vùng miền",
            "Ngày sinh",
            "Loại hợp đồng",
            "Ngày vào công ty",
            "Ngày ký hợp đồng chính thức đầu tiên",
            "Ngày nghỉ việc",
            "Thời gian tạo",
        ]
        ws.append(headers)
        ws.append([
            1,
            "10001",
            "E001",
            "Nhan Vien",
            "recipient@f88.vn",
            "84900000000",
            "Đang hoạt động",
            "Admin",
            "F88",
            "Chuyên viên",
            "Tập đoàn F88 || Khối Vận hành || Phòng Hành chính",
            "",
            "",
            "",
            "",
            "",
            "",
            "24/03/2026 10:00",
        ])
        ws.append([
            2,
            None,
            None,
            None,
            "recipient2@f88.vn",
            None,
            None,
            None,
            "F88",
            "Chuyên viên",
            "Tập đoàn F88 || Khối Vận hành || Phòng Hành chính",
            "",
            "",
            "",
            "",
            "",
            "",
            "24/03/2026 10:05",
        ])
        ws.merge_cells("B2:B3")
        ws.merge_cells("C2:C3")
        ws.merge_cells("D2:D3")
        ws.merge_cells("F2:F3")
        ws.merge_cells("G2:G3")
        ws.merge_cells("H2:H3")

        content = BytesIO()
        workbook.save(content)
        content.seek(0)

        self.client.force_login(superuser)
        response = self.client.post(
            reverse("admindocuments:parcel_recipient_directory"),
            data={
                "action": "upload",
                "file": SimpleUploadedFile(
                    "recipient_upload.xlsx",
                    content.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )

        self.assertEqual(response.status_code, 302)
        current_batch = AdmParcelRecipientImportBatch.objects.get(is_current=True)
        imported = list(
            current_batch.recipients.order_by("row_number").values(
                "gapo_user_id", "employee_code", "full_name", "department_name"
            )
        )
        self.assertGreaterEqual(len(imported), 2)
        self.assertEqual(imported[0]["department_name"], "Phòng Hành chính")
        self.assertEqual(imported[1]["gapo_user_id"], "10001")
        self.assertEqual(imported[1]["employee_code"], "E001")
        self.assertEqual(imported[1]["full_name"], "Nhan Vien")

    def test_recipient_directory_requires_superuser(self):
        self.client.login(username="viewer", password="secret")
        response = self.client.get(reverse("admindocuments:parcel_recipient_directory"))

        self.assertEqual(response.status_code, 403)

    def test_recipient_search_by_full_phone_masks_sensitive_data_for_viewer(self):
        self.client.login(username="viewer", password="secret")
        url = reverse("admindocuments:parcel_recipient_search")

        partial = self.client.get(url, {"q": "849000"})
        exact = self.client.get(url, {"q": "84900000000"})

        self.assertEqual(partial.status_code, 200)
        self.assertEqual(partial.json()["results"], [])
        self.assertEqual(exact.status_code, 200)
        results = exact.json()["results"]
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["full_name"], "Nhan Vien")
        self.assertEqual(results[0]["phone_number"], "849*****000")
        self.assertTrue(results[0]["email"].startswith("re"))

    def test_recipient_search_masks_sensitive_data_for_superuser_too(self):
        self.client.login(username="admin", password="secret")
        url = reverse("admindocuments:parcel_recipient_search")

        response = self.client.get(url, {"q": "84900000000"})

        self.assertEqual(response.status_code, 200)
        results = response.json()["results"]
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["phone_number"], "849*****000")
        self.assertNotEqual(results[0]["phone_number"], "84900000000")
