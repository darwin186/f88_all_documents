import logging
import os
import uuid
from datetime import datetime, date, timedelta
from io import BytesIO
from functools import wraps
from urllib.parse import urlencode, urljoin, urlparse, urlunparse
import qrcode
from qrcode.image.svg import SvgPathImage

from django.contrib import messages
from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.db import IntegrityError, transaction, models
from django.db.models import F
from django.db.models import Count, Q, Max
from django.db.models.functions import ExtractYear
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.templatetags.static import static
from django.utils import timezone
from django.conf import settings
from django.http import JsonResponse

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app_documents.models import GapoScheduledMessage, Shop, Region, UserProfile
from app_documents.tasks import send_gapo_scheduled_message

from .forms import (
    AdmAdministrativeDocumentForm,
    AdmAdministrativeDocumentSplitForm,
    AdmAdministrativeDocumentUpdateForm,
    AdmAdministrativeDocumentEditPopupForm,
    AdmIncomingDispatchForm,
    AdmParcelReceiveLocationForm,
    AdmParcelAutoNotifySettingForm,
    AdmParcelReceiptForm,
    AdmPaperDocumentForm,
    AdmDocumentTypeForm,
    AdmContentTypeForm,
    AdmSignerRoleForm,
    AdmDocumentStatusForm,
    AdmDocumentStatusColorsForm,
    AdmCompanyForm,
    AdmDepartmentForm,
    AdmPaperTypeForm,
    AdmCourierCompanyForm,
)
from .models import (
    AdmAdministrativeDocument,
    AdmCompany,
    AdmContentType,
    AdmCourierCompany,
    AdmDepartment,
    AdmDocumentStatus,
    AdmDocumentType,
    AdmDocumentAttachment,
    AdmDocumentCounter,
    AdmIncomingDispatchStatus,
    AdmIncomingDispatchType,
    AdmIncomingDispatch,
    AdmIncomingDispatchImage,
    AdmIncomingDispatchStatusLog,
    AdmParcelAutoNotifySetting,
    AdmParcelDynamicTemplate,
    AdmParcelNotificationBatch,
    AdmParcelRecipientCatalog,
    AdmParcelRecipientImportBatch,
    AdmParcelReceiveLocation,
    AdmParcelReceipt,
    AdmParcelReceiptImage,
    AdmParcelReceiptLog,
    AdmParcelSenderSuggestion,
    AdmPaperCounter,
    AdmPaperDocument,
    AdmPaperType,
    AdmSignerRole,
    AdmAdministrativeDocumentHistory,
    administrative_number_year,
)
from .services import (
    NotificationSendError,
    allocate_paper_running_number,
    allocate_running_number,
    send_via_gapo,
    create_split_document,
    next_split_suffix,
    split_number_group,
)
from .access import has_admin_docs_access


def _has_admin_docs_access(user) -> bool:
    return has_admin_docs_access(user)


def _has_incoming_dispatch_create_access(user) -> bool:
    return bool(getattr(user, "is_authenticated", False))


def _has_incoming_dispatch_status_edit_access(user) -> bool:
    if not getattr(user, "is_authenticated", False):
        return False
    if getattr(user, "is_superuser", False):
        return True
    if user.has_perm("app_admindocuments.change_admincomingdispatch"):
        return True
    return _has_admin_docs_access(user)


logger = logging.getLogger(__name__)

INCOMING_TYPE_DOC = "cong_van"
INCOMING_TYPE_PARCEL = "buu_pham_buu_kien"

INCOMING_DOC_FLOW = [
    ("doc_received", "Đã tiếp nhận"),
    ("doc_at_clerical", "Đang ở Văn thư"),
    ("doc_at_assistant", "Đang ở Ban trợ lý"),
    ("doc_archived", "Lưu trữ"),
]

LEGACY_INCOMING_DOC_STATUS_MAP = {
    "doc_pending_signer": "doc_at_assistant",
    "in_progress": "doc_at_clerical",
    "to_btl": "doc_at_assistant",
    "to_pc": "doc_at_assistant",
    "received": "doc_received",
    "at_clerical": "doc_at_clerical",
    "at_assistant": "doc_at_assistant",
    "pending_signer": "doc_at_assistant",
    "done": "doc_archived",
    "archived": "doc_archived",
}

INCOMING_PARCEL_FLOW = [
    ("pkg_received", "Lễ tân tiếp nhận"),
    ("pkg_at_clerical", "Đã thông báo"),
    ("pkg_done", "Hoàn tất"),
]

INCOMING_ARCHIVED_STATUS_CODES = {"doc_archived", "pkg_archived", "archived"}
PARCEL_HIDDEN_LIST_STATUS_CODES = INCOMING_ARCHIVED_STATUS_CODES | {"pkg_processing", "pkg_done"}
INCOMING_VISIBLE_STATUS_CODES = [
    code
    for code, _ in (INCOMING_DOC_FLOW + INCOMING_PARCEL_FLOW)
    if code not in INCOMING_ARCHIVED_STATUS_CODES
]

PARCEL_RECIPIENT_IMPORT_HEADERS = {
    "stt": "stt",
    "user id": "gapo_user_id",
    "gapo user id": "gapo_user_id",
    "mã nhân viên": "employee_code",
    "tên thành viên": "full_name",
    "email": "email",
    "số điện thoại": "phone_number",
    "trạng thái": "employment_status",
    "quyền": "permission_name",
    "sơ đồ tổ chức": "org_chart",
    "chức vụ": "position_name",
    "phòng ban đầy đủ": "department_full",
    "vùng miền": "region_name",
    "ngày sinh": "birth_date",
    "ngày vào công ty": "company_join_date",
    "ngày ký hợp đồng chính thức đầu tiên": "contract_start_date",
    "ngày nghỉ việc": "leave_date",
    "thời gian tạo": "source_created_at",
}

INCOMING_SCREEN_META = {
    INCOMING_TYPE_DOC: {
        "route_name": "incoming_document_list",
        "page_title": "Công văn đến",
        "page_description": "Tiếp nhận và theo dõi công văn đến.",
        "create_button_label": "Tiếp nhận Công văn đến",
        "modal_title": "Tiếp nhận công văn đến",
        "success_message": "Đã nhập công văn đến thành công.",
        "empty_message": "Chưa có công văn đến nào.",
        "search_placeholder": "Số hiệu hoặc tên trích yếu...",
        "auto_status_note": 'Người phụ trách, ngày nhận và trạng thái "Đã tiếp nhận" được hệ thống tự động ghi nhận.',
    },
    INCOMING_TYPE_PARCEL: {
        "route_name": "parcel_receipt_list",
        "page_title": "Bưu phẩm/Bưu kiện",
        "page_description": "Tiếp nhận và theo dõi bưu phẩm, bưu kiện.",
        "create_button_label": "Tiếp nhận Bưu phẩm/Bưu kiện",
        "modal_title": "Tiếp nhận bưu phẩm/bưu kiện",
        "success_message": "Đã tiếp nhận bưu phẩm/bưu kiện thành công.",
        "empty_message": "Chưa có bưu phẩm hoặc bưu kiện nào.",
        "search_placeholder": "Số hiệu hoặc nội dung...",
        "auto_status_note": 'Người phụ trách, ngày nhận và trạng thái khởi tạo được hệ thống tự động ghi nhận.',
    },
}


def _incoming_dispatch_route_name(item_type_code):
    meta = INCOMING_SCREEN_META.get(item_type_code) or INCOMING_SCREEN_META[INCOMING_TYPE_DOC]
    return f"admindocuments:{meta['route_name']}"


def _incoming_dispatch_list_url(item_type_code):
    return reverse(_incoming_dispatch_route_name(item_type_code))


def _exclude_hidden_parcels(queryset):
    return queryset.exclude(
        Q(status_id__in=PARCEL_HIDDEN_LIST_STATUS_CODES)
        | Q(completed_at__isnull=False)
        | Q(confirmed_at__isnull=False)
    )


def _completed_parcels_queryset(queryset):
    return queryset.filter(
        Q(status_id="pkg_done")
        | Q(completed_at__isnull=False)
        | Q(confirmed_at__isnull=False)
    )


def _save_incoming_dispatch_images(dispatch, files, user):
    saved_count = 0
    for image_file in files:
        AdmIncomingDispatchImage.objects.create(
            dispatch=dispatch,
            image=image_file,
            uploaded_by=user,
        )
        saved_count += 1
    return saved_count


def _save_parcel_receipt_images(parcel_receipt, files, user):
    for image_file in files:
        AdmParcelReceiptImage.objects.create(
            parcel_receipt=parcel_receipt,
            image=image_file,
            uploaded_by=user,
        )


def _validate_parcel_images(files):
    if len(files) > 5:
        return "Chỉ được tải tối đa 5 tệp."
    for attachment in files:
        if getattr(attachment, "size", 0) > 30 * 1024 * 1024:
            return f"Tệp '{attachment.name}' vượt quá 30MB."
    return ""


def _validate_incoming_dispatch_files(files):
    for attachment in files:
        if getattr(attachment, "size", 0) > 30 * 1024 * 1024:
            return f"Tệp '{attachment.name}' vượt quá 30MB."
    return ""


def _log_parcel_event(parcel_receipt, action, actor=None, from_status=None, to_status=None, note="", metadata=None):
    AdmParcelReceiptLog.objects.create(
        parcel_receipt=parcel_receipt,
        action=action,
        actor=actor,
        from_status_id=from_status,
        to_status_id=to_status,
        note=note,
        metadata=metadata or {},
    )


def _legacy_recipient_display_name(user):
    profile = getattr(user, "userprofile", None)
    full_name = user.get_full_name() or user.username
    employee_code = (getattr(profile, "employee_code", "") or "").strip()
    if employee_code:
        return f"{full_name} - {employee_code}"
    return full_name


def _parcel_recipient_display(parcel_receipt):
    full_name = (getattr(parcel_receipt, "recipient_name", "") or "").strip()
    employee_code = (getattr(parcel_receipt, "recipient_employee_code", "") or "").strip()
    if full_name:
        return f"{full_name} - {employee_code}" if employee_code else full_name
    if getattr(parcel_receipt, "recipient_user_id", None):
        return _legacy_recipient_display_name(parcel_receipt.recipient_user)
    return "Chưa gán người nhận"


def _parcel_is_unassigned(parcel_receipt):
    recipient_name = (getattr(parcel_receipt, "recipient_name", "") or "").strip()
    return not getattr(parcel_receipt, "recipient_directory_id", None) and recipient_name in {
        "",
        AdmParcelReceiptForm.UNKNOWN_RECIPIENT_LABEL,
    }


def _parcel_group_key(parcel_receipt):
    if getattr(parcel_receipt, "recipient_directory_id", None):
        return f"dir-{parcel_receipt.recipient_directory_id}"
    fallback_name = (
        getattr(parcel_receipt, "recipient_name", "")
        or _parcel_recipient_display(parcel_receipt)
        or "unknown"
    )
    return f"name-{fallback_name.strip()}"


def _parcel_recipient_receiver_id(parcel_receipt):
    gapo_user_id = (getattr(parcel_receipt, "recipient_gapo_user_id", "") or "").strip()
    if gapo_user_id:
        return gapo_user_id
    profile = getattr(parcel_receipt.recipient_user, "userprofile", None)
    return (getattr(profile, "gapo_user_id", None) or "").strip()


def _build_public_absolute_url(request, path):
    if not path:
        return ""
    if str(path).startswith(("http://", "https://")):
        parsed = urlparse(str(path))
    else:
        base_url = (getattr(settings, "PUBLIC_APP_BASE_URL", "") or "").strip().rstrip("/")
        if base_url:
            parsed = urlparse(urljoin(f"{base_url}/", str(path).lstrip("/")))
        else:
            parsed = urlparse(request.build_absolute_uri(path))

    hostname = (parsed.hostname or "").lower()
    if parsed.scheme == "http" and hostname not in {"127.0.0.1", "localhost", "::1"}:
        parsed = parsed._replace(scheme="https")

    return urlunparse(parsed)


def _build_parcel_confirmation_url(request, token):
    return _build_public_absolute_url(
        request,
        reverse("admindocuments:parcel_receipt_confirm_short", args=[token]),
    )


def _build_parcel_batch_confirmation_url(request, token):
    return _build_public_absolute_url(
        request,
        reverse("admindocuments:parcel_batch_confirm_short", args=[token]),
    )


def _build_parcel_confirmation_qr_svg_url(token):
    return reverse("admindocuments:parcel_receipt_confirm_qr", args=[token])


def _build_parcel_batch_confirmation_qr_svg_url(token):
    return reverse("admindocuments:parcel_batch_confirm_qr", args=[token])


def _build_qr_svg_response(target_url):
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=8,
        border=2,
    )
    qr.add_data(target_url)
    qr.make(fit=True)
    image = qr.make_image(image_factory=SvgPathImage)
    buffer = BytesIO()
    image.save(buffer)
    return HttpResponse(buffer.getvalue(), content_type="image/svg+xml")


def _build_parcel_dynamic_image_url(request):
    relative_path = "admindocuments/parcel-notify-card.svg"
    try:
        asset_url = static(relative_path)
    except ValueError:
        asset_url = f"{settings.STATIC_URL.rstrip('/')}/{relative_path}"
    return _build_public_absolute_url(request, asset_url)


def _get_active_parcel_dynamic_template():
    return AdmParcelDynamicTemplate.objects.filter(
        template_type=AdmParcelDynamicTemplate.TEMPLATE_PARCEL_NOTIFY_CONFIRM,
        is_active=True,
    ).first()


def _normalize_gapo_hex(value, *, fallback):
    text = (value or fallback or "").strip()
    if not text:
        return fallback
    return text if text.startswith("#") else f"#{text}"


def _build_parcel_dynamic_context(parcels, first_parcel, confirm_url):
    company_name = getattr(getattr(first_parcel, "receiving_company", None), "name", "") or ""
    return {
        "recipient_name": _parcel_recipient_display(first_parcel),
        "parcel_count": str(len(parcels)),
        "primary_sender": first_parcel.sender_unit or "Lễ tân tòa nhà",
        "company_name": company_name,
        "confirm_url": confirm_url,
    }


def _ensure_parcel_confirm_url_in_body(body_text, template_text, confirm_url):
    resolved_body = (body_text or "").strip()
    if not resolved_body:
        return "Bạn có bưu kiện mới. Liên hệ lễ tân để nhận và bấm nút bên dưới để xác nhận nhận hàng."
    if "xác nhận nhận hàng" not in resolved_body.lower():
        resolved_body = resolved_body.replace(
            "bấm nút bên dưới để xác nhận.",
            "bấm nút bên dưới để xác nhận nhận hàng.",
        )
    return resolved_body


def _gapo_markdown_link(url, label="tại đây"):
    return f"[{label}]({url})"


def _send_parcel_web_url_button_fallback(receiver_id, confirm_url, request):
    """Send a carousel button using web_url to improve browser open compatibility."""
    image_url = _build_parcel_dynamic_image_url(request)
    send_via_gapo(
        str(receiver_id),
        "Xác nhận nhận hàng tại đây",
        target_type="receiver",
        body_type="carousel",
        body_metadata={
            "metadata": {
                "carousel_cards": [
                    {
                        "title": "Xác nhận nhận bưu kiện",
                        "subtitle": "Bấm nút để mở trang xác nhận trên trình duyệt.",
                        "image_url": image_url,
                        "buttons": [
                            {
                                "title": "Bấm tại đây",
                                "type": "web_url",
                                "payload": confirm_url,
                            }
                        ],
                    }
                ]
            }
        },
    )


def _send_parcel_clickable_link_fallback(receiver_id, confirm_url):
    """Send an explicit markdown plaintext follow-up for clickable link support."""
    send_via_gapo(
        str(receiver_id),
        f"Xác nhận nhận hàng {_gapo_markdown_link(confirm_url)}",
        target_type="receiver",
        body_type="text",
        body_metadata={"is_markdown_text": True},
    )


def _get_active_parcel_auto_notify_setting():
    return (
        AdmParcelAutoNotifySetting.objects.filter(is_active=True)
        .order_by("updated_at", "id")
        .last()
    )


def _next_parcel_reminder_at(base_time, setting):
    threshold = base_time + timedelta(hours=24)
    skip_weekends = getattr(setting, "skip_weekends", True)
    send_slots = getattr(setting, "get_reminder_time_slots", lambda: [])()
    if not send_slots:
        send_slots = [datetime.strptime("09:00", "%H:%M").time()]

    candidate_date = threshold.date()
    while True:
        if skip_weekends and candidate_date.weekday() >= 5:
            candidate_date += timedelta(days=1)
            continue
        for send_time in send_slots:
            candidate = datetime.combine(candidate_date, send_time)
            if timezone.is_aware(threshold):
                candidate = timezone.make_aware(candidate, timezone.get_current_timezone())
            if candidate > threshold:
                return candidate
        candidate_date += timedelta(days=1)


def _cancel_batch_reminder(batch):
    reminder = getattr(batch, "reminder_schedule", None)
    if reminder and reminder.status == GapoScheduledMessage.Status.PENDING:
        reminder.status = GapoScheduledMessage.Status.CANCELLED
        reminder.save(update_fields=["status", "updated_at"])


def _enqueue_gapo_schedule(schedule, *, eta):
    try:
        send_gapo_scheduled_message.apply_async(args=[schedule.id], eta=eta)
    except Exception as exc:
        schedule.status = GapoScheduledMessage.Status.FAILED
        schedule.last_error = str(exc)
        schedule.save(update_fields=["status", "last_error", "updated_at"])
        raise RuntimeError(
            "Không kết nối được hàng đợi nhắc lại. Kiểm tra Redis/Celery."
        ) from exc


def _schedule_parcel_batch_reminder(batch, parcels, request_user, request, *, confirm_url):
    setting = _get_active_parcel_auto_notify_setting()
    if not setting:
        return None

    first = parcels[0]
    receiver_id = _parcel_recipient_receiver_id(first)
    if not receiver_id:
        return None

    schedule_at = _next_parcel_reminder_at(timezone.now(), setting)
    template = _get_active_parcel_dynamic_template() or AdmParcelDynamicTemplate()
    context = _build_parcel_dynamic_context(parcels, first, confirm_url)
    title_text = template.render_text(template.title_template, context)
    body_text = template.render_text(template.body_template, context)
    body_text = _ensure_parcel_confirm_url_in_body(
        body_text,
        template.body_template,
        confirm_url,
    )
    button_text = template.render_text(template.button_text, context)
    image_url = template.hero_image_url or _build_parcel_dynamic_image_url(request)
    reminder_message = (
        f"Nhắc lại: bạn có {len(parcels)} kiện hàng chưa được nhận. "
        f"Liên hệ lễ tân tòa nhà để được hỗ trợ. Xác nhận {_gapo_markdown_link(confirm_url)}"
    )
    body_metadata = {
        "metadata": {
            "layout": {
                "type": "container",
                "direction": "vertical",
                "width": 0,
                "height": 0,
                "alignment": "fill",
                "background": "",
                "item_spacing": 0,
                "insets": "",
                "deep_link": confirm_url,
                "border": {
                    "color": _normalize_gapo_hex(template.card_border_color, fallback="#DADDE1"),
                    "corner_radius": 16,
                    "width": 1,
                },
                "photo_url": "",
                "children": [
                    {
                        "type": "photo",
                        "photo_url": image_url,
                        "height": 200,
                        "width": 0,
                    },
                    {
                        "type": "container",
                        "direction": "vertical",
                        "width": 0,
                        "height": 0,
                        "alignment": "fill",
                        "background": "",
                        "item_spacing": 8,
                        "insets": "16,16,12,16",
                        "children": [
                            {
                                "type": "text",
                                "photo_url": "",
                                "text_object": {
                                    "text": title_text,
                                    "font": {"name": "SF Pro Text", "style": "semibold", "size": 16},
                                    "number_of_lines": 0,
                                    "color": "#10203A",
                                },
                                "width": 0,
                                "height": 0,
                                "item_spacing": 0,
                            },
                            {
                                "type": "text",
                                "photo_url": "",
                                "text_object": {
                                    "text": f"Nhắc lại: {body_text}",
                                    "font": {"name": "SF Pro Text", "style": "regular", "size": 14},
                                    "number_of_lines": 0,
                                    "color": "#5B667A",
                                },
                                "width": 0,
                                "height": 0,
                                "item_spacing": 0,
                            },
                        ],
                    },
                    {
                        "type": "container",
                        "direction": "vertical",
                        "width": 0,
                        "height": 0,
                        "alignment": "fill",
                        "background": "",
                        "item_spacing": 8,
                        "insets": "0,16,16,16",
                        "children": [
                            {
                                "type": "button",
                                "direction": "vertical",
                                "deep_link": confirm_url,
                                "text_object": {
                                    "text": button_text,
                                    "color": _normalize_gapo_hex(
                                        template.button_text_color, fallback="#FFFFFF"
                                    ).lstrip("#"),
                                    "font": {"name": "SF Pro Text", "style": "semibold", "size": 16},
                                    "number_of_lines": 0,
                                },
                                "alignment": "fill",
                                "height": 44,
                                "width": 0,
                                "border": {"width": 0, "color": "", "corner_radius": 10},
                                "background": _normalize_gapo_hex(
                                    template.button_bg_color, fallback="#16A34A"
                                ).lstrip("#"),
                                "photo_url": "",
                            }
                        ],
                    },
                ],
            }
        }
    }
    schedule = GapoScheduledMessage.objects.create(
        receiver_id=str(receiver_id),
        message=reminder_message,
        body_type=GapoScheduledMessage.BodyType.DYNAMIC,
        body_metadata=body_metadata,
        schedule_at=schedule_at,
        created_by=request_user,
    )
    _enqueue_gapo_schedule(schedule, eta=schedule_at)
    batch.reminder_schedule = schedule
    batch.reminder_scheduled_at = schedule_at
    batch.save(update_fields=["reminder_schedule", "reminder_scheduled_at", "updated_at"])
    return schedule


def _parcel_batch_group_key(batch):
    if getattr(batch, "recipient_directory_id", None):
        return f"dir-{batch.recipient_directory_id}"
    return f"name-{(getattr(batch, 'recipient_name', '') or 'unknown').strip()}"


def _cancel_parcel_reminder(parcel_receipt):
    reminder = parcel_receipt.reminder_schedule
    if reminder and reminder.status == GapoScheduledMessage.Status.PENDING:
        reminder.status = GapoScheduledMessage.Status.CANCELLED
        reminder.save(update_fields=["status", "updated_at"])


def _build_parcel_notification_message(parcel_receipt, request):
    receiver_id = _parcel_recipient_receiver_id(parcel_receipt)
    if not receiver_id:
        raise ValueError("Người nhận chưa có GAPO ID.")

    confirm_url = _build_parcel_confirmation_url(request, parcel_receipt.confirmation_token)
    recipient_name = _parcel_recipient_display(parcel_receipt)
    parcel_type_label = parcel_receipt.get_parcel_type_display()
    body = (
        f"Bạn có {parcel_type_label.lower()} mới từ '{parcel_receipt.sender_unit}'. "
        f"Người nhận: {recipient_name}. "
        f"Xác nhận nhận hàng {_gapo_markdown_link(confirm_url)}"
    )
    return str(receiver_id), body, confirm_url


def _schedule_parcel_reminder(parcel_receipt, request_user, request, *, confirm_url=None):
    receiver_id, _, resolved_confirm_url = _build_parcel_notification_message(parcel_receipt, request)
    confirm_url = confirm_url or resolved_confirm_url
    now = timezone.now()
    reminder_schedule = GapoScheduledMessage.objects.create(
        receiver_id=str(receiver_id),
        message=(
            f"Nhắc lại: bạn chưa xác nhận bưu phẩm từ '{parcel_receipt.sender_unit}'. "
            f"Vui lòng xác nhận {_gapo_markdown_link(confirm_url)}"
        ),
        schedule_at=now + timedelta(hours=24),
        created_by=request_user,
    )
    _enqueue_gapo_schedule(reminder_schedule, eta=reminder_schedule.schedule_at)
    parcel_receipt.reminder_schedule = reminder_schedule
    parcel_receipt.reminder_scheduled_at = reminder_schedule.schedule_at
    parcel_receipt.save(update_fields=["reminder_schedule", "reminder_scheduled_at", "updated_at"])
    return reminder_schedule


def _create_parcel_notification_batch(parcels, request_user):
    first = parcels[0]
    batch = AdmParcelNotificationBatch.objects.create(
        recipient_directory=first.recipient_directory,
        recipient_name=(first.recipient_name or _parcel_recipient_display(first) or "").strip(),
        recipient_employee_code=(first.recipient_employee_code or "").strip(),
        recipient_gapo_user_id=_parcel_recipient_receiver_id(first),
        recipient_department=(first.recipient_department or "").strip(),
        parcel_count=len(parcels),
        created_by=request_user,
        updated_by=request_user,
    )
    batch.parcels.set([parcel.id for parcel in parcels])
    return batch


def _serialize_selected_parcel_ids(value):
    ids = []
    for raw in (value or "").split(","):
        raw = raw.strip()
        if raw.isdigit():
            ids.append(int(raw))
    return ids


def _selected_parcels_from_request(request):
    selected_ids = _serialize_selected_parcel_ids(request.POST.get("selected_parcels"))
    if not selected_ids:
        raise ValueError("Cần chọn ít nhất một bưu kiện.")
    parcels = list(
        AdmParcelReceipt.objects.select_related(
            "recipient_directory",
            "recipient_user",
            "recipient_user__userprofile",
            "status",
        ).filter(pk__in=selected_ids)
    )
    if len(parcels) != len(set(selected_ids)):
        raise ValueError("Có bưu kiện không còn tồn tại.")
    return parcels


def _selected_parcels_for_group(request):
    parcels = _selected_parcels_from_request(request)
    group_keys = {_parcel_group_key(parcel) for parcel in parcels}
    if len(group_keys) != 1:
        raise ValueError("Chỉ được thao tác các bưu kiện của cùng một người nhận.")
    return parcels


def _send_parcel_group_notification_now(parcels, request_user, request):
    if not parcels:
        raise ValueError("Không có bưu kiện để gửi thông báo.")
    invalid_statuses = [
        parcel for parcel in parcels if parcel.status_id not in {"pkg_received", "pkg_at_clerical"}
    ]
    if invalid_statuses:
        raise ValueError("Chỉ gửi thông báo cho các bưu kiện chưa xác nhận.")

    first = parcels[0]
    receiver_id = _parcel_recipient_receiver_id(first)
    if not receiver_id:
        raise ValueError("Người nhận chưa có GAPO ID.")

    batch = _create_parcel_notification_batch(parcels, request_user)
    confirm_url = _build_parcel_batch_confirmation_url(request, batch.token)
    parcel_label = "kiện hàng" if len(parcels) > 1 else "kiện hàng"
    message = (
        f"Bạn có {len(parcels)} {parcel_label} chưa được nhận. "
        f"Liên hệ lễ tân tòa nhà để được hỗ trợ. Xác nhận {_gapo_markdown_link(confirm_url)}"
    )
    template = _get_active_parcel_dynamic_template() or AdmParcelDynamicTemplate()
    context = _build_parcel_dynamic_context(parcels, first, confirm_url)
    image_url = template.hero_image_url or _build_parcel_dynamic_image_url(request)
    title_text = template.render_text(template.title_template, context)
    body_text = template.render_text(template.body_template, context)
    body_text = _ensure_parcel_confirm_url_in_body(
        body_text,
        template.body_template,
        confirm_url,
    )
    button_text = template.render_text(template.button_text, context)
    card_border_color = _normalize_gapo_hex(
        template.card_border_color, fallback="#DADDE1"
    )
    button_bg_color = _normalize_gapo_hex(
        template.button_bg_color, fallback="#16A34A"
    )
    button_text_color = _normalize_gapo_hex(
        template.button_text_color, fallback="#FFFFFF"
    )
    body_metadata = {
        "metadata": {
            "layout": {
                "type": "container",
                "direction": "vertical",
                "width": 0,
                "height": 0,
                "alignment": "fill",
                "background": "",
                "item_spacing": 0,
                "insets": "",
                "deep_link": confirm_url,
                "border": {
                    "color": card_border_color,
                    "corner_radius": 16,
                    "width": 1,
                },
                "photo_url": "",
                "children": [
                    {
                        "type": "photo",
                        "photo_url": image_url,
                        "height": 200,
                        "width": 0,
                    },
                    {
                        "type": "container",
                        "direction": "vertical",
                        "width": 0,
                        "height": 0,
                        "alignment": "fill",
                        "background": "",
                        "item_spacing": 8,
                        "insets": "16,16,12,16",
                        "children": [
                            {
                                "type": "text",
                                "photo_url": "",
                                "text_object": {
                                    "text": title_text,
                                    "font": {
                                        "name": "SF Pro Text",
                                        "style": "semibold",
                                        "size": 16,
                                    },
                                    "number_of_lines": 0,
                                    "color": "#10203A",
                                },
                                "width": 0,
                                "height": 0,
                                "item_spacing": 0,
                            },
                            {
                                "type": "text",
                                "photo_url": "",
                                "text_object": {
                                    "text": body_text,
                                    "font": {
                                        "name": "SF Pro Text",
                                        "style": "regular",
                                        "size": 14,
                                    },
                                    "number_of_lines": 0,
                                    "color": "#5B667A",
                                },
                                "width": 0,
                                "height": 0,
                                "item_spacing": 0,
                            },
                        ],
                    },
                    {
                        "type": "container",
                        "direction": "vertical",
                        "width": 0,
                        "height": 0,
                        "alignment": "fill",
                        "background": "",
                        "item_spacing": 8,
                        "insets": "0,16,16,16",
                        "children": [
                            {
                                "type": "button",
                                "direction": "vertical",
                                "deep_link": confirm_url,
                                "text_object": {
                                    "text": button_text,
                                    "color": button_text_color.lstrip("#"),
                                    "font": {
                                        "name": "SF Pro Text",
                                        "style": "semibold",
                                        "size": 16,
                                    },
                                    "number_of_lines": 0,
                                },
                                "alignment": "fill",
                                "height": 44,
                                "width": 0,
                                "border": {
                                    "width": 0,
                                    "color": "",
                                    "corner_radius": 10,
                                },
                                "background": button_bg_color.lstrip("#"),
                                "photo_url": "",
                            }
                        ],
                    },
                ],
            }
        }
    }
    try:
        _send_parcel_clickable_link_fallback(receiver_id, confirm_url)
    except NotificationSendError:
        pass

    response = send_via_gapo(
        str(receiver_id),
        message,
        target_type="receiver",
        body_type="dynamic",
        body_metadata=body_metadata,
    )
    try:
        _send_parcel_web_url_button_fallback(receiver_id, confirm_url, request)
    except NotificationSendError:
        # Keep dynamic message flow resilient even when web_url fallback fails.
        pass
    now = timezone.now()
    batch.message_text = message
    batch.status = AdmParcelNotificationBatch.Status.SENT
    batch.notified_at = now
    batch.notification_send_count = 1
    batch.updated_by = request_user
    batch.save(
        update_fields=[
            "message_text",
            "status",
            "notified_at",
            "notification_send_count",
            "updated_by",
            "updated_at",
        ]
    )
    reminder_error = ""
    try:
        _schedule_parcel_batch_reminder(batch, parcels, request_user, request, confirm_url=confirm_url)
    except Exception as exc:
        reminder_error = str(exc)

    for parcel in parcels:
        previous_status = parcel.status_id
        parcel.notified_at = now
        if parcel.status_id == "pkg_received":
            parcel.status_id = "pkg_at_clerical"
        parcel.updated_by = request_user
        parcel.save(update_fields=["notified_at", "status", "updated_by", "updated_at"])
        _log_parcel_event(
            parcel,
            AdmParcelReceiptLog.ACTION_NOTIFIED,
            actor=request_user,
            from_status=previous_status,
            to_status=parcel.status_id,
            note=f"Đã gửi thông báo nhận theo lô #{batch.id}.",
            metadata={"batch_id": batch.id, "parcel_count": len(parcels)},
        )
    return batch, response, reminder_error


def _send_parcel_notification_now(parcel_receipt, request_user, request):
    receiver_id, body, confirm_url = _build_parcel_notification_message(parcel_receipt, request)
    try:
        _send_parcel_clickable_link_fallback(receiver_id, confirm_url)
    except NotificationSendError:
        pass
    response = send_via_gapo(str(receiver_id), body, target_type="receiver")
    try:
        _send_parcel_web_url_button_fallback(receiver_id, confirm_url, request)
    except NotificationSendError:
        pass
    now = timezone.now()
    notification_schedule = parcel_receipt.notification_schedule
    if notification_schedule:
        notification_schedule.receiver_id = str(receiver_id)
        notification_schedule.message = body
        notification_schedule.schedule_at = now
        notification_schedule.status = GapoScheduledMessage.Status.SENT
        notification_schedule.sent_at = now
        notification_schedule.last_error = ""
        notification_schedule.created_by = request_user
        notification_schedule.save(
            update_fields=[
                "receiver_id",
                "message",
                "schedule_at",
                "status",
                "sent_at",
                "last_error",
                "created_by",
                "updated_at",
            ]
        )
    else:
        notification_schedule = GapoScheduledMessage.objects.create(
            receiver_id=str(receiver_id),
            message=body,
            schedule_at=now,
            status=GapoScheduledMessage.Status.SENT,
            sent_at=now,
            created_by=request_user,
        )

    reminder_error = ""
    try:
        _cancel_parcel_reminder(parcel_receipt)
        _schedule_parcel_reminder(
            parcel_receipt,
            request_user,
            request,
            confirm_url=confirm_url,
        )
    except Exception as exc:
        reminder_error = str(exc)
    parcel_receipt.notification_schedule = notification_schedule
    parcel_receipt.notified_at = now
    from_status = parcel_receipt.status_id
    parcel_receipt.status_id = "pkg_at_clerical"
    parcel_receipt.updated_by = request_user
    parcel_receipt.save(
        update_fields=[
            "notification_schedule",
            "notified_at",
            "status",
            "updated_by",
            "updated_at",
        ]
    )
    _log_parcel_event(
        parcel_receipt,
        AdmParcelReceiptLog.ACTION_NOTIFIED,
        actor=request_user,
        from_status=from_status,
        to_status=parcel_receipt.status_id,
        note="Đã gửi thông báo GAPO cho người nhận.",
    )
    return response, reminder_error


def _incoming_flow_by_type(item_type_code, signer_label="Người ký"):
    if item_type_code == INCOMING_TYPE_DOC:
        return list(INCOMING_DOC_FLOW)
    if item_type_code == INCOMING_TYPE_PARCEL:
        return list(INCOMING_PARCEL_FLOW)
    return list(INCOMING_DOC_FLOW)


def _resolve_initial_incoming_status_code(item_type_code):
    flow = _incoming_flow_by_type(item_type_code)
    if flow:
        first_code = flow[0][0]
        if AdmIncomingDispatchStatus.objects.filter(code=first_code, is_active=True).exists():
            return first_code
    fallback = (
        AdmIncomingDispatchStatus.objects.filter(is_active=True)
        .order_by("sort_order", "name")
        .values_list("code", flat=True)
        .first()
    )
    return fallback


def _normalize_incoming_dispatch_status_code(item_type_code, status_code):
    if item_type_code == INCOMING_TYPE_DOC:
        return LEGACY_INCOMING_DOC_STATUS_MAP.get(status_code, status_code)
    return status_code


def _incoming_flow_label(item_type_code, status_code, signer_label="Người ký"):
    normalized_code = _normalize_incoming_dispatch_status_code(item_type_code, status_code)
    for code, label in _incoming_flow_by_type(item_type_code, signer_label=signer_label):
        if code == normalized_code:
            return label
    return ""


def _incoming_status_badge_tokens(status):
    return {
        "bg": (getattr(status, "badge_bg_color", "") or "#F3F4F6").strip(),
        "text": (getattr(status, "badge_text_color", "") or "#374151").strip(),
    }


def _allowed_status_codes_for_dispatch(dispatch):
    signer_label = (dispatch.signer_name or "").strip() or "Người ký"
    flow = _incoming_flow_by_type(dispatch.incoming_item_type_id, signer_label=signer_label)
    codes = [code for code, _ in flow]
    initial_code = flow[0][0] if flow else ""
    if dispatch.incoming_item_type_id == INCOMING_TYPE_DOC and initial_code:
        codes = [code for code in codes if code != initial_code]
    normalized_current = _normalize_incoming_dispatch_status_code(
        dispatch.incoming_item_type_id,
        dispatch.status_id,
    )
    if normalized_current and normalized_current not in codes and normalized_current != initial_code:
        codes.append(normalized_current)
    return codes


def _allowed_status_codes_for_parcel(parcel_receipt):
    codes = [code for code, _ in INCOMING_PARCEL_FLOW]
    if parcel_receipt.status_id and parcel_receipt.status_id not in codes:
        codes.append(parcel_receipt.status_id)
    return codes



def _build_incoming_workflow_steps(dispatch):
    current_status = _normalize_incoming_dispatch_status_code(
        dispatch.incoming_item_type_id,
        dispatch.status_id,
    )
    signer_label = (dispatch.signer_name or "").strip() or "Người ký"
    flow = _incoming_flow_by_type(dispatch.incoming_item_type_id, signer_label=signer_label)
    logs = list(
        getattr(dispatch, "status_logs_display", None)
        or dispatch.status_logs.select_related("to_status").all()
    )
    first_code = flow[0][0] if flow else ""
    received_timestamp = getattr(dispatch, "created_at", None)
    if received_timestamp is None and getattr(dispatch, "received_date", None):
        received_timestamp = datetime.combine(dispatch.received_date, datetime.min.time())
    timestamp_map = {first_code: received_timestamp} if first_code else {}
    for log in reversed(logs):
        normalized_to = _normalize_incoming_dispatch_status_code(
            dispatch.incoming_item_type_id,
            log.to_status_id,
        )
        if normalized_to and normalized_to not in timestamp_map:
            timestamp_map[normalized_to] = log.changed_at
    if current_status and current_status not in timestamp_map and getattr(dispatch, "updated_at", None):
        timestamp_map[current_status] = dispatch.updated_at
    index_map = {code: idx for idx, (code, _) in enumerate(flow)}
    current_index = index_map.get(current_status, -1)
    steps = []
    for idx, (code, label) in enumerate(flow):
        steps.append(
            {
                "code": code,
                "label": label,
                "is_current": idx == current_index,
                "is_done": current_index > idx,
                "timestamp": timestamp_map.get(code),
            }
        )
    return steps


def _build_parcel_workflow_steps(parcel_receipt):
    current_status = "pkg_done" if parcel_receipt.status_id in {"pkg_processing", "pkg_done"} else parcel_receipt.status_id
    flow = list(INCOMING_PARCEL_FLOW)
    index_map = {code: idx for idx, (code, _) in enumerate(flow)}
    current_index = index_map.get(current_status, -1)
    timestamp_map = {
        "pkg_received": parcel_receipt.received_at,
        "pkg_at_clerical": parcel_receipt.notified_at,
        "pkg_done": parcel_receipt.completed_at or parcel_receipt.confirmed_at,
    }
    steps = []
    for idx, (code, label) in enumerate(flow):
        steps.append(
            {
                "code": code,
                "label": label,
                "is_current": idx == current_index,
                "is_done": current_index > idx,
                "timestamp": timestamp_map.get(code),
            }
        )
    return steps


def _parcel_status_tone(parcel_receipt):
    normalized_status = "pkg_done" if parcel_receipt.status_id in {"pkg_processing", "pkg_done"} else parcel_receipt.status_id
    if normalized_status == "pkg_received":
        return "received"
    if normalized_status == "pkg_done":
        return "done"
    return "notified"


def _parcel_actual_receiver_display(parcel_receipt):
    actual_name = (parcel_receipt.actual_receiver_name or "").strip()
    actual_code = (parcel_receipt.actual_receiver_employee_code or "").strip()
    if not actual_name:
        return _parcel_recipient_display(parcel_receipt) or "Chưa xác nhận"
    intended_name = (parcel_receipt.recipient_name or "").strip()
    intended_code = (parcel_receipt.recipient_employee_code or "").strip()
    is_proxy = (
        (actual_code and intended_code and actual_code != intended_code)
        or (actual_name and intended_name and actual_name != intended_name)
        or (bool(parcel_receipt.proxy_receiver_name) and actual_name == parcel_receipt.proxy_receiver_name)
    )
    label = actual_name
    if actual_code:
        label = f"{label} - {actual_code}"
    if is_proxy:
        label = f"{label} (nhận hộ)"
    return label


def _parcel_extra_audit_logs(parcel_receipt):
    duplicate_actions = {
        AdmParcelReceiptLog.ACTION_CREATED,
        AdmParcelReceiptLog.ACTION_NOTIFIED,
        AdmParcelReceiptLog.ACTION_CONFIRMED,
        AdmParcelReceiptLog.ACTION_HANDED_OVER,
    }
    return [log for log in parcel_receipt.audit_logs.all() if log.action not in duplicate_actions]


def _parcel_type_tone(parcel_receipt):
    if parcel_receipt.parcel_type == AdmParcelReceipt.ParcelType.DOSSIER:
        return "document"
    if parcel_receipt.parcel_type == AdmParcelReceipt.ParcelType.GOODS:
        return "goods"
    return "other"


def _parcel_reminder_meta(reminder_schedule, reminder_scheduled_at, *, notification_send_count=0):
    meta = {
        "badge_label": "",
        "badge_tone": "",
        "status_label": "Chưa lên lịch nhắc",
        "scheduled_at": reminder_scheduled_at,
        "scheduled_label": reminder_scheduled_at.strftime("%d/%m/%Y %H:%M")
        if reminder_scheduled_at
        else "",
        "last_error": "",
    }
    if not reminder_schedule:
        if notification_send_count:
            meta["badge_label"] = "Chưa lên lịch nhắc"
            meta["badge_tone"] = "warning"
        return meta

    meta["last_error"] = (reminder_schedule.last_error or "").strip()
    if reminder_schedule.status == GapoScheduledMessage.Status.PENDING:
        meta["badge_label"] = "Đã lên lịch nhắc"
        meta["badge_tone"] = "scheduled"
        meta["status_label"] = "Đã lên lịch nhắc"
    elif reminder_schedule.status == GapoScheduledMessage.Status.SENT:
        meta["badge_label"] = "Đã nhắc lại"
        meta["badge_tone"] = "done"
        meta["status_label"] = "Đã gửi nhắc lại"
    elif reminder_schedule.status == GapoScheduledMessage.Status.FAILED:
        meta["badge_label"] = "Chưa lên lịch nhắc"
        meta["badge_tone"] = "warning"
        meta["status_label"] = "Lên lịch nhắc lỗi"
    elif reminder_schedule.status == GapoScheduledMessage.Status.CANCELLED:
        meta["badge_label"] = "Đã hủy lịch nhắc"
        meta["badge_tone"] = "muted"
        meta["status_label"] = "Đã hủy lịch nhắc"
    return meta


def _prepare_parcel_receipt_view_state(parcel_receipt, request=None):
    parcel_receipt.workflow_steps = _build_parcel_workflow_steps(parcel_receipt)
    parcel_receipt.extra_audit_logs = _parcel_extra_audit_logs(parcel_receipt)
    parcel_receipt.actual_receiver_display = _parcel_actual_receiver_display(parcel_receipt)
    parcel_receipt.status_tone = _parcel_status_tone(parcel_receipt)
    parcel_receipt.recipient_label = _parcel_recipient_display(parcel_receipt)
    parcel_receipt.is_unassigned = _parcel_is_unassigned(parcel_receipt)
    parcel_receipt.recipient_group_key = _parcel_group_key(parcel_receipt)
    parcel_receipt.parcel_type_tone = _parcel_type_tone(parcel_receipt)
    latest_batch = None
    if hasattr(parcel_receipt, "_prefetched_objects_cache") and "notification_batches" in parcel_receipt._prefetched_objects_cache:
        prefetched_batches = list(parcel_receipt.notification_batches.all())
        latest_batch = prefetched_batches[0] if prefetched_batches else None
    parcel_receipt.latest_notification_batch = latest_batch
    if latest_batch:
        parcel_receipt.reminder_meta = _parcel_reminder_meta(
            latest_batch.reminder_schedule,
            latest_batch.reminder_scheduled_at,
            notification_send_count=latest_batch.notification_send_count or 0,
        )
    else:
        parcel_receipt.reminder_meta = _parcel_reminder_meta(
            parcel_receipt.reminder_schedule,
            parcel_receipt.reminder_scheduled_at,
            notification_send_count=1 if parcel_receipt.notified_at else 0,
        )
    if request and parcel_receipt.confirmation_token:
        parcel_receipt.confirmation_url = _build_parcel_confirmation_url(
            request, parcel_receipt.confirmation_token
        )
        parcel_receipt.confirmation_qr_svg_url = _build_parcel_confirmation_qr_svg_url(
            parcel_receipt.confirmation_token
        )
    else:
        parcel_receipt.confirmation_url = ""
        parcel_receipt.confirmation_qr_svg_url = ""
    return parcel_receipt


def _build_parcel_recipient_groups(parcel_receipts, *, split_by_day=False, request=None):
    recipient_group_map = {}
    recipient_groups = []
    grouping_mode = "day" if split_by_day else "recipient"
    batch_queryset = (
        AdmParcelNotificationBatch.objects.filter(parcels__in=parcel_receipts)
        .select_related("reminder_schedule")
        .prefetch_related("parcels")
        .distinct()
        .order_by("-created_at")
    )

    for parcel_receipt in parcel_receipts:
        recipient_key = _parcel_group_key(parcel_receipt)
        day_key = (
            parcel_receipt.received_at.date().isoformat()
            if parcel_receipt.received_at
            else "unknown"
        )
        group_key = day_key if split_by_day else recipient_key
        group = recipient_group_map.get(group_key)
        if not group:
            group = {
                "key": group_key,
                "recipient_label": _parcel_recipient_display(parcel_receipt),
                "department": (parcel_receipt.recipient_department or "").strip() or "Chưa có phòng ban",
                "parcels": [],
                "received_day": parcel_receipt.received_at.date() if parcel_receipt.received_at else None,
                "recipient_keys": set(),
            }
            recipient_group_map[group_key] = group
            recipient_groups.append(group)
        group["parcels"].append(parcel_receipt)
        group["recipient_keys"].add(recipient_key)

    for group in recipient_groups:
        group["parcel_ids"] = [parcel.id for parcel in group["parcels"]]
        group["count"] = len(group["parcels"])
        group["parcel_ids_set"] = set(group["parcel_ids"])
        group["recipient_count"] = len(group["recipient_keys"])
        group["notified_count"] = sum(
            1 for parcel in group["parcels"] if parcel.status_id == "pkg_at_clerical"
        )
        group["confirmed_count"] = sum(
            1 for parcel in group["parcels"] if parcel.status_id in {"pkg_processing", "pkg_done"}
        )
        group["notifiable_parcel_ids"] = [
            parcel.id
            for parcel in group["parcels"]
            if parcel.status_id in {"pkg_received", "pkg_at_clerical"}
        ]
        group["notifiable_count"] = len(group["notifiable_parcel_ids"])
        group["status_label"] = (
            "Đã thông báo"
            if group["notified_count"] and group["notifiable_count"] == 0
            else "Chưa nhận"
        )
        group["status_tone"] = "notified" if group["notified_count"] else "received"
        group["supports_batch_actions"] = True
        group["has_unassigned"] = any(parcel.is_unassigned for parcel in group["parcels"])
        if grouping_mode == "day":
            group["recipient_label"] = (
                f"Nhận ngày {group['received_day']:%d/%m/%Y}"
                if group["received_day"]
                else "Chưa rõ ngày nhận"
            )
            group["department"] = f"{group['recipient_count']} người nhận · {group['count']} kiện"

        related_batches = []
        for batch in batch_queryset:
            batch_parcel_ids = {parcel.id for parcel in batch.parcels.all()}
            if batch_parcel_ids & group["parcel_ids_set"]:
                related_batches.append(batch)
            if len(related_batches) >= 5:
                break
        group["notification_batches"] = related_batches
        group["notification_send_count"] = sum(
            batch.notification_send_count or 1 for batch in related_batches
        )
        group["latest_notified_at"] = (
            group["notification_batches"][0].notified_at
            if group["notification_batches"]
            else None
        )
        latest_batch = group["notification_batches"][0] if group["notification_batches"] else None
        if latest_batch:
            group["reminder_meta"] = _parcel_reminder_meta(
                latest_batch.reminder_schedule,
                latest_batch.reminder_scheduled_at,
                notification_send_count=latest_batch.notification_send_count or 0,
            )
        else:
            group["reminder_meta"] = _parcel_reminder_meta(
                None,
                None,
                notification_send_count=group["notification_send_count"],
            )
        if request and group["notification_batches"] and group["recipient_count"] == 1:
            group["latest_confirm_url"] = _build_parcel_batch_confirmation_url(
                request, latest_batch.token
            )
            group["latest_confirm_qr_svg_url"] = _build_parcel_batch_confirmation_qr_svg_url(
                latest_batch.token
            )
        else:
            group["latest_confirm_url"] = ""
            group["latest_confirm_qr_svg_url"] = ""
    return recipient_groups


def _create_attachment_version(document, user, uploaded_file=None, link=None, note=None):
    """Create a new attachment version (file or link) for a document."""
    if not uploaded_file and not link:
        return None

    if uploaded_file:
        uploaded_file.seek(0)

    current_max = (
        document.attachments.aggregate(max_version=Max("version")).get("max_version")
        or 0
    )
    new_version = current_max + 1
    document.attachments.filter(is_latest=True).update(is_latest=False)
    attachment = AdmDocumentAttachment.objects.create(
        document=document,
        version=new_version,
        file=uploaded_file if uploaded_file else None,
        original_name=getattr(uploaded_file, "name", None),
        link=link or None,
        is_latest=True,
        is_deleted=False,
        created_by=user,
        note=note,
    )
    if uploaded_file:
        # keep compatibility field pointing to latest file
        document.attachment = attachment.file
        document.save(update_fields=["attachment"])
    return attachment


PAPER_IMPORT_HEADERS = (
    "Số hiệu (để trống nếu muốn hệ thống sinh)",
    "Loại giấy", "Miền", "Phòng giao dịch", "Phòng ban nội bộ", "Người phụ trách",
    "Nội dung", "Đơn vị CPN", "Mã vận đơn", "Mã Ticket", "Ghi chú", "Ngày tạo (yyyy-mm-dd)",
)


def _parse_paper_rows(ws, request_user):
    """Parse by column name, supporting new templates and optional legacy status."""
    required_headers = {header.casefold() for header in PAPER_IMPORT_HEADERS} - {"mã ticket"}
    allowed_headers = required_headers | {"mã ticket", "tình trạng"}
    header_row = [str(value).strip().casefold() if value is not None else ""
                  for value in next(ws.iter_rows(max_row=1, values_only=True), ())]
    while header_row and not header_row[-1]:
        header_row.pop()
    if (required_headers - set(header_row) or set(header_row) - allowed_headers
            or len(header_row) != len(set(header_row))):
        return ["Header không đúng định dạng mẫu, vui lòng tải file mẫu mới nhất."], []

    paper_type_map = {p.name.lower(): p for p in AdmPaperType.objects.all()}
    shop_map = {s.shop_name.lower(): s for s in Shop.objects.all()}
    dept_map = {d.name.lower(): d for d in AdmDepartment.objects.all()}
    courier_map = {c.name.lower(): c for c in AdmCourierCompany.objects.all()}
    region_map = {r.region_name.lower(): r for r in Region.objects.all()}

    today = timezone.localdate() if settings.USE_TZ else date.today()

    errors = []
    rows_data = []

    for idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
            continue
        columns = dict(zip(header_row, values))
        number_full = columns.get("số hiệu (để trống nếu muốn hệ thống sinh)")
        paper_type_name = columns.get("loại giấy")
        region = columns.get("miền")
        requested_dept_name = columns.get("phòng giao dịch")
        internal_dept_name = columns.get("phòng ban nội bộ")
        responsible = columns.get("người phụ trách")
        summary = columns.get("nội dung")
        courier_name = columns.get("đơn vị cpn")
        tracking_code = columns.get("mã vận đơn")
        ticket_value = columns.get("mã ticket")
        ticket_code = str(ticket_value).strip() if ticket_value is not None else ""
        status_text = columns.get("tình trạng")
        note = columns.get("ghi chú")
        created_str = columns.get("ngày tạo (yyyy-mm-dd)")

        row_errors = []
        if any(value is not None and str(value).strip() for value in values[len(header_row):]):
            row_errors.append("Có dữ liệu ở cột không có tiêu đề.")
        if len(ticket_code) > AdmPaperDocument._meta.get_field("ticket_code").max_length:
            row_errors.append("Mã Ticket không được vượt quá 200 ký tự.")

        pt_key = str(paper_type_name).strip().lower() if paper_type_name else ""
        paper_type = paper_type_map.get(pt_key)
        if not paper_type:
            row_errors.append("Loại giấy không hợp lệ/để trống.")

        region_val = str(region).strip() if region else ""
        region_obj = None
        if not region_val:
            row_errors.append("Miền không được để trống.")
        else:
            region_obj = region_map.get(region_val.lower())
            if not region_obj:
                row_errors.append(f"Miền '{region_val}' không tồn tại.")

        requested_dept = None
        if requested_dept_name:
            requested_dept = shop_map.get(str(requested_dept_name).strip().lower())
            if not requested_dept:
                row_errors.append(f"Phòng giao dịch '{requested_dept_name}' không tồn tại.")

        internal_dept = None
        if internal_dept_name:
            internal_dept = dept_map.get(str(internal_dept_name).strip().lower())
            if not internal_dept:
                row_errors.append(f"Phòng ban nội bộ '{internal_dept_name}' không tồn tại.")

        if not requested_dept and not internal_dept:
            row_errors.append("Cần chọn Phòng giao dịch hoặc Phòng ban nội bộ (có thể chọn cả hai).")

        courier = None
        if courier_name:
            courier = courier_map.get(str(courier_name).strip().lower())
            if not courier:
                row_errors.append(f"Đơn vị CPN '{courier_name}' không tồn tại.")

        try:
            created_date = (
                datetime.strptime(str(created_str), "%Y-%m-%d").date() if created_str else today
            )
        except ValueError:
            row_errors.append("Ngày tạo không đúng định dạng yyyy-mm-dd.")

        summary_text = str(summary).strip() if summary else ""
        if not summary_text:
            row_errors.append("Nội dung không được để trống.")

        if row_errors:
            errors.append(f"Dòng {idx}: " + "; ".join(row_errors))
            continue

        rows_data.append(
            {
                "paper_type_id": paper_type.id,
                "paper_type_code": paper_type.code,
                "region": region_val,
                "requested_dept_id": requested_dept.pk if requested_dept else None,
                "internal_dept_id": internal_dept.id if internal_dept else None,
                "responsible": str(responsible).strip()
                if responsible
                else (request_user.get_full_name() or request_user.username),
                "summary": summary_text,
                "courier_id": courier.id if courier else None,
                "tracking_code": str(tracking_code).strip() if tracking_code else None,
                "ticket_code": ticket_code,
                "status": str(status_text).strip() if status_text else "",
                "note": str(note).strip() if note else "",
                "created_date": created_date,
                "number_full": str(number_full).strip() if number_full else None,
            }
        )

    return errors, rows_data


def admin_staff_required(view_func):
    @wraps(view_func)
    def _wrapped(request, *args, **kwargs):
        if not request.user.is_authenticated:
            raise PermissionDenied
        if _has_admin_docs_access(request.user):
            return view_func(request, *args, **kwargs)
        messages.error(request, "Bạn không có quyền truy cập mục này.")
        return render(request, "403.html", status=403)

    return _wrapped


def superuser_required(view_func):
    @wraps(view_func)
    def _wrapped(request, *args, **kwargs):
        if not request.user.is_authenticated or not request.user.is_superuser:
            messages.error(request, "Chỉ super admin mới có quyền truy cập mục này.")
            return render(request, "403.html", status=403)
        return view_func(request, *args, **kwargs)

    return _wrapped


def _safe_sheet_cell_value(ws, row_idx, col_idx, merged_lookup):
    cell = ws.cell(row_idx, col_idx)
    if cell.value is not None:
        return cell.value
    anchor = merged_lookup.get((row_idx, col_idx))
    if anchor:
        return ws.cell(*anchor).value
    return None


def _normalize_excel_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    text = str(value).strip()
    if text.lower() == "none":
        return ""
    return text


def _normalize_excel_header(value):
    """Normalize headers while retaining Vietnamese accents for exact aliases."""
    return " ".join(
        _normalize_excel_text(value).replace("\ufeff", "").split()
    ).casefold()


def _normalize_phone_number(value):
    return "".join(ch for ch in _normalize_excel_text(value) if ch.isdigit())


def _parcel_department_leaf(department_full):
    parts = [part.strip() for part in (department_full or "").split("||") if part.strip()]
    return parts[-1] if parts else ""


def _parse_parcel_recipient_workbook(file_path):
    # Own the file handle so Windows can release the uploaded snapshot immediately.
    if hasattr(file_path, "read"):
        workbook = load_workbook(file_path, data_only=True)
    else:
        with open(file_path, "rb") as source:
            workbook = load_workbook(source, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    merged_lookup = {}
    for cell_range in worksheet.merged_cells.ranges:
        anchor = (cell_range.min_row, cell_range.min_col)
        for row_idx in range(cell_range.min_row, cell_range.max_row + 1):
            for col_idx in range(cell_range.min_col, cell_range.max_col + 1):
                merged_lookup[(row_idx, col_idx)] = anchor

    # IT exports have changed the user-id label and can prepend report-title rows.
    # Detect the actual header row instead of assuming it is always the first row.
    header_row = 1
    raw_headers = []
    best_header_count = -1
    for candidate_row in range(1, min(worksheet.max_row, 10) + 1):
        candidate_headers = [
            _normalize_excel_header(
                _safe_sheet_cell_value(
                    worksheet, candidate_row, col_idx, merged_lookup
                )
            )
            for col_idx in range(1, worksheet.max_column + 1)
        ]
        recognized_count = sum(
            header in PARCEL_RECIPIENT_IMPORT_HEADERS
            for header in candidate_headers
        )
        if recognized_count > best_header_count:
            header_row = candidate_row
            raw_headers = candidate_headers
            best_header_count = recognized_count

    header_map = {}
    stt_col = None
    for idx, header in enumerate(raw_headers, 1):
        key = PARCEL_RECIPIENT_IMPORT_HEADERS.get(header)
        if key:
            header_map[idx] = key
            if key == "stt":
                stt_col = idx

    required_fields = {"gapo_user_id", "full_name", "department_full"}
    missing_fields = required_fields.difference(header_map.values())
    if missing_fields:
        missing_labels = ", ".join(sorted(missing_fields))
        raise ValueError(
            "File Excel không đúng cấu trúc danh bạ người nhận; "
            f"không nhận diện được các cột: {missing_labels}."
        )

    rows = []
    stats = {
        "merged_ranges": len(list(worksheet.merged_cells.ranges)),
        "missing_gapo": 0,
        "missing_department": 0,
        "inactive_rows": 0,
        "header_row": header_row,
        "recognized_headers": len(header_map),
    }
    for row_idx in range(header_row + 1, worksheet.max_row + 1):
        # Skip continuation rows (users with multiple departments have merged cells;
        # raw STT cell is None for those extra rows).
        if stt_col and worksheet.cell(row_idx, stt_col).value is None:
            continue

        record = {}
        for col_idx, field_name in header_map.items():
            record[field_name] = _normalize_excel_text(
                _safe_sheet_cell_value(worksheet, row_idx, col_idx, merged_lookup)
            )

        if not any(record.values()):
            continue

        full_name = record.get("full_name", "")
        gapo_user_id = record.get("gapo_user_id", "")
        department_full = record.get("department_full", "")
        employment_status = record.get("employment_status", "")
        is_active_member = employment_status.lower() == "đang hoạt động"
        department_name = _parcel_department_leaf(department_full)

        if not gapo_user_id:
            stats["missing_gapo"] += 1
        if not department_name:
            stats["missing_department"] += 1
        if not is_active_member:
            stats["inactive_rows"] += 1

        if not full_name and not gapo_user_id:
            continue

        record["department_name"] = department_name
        record["is_active_member"] = is_active_member
        record["row_number"] = row_idx
        rows.append(record)

    stats["department_rows"] = sum(
        bool(row.get("department_name")) for row in rows
    )
    stats["department_count"] = len(
        {
            row["department_name"]
            for row in rows
            if row.get("department_name")
        }
    )

    return {
        "sheet_name": worksheet.title,
        "total_rows": max(worksheet.max_row - header_row, 0),
        "rows": rows,
        "summary": stats,
    }


@transaction.atomic
def _activate_parcel_recipient_batch(batch):
    AdmParcelRecipientImportBatch.objects.filter(is_current=True).exclude(
        pk=batch.pk
    ).update(is_current=False)
    batch.is_current = True
    batch.activated_at = timezone.now()
    batch.save(update_fields=["is_current", "activated_at", "updated_at"])


@transaction.atomic
def _import_parcel_recipient_batch(batch):
    parsed = _parse_parcel_recipient_workbook(batch.file.path)
    rows = parsed["rows"]
    batch.sheet_name = parsed["sheet_name"]
    batch.total_rows = parsed["total_rows"]
    batch.imported_rows = len(rows)
    batch.active_rows = sum(1 for row in rows if row["is_active_member"])
    batch.has_errors = False
    batch.summary = parsed["summary"]
    batch.save(
        update_fields=[
            "sheet_name",
            "total_rows",
            "imported_rows",
            "active_rows",
            "has_errors",
            "summary",
            "updated_at",
        ]
    )
    batch.recipients.all().delete()
    AdmParcelRecipientCatalog.objects.bulk_create(
        [
            AdmParcelRecipientCatalog(
                import_batch=batch,
                row_number=row["row_number"],
                gapo_user_id=row.get("gapo_user_id", ""),
                employee_code=row.get("employee_code", ""),
                full_name=row.get("full_name", ""),
                email=row.get("email", ""),
                phone_number=row.get("phone_number", ""),
                phone_number_normalized=_normalize_phone_number(row.get("phone_number", "")),
                employment_status=row.get("employment_status", ""),
                permission_name=row.get("permission_name", ""),
                org_chart=row.get("org_chart", ""),
                position_name=row.get("position_name", ""),
                department_full=row.get("department_full", ""),
                department_name=row.get("department_name", ""),
                region_name=row.get("region_name", ""),
                birth_date=row.get("birth_date", ""),
                company_join_date=row.get("company_join_date", ""),
                contract_start_date=row.get("contract_start_date", ""),
                leave_date=row.get("leave_date", ""),
                source_created_at=row.get("source_created_at", ""),
                is_active_member=row.get("is_active_member", False),
                raw_data=row,
            )
            for row in rows
        ],
        batch_size=500,
    )
    _activate_parcel_recipient_batch(batch)
    return batch


@login_required
def workflow_hub(request):
    can_manage_documents = _has_admin_docs_access(request.user)
    workflows = [
        {"key": "document", "title": "Văn bản hành chính",
         "description": "Tạo văn bản, cấp số hiệu và lưu trữ tệp đính kèm theo phiên bản.",
         "route": "admindocuments:admindocuments_list", "allowed": can_manage_documents},
        {"key": "paper", "title": "Giấy tờ hành chính",
         "description": "Tạo và quản lý các loại giấy tờ hành chính theo danh mục.",
         "route": "admindocuments:paper_document_list", "allowed": can_manage_documents},
        {"key": "incoming", "title": "Công văn đến",
         "description": "Tiếp nhận công văn và theo dõi quá trình chuyển giao, xử lý.",
         "route": "admindocuments:incoming_document_list", "allowed": True},
        {"key": "parcel", "title": "Bưu kiện đến",
         "description": "Tiếp nhận bưu kiện, xác định người nhận và theo dõi bàn giao.",
         "route": "admindocuments:parcel_receipt_list", "allowed": True},
    ]
    for workflow in workflows:
        url = reverse(workflow["route"])
        workflow["create_url"] = f"{url}?create=1"
        workflow["list_url"] = reverse("admindocuments:parcel_receipt_completed_list") if workflow["key"] == "parcel" else url
        workflow["list_allowed"] = workflow["allowed"] and (workflow["key"] != "parcel" or can_manage_documents)
        if workflow["key"] in {"document", "paper"}:
            tab = "admin" if workflow["key"] == "document" else "paper"
            workflow["counter_url"] = f"{reverse('admindocuments:admindocuments_counters')}?tab={tab}"
            workflow["counter_label"] = "Số hiệu văn bản hành chính" if tab == "admin" else "Số hiệu giấy tờ hành chính"
    return render(request, "admindocuments/workflow_hub.html", {
        "workflows": workflows,
        "has_admin_docs_access": can_manage_documents,
        "is_super_admin": request.user.is_superuser,
    })


def _incoming_dispatch_list(request, item_type_code):
    can_create = _has_incoming_dispatch_create_access(request.user)
    can_edit_status = _has_incoming_dispatch_status_edit_access(request.user)
    screen_meta = INCOMING_SCREEN_META.get(item_type_code, INCOMING_SCREEN_META[INCOMING_TYPE_DOC])

    query = request.GET.get("q", "").strip()
    status_filter = request.GET.get("status", "")
    department_filter = request.GET.get("department", "")
    company_filter = request.GET.get("company", "")
    receive_company = (request.GET.get("receive_company") or "").strip()
    open_dispatch_id = (request.GET.get("open") or "").strip()
    companies = AdmCompany.objects.filter(is_active=True).order_by("name")
    prefill_company = None

    if receive_company:
        prefill_company = companies.filter(code=receive_company).first()

    if request.method == "POST":
        form = AdmIncomingDispatchForm(request.POST, item_type_code=item_type_code)
        receive_company = (request.POST.get("receiving_company") or "").strip()
        prefill_company = companies.filter(code=receive_company).first()
        if form.is_valid():
            dispatch = form.save(commit=False)
            dispatch.responsible_user = request.user
            dispatch.created_by = request.user
            dispatch.updated_by = request.user
            initial_status = _resolve_initial_incoming_status_code(
                dispatch.incoming_item_type_id
            )
            if not initial_status:
                messages.error(
                    request,
                    "Chưa có cấu hình trạng thái khởi tạo cho tiếp nhận thư từ.",
                )
                return redirect(_incoming_dispatch_route_name(item_type_code))
            dispatch.status_id = initial_status
            dispatch.save()
            form.save_m2m()
            files = request.FILES.getlist("images")
            file_error = _validate_incoming_dispatch_files(files) if files else ""
            if file_error:
                messages.error(request, file_error)
                return redirect(_incoming_dispatch_route_name(item_type_code))
            _save_incoming_dispatch_images(dispatch, files, request.user)
            messages.success(request, screen_meta["success_message"])
            return redirect(_incoming_dispatch_route_name(item_type_code))
        messages.error(request, "Dữ liệu không hợp lệ, vui lòng kiểm tra lại.")
    else:
        form = AdmIncomingDispatchForm(item_type_code=item_type_code)
        if prefill_company:
            form.initial["receiving_company"] = prefill_company.code

    dispatches_qs = (
        AdmIncomingDispatch.objects.select_related(
            "responsible_user",
            "receiving_company",
            "incoming_item_type",
            "status",
        )
        .prefetch_related(
            "processing_departments",
            "processing_departments__company",
            "images",
        )
        .order_by("-created_at")
    )
    dispatches_qs = dispatches_qs.filter(incoming_item_type_id=item_type_code)
    if item_type_code == INCOMING_TYPE_DOC:
        if status_filter:
            dispatches_qs = dispatches_qs.filter(status=status_filter)
        else:
            dispatches_qs = dispatches_qs.exclude(status_id__in=INCOMING_ARCHIVED_STATUS_CODES)
    else:
        dispatches_qs = dispatches_qs.exclude(status_id__in=PARCEL_HIDDEN_LIST_STATUS_CODES)
    if query:
        dispatches_qs = dispatches_qs.filter(
            Q(document_number__icontains=query) | Q(summary__icontains=query)
        )
    if status_filter and item_type_code != INCOMING_TYPE_DOC:
        dispatches_qs = dispatches_qs.filter(status=status_filter)
    if department_filter:
        dispatches_qs = dispatches_qs.filter(
            processing_departments__id=department_filter
        )
    if company_filter:
        dispatches_qs = dispatches_qs.filter(receiving_company_id=company_filter)

    dispatches_qs = dispatches_qs.distinct()
    paginator = Paginator(dispatches_qs, 25)
    page_number = request.GET.get("page", "1")
    try:
        dispatches = paginator.page(page_number)
    except PageNotAnInteger:
        dispatches = paginator.page(1)
    except EmptyPage:
        dispatches = paginator.page(paginator.num_pages)
    active_status_map = {
        status.code: status
        for status in AdmIncomingDispatchStatus.objects.filter(is_active=True)
    }
    for dispatch in dispatches.object_list:
        dispatch.normalized_status_id = _normalize_incoming_dispatch_status_code(
            dispatch.incoming_item_type_id,
            dispatch.status_id,
        )
        normalized_flow_label = _incoming_flow_label(
            dispatch.incoming_item_type_id,
            dispatch.status_id,
            signer_label=(dispatch.signer_name or "").strip() or "Người ký",
        )
        dispatch.normalized_status_name = (
            active_status_map.get(dispatch.normalized_status_id).name
            if dispatch.normalized_status_id in active_status_map
            else normalized_flow_label or dispatch.status.name
        )
        badge_source = active_status_map.get(dispatch.normalized_status_id) or dispatch.status
        badge_tokens = _incoming_status_badge_tokens(badge_source)
        dispatch.normalized_status_badge_bg_color = badge_tokens["bg"]
        dispatch.normalized_status_badge_text_color = badge_tokens["text"]
        dispatch.available_statuses = [
            active_status_map[code]
            for code in _allowed_status_codes_for_dispatch(dispatch)
            if code in active_status_map
        ]

    departments = AdmDepartment.objects.select_related("company").filter(
        is_active=True
    ).order_by("company__code", "name")
    if item_type_code == INCOMING_TYPE_DOC:
        allowed_codes = [code for code, _ in _incoming_flow_by_type(item_type_code)]
    else:
        allowed_codes = [
            code
            for code, _ in _incoming_flow_by_type(item_type_code)
            if code not in INCOMING_ARCHIVED_STATUS_CODES
        ]
    status_options = AdmIncomingDispatchStatus.objects.filter(
        is_active=True, code__in=allowed_codes
    ).order_by("sort_order", "name")
    selected_dispatch = None
    selected_workflow_steps = []
    selected_recipient_group_key = ""
    if open_dispatch_id.isdigit():
        selected_dispatch = (
            dispatches_qs.filter(pk=int(open_dispatch_id))
            .select_related(
                "responsible_user",
                "receiving_company",
                "incoming_item_type",
                "status",
            )
            .prefetch_related(
                "processing_departments",
                "processing_departments__company",
                "images",
            )
            .first()
        )
        if selected_dispatch:
            selected_dispatch.normalized_status_id = _normalize_incoming_dispatch_status_code(
                selected_dispatch.incoming_item_type_id,
                selected_dispatch.status_id,
            )
            selected_dispatch.normalized_status_name = (
                active_status_map.get(selected_dispatch.normalized_status_id).name
                if selected_dispatch.normalized_status_id in active_status_map
                else _incoming_flow_label(
                    selected_dispatch.incoming_item_type_id,
                    selected_dispatch.status_id,
                    signer_label=(selected_dispatch.signer_name or "").strip() or "Người ký",
                ) or selected_dispatch.status.name
            )
            badge_source = active_status_map.get(selected_dispatch.normalized_status_id) or selected_dispatch.status
            badge_tokens = _incoming_status_badge_tokens(badge_source)
            selected_dispatch.normalized_status_badge_bg_color = badge_tokens["bg"]
            selected_dispatch.normalized_status_badge_text_color = badge_tokens["text"]
            selected_dispatch.available_statuses = [
                active_status_map[code]
                for code in _allowed_status_codes_for_dispatch(selected_dispatch)
                if code in active_status_map
            ]
            selected_dispatch.current_processing_department_id = (
                selected_dispatch.processing_departments.values_list("id", flat=True).first()
            )
            selected_dispatch.status_logs_display = list(
                selected_dispatch.status_logs.select_related(
                    "from_status",
                    "to_status",
                    "changed_by",
                ).all()
            )
            selected_workflow_steps = _build_incoming_workflow_steps(selected_dispatch)

    context = {
        "dispatches": dispatches,
        "page_obj": dispatches,
        "paginator": paginator,
        "is_paginated": paginator.num_pages > 1,
        "form": form,
        "incoming_item_type_field": form["incoming_item_type"] if "incoming_item_type" in form.fields else None,
        "q": query,
        "selected_status": status_filter,
        "selected_department": department_filter,
        "selected_company": company_filter,
        "status_options": status_options,
        "companies": companies,
        "departments": departments,
        "prefill_company": prefill_company,
        "selected_dispatch": selected_dispatch,
        "selected_workflow_steps": selected_workflow_steps,
        "can_create": can_create,
        "can_edit_status": can_edit_status,
        "has_admin_docs_access": _has_admin_docs_access(request.user),
        "screen_key": screen_meta["route_name"],
        "page_title": screen_meta["page_title"],
        "page_description": screen_meta["page_description"],
        "create_button_label": screen_meta["create_button_label"],
        "modal_title": screen_meta["modal_title"],
        "empty_message": screen_meta["empty_message"],
        "search_placeholder": screen_meta["search_placeholder"],
        "auto_status_note": screen_meta["auto_status_note"],
        "current_route_name": screen_meta["route_name"],
        "current_url_name": _incoming_dispatch_route_name(item_type_code),
        "status_action_name": "admindocuments:incoming_dispatch_change_status",
        "upload_action_name": "admindocuments:incoming_dispatch_upload_images",
        "item_type_code": item_type_code,
        "item_type_name": (
            AdmIncomingDispatchType.objects.filter(code=item_type_code)
            .values_list("name", flat=True)
            .first()
            or screen_meta["page_title"]
        ),
    }
    return render(request, "admindocuments/incoming_dispatch_list.html", context)


@login_required
def incoming_dispatch_list(request):
    return redirect("admindocuments:incoming_document_list")


@login_required
def incoming_document_list(request):
    return _incoming_dispatch_list(request, INCOMING_TYPE_DOC)


@login_required
def parcel_receipt_list(request):
    can_create = _has_incoming_dispatch_create_access(request.user)
    can_edit_status = _has_incoming_dispatch_status_edit_access(request.user)
    can_notify = _has_incoming_dispatch_create_access(request.user)
    screen_meta = INCOMING_SCREEN_META[INCOMING_TYPE_PARCEL]

    query = request.GET.get("q", "").strip()
    status_filter = request.GET.get("status", "")
    department_filter = request.GET.get("department", "").strip()
    company_filter = request.GET.get("company", "")
    recipient_filter = request.GET.get("recipient", "").strip()
    unnotified_only = request.GET.get("unnotified") == "1"
    unassigned_only = request.GET.get("unassigned") == "1"
    split_by_day = request.GET.get("split_by_day", "1") != "0"
    start_date = (request.GET.get("start_date") or "").strip()
    end_date = (request.GET.get("end_date") or "").strip()
    receive_location_filter = (request.GET.get("receive_location") or "").strip()
    receive_company = (request.GET.get("receive_company") or "").strip()
    open_dispatch_id = (request.GET.get("open") or "").strip()
    highlight_dispatch_id = (request.GET.get("highlight") or "").strip()
    companies = AdmCompany.objects.filter(is_active=True).order_by("name")
    receive_locations = AdmParcelReceiveLocation.objects.filter(is_active=True).order_by("name")
    prefill_company = None
    default_receive_location_id = (
        UserProfile.objects.filter(user=request.user)
        .values_list("default_receive_location_id", flat=True)
        .first()
    )
    if not receive_location_filter and default_receive_location_id:
        receive_location_filter = str(default_receive_location_id)

    if receive_company:
        prefill_company = companies.filter(code=receive_company).first()

    if request.method == "POST":
        form = AdmParcelReceiptForm(request.POST, user=request.user)
        receive_company = (request.POST.get("receiving_company") or "").strip()
        prefill_company = companies.filter(code=receive_company).first()
        if form.is_valid():
            image_error = _validate_parcel_images(request.FILES.getlist("images"))
            if image_error:
                form.add_error(None, image_error)
            else:
                parcel_receipt = None
                try:
                    with transaction.atomic():
                        parcel_receipt = form.save(commit=False)
                        parcel_receipt.received_by = request.user
                        parcel_receipt.created_by = request.user
                        parcel_receipt.updated_by = request.user
                        parcel_receipt.received_at = timezone.now()
                        initial_status = _resolve_initial_incoming_status_code(INCOMING_TYPE_PARCEL)
                        if not initial_status:
                            raise ValueError("Chưa có cấu hình trạng thái khởi tạo cho bưu phẩm.")
                        parcel_receipt.status_id = initial_status
                        parcel_receipt.save()
                        sender_name = (parcel_receipt.sender_unit or "").strip()
                        if sender_name and not AdmParcelSenderSuggestion.objects.filter(
                            name__iexact=sender_name
                        ).exists():
                            AdmParcelSenderSuggestion.objects.create(name=sender_name)
                        _save_parcel_receipt_images(
                            parcel_receipt,
                            request.FILES.getlist("images"),
                            request.user,
                        )
                        _log_parcel_event(
                            parcel_receipt,
                            AdmParcelReceiptLog.ACTION_CREATED,
                            actor=request.user,
                            to_status=parcel_receipt.status_id,
                            note="Lễ tân tiếp nhận.",
                        )
                except Exception as exc:
                    messages.error(request, f"Lưu tiếp nhận thất bại: {exc}")
                else:
                    messages.success(
                        request,
                        "Đã tiếp nhận bưu phẩm/bưu kiện.",
                    )
                    return redirect(
                        f"{reverse('admindocuments:parcel_receipt_list')}?open={parcel_receipt.id}&highlight={parcel_receipt.id}&pane=list"
                    )
        if form.errors:
            error_parts = []
            for field_errors in form.errors.values():
                if isinstance(field_errors, (list, tuple)):
                    error_parts.extend(str(item) for item in field_errors if str(item))
            if error_parts:
                messages.error(request, "Dữ liệu không hợp lệ: " + " | ".join(error_parts))
            else:
                messages.error(request, "Dữ liệu không hợp lệ, vui lòng kiểm tra lại.")
        else:
            messages.error(request, "Dữ liệu không hợp lệ, vui lòng kiểm tra lại.")
    else:
        form = AdmParcelReceiptForm(user=request.user)
        if prefill_company:
            form.initial["receiving_company"] = prefill_company.code

    dispatches_qs = (
        AdmParcelReceipt.objects.select_related(
            "received_by",
            "receiving_company",
            "status",
            "recipient_directory",
            "recipient_user",
            "recipient_user__userprofile",
        )
        .prefetch_related(
            "images",
            "audit_logs",
            "audit_logs__actor",
            models.Prefetch(
                "notification_batches",
                queryset=AdmParcelNotificationBatch.objects.select_related(
                    "reminder_schedule"
                ).order_by("-created_at"),
            ),
        )
        .order_by("-received_at", "-id")
    )
    dispatches_qs = _exclude_hidden_parcels(dispatches_qs)
    if query:
        dispatches_qs = dispatches_qs.filter(
            Q(sender_unit__icontains=query)
            | Q(tracking_code__icontains=query)
            | Q(recipient_name__icontains=query)
            | Q(recipient_employee_code__icontains=query)
            | Q(recipient_directory__full_name__icontains=query)
            | Q(recipient_directory__employee_code__icontains=query)
            | Q(recipient_user__first_name__icontains=query)
            | Q(recipient_user__last_name__icontains=query)
            | Q(recipient_user__username__icontains=query)
            | Q(recipient_user__userprofile__employee_code__icontains=query)
        )
    if status_filter:
        dispatches_qs = dispatches_qs.filter(status=status_filter)
    if department_filter:
        dispatches_qs = dispatches_qs.filter(
            Q(recipient_department__icontains=department_filter)
            | Q(recipient_directory__department_name__icontains=department_filter)
        )
    if recipient_filter:
        dispatches_qs = dispatches_qs.filter(
            Q(recipient_name__icontains=recipient_filter)
            | Q(recipient_employee_code__icontains=recipient_filter)
            | Q(recipient_directory__full_name__icontains=recipient_filter)
            | Q(recipient_directory__employee_code__icontains=recipient_filter)
            | Q(recipient_directory__email__icontains=recipient_filter)
            | Q(recipient_user__first_name__icontains=recipient_filter)
            | Q(recipient_user__last_name__icontains=recipient_filter)
            | Q(recipient_user__username__icontains=recipient_filter)
        )
    if unnotified_only:
        dispatches_qs = dispatches_qs.filter(notified_at__isnull=True)
    if unassigned_only:
        dispatches_qs = dispatches_qs.filter(
            Q(recipient_directory__isnull=True)
            & (Q(recipient_name__exact="") | Q(recipient_name=AdmParcelReceiptForm.UNKNOWN_RECIPIENT_LABEL))
        )
    if start_date:
        for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                dispatches_qs = dispatches_qs.filter(received_at__date__gte=datetime.strptime(start_date, fmt).date())
                break
            except ValueError:
                continue
    if end_date:
        for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                dispatches_qs = dispatches_qs.filter(received_at__date__lte=datetime.strptime(end_date, fmt).date())
                break
            except ValueError:
                continue
    if receive_location_filter and receive_location_filter != "all":
        dispatches_qs = dispatches_qs.filter(receive_location_id=receive_location_filter)

    dispatches_qs = dispatches_qs.distinct()
    paginator = Paginator(dispatches_qs, 25)
    page_number = request.GET.get("page", "1")
    try:
        dispatches = paginator.page(page_number)
    except PageNotAnInteger:
        dispatches = paginator.page(1)
    except EmptyPage:
        dispatches = paginator.page(paginator.num_pages)

    active_status_map = {
        status.code: status
        for status in AdmIncomingDispatchStatus.objects.filter(is_active=True)
    }
    for parcel_receipt in dispatches.object_list:
        parcel_receipt.available_statuses = [
            active_status_map[code]
            for code in _allowed_status_codes_for_parcel(parcel_receipt)
            if code in active_status_map
        ]
        _prepare_parcel_receipt_view_state(parcel_receipt, request=request)

    grouping_mode = "day" if split_by_day else "recipient"
    recipient_groups = _build_parcel_recipient_groups(
        dispatches.object_list,
        split_by_day=split_by_day,
        request=request,
    )

    current_recipient_batch = (
        AdmParcelRecipientImportBatch.objects.filter(is_current=True)
        .order_by("-created_at")
        .first()
    )
    recipient_options = AdmParcelRecipientCatalog.objects.none()
    departments = []
    if current_recipient_batch:
        recipient_options = current_recipient_batch.recipients.filter(
            is_active_member=True
        ).exclude(department_name__exact="").order_by(
            "department_name", "full_name", "employee_code"
        )
        departments = (
            current_recipient_batch.recipients.filter(is_active_member=True)
            .exclude(department_name__exact="")
            .values_list("department_name", flat=True)
            .distinct()
            .order_by("department_name")
        )
    allowed_codes = [
        code
        for code, _ in INCOMING_PARCEL_FLOW
        if code not in PARCEL_HIDDEN_LIST_STATUS_CODES
    ]
    status_options = AdmIncomingDispatchStatus.objects.filter(
        is_active=True, code__in=allowed_codes
    ).order_by("sort_order", "name")
    selected_dispatch = None
    selected_workflow_steps = []
    selected_recipient_group_key = ""
    if open_dispatch_id.isdigit():
        selected_dispatch = (
            dispatches_qs.filter(pk=int(open_dispatch_id))
            .select_related(
                "received_by",
                "receiving_company",
                "status",
                "recipient_directory",
                "recipient_user",
                "recipient_user__userprofile",
                "notification_schedule",
                "reminder_schedule",
            )
            .first()
        )
        if selected_dispatch:
            _prepare_parcel_receipt_view_state(selected_dispatch, request=request)
            selected_workflow_steps = selected_dispatch.workflow_steps
            selected_recipient_group_key = (
                selected_dispatch.received_at.date().isoformat()
                if split_by_day and selected_dispatch.received_at
                else _parcel_group_key(selected_dispatch)
            )

    auto_notify_setting = _get_active_parcel_auto_notify_setting()

    context = {
        "dispatches": dispatches,
        "page_obj": dispatches,
        "paginator": paginator,
        "is_paginated": paginator.num_pages > 1,
        "form": form,
        "incoming_item_type_field": form["incoming_item_type"] if "incoming_item_type" in form.fields else None,
        "q": query,
        "selected_status": status_filter,
        "selected_department": department_filter,
        "selected_company": "",
        "selected_recipient": recipient_filter,
        "selected_unnotified": unnotified_only,
        "selected_unassigned": unassigned_only,
        "split_by_day": split_by_day,
        "grouping_mode": grouping_mode,
        "selected_start_date": start_date,
        "selected_end_date": end_date,
        "selected_receive_location": receive_location_filter,
        "status_options": status_options,
        "companies": companies,
        "receive_locations": receive_locations,
        "departments": departments,
        "recipient_options": recipient_options,
        "recipient_groups": recipient_groups,
        "current_recipient_batch": current_recipient_batch,
        "sender_suggestions": getattr(form, "sender_suggestions", []),
        "auto_notify_setting": auto_notify_setting,
        "prefill_company": prefill_company,
        "selected_dispatch": selected_dispatch,
        "selected_recipient_group_key": selected_recipient_group_key,
        "highlight_dispatch_id": highlight_dispatch_id or open_dispatch_id,
        "selected_workflow_steps": selected_workflow_steps,
        "can_create": can_create,
        "can_edit_status": can_edit_status,
        "can_notify": can_notify,
        "has_admin_docs_access": _has_admin_docs_access(request.user),
        "screen_key": screen_meta["route_name"],
        "page_title": screen_meta["page_title"],
        "page_description": screen_meta["page_description"],
        "create_button_label": screen_meta["create_button_label"],
        "modal_title": screen_meta["modal_title"],
        "empty_message": screen_meta["empty_message"],
        "search_placeholder": screen_meta["search_placeholder"],
        "auto_status_note": screen_meta["auto_status_note"],
        "current_route_name": screen_meta["route_name"],
        "current_url_name": "admindocuments:parcel_receipt_list",
        "status_action_name": "admindocuments:parcel_receipt_change_status",
        "notify_action_name": "admindocuments:parcel_receipt_send_notification",
        "upload_action_name": "admindocuments:parcel_receipt_upload_images",
        "parcel_confirm_base_url": _build_public_absolute_url(
            request,
            reverse("admindocuments:parcel_receipt_confirm_short", args=["TOKEN"]),
        ).replace("TOKEN", ""),
        "item_type_code": INCOMING_TYPE_PARCEL,
        "item_type_name": screen_meta["page_title"],
    }
    return render(request, "admindocuments/parcel_receipt_list.html", context)


def _completed_parcel_filters(request):
    """Áp dụng bộ lọc màn hình "Bưu kiện đã nhận" (dùng chung cho danh sách và xuất Excel)."""
    values = {
        "q": (request.GET.get("q") or "").strip(),
        "recipient": (request.GET.get("recipient") or "").strip(),
        "employee_code": (request.GET.get("employee_code") or "").strip(),
        "department": (request.GET.get("department") or "").strip(),
        "company": (request.GET.get("company") or "").strip(),
        "parcel_type": (request.GET.get("parcel_type") or "").strip(),
        "status": (request.GET.get("status") or "").strip(),
        "start_date": (request.GET.get("start_date") or "").strip(),
        "end_date": (request.GET.get("end_date") or "").strip(),
    }

    parcels_qs = (
        AdmParcelReceipt.objects.select_related(
            "received_by",
            "receiving_company",
            "status",
            "receive_location",
            "recipient_directory",
            "recipient_user",
            "recipient_user__userprofile",
        )
        .prefetch_related("images")
        .order_by("-completed_at", "-confirmed_at", "-received_at", "-id")
    )
    base_qs = _completed_parcels_queryset(parcels_qs)
    parcels_qs = base_qs

    if values["q"]:
        query = values["q"]
        parcels_qs = parcels_qs.filter(
            Q(tracking_code__icontains=query)
            | Q(sender_unit__icontains=query)
            | Q(content__icontains=query)
            | Q(recipient_name__icontains=query)
            | Q(recipient_employee_code__icontains=query)
            | Q(actual_receiver_name__icontains=query)
            | Q(actual_receiver_employee_code__icontains=query)
        )
    if values["recipient"]:
        recipient_filter = values["recipient"]
        parcels_qs = parcels_qs.filter(
            Q(recipient_name__icontains=recipient_filter)
            | Q(recipient_employee_code__icontains=recipient_filter)
            | Q(actual_receiver_name__icontains=recipient_filter)
            | Q(actual_receiver_employee_code__icontains=recipient_filter)
            | Q(recipient_directory__full_name__icontains=recipient_filter)
            | Q(recipient_directory__employee_code__icontains=recipient_filter)
        )
    if values["employee_code"]:
        code = values["employee_code"]
        parcels_qs = parcels_qs.filter(
            Q(recipient_employee_code__icontains=code)
            | Q(actual_receiver_employee_code__icontains=code)
            | Q(recipient_directory__employee_code__icontains=code)
        )
    if values["department"]:
        parcels_qs = parcels_qs.filter(
            Q(recipient_department__icontains=values["department"])
            | Q(recipient_directory__department_name__icontains=values["department"])
        )
    if values["company"]:
        parcels_qs = parcels_qs.filter(receiving_company_id=values["company"])
    if values["parcel_type"] in AdmParcelReceipt.ParcelType.values:
        parcels_qs = parcels_qs.filter(parcel_type=values["parcel_type"])
    if values["status"]:
        parcels_qs = parcels_qs.filter(status_id=values["status"])

    def parse_date(value):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
        return None

    start = parse_date(values["start_date"])
    end = parse_date(values["end_date"])
    if start and end and start > end:
        start, end = end, start
    if start:
        parcels_qs = parcels_qs.filter(received_at__date__gte=start)
    if end:
        parcels_qs = parcels_qs.filter(received_at__date__lte=end)
    values["start_date"] = start.strftime("%Y-%m-%d") if start else ""
    values["end_date"] = end.strftime("%Y-%m-%d") if end else ""

    options = {
        "companies": AdmCompany.objects.filter(code__in=base_qs.values("receiving_company_id")).order_by("name"),
        "parcel_types": [{"code": code, "name": label} for code, label in AdmParcelReceipt.ParcelType.choices],
        "statuses": AdmIncomingDispatchStatus.objects.filter(code__in=base_qs.values("status_id")).order_by("name"),
    }
    return parcels_qs.distinct(), values, options


@login_required
def parcel_receipt_completed_list(request):
    if not _has_admin_docs_access(request.user):
        return HttpResponseForbidden("Bạn không có quyền xem dữ liệu bưu kiện đã nhận.")

    parcels_qs, filters, options = _completed_parcel_filters(request)
    paginator = Paginator(parcels_qs, 25)
    page_number = request.GET.get("page", "1")
    try:
        parcels = paginator.page(page_number)
    except PageNotAnInteger:
        parcels = paginator.page(1)
    except EmptyPage:
        parcels = paginator.page(paginator.num_pages)

    for parcel in parcels.object_list:
        parcel.recipient_display = _parcel_recipient_display(parcel) or "Chưa xác định"
        parcel.actual_receiver_display = _parcel_actual_receiver_display(parcel)
        parcel.finished_at = parcel.completed_at or parcel.confirmed_at

    export_params = request.GET.copy()
    export_params.pop("page", None)
    context = {
        "active": "parcel_receipt_completed_list",
        "screen_key": "parcel_receipt_completed_list",
        "page_title": "Bưu kiện đã nhận",
        "page_description": "Tra cứu lại bưu phẩm, bưu kiện đã được người nhận xác nhận.",
        "parcels": parcels,
        "page_obj": parcels,
        "paginator": paginator,
        "is_paginated": paginator.num_pages > 1,
        "export_query": export_params.urlencode(),
        "company_options": options["companies"],
        "type_options": options["parcel_types"],
        "status_options": options["statuses"],
        "has_admin_docs_access": _has_admin_docs_access(request.user),
        **filters,
    }
    return render(request, "admindocuments/parcel_receipt_completed_list.html", context)


def _excel_text(value):
    """Chặn công thức Excel (=, +, -, @) trong dữ liệu người dùng nhập."""
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@") else text


@login_required
def parcel_receipt_completed_export(request):
    if not _has_admin_docs_access(request.user):
        return HttpResponseForbidden("Bạn không có quyền xuất dữ liệu bưu kiện đã nhận.")

    parcels_qs, _filters, _options = _completed_parcel_filters(request)
    local = timezone.get_current_timezone()

    def when(value):
        if not value:
            return ""
        if timezone.is_aware(value):
            value = value.astimezone(local)
        return value.strftime("%d/%m/%Y %H:%M")

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Bưu kiện đã nhận")
    headers = [
        ("STT", 6), ("Tiếp nhận", 17), ("Hoàn tất", 17), ("Người nhận", 26), ("Mã NV người nhận", 15),
        ("Phòng ban", 22), ("Người nhận thực tế", 26), ("Mã NV người nhận thực tế", 17), ("Đơn vị gửi", 26),
        ("Mã vận đơn", 20), ("Nội dung", 30), ("Nơi nhận", 24), ("Địa chỉ nơi nhận", 30), ("Loại", 12),
        ("Công ty", 16), ("Lễ tân", 20), ("Trạng thái", 14),
    ]
    for index, (_title, width) in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    header_cells = []
    for title, _width in headers:
        cell = WriteOnlyCell(sheet, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="00844A", end_color="00844A")
        header_cells.append(cell)
    sheet.append(header_cells)

    for index, parcel in enumerate(parcels_qs.iterator(chunk_size=500), start=1):
        directory = parcel.recipient_directory
        recipient_name = (parcel.recipient_name or "").strip() or (directory.full_name if directory else "")
        recipient_code = (parcel.recipient_employee_code or "").strip() or (directory.employee_code if directory else "")
        department = parcel.recipient_department or (directory.department_name if directory else "")
        location = parcel.receive_location
        receptionist = ""
        if parcel.received_by_id:
            receptionist = parcel.received_by.get_full_name() or parcel.received_by.username
        sheet.append([
            index,
            when(parcel.received_at),
            when(parcel.completed_at or parcel.confirmed_at),
            _excel_text(recipient_name or _parcel_recipient_display(parcel)),
            _excel_text(recipient_code),
            _excel_text(department),
            _excel_text(parcel.actual_receiver_name or ""),
            _excel_text(parcel.actual_receiver_employee_code or ""),
            _excel_text(parcel.sender_unit),
            _excel_text(parcel.tracking_code),
            _excel_text(parcel.content),
            _excel_text(location.name if location else ""),
            _excel_text(location.address if location else ""),
            parcel.get_parcel_type_display(),
            _excel_text(parcel.receiving_company.name if parcel.receiving_company_id else ""),
            _excel_text(receptionist),
            _excel_text(parcel.status.name if parcel.status_id else ""),
        ])

    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="buu-kien-da-nhan-{date.today():%Y%m%d}.xlsx"'
    return response


@login_required
@superuser_required
def parcel_recipient_directory(request):
    """Keep old upload links working while the directory lives in Master Data."""
    if request.method == "GET":
        query = request.GET.urlencode()
        url = reverse("master_data_gapo_directory")
        return redirect(f"{url}?{query}" if query else url)
    from app_documents.master_data import gapo_directory_page
    return gapo_directory_page(request)


@login_required
def parcel_auto_notify_settings(request):
    if not _has_admin_docs_access(request.user):
        return render(request, "403.html", status=403)

    setting, _ = AdmParcelAutoNotifySetting.objects.get_or_create(
        code="default",
        defaults={"name": "Nhắc lại bưu kiện"},
    )
    if request.method == "POST":
        form = AdmParcelAutoNotifySettingForm(request.POST, instance=setting)
        if form.is_valid():
            setting = form.save(commit=False)
            setting.updated_by = request.user
            setting.save()
            messages.success(request, "Đã lưu cấu hình nhắc lại bưu kiện.")
            return redirect("admindocuments:parcel_auto_notify_settings")
        messages.error(request, "Cấu hình không hợp lệ.")
    else:
        form = AdmParcelAutoNotifySettingForm(instance=setting)

    return render(
        request,
        "admindocuments/parcel_auto_notify_settings.html",
        {
            "form": form,
            "setting": setting,
            "has_admin_docs_access": _has_admin_docs_access(request.user),
        },
    )


@login_required
def parcel_receipt_item_partial(request, doc_id: int):
    split_by_day = request.GET.get("split_by_day", "1") != "0"
    parcel_receipt = get_object_or_404(
        AdmParcelReceipt.objects.select_related(
            "received_by",
            "receiving_company",
            "status",
            "recipient_directory",
            "recipient_user",
            "recipient_user__userprofile",
        ).prefetch_related(
            "images",
            "audit_logs",
            "audit_logs__actor",
            models.Prefetch(
                "notification_batches",
                queryset=AdmParcelNotificationBatch.objects.select_related(
                    "reminder_schedule"
                ).order_by("-created_at"),
            ),
        ),
        pk=doc_id,
    )
    _prepare_parcel_receipt_view_state(parcel_receipt, request=request)
    group_queryset = AdmParcelReceipt.objects.select_related(
        "received_by",
        "receiving_company",
        "status",
        "recipient_directory",
        "recipient_user",
        "recipient_user__userprofile",
    ).prefetch_related(
        "images",
        "audit_logs",
        "audit_logs__actor",
        models.Prefetch(
            "notification_batches",
            queryset=AdmParcelNotificationBatch.objects.select_related(
                "reminder_schedule"
            ).order_by("-created_at"),
        ),
    )
    if split_by_day and parcel_receipt.received_at:
        group_queryset = group_queryset.filter(received_at__date=parcel_receipt.received_at.date())
    elif parcel_receipt.recipient_directory_id:
        group_queryset = group_queryset.filter(recipient_directory_id=parcel_receipt.recipient_directory_id)
    else:
        recipient_name = (parcel_receipt.recipient_name or "").strip()
        if recipient_name:
            group_queryset = group_queryset.filter(
                recipient_directory__isnull=True,
                recipient_name=recipient_name,
            )
        else:
            group_queryset = group_queryset.filter(recipient_directory__isnull=True).filter(
                Q(recipient_name__exact="")
                | Q(recipient_name=AdmParcelReceiptForm.UNKNOWN_RECIPIENT_LABEL)
            )
    group_queryset = _exclude_hidden_parcels(group_queryset).order_by("-received_at", "-id")
    for sibling in group_queryset:
        _prepare_parcel_receipt_view_state(sibling, request=request)
    groups = _build_parcel_recipient_groups(
        group_queryset,
        split_by_day=split_by_day,
        request=request,
    )
    group_html = ""
    if groups:
        group_html = render_to_string(
            "admindocuments/includes/parcel_recipient_group_card.html",
            {
                "group": groups[0],
                "grouping_mode": "day" if split_by_day else "recipient",
                "highlight_dispatch_id": "",
                "upload_action_name": "admindocuments:parcel_receipt_upload_images",
                "request": request,
            },
            request=request,
        )
    if (
        parcel_receipt.status_id in PARCEL_HIDDEN_LIST_STATUS_CODES
        or parcel_receipt.confirmed_at is not None
        or parcel_receipt.completed_at is not None
    ):
        return JsonResponse({"remove": True, "group_html": group_html, "status": parcel_receipt.status.name})
    html = render_to_string(
        "admindocuments/includes/parcel_receipt_list_item.html",
        {
            "doc": parcel_receipt,
            "highlight_dispatch_id": "",
            "grouping_mode": "day" if split_by_day else "recipient",
            "upload_action_name": "admindocuments:parcel_receipt_upload_images",
            "request": request,
        },
        request=request,
    )
    return JsonResponse({"html": html, "group_html": group_html, "status": parcel_receipt.status.name})


@login_required
def parcel_receipt_confirm_qr(request, token: str):
    parcel_receipt = get_object_or_404(AdmParcelReceipt, confirmation_token=token)
    confirm_url = _build_parcel_confirmation_url(request, parcel_receipt.confirmation_token)
    return _build_qr_svg_response(confirm_url)


@login_required
def parcel_batch_confirm_qr(request, token: str):
    batch = get_object_or_404(AdmParcelNotificationBatch, token=token)
    confirm_url = _build_parcel_batch_confirmation_url(request, batch.token)
    return _build_qr_svg_response(confirm_url)


@login_required
def parcel_recipient_search(request):
    current_batch = (
        AdmParcelRecipientImportBatch.objects.filter(is_current=True)
        .order_by("-created_at")
        .first()
    )
    if not current_batch:
        return JsonResponse({"results": []})

    query = (request.GET.get("q") or "").strip()
    department = (request.GET.get("department") or "").strip()
    selected_id = (request.GET.get("selected_id") or "").strip()
    queryset = current_batch.recipients.filter(is_active_member=True)

    if selected_id.isdigit():
        queryset = queryset.filter(pk=int(selected_id))
    else:
        if department:
            queryset = queryset.filter(department_name=department)
        if not query:
            return JsonResponse({"results": []})

        phone_query = "".join(ch for ch in query if ch.isdigit())
        looks_like_phone = phone_query and all(
            ch.isdigit() or ch in " +-.()"
            for ch in query
        )
        if looks_like_phone:
            queryset = queryset.filter(phone_number_normalized=phone_query)
        else:
            queryset = queryset.filter(
                Q(full_name__icontains=query)
                | Q(employee_code__icontains=query)
                | Q(email__icontains=query)
            )

    queryset = queryset.order_by("full_name", "employee_code")[:20]
    results = []
    for recipient in queryset:
        # Parcel search is an operational UI, so sensitive fields stay masked
        # even when the current session belongs to a superuser.
        contact_phone = recipient.mask_phone_number(recipient.phone_number)
        contact_email = recipient.mask_email(recipient.email)
        birth_date = recipient.mask_birth_date(recipient.birth_date)
        results.append(
            {
                "id": recipient.id,
                "full_name": recipient.full_name,
                "employee_code": recipient.employee_code,
                "department_name": recipient.department_name,
                "phone_number": contact_phone,
                "email": contact_email,
                "birth_date": birth_date,
                "label": " - ".join(
                    [
                        part
                        for part in [
                            recipient.full_name,
                            recipient.employee_code,
                            contact_phone,
                            contact_email,
                        ]
                        if part
                    ]
                ),
            }
        )
    return JsonResponse({"results": results})


@login_required
def incoming_dispatch_change_status(request, doc_id: int):
    if request.method != "POST":
        return redirect("admindocuments:incoming_document_list")
    if not _has_incoming_dispatch_status_edit_access(request.user):
        messages.error(request, "Bạn không có quyền thay đổi tình trạng.")
        return render(request, "403.html", status=403)

    dispatch = get_object_or_404(AdmIncomingDispatch, pk=doc_id)
    new_status = (request.POST.get("status") or "").strip()
    valid_statuses = set(
        AdmIncomingDispatchStatus.objects.filter(is_active=True).values_list(
            "code", flat=True
        )
    )
    allowed_statuses = set(_allowed_status_codes_for_dispatch(dispatch))
    if new_status and new_status not in allowed_statuses:
        messages.error(request, "Bước xử lý không hợp lệ cho loại tiếp nhận này.")
        return redirect(_incoming_dispatch_route_name(dispatch.incoming_item_type_id))
    if new_status not in valid_statuses:
        messages.error(request, "Tình trạng không hợp lệ.")
        return redirect(_incoming_dispatch_route_name(dispatch.incoming_item_type_id))
    if dispatch.status_id == new_status:
        messages.info(request, "Tình trạng đã ở giá trị này.")
        return redirect(_incoming_dispatch_route_name(dispatch.incoming_item_type_id))

    previous_status = dispatch.status_id
    dispatch.status_id = new_status
    dispatch.updated_by = request.user
    dispatch.save(update_fields=["status", "updated_by", "updated_at"])
    AdmIncomingDispatchStatusLog.objects.create(
        dispatch=dispatch,
        from_status_id=previous_status,
        to_status_id=new_status,
        changed_by=request.user,
    )
    messages.success(request, "Cập nhật tình trạng thành công.")

    next_url = request.POST.get("next", "")
    if next_url.startswith("/"):
        return redirect(next_url)
    return redirect(_incoming_dispatch_route_name(dispatch.incoming_item_type_id))


@login_required
def incoming_dispatch_update_department(request, doc_id: int):
    if request.method != "POST":
        return redirect("admindocuments:incoming_document_list")
    if not _has_incoming_dispatch_status_edit_access(request.user):
        messages.error(request, "Bạn không có quyền cập nhật phòng ban xử lý.")
        return render(request, "403.html", status=403)

    dispatch = get_object_or_404(
        AdmIncomingDispatch.objects.prefetch_related("processing_departments"),
        pk=doc_id,
    )
    department_id = (request.POST.get("processing_department") or "").strip()
    if department_id:
        department = get_object_or_404(
            AdmDepartment.objects.filter(is_active=True),
            pk=department_id,
        )
        dispatch.processing_departments.set([department])
        messages.success(request, "Đã cập nhật phòng ban xử lý.")
    else:
        dispatch.processing_departments.clear()
        messages.success(request, "Đã xóa phòng ban xử lý.")

    dispatch.updated_by = request.user
    dispatch.save(update_fields=["updated_by", "updated_at"])

    next_url = request.POST.get("next", "")
    if next_url.startswith("/"):
        return redirect(next_url)
    return redirect(f"{_incoming_dispatch_list_url(dispatch.incoming_item_type_id)}?open={dispatch.id}")


@login_required
def incoming_dispatch_upload_images(request, doc_id: int):
    if request.method != "POST":
        return redirect("admindocuments:incoming_document_list")

    dispatch = get_object_or_404(AdmIncomingDispatch, pk=doc_id)
    files = request.FILES.getlist("images")
    if not files:
        messages.warning(request, "Bạn chưa chọn tệp để tải lên.")
    else:
        error = _validate_incoming_dispatch_files(files)
        if error:
            messages.error(request, error)
        else:
            saved_count = _save_incoming_dispatch_images(dispatch, files, request.user)
            if saved_count:
                messages.success(request, "Đã tải tệp đính kèm.")
            else:
                messages.warning(request, "Không có tệp hợp lệ để tải lên.")

    next_url = request.POST.get("next", "")
    if next_url.startswith("/"):
        return redirect(next_url)
    return redirect(f"{_incoming_dispatch_list_url(dispatch.incoming_item_type_id)}?open={dispatch.id}")


@login_required
def parcel_receipt_change_status(request, doc_id: int):
    if request.method != "POST":
        return redirect("admindocuments:parcel_receipt_list")
    if not _has_incoming_dispatch_status_edit_access(request.user):
        messages.error(request, "Bạn không có quyền thay đổi tình trạng.")
        return render(request, "403.html", status=403)

    parcel_receipt = get_object_or_404(AdmParcelReceipt, pk=doc_id)
    new_status = (request.POST.get("status") or "").strip()
    valid_statuses = set(
        AdmIncomingDispatchStatus.objects.filter(is_active=True).values_list(
            "code", flat=True
        )
    )
    allowed_statuses = set(_allowed_status_codes_for_parcel(parcel_receipt))
    if new_status and new_status not in allowed_statuses:
        messages.error(request, "Bước xử lý không hợp lệ cho bưu phẩm/bưu kiện.")
        return redirect("admindocuments:parcel_receipt_list")
    if new_status not in valid_statuses:
        messages.error(request, "Tình trạng không hợp lệ.")
        return redirect("admindocuments:parcel_receipt_list")
    if parcel_receipt.status_id == new_status:
        messages.info(request, "Tình trạng đã ở giá trị này.")
        return redirect("admindocuments:parcel_receipt_list")

    from_status = parcel_receipt.status_id
    parcel_receipt.status_id = new_status
    parcel_receipt.updated_by = request.user
    update_fields = ["status", "updated_by", "updated_at"]
    if new_status == "pkg_processing" and parcel_receipt.confirmed_at is None:
        parcel_receipt.confirmed_at = timezone.now()
        update_fields.append("confirmed_at")
    if new_status in {"pkg_processing", "pkg_done"}:
        # Treat legacy processing/done states as received so pending reminders stop.
        _cancel_parcel_reminder(parcel_receipt)
    if new_status == "pkg_done" and parcel_receipt.completed_at is None:
        parcel_receipt.completed_at = timezone.now()
        update_fields.append("completed_at")
    parcel_receipt.save(update_fields=update_fields)
    _log_parcel_event(
        parcel_receipt,
        AdmParcelReceiptLog.ACTION_UPDATED,
        actor=request.user,
        from_status=from_status,
        to_status=new_status,
        note="Cập nhật tình trạng thủ công.",
    )
    messages.success(request, "Cập nhật tình trạng thành công.")

    next_url = request.POST.get("next", "")
    if next_url.startswith("/"):
        return redirect(next_url)
    return redirect("admindocuments:parcel_receipt_list")


@login_required
def parcel_receipt_send_notification(request, doc_id: int):
    if request.method != "POST":
        return redirect("admindocuments:parcel_receipt_list")
    if not _has_incoming_dispatch_create_access(request.user):
        messages.error(request, "Bạn không có quyền gửi thông báo.")
        return render(request, "403.html", status=403)

    parcel_receipt = get_object_or_404(AdmParcelReceipt, pk=doc_id)
    if parcel_receipt.status_id != "pkg_received":
        messages.info(request, "Bưu kiện này đã được gửi thông báo trước đó.")
        next_url = request.POST.get("next", "")
        if next_url.startswith("/"):
            return redirect(next_url)
        return redirect("admindocuments:parcel_receipt_list")

    try:
        _, reminder_error = _send_parcel_notification_now(parcel_receipt, request.user, request)
    except (ValueError, NotificationSendError) as exc:
        messages.error(request, f"Gửi thông báo nhận thất bại: {exc}")
    except Exception as exc:
        messages.error(request, f"Gửi thông báo nhận thất bại: {exc}")
    else:
        messages.success(request, "Đã gửi thông báo nhận thành công.")
        if reminder_error:
            messages.warning(request, f"Đã gửi thông báo nhưng chưa lên lịch nhắc lại: {reminder_error}")

    next_url = request.POST.get("next", "")
    if next_url.startswith("/"):
        return redirect(next_url)
    return redirect("admindocuments:parcel_receipt_list")


@login_required
def parcel_receipt_send_group_notification(request):
    if request.method != "POST":
        return redirect("admindocuments:parcel_receipt_list")
    if not _has_incoming_dispatch_create_access(request.user):
        messages.error(request, "Bạn không có quyền gửi thông báo.")
        return render(request, "403.html", status=403)
    try:
        parcels = _selected_parcels_from_request(request)
        grouped_parcels = {}
        for parcel in parcels:
            grouped_parcels.setdefault(_parcel_group_key(parcel), []).append(parcel)
        batches = []
        errors = []
        reminder_warnings = []
        for group_parcels in grouped_parcels.values():
            try:
                batch, _, reminder_error = _send_parcel_group_notification_now(
                    group_parcels, request.user, request
                )
            except (ValueError, NotificationSendError) as exc:
                errors.append(str(exc))
            else:
                batches.append(batch)
                if reminder_error:
                    reminder_warnings.append(reminder_error)
    except Exception as exc:
        messages.error(request, f"Gửi thông báo nhận thất bại: {exc}")
    else:
        if batches:
            total_parcels = sum(batch.parcel_count for batch in batches)
            if len(batches) == 1:
                batch = batches[0]
                messages.success(
                    request,
                    f"Đã gửi thông báo cho {batch.parcel_count} bưu kiện của {batch.recipient_name or 'người nhận'}.",
                )
            else:
                messages.success(
                    request,
                    f"Đã gửi thông báo cho {total_parcels} bưu kiện của {len(batches)} người nhận.",
                )
        if reminder_warnings:
            messages.warning(
                request,
                "Đã gửi thông báo nhưng chưa lên lịch nhắc lại: " + " | ".join(reminder_warnings[:3]),
            )
        if errors:
            messages.warning(request, "Một số nhóm chưa gửi được: " + " | ".join(errors[:3]))
        if not batches and errors:
            next_url = request.POST.get("next", "")
            if next_url.startswith("/"):
                return redirect(next_url)
            return redirect("admindocuments:parcel_receipt_list")

    next_url = request.POST.get("next", "")
    if next_url.startswith("/"):
        return redirect(next_url)
    return redirect("admindocuments:parcel_receipt_list")


@login_required
def parcel_receipt_mark_handed_over(request):
    if request.method != "POST":
        return redirect("admindocuments:parcel_receipt_list")
    if not _has_incoming_dispatch_create_access(request.user):
        messages.error(request, "Bạn không có quyền cập nhật bàn giao.")
        return render(request, "403.html", status=403)

    messages.info(request, "Bước bàn giao đã được gộp. Khi người nhận xác nhận, bưu kiện sẽ tự hoàn tất.")

    next_url = request.POST.get("next", "")
    if next_url.startswith("/"):
        return redirect(next_url)
    return redirect("admindocuments:parcel_receipt_list")


@login_required
def parcel_receipt_assign_recipient(request, doc_id: int):
    if request.method != "POST":
        return redirect("admindocuments:parcel_receipt_list")
    if not _has_incoming_dispatch_create_access(request.user):
        messages.error(request, "Bạn không có quyền cập nhật người nhận.")
        return render(request, "403.html", status=403)

    parcel_receipt = get_object_or_404(AdmParcelReceipt, pk=doc_id)
    recipient_id = (request.POST.get("recipient_directory_id") or "").strip()
    if not recipient_id.isdigit():
        messages.error(request, "Cần chọn người nhận từ danh bạ.")
    else:
        recipient = get_object_or_404(
            AdmParcelRecipientCatalog,
            pk=int(recipient_id),
            is_active_member=True,
        )
        previous_recipient = _parcel_recipient_display(parcel_receipt)
        parcel_receipt.recipient_directory = recipient
        parcel_receipt.recipient_department = (recipient.department_name or "").strip()
        parcel_receipt.recipient_name = recipient.full_name or ""
        parcel_receipt.recipient_employee_code = recipient.employee_code or ""
        parcel_receipt.recipient_gapo_user_id = recipient.gapo_user_id or ""
        parcel_receipt.updated_by = request.user
        parcel_receipt.save(
            update_fields=[
                "recipient_directory",
                "recipient_department",
                "recipient_name",
                "recipient_employee_code",
                "recipient_gapo_user_id",
                "updated_by",
                "updated_at",
            ]
        )
        _log_parcel_event(
            parcel_receipt,
            AdmParcelReceiptLog.ACTION_UPDATED,
            actor=request.user,
            to_status=parcel_receipt.status_id,
            note="Cập nhật người nhận.",
            metadata={
                "from_recipient": previous_recipient,
                "to_recipient": _parcel_recipient_display(parcel_receipt),
            },
        )
        messages.success(request, "Đã cập nhật người nhận.")

    next_url = request.POST.get("next", "")
    if next_url.startswith("/"):
        return redirect(next_url)
    return redirect(f"{reverse('admindocuments:parcel_receipt_list')}?open={parcel_receipt.id}")


@login_required
def parcel_receipt_upload_images(request, doc_id: int):
    if request.method != "POST":
        return redirect("admindocuments:parcel_receipt_list")

    parcel_receipt = get_object_or_404(AdmParcelReceipt, pk=doc_id)
    files = request.FILES.getlist("images")
    if not files:
        messages.warning(request, "Bạn chưa chọn ảnh để tải lên.")
    else:
        _save_parcel_receipt_images(parcel_receipt, files, request.user)
        messages.success(request, "Đã tải ảnh đính kèm.")

    next_url = request.POST.get("next", "")
    if next_url.startswith("/"):
        return redirect(next_url)
    return redirect(f"{reverse('admindocuments:parcel_receipt_list')}?open={parcel_receipt.id}")


def parcel_receipt_confirm(request, token: str):
    parcel_receipt = get_object_or_404(
        AdmParcelReceipt.objects.select_related(
            "recipient_directory",
            "recipient_user",
            "recipient_user__userprofile",
            "receiving_company",
            "status",
        ).prefetch_related("images"),
        confirmation_token=token,
    )
    if request.method == "POST":
        if parcel_receipt.status_id in {"pkg_processing", "pkg_done"}:
            messages.info(request, "Bưu phẩm đã được xác nhận trước đó.")
            return redirect(request.path)
        confirmed = request.POST.get("confirmed") == "1"
        claim_mode = (request.POST.get("claim_mode") or "self").strip()
        actual_receiver_name = (request.POST.get("actual_receiver_name") or "").strip()
        actual_receiver_employee_code = (request.POST.get("actual_receiver_employee_code") or "").strip()
        if not confirmed:
            messages.error(request, "Cần xác nhận đã nhận hàng.")
        elif claim_mode == "proxy" and not actual_receiver_name:
            messages.error(request, "Cần nhập tên người nhận hộ.")
        else:
            from_status = parcel_receipt.status_id
            if claim_mode != "proxy" and not actual_receiver_name:
                actual_receiver_name = _parcel_recipient_display(parcel_receipt)
            parcel_receipt.actual_receiver_name = actual_receiver_name
            parcel_receipt.actual_receiver_employee_code = actual_receiver_employee_code
            now = timezone.now()
            parcel_receipt.status_id = "pkg_done"
            parcel_receipt.confirmed_at = now
            parcel_receipt.completed_at = now
            parcel_receipt.save(
                update_fields=[
                    "actual_receiver_name",
                    "actual_receiver_employee_code",
                    "status",
                    "confirmed_at",
                    "completed_at",
                    "updated_at",
                ]
            )
            _cancel_parcel_reminder(parcel_receipt)
            _log_parcel_event(
                parcel_receipt,
                AdmParcelReceiptLog.ACTION_CONFIRMED,
                from_status=from_status,
                to_status=parcel_receipt.status_id,
                note="Người nhận xác nhận đã nhận hàng từ lễ tân.",
                metadata={"actual_receiver_name": actual_receiver_name},
            )
            messages.success(request, "Đã xác nhận nhận hàng và hoàn tất bưu kiện.")
            return redirect(request.path)
    return render(
        request,
        "admindocuments/parcel_receipt_confirm.html",
        {"parcel": parcel_receipt},
    )


def parcel_batch_confirm(request, token: str):
    batch = get_object_or_404(
        AdmParcelNotificationBatch.objects.prefetch_related(
            "parcels",
            "parcels__status",
            "parcels__receiving_company",
            "parcels__images",
        ),
        token=token,
    )
    parcels = list(batch.parcels.all().order_by("-received_at", "-id"))
    if request.method == "POST":
        confirmed = request.POST.get("confirmed") == "1"
        claim_mode = (request.POST.get("claim_mode") or "self").strip()
        actual_receiver_name = (request.POST.get("actual_receiver_name") or "").strip()
        actual_receiver_employee_code = (request.POST.get("actual_receiver_employee_code") or "").strip()
        if not confirmed:
            messages.error(request, "Cần xác nhận đã nhận hàng.")
        else:
            if claim_mode == "proxy" and not actual_receiver_name:
                messages.error(request, "Cần nhập tên người nhận hộ.")
                return render(
                    request,
                    "admindocuments/parcel_batch_confirm.html",
                    {"batch": batch, "parcels": parcels},
                )
            if claim_mode != "proxy" and not actual_receiver_name:
                actual_receiver_name = batch.recipient_name or "Người nhận"
            now = timezone.now()
            with transaction.atomic():
                batch.status = AdmParcelNotificationBatch.Status.CONFIRMED
                batch.confirmed_at = now
                batch.actual_receiver_name = actual_receiver_name
                batch.actual_receiver_employee_code = actual_receiver_employee_code
                batch.save(
                    update_fields=[
                        "status",
                        "confirmed_at",
                        "actual_receiver_name",
                        "actual_receiver_employee_code",
                        "updated_at",
                    ]
                )
                for parcel in parcels:
                    if parcel.status_id == "pkg_done":
                        continue
                    from_status = parcel.status_id
                    parcel.actual_receiver_name = actual_receiver_name
                    parcel.actual_receiver_employee_code = actual_receiver_employee_code
                    parcel.status_id = "pkg_done"
                    parcel.confirmed_at = now
                    parcel.completed_at = now
                    parcel.save(
                        update_fields=[
                            "actual_receiver_name",
                            "actual_receiver_employee_code",
                            "status",
                            "confirmed_at",
                            "completed_at",
                            "updated_at",
                        ]
                    )
                    _cancel_parcel_reminder(parcel)
                    _log_parcel_event(
                        parcel,
                        AdmParcelReceiptLog.ACTION_CONFIRMED,
                        from_status=from_status,
                        to_status=parcel.status_id,
                        note=f"Người nhận xác nhận theo lô #{batch.id} và hoàn tất bưu kiện.",
                        metadata={
                            "batch_id": batch.id,
                            "actual_receiver_name": actual_receiver_name,
                        },
                    )
                _cancel_batch_reminder(batch)
            messages.success(request, "Đã xác nhận nhận hàng và hoàn tất các bưu kiện đã chọn.")
            return redirect(request.path)
    return render(
        request,
        "admindocuments/parcel_batch_confirm.html",
        {"batch": batch, "parcels": parcels},
    )


@login_required
def parcel_receipt_register_proxy(request, doc_id: int):
    if request.method != "POST":
        return redirect("admindocuments:parcel_receipt_list")
    parcel_receipt = get_object_or_404(AdmParcelReceipt, pk=doc_id)
    claim_mode = (request.POST.get("claim_mode") or "self").strip()
    proxy_receiver_name = (request.POST.get("proxy_receiver_name") or "").strip()
    if claim_mode == "proxy" and not proxy_receiver_name:
        messages.error(request, "Cần nhập tên người nhận hộ.")
        return redirect(f"{reverse('admindocuments:parcel_receipt_list')}?open={parcel_receipt.id}")
    if claim_mode == "proxy":
        if not parcel_receipt.proxy_qr_token:
            parcel_receipt.proxy_qr_token = uuid.uuid4().hex
        parcel_receipt.proxy_receiver_name = proxy_receiver_name
        parcel_receipt.proxy_receiver_employee_code = ""
    else:
        parcel_receipt.proxy_receiver_name = ""
        parcel_receipt.proxy_receiver_employee_code = ""
    parcel_receipt.updated_by = request.user
    parcel_receipt.save(
        update_fields=[
            "proxy_receiver_name",
            "proxy_receiver_employee_code",
            "proxy_qr_token",
            "updated_by",
            "updated_at",
        ]
    )
    _log_parcel_event(
        parcel_receipt,
        AdmParcelReceiptLog.ACTION_PROXY_REGISTERED,
        actor=request.user,
        to_status=parcel_receipt.status_id,
        note="Cập nhật hình thức nhận hàng.",
        metadata={
            "proxy_receiver_name": proxy_receiver_name,
            "claim_mode": claim_mode,
        },
    )
    if claim_mode == "proxy":
        messages.success(request, "Đã đăng ký người nhận hộ.")
    else:
        messages.success(request, "Đã cập nhật nhận chính chủ.")
    return redirect(f"{reverse('admindocuments:parcel_receipt_list')}?open={parcel_receipt.id}")


@login_required
def parcel_receipt_proxy_claim(request, token: str):
    parcel_receipt = get_object_or_404(
        AdmParcelReceipt.objects.select_related(
            "recipient_directory", "recipient_user", "recipient_user__userprofile"
        ),
        proxy_qr_token=token,
    )
    if request.method == "POST":
        actual_receiver_name = (request.POST.get("actual_receiver_name") or "").strip()
        actual_receiver_employee_code = (request.POST.get("actual_receiver_employee_code") or "").strip()
        if not actual_receiver_name:
            messages.error(request, "Cần nhập họ tên người thực tế đến nhận.")
        else:
            from_status = parcel_receipt.status_id
            parcel_receipt.actual_receiver_name = actual_receiver_name
            parcel_receipt.actual_receiver_employee_code = actual_receiver_employee_code
            now = timezone.now()
            parcel_receipt.status_id = "pkg_done"
            parcel_receipt.confirmed_at = now
            parcel_receipt.completed_at = now
            parcel_receipt.updated_by = request.user
            parcel_receipt.save(
                update_fields=[
                    "actual_receiver_name",
                    "actual_receiver_employee_code",
                    "status",
                    "confirmed_at",
                    "completed_at",
                    "updated_by",
                    "updated_at",
                ]
            )
            _cancel_parcel_reminder(parcel_receipt)
            _log_parcel_event(
                parcel_receipt,
                AdmParcelReceiptLog.ACTION_CONFIRMED,
                actor=request.user,
                from_status=from_status,
                to_status=parcel_receipt.status_id,
                note="Xác nhận nhận hộ bằng mã QR/token.",
                metadata={"actual_receiver_name": actual_receiver_name, "proxy_receiver_name": parcel_receipt.proxy_receiver_name},
            )
            messages.success(request, "Đã ghi nhận nhận hộ thành công.")
            return redirect("admindocuments:parcel_receipt_list")
    return render(
        request,
        "admindocuments/parcel_receipt_proxy_claim.html",
        {"parcel": parcel_receipt},
    )


def _administrative_document_date_filters(queryset, request, *, include_issue_date=True):
    def parse(value):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value or "", fmt).date()
            except (TypeError, ValueError):
                continue
        return None

    start = parse(request.GET.get("start_date"))
    end = parse(request.GET.get("end_date"))
    issue = parse(request.GET.get("issue_date")) if include_issue_date else None
    all_created = request.GET.get("created_all") == "1" and not start and not end
    if not start and not end and not all_created:
        today = timezone.localdate() if settings.USE_TZ else date.today()
        start, end = date(today.year, 1, 1), date(today.year, 12, 31)
    if start and end and start > end:
        start, end = end, start
    if start:
        queryset = queryset.filter(created_at__date__gte=start)
    if end:
        queryset = queryset.filter(created_at__date__lte=end)
    if issue:
        queryset = queryset.filter(issue_date=issue)
    return queryset, {
        "start_date": start.strftime("%d/%m/%Y") if start else "",
        "end_date": end.strftime("%d/%m/%Y") if end else "",
        "issue_date": issue.strftime("%d/%m/%Y") if issue else "",
        "created_all": "1" if all_created else "",
    }


@login_required
@admin_staff_required
def document_list(request):
    query = request.GET.get("q", "").strip()
    company_filter = request.GET.get("company")
    content_type_filter = request.GET.get("ctype")
    doc_type_filter = request.GET.get("dtype")
    sort = request.GET.get("sort", "created")
    direction = request.GET.get("dir", "desc")
    new_doc_id = request.GET.get("new")

    sort_map = {
        "number": "document_number_full",
        "title": "title",
        "type": "doc_type__name",
        "status": "status__code",
        "company": "issuing_company__name",
        "issue": "issue_date",
        "effective": "effective_date",
        "expiry": "expiry_date",
        "created": "created_at",
    }
    order_field = sort_map.get(sort, "created_at")
    if direction == "desc":
        order_field = f"-{order_field}"

    documents_qs = AdmAdministrativeDocument.objects.select_related(
        "doc_type", "signer_role", "status", "issuing_company", "created_by"
    ).filter(is_void=False)
    if query:
        documents_qs = documents_qs.filter(
            Q(title__icontains=query) | Q(document_number_full__icontains=query)
        )
    if company_filter:
        documents_qs = documents_qs.filter(issuing_company_id=company_filter)
    if content_type_filter:
        documents_qs = documents_qs.filter(content_type_id=content_type_filter)
    if doc_type_filter:
        documents_qs = documents_qs.filter(doc_type_id=doc_type_filter)
    documents_qs, date_values = _administrative_document_date_filters(documents_qs, request)
    documents_qs = documents_qs.order_by(order_field)

    page = request.GET.get("page", "1")
    paginator = Paginator(documents_qs, 25)
    try:
        documents = paginator.page(page)
    except PageNotAnInteger:
        documents = paginator.page(1)
    except EmptyPage:
        documents = paginator.page(paginator.num_pages)

    form = AdmAdministrativeDocumentForm()
    doc_types = AdmDocumentType.objects.all().order_by("name")
    content_types = AdmContentType.objects.all().order_by("name")
    signer_roles = AdmSignerRole.objects.all().order_by("title")
    statuses = AdmDocumentStatus.objects.all().order_by("name")
    companies = AdmCompany.objects.all().order_by("name")
    departments = (
        AdmDepartment.objects.select_related("company").all().order_by("name")
    )
    reference_docs = AdmAdministrativeDocument.objects.all().order_by("-created_at")
    ref_prefill_id = request.GET.get("ref", "")
    default_status = AdmDocumentStatus.objects.filter(
        code=AdmDocumentStatus.CODE_DRAFT
    ).first()

    filter_values = {
        "q": query, "company": company_filter or "", "ctype": content_type_filter or "",
        "dtype": doc_type_filter or "", **date_values,
    }
    filter_values = {key: value for key, value in filter_values.items() if value}
    edit_document_id = request.GET.get("edit", "")
    edit_document_id = int(edit_document_id) if edit_document_id.isdecimal() and len(edit_document_id) <= 18 else ""

    return render(
        request,
        "admindocuments/document_list.html",
        {
            "documents": documents,
            "paginator": paginator,
            "page_obj": documents,
            "page_numbers": paginator.get_elided_page_range(documents.number, on_each_side=1, on_ends=1),
            "filter_query": urlencode(filter_values),
            "list_query": urlencode({**filter_values, "sort": sort, "dir": direction}),
            "is_paginated": paginator.num_pages > 1,
            "sort": sort,
            "dir": direction,
            "q": query,
            "active_company_id": company_filter or "",
            "form": form,
            "doc_types": doc_types,
            "content_types": content_types,
            "signer_roles": signer_roles,
            "statuses": statuses,
            "companies": companies,
            "departments": departments,
            "default_status": default_status,
            "has_admin_docs_access": _has_admin_docs_access(request.user),
            "STATUS_DRAFT": AdmDocumentStatus.CODE_DRAFT,
            "STATUS_PENDING": AdmDocumentStatus.CODE_PENDING,
            "STATUS_ISSUED": AdmDocumentStatus.CODE_ISSUED,
            "STATUS_EXPIRED": AdmDocumentStatus.CODE_EXPIRED,
            "new_doc_id": new_doc_id or "",
            "active_content_type_id": content_type_filter or "",
            "active_doc_type_id": doc_type_filter or "",
            "reference_docs": reference_docs,
            "ref_prefill_id": ref_prefill_id,
            "edit_document_id": edit_document_id,
            **date_values,
        },
    )


@login_required
@admin_staff_required
def document_list_export(request):
    query = request.GET.get("q", "").strip()
    company_filter = request.GET.get("company")
    content_type_filter = request.GET.get("ctype")
    doc_type_filter = request.GET.get("dtype")
    sort = request.GET.get("sort", "created")
    direction = request.GET.get("dir", "desc")

    sort_map = {
        "number": "document_number_full",
        "title": "title",
        "type": "doc_type__name",
        "status": "status__code",
        "company": "issuing_company__name",
        "issue": "issue_date",
        "effective": "effective_date",
        "expiry": "expiry_date",
        "created": "created_at",
    }
    order_field = sort_map.get(sort, "created_at")
    if direction == "desc":
        order_field = f"-{order_field}"

    documents_qs = AdmAdministrativeDocument.objects.select_related(
        "doc_type", "signer_role", "status", "issuing_company", "issuing_department"
    ).filter(is_void=False)
    if query:
        documents_qs = documents_qs.filter(
            Q(title__icontains=query) | Q(document_number_full__icontains=query)
        )
    if company_filter:
        documents_qs = documents_qs.filter(issuing_company_id=company_filter)
    if content_type_filter:
        documents_qs = documents_qs.filter(content_type_id=content_type_filter)
    if doc_type_filter:
        documents_qs = documents_qs.filter(doc_type_id=doc_type_filter)
    documents_qs, _ = _administrative_document_date_filters(documents_qs, request)
    documents_qs = documents_qs.order_by(order_field)

    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách văn bản"
    headers = [
        "STT",
        "Số hiệu",
        "Tên văn bản",
        "Loại văn bản",
        "Loại nội dung",
        "Trạng thái",
        "Công ty ban hành",
        "Phòng ban ban hành",
        "Người ký",
        "Số hiệu tham chiếu",
        "Tham chiếu tới",
        "Ngày ban hành",
        "Ngày hiệu lực",
        "Ngày hết hiệu lực",
        "Mã ticket",
        "Người tạo",
        "Ngày tạo",
        "Ghi chú",
    ]
    ws.append(headers)
    for idx, d in enumerate(documents_qs, start=1):
        ws.append(
            [
                idx,
                d.document_number_full,
                d.title,
                d.doc_type.name if d.doc_type_id else "",
                d.content_type.name if d.content_type_id else "",
                d.status.name if d.status_id else "",
                d.issuing_company.name if d.issuing_company_id else "",
                d.issuing_department.name if d.issuing_department_id else "",
                d.signer_role.title if d.signer_role_id else "",
                d.reference_number or "",
                d.reference_document.document_number_full
                if d.reference_document_id
                else "",
                d.issue_date.strftime("%d/%m/%Y") if d.issue_date else "",
                d.effective_date.strftime("%d/%m/%Y") if d.effective_date else "",
                d.expiry_date.strftime("%d/%m/%Y") if d.expiry_date else "",
                d.ticket_code or "",
                d.created_by.get_full_name() if d.created_by_id else "",
                d.created_at.astimezone(timezone.get_current_timezone()).strftime(
                    "%d/%m/%Y %H:%M"
                )
                if d.created_at
                else "",
                d.note or "",
            ]
        )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    now_str = date.today().strftime("%Y%m%d")
    filename = f"danhsachvanbanhanhchin_{now_str}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _paper_document_queryset(request):
    values = {key: request.GET.get(key, "").strip() for key in
              ("q", "paper_type", "region", "requested_department")}
    qs = AdmPaperDocument.objects.select_related(
        "paper_type", "requested_department", "department", "courier_company", "created_by", "updated_by"
    ).filter(is_deleted=False)
    if values["q"]:
        qs = qs.filter(Q(document_number_full__icontains=values["q"]) |
                       Q(courier_tracking_code__icontains=values["q"]) | Q(summary__icontains=values["q"]) |
                       Q(ticket_code__icontains=values["q"]))
    for key in ("paper_type", "requested_department"):
        if values[key].isdecimal() and len(values[key]) <= 18:
            qs = qs.filter(**{f"{key}_id": values[key]})
        elif values[key]:
            values[key] = ""
    if values["region"]:
        qs = qs.filter(region=values["region"])
    qs, date_values = _administrative_document_date_filters(qs, request, include_issue_date=False)
    values.update({key: value for key, value in date_values.items() if key != "issue_date"})
    sort_map = {"code": "document_number_full", "running": "running_number", "type": "paper_type__name",
                "region": "region", "responsible": "responsible_person", "department": "requested_department__shop_name",
                "ticket": "ticket_code", "status": "status", "created": "created_at"}
    sort = request.GET.get("sort", "created")
    if sort not in sort_map:
        sort = "created"
    direction = "asc" if request.GET.get("dir") == "asc" else "desc"
    order = sort_map[sort]
    return qs.order_by(order if direction == "asc" else f"-{order}", "-pk"), values, sort, direction


def _paper_edit_context(doc, form=None):
    form = form if form is not None else AdmPaperDocumentForm(instance=doc, prefix="paper-edit")
    names = list(form.fields)
    return {"doc": doc, "form": form, "edit_columns": [[form[name] for name in names[:5]],
                                                               [form[name] for name in names[5:]]]}


@login_required
@admin_staff_required
def paper_document_list(request):
    documents_qs, filters, sort, direction = _paper_document_queryset(request)
    edit_id = request.GET.get("edit", "")
    edit_id = int(edit_id) if edit_id.isdecimal() and len(edit_id) <= 18 else ""
    form = AdmPaperDocumentForm()
    legacy_edit_id = ""
    if request.method == "POST":
        legacy_edit_id = request.POST.get("edit_id", "")
        edit_instance = get_object_or_404(AdmPaperDocument, pk=legacy_edit_id, is_deleted=False) if legacy_edit_id else None
        form = AdmPaperDocumentForm(request.POST, instance=edit_instance)
        if form.is_valid():
            paper_doc = form.save(commit=False)
            if edit_instance:
                paper_doc.updated_by = request.user
                paper_doc.save()
                messages.success(request, f"Đã cập nhật giấy hành chính: {paper_doc.document_number_full}")
            else:
                paper_doc.created_by = request.user
                paper_doc.responsible_person = (request.user.get_full_name() or "").strip() or request.user.username
                for attempt in range(3):
                    try:
                        with transaction.atomic():
                            paper_doc.running_number, paper_doc.document_number_full = allocate_paper_running_number(
                                paper_doc.paper_type_id, paper_doc.paper_type.code, timezone.now().year)
                            paper_doc.save()
                        break
                    except IntegrityError:
                        if attempt == 2:
                            messages.error(request, "Không thể sinh số hiệu duy nhất, vui lòng thử lại.")
                            return redirect("admindocuments:paper_document_list")
                request.session["paper_new_ids"] = [*request.session.get("paper_new_ids", []), paper_doc.pk]
                request.session["paper_new_numbers"] = [*request.session.get("paper_new_numbers", []), paper_doc.document_number_full]
                messages.success(request, f"Tạo giấy hành chính thành công: {paper_doc.document_number_full}")
            params = request.GET.copy()
            params.pop("edit", None)
            params.pop("create", None)
            target = reverse("admindocuments:paper_document_list")
            return redirect(f"{target}?{params.urlencode()}" if params else target)
        messages.error(request, "Dữ liệu không hợp lệ:\n" + "\n".join(
            f"{form.fields[field].label if field in form.fields else field}: {', '.join(errors)}"
            for field, errors in form.errors.items()))

    paginator = Paginator(documents_qs, 25)
    documents = paginator.get_page(request.GET.get("page", "1"))
    query_values = {key: value for key, value in filters.items() if value}
    choices_qs = AdmPaperDocument.objects.filter(is_deleted=False).order_by()
    context = {
        **filters, "documents": documents, "paginator": paginator, "form": form,
        "paper_types": AdmPaperType.objects.all().order_by("name"),
        "region_options": [{"pk": value, "name": value} for value in choices_qs.exclude(region__isnull=True).exclude(region="").values_list("region", flat=True).distinct().order_by("region")],
        "shop_options": [{"pk": shop.pk, "name": shop.shop_name} for shop in form.fields["requested_department"].queryset],
        "sort": sort, "dir": direction, "filter_query": urlencode(query_values),
        "list_query": urlencode({**query_values, "sort": sort, "dir": direction}),
        "edit_document_id": edit_id, "legacy_edit_id": legacy_edit_id,
        "create_modal_open": request.method == "POST", "paper_documents_nav": True,
        "has_admin_docs_access": _has_admin_docs_access(request.user),
        "create_columns": [[form[name] for name in list(form.fields)[:5]], [form[name] for name in list(form.fields)[5:]]],
        "new_ids": request.session.pop("paper_new_ids", []),
        "new_numbers": request.session.pop("paper_new_numbers", []),
    }
    return render(request, "admindocuments/paper_document_list.html", context)


@login_required
@admin_staff_required
def paper_document_export(request):
    documents, _, _, _ = _paper_document_queryset(request)
    wb = Workbook()
    ws = wb.active
    ws.title = "Giấy hành chính"
    ws.append(["STT", "Số hiệu", "Loại giấy", "Miền", "Người phụ trách", "Phòng giao dịch", "Phòng ban nội bộ",
               "Nội dung", "Đơn vị CPN", "Mã vận đơn", "Tình trạng", "Mã ticket", "Ngày tạo", "Ghi chú"])
    for index, doc in enumerate(documents.iterator(chunk_size=500), 1):
        row = [index, doc.document_number_full, doc.paper_type.name, doc.region or "", doc.responsible_person,
                   doc.requested_department.shop_name if doc.requested_department_id else "",
                   doc.department.name if doc.department_id else "", doc.summary,
                   doc.courier_company.name if doc.courier_company_id else "", doc.courier_tracking_code or "",
                   doc.status or "", doc.ticket_code or "", doc.created_at.strftime("%d/%m/%Y %H:%M"), doc.note or ""]
        ws.append(row)
        for column, value in enumerate(row, 1):
            if isinstance(value, str):
                ws.cell(row=index + 1, column=column).data_type = "s"
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="giay-hanh-chinh-{date.today():%Y%m%d}.xlsx"'
    return response


@login_required
@admin_staff_required
def paper_document_detail(request, doc_id: int):
    doc = get_object_or_404(AdmPaperDocument.objects.select_related(
        "paper_type", "requested_department", "courier_company", "department", "created_by", "updated_by"),
        pk=doc_id, is_deleted=False)
    if request.method == "POST":
        form = AdmPaperDocumentForm(request.POST, instance=doc)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.updated_by = request.user
            obj.save()
            messages.success(request, f"Đã cập nhật giấy hành chính: {obj.document_number_full}")
            return redirect(f"{reverse('admindocuments:paper_document_list')}?edit={doc.pk}")
        return render(request, "admindocuments/includes/paper_document_edit_form.html",
                      _paper_edit_context(doc, form), status=422)
    if request.GET.get("edit") == "1":
        return render(request, "admindocuments/includes/paper_document_edit_form.html", _paper_edit_context(doc))
    if request.GET.get("drawer") == "1":
        return render(request, "admindocuments/paper_document_detail_drawer.html", {"doc": doc})
    return redirect(f"{reverse('admindocuments:paper_document_list')}?edit={doc.pk}")


@login_required
@admin_staff_required
@transaction.atomic
def paper_document_update_field(request, doc_id: int):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "Method not allowed."}, status=405)
    doc = get_object_or_404(AdmPaperDocument.objects.select_for_update(of=("self",)), pk=doc_id, is_deleted=False)
    field = request.POST.get("field", "")
    allowed = AdmPaperDocumentForm.Meta.fields
    if field not in allowed:
        return JsonResponse({"ok": False, "error": "Không được chỉnh sửa trường này."}, status=400)
    data = {name: getattr(doc, f"{name}_id" if doc._meta.get_field(name).is_relation else name) or "" for name in allowed}
    data[field] = request.POST.get("value", "")
    form = AdmPaperDocumentForm(data, instance=doc)
    if not form.is_valid():
        return JsonResponse({"ok": False, "error": "\n".join(str(error) for errors in form.errors.values() for error in errors)}, status=400)
    value = form.cleaned_data[field]
    doc.refresh_from_db()
    setattr(doc, field, value)
    doc.updated_by = request.user
    doc.save(update_fields=[field, "updated_by", "updated_at"])
    display_names = {"paper_type": "name", "requested_department": "shop_name", "department": "name", "courier_company": "name"}
    display = getattr(value, display_names[field]) if value and field in display_names else (value or "—")
    payload = {"ok": True, "field": field, "value": value.pk if hasattr(value, "pk") else value or "",
               "display": display, "document_id": doc.pk}
    if field == "status":
        payload["status"] = {"name": display, "text_color": "#6b7280", "bg_color": "#f3f4f6"}
    return JsonResponse(payload)


@login_required
@admin_staff_required
def paper_document_hide(request, doc_id: int):
    if request.method != "POST":
        return redirect("admindocuments:paper_document_list")
    paper_doc = get_object_or_404(AdmPaperDocument, pk=doc_id, is_deleted=False)
    paper_doc.document_number_full = f"{paper_doc.document_number_full}#VOID#{uuid.uuid4().hex[:6]}"
    paper_doc.is_deleted = True
    paper_doc.deleted_at = timezone.now()
    paper_doc.save(update_fields=["document_number_full", "is_deleted", "deleted_at"])
    messages.success(request, "Đã ẩn giấy và thu hồi số hiệu.")
    return redirect("admindocuments:paper_document_list")


@login_required
@admin_staff_required
def paper_document_import(request):
    if request.method != "POST" or "file" not in request.FILES:
        messages.error(request, "Vui lòng chọn file .xlsx để tải lên.")
        return redirect("admindocuments:paper_document_list")

    upload = request.FILES["file"]
    try:
        wb = load_workbook(upload)
        ws = wb.active
    except Exception:
        messages.error(request, "File không hợp lệ hoặc không thể đọc.")
        return redirect("admindocuments:paper_document_list")

    errors, rows_data = _parse_paper_rows(ws, request.user)
    if errors:
        messages.error(request, "Không nhập dữ liệu. Lỗi:\n" + "\n".join(errors))
        return redirect("admindocuments:paper_document_list")

    created_count = 0
    with transaction.atomic():
        for item in rows_data:
            year_now = item["created_date"].year
            running_number, document_number_full = allocate_paper_running_number(
                item["paper_type_id"], item["paper_type_code"], year_now
            )
            obj = AdmPaperDocument(
                paper_type_id=item["paper_type_id"],
                running_number=running_number,
                region=item["region"],
                requested_department_id=item["requested_dept_id"],
                department_id=item["internal_dept_id"],
                responsible_person=item["responsible"],
                document_number_full=document_number_full,
                summary=item["summary"],
                courier_company_id=item["courier_id"],
                courier_tracking_code=item["tracking_code"],
                ticket_code=item.get("ticket_code", ""),
                status=item["status"],
                note=item["note"],
                created_by=request.user,
                created_at=datetime.combine(item["created_date"], datetime.min.time()).replace(
                    tzinfo=timezone.get_current_timezone()
                ),
            )
            obj.save()
            created_count += 1

    messages.success(request, f"Đã nhập {created_count} dòng thành công.")
    return redirect("admindocuments:paper_document_list")

@login_required
@admin_staff_required
def paper_document_template(request):
    """Generate Excel template for paper documents import."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Template"
    headers = PAPER_IMPORT_HEADERS
    ws.append(headers)

    sample = AdmPaperDocument.objects.select_related(
        "paper_type", "requested_department", "department", "courier_company"
    ).order_by("-created_at").first()
    if sample:
        ws.append(
            [
                "",
                sample.paper_type.name if sample.paper_type_id else "",
                sample.region or "",
                sample.requested_department.shop_name if sample.requested_department_id else "",
                sample.department.name if sample.department_id else "",
                sample.responsible_person or "",
                sample.summary or "",
                sample.courier_company.name if sample.courier_company_id else "",
                sample.courier_tracking_code or "",
                sample.ticket_code or "",
                sample.note or "",
                sample.created_at.date().isoformat(),
            ]
        )
    else:
        ws.append(
            [
                "",
                "Công văn",
                "Miền Bắc",
                "PGD Hà Nội",
                "Phòng Kế toán",
                "Nguyễn Văn A",
                "Nội dung ví dụ",
                "VNPost",
                "ABC123456",
                "TICKET-001",
                "Ghi chú ví dụ",
                date.today().isoformat(),
            ]
        )
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response[
        "Content-Disposition"
    ] = 'attachment; filename="paper_document_template.xlsx"'
    return response


@login_required
@admin_staff_required
def paper_document_import_url(request):
    if request.method != "POST":
        return JsonResponse({"error": "Method not allowed"}, status=405)
    if "file" not in request.FILES:
        return JsonResponse({"error": "Thiếu file .xlsx"}, status=400)
    upload = request.FILES["file"]
    try:
        wb = load_workbook(upload)
        ws = wb.active
    except Exception:
        return JsonResponse({"error": "File không hợp lệ hoặc không thể đọc."}, status=400)

    errors, rows_data = _parse_paper_rows(ws, request.user)
    if errors:
        return JsonResponse({"error": "Lỗi dữ liệu", "details": errors}, status=400)

    token = uuid.uuid4().hex
    serialized = []
    for item in rows_data:
        s = item.copy()
        s["created_date"] = s["created_date"].isoformat()
        serialized.append(s)
    request.session[f"paper_import_{token}"] = serialized
    request.session.modified = True
    return JsonResponse({"status": "ok", "token": token, "rows": len(rows_data)})


@login_required
@admin_staff_required
def paper_document_import_commit(request):
    if request.method != "POST":
        return JsonResponse({"error": "Method not allowed"}, status=405)
    token = request.POST.get("token")
    if not token:
        return JsonResponse({"error": "Thiếu token"}, status=400)
    rows_data = request.session.get(f"paper_import_{token}")
    if not rows_data:
        return JsonResponse({"error": "Token không hợp lệ hoặc đã hết hạn"}, status=400)

    created_count = 0
    created_ids = []
    created_numbers = []
    with transaction.atomic():
        for item in rows_data:
            created_date = datetime.fromisoformat(item["created_date"]).date()
            year_now = created_date.year
            running_number, document_number_full = allocate_paper_running_number(
                item["paper_type_id"], item["paper_type_code"], year_now
            )
            obj = AdmPaperDocument(
                paper_type_id=item["paper_type_id"],
                running_number=running_number,
                region=item["region"],
                requested_department_id=item["requested_dept_id"],
                department_id=item["internal_dept_id"],
                responsible_person=item["responsible"],
                document_number_full=document_number_full,
                summary=item["summary"],
                courier_company_id=item["courier_id"],
                courier_tracking_code=item["tracking_code"],
                ticket_code=item.get("ticket_code", ""),
                status=item["status"],
                note=item["note"],
                created_by=request.user,
                created_at=datetime.combine(created_date, datetime.min.time()).replace(
                    tzinfo=timezone.get_current_timezone()
                ),
            )
            obj.save()
            created_count += 1
            created_ids.append(obj.id)
            created_numbers.append(document_number_full)
    try:
        del request.session[f"paper_import_{token}"]
        request.session.modified = True
    except KeyError:
        pass
    if created_ids:
        request.session["paper_new_ids"] = created_ids
        request.session["paper_new_numbers"] = created_numbers
        request.session.modified = True
    messages.success(request, f"Đã nhập {created_count} giấy hành chính từ file Excel thành công.")
    return JsonResponse({"status": "ok", "created": created_count})


def counter_status(next_number, used_max, used_count):
    """Trạng thái hiển thị của bộ đếm số hiệu.

    Số trống (chưa dùng hoặc đã thu hồi) luôn có thể được cấp lại bằng cách nhập số đó vào
    "Số tiếp theo", nên đây là thông tin chứ không phải cảnh báo.
    used_count: số lượng số gốc đang được dùng (không tính số thu hồi).
    """
    free_below = max(used_max - used_count, 0)  # số trống nằm dưới số đã dùng cao nhất
    free_ahead = max(next_number - used_max - 1, 0)  # số bị bỏ qua giữa số đã dùng cao nhất và số tiếp theo
    free_total = free_below + free_ahead
    if next_number <= used_max:
        return {"code": "fill", "label": "Đang cấp số trống", "hint": f"Số tiếp theo ({next_number}) thấp hơn số đã dùng cao nhất ({used_max}): đang cấp lại số còn trống."}
    if free_total:
        parts = []
        if free_below:
            parts.append(f"{free_below} số trống nằm dưới {used_max}")
        if free_ahead:
            parts.append(f"{free_ahead} số bỏ qua từ {used_max + 1} đến {next_number - 1}")
        return {"code": "free", "label": f"Còn {free_total} số trống", "hint": "; ".join(parts).capitalize() + ". Nhập số đó vào \"Số tiếp theo\" để cấp lại."}
    return {"code": "ok", "label": "Liên tục", "hint": "Không có số trống."}


@login_required
@admin_staff_required
def document_counter_manage(request):
    doc_types = AdmDocumentType.objects.all().order_by("name")
    companies = AdmCompany.objects.all().order_by("name")
    paper_types = AdmPaperType.objects.all().order_by("name")
    year_now = timezone.now().year
    active_tab = "paper" if request.GET.get("tab") == "paper" else "admin"

    def counters_url(tab):
        base = reverse("admindocuments:admindocuments_counters")
        return f"{base}?tab=paper" if tab == "paper" else base

    if request.method == "POST" and request.POST.get("target", "admin") == "admin":
        doc_type_id = request.POST.get("doc_type")
        company_id = request.POST.get("company")
        year = request.POST.get("year") or year_now
        next_number = request.POST.get("next_number")

        try:
            year = int(year)
            next_number = int(next_number)
        except (TypeError, ValueError):
            messages.error(request, "Năm và số tiếp theo phải là số.")
            return redirect(counters_url("admin"))

        if not doc_type_id or not company_id or next_number < 1 or not 1 <= year <= 9999:
            messages.error(request, "Vui lòng chọn đủ loại văn bản và công ty.")
            return redirect(counters_url("admin"))

        with transaction.atomic():
            counter, _ = AdmDocumentCounter.objects.select_for_update().get_or_create(
                doc_type_id=doc_type_id, company_id=company_id, year=year,
                defaults={"next_number": 1},
            )
            used_qs = AdmAdministrativeDocument.objects.filter(
                doc_type_id=doc_type_id,
                issuing_company_id=company_id,
                number_year=year,
                number_suffix="",
                is_void=False,
            )
            max_used = used_qs.aggregate(mx=Max("running_number"))["mx"] or 0
            exists_number = used_qs.filter(running_number=next_number).exists()
            if exists_number:
                messages.error(
                    request,
                    f"Số {next_number} đã được sử dụng (đã dùng tới {max_used}).",
                )
                return redirect(counters_url("admin"))

            counter.next_number = next_number
            counter.save(update_fields=["next_number", "updated_at"])
        dt_name = doc_types.filter(pk=doc_type_id).first()
        dt_name = dt_name.name if dt_name else doc_type_id
        comp_name = companies.filter(pk=company_id).first()
        comp_code = comp_name.code if comp_name else company_id
        messages.success(request, f"Cập nhật số tiếp theo cho '{dt_name}' - {comp_code} năm {year} -> {next_number}.")
        return redirect(counters_url("admin"))

    # prepare tracking
    used_map = {
        (row["doc_type_id"], row["issuing_company_id"], row["number_year"]): {
            "used_max": row["used_max"],
            "total": row["total"],
            "base_total": row["base_total"],
        }
        for row in AdmAdministrativeDocument.objects.filter(is_void=False)
        .values("doc_type_id", "issuing_company_id", "number_year")
        .annotate(used_max=Max("running_number"), total=Count("id"), base_total=Count("id", filter=Q(number_suffix="")))
    }

    counters = AdmDocumentCounter.objects.select_related("doc_type", "company").order_by(
        "-year", "doc_type__name", "company__name"
    )
    counter_rows = []
    for c in counters:
            key = (c.doc_type_id, c.company_id, c.year)
            used_info = used_map.get(key, {"used_max": 0, "total": 0, "base_total": 0})
            gap_from = (used_info["used_max"] or 0) + 1
            counter_rows.append(
                {
                "doc_type": c.doc_type,
                "doc_type_id": c.doc_type_id,
                "company": c.company,
                "company_id": c.company_id,
                "year": c.year,
                "next_number": c.next_number,
                "used_max": used_info["used_max"] or 0,
                "total": used_info["total"] or 0,
                "base_total": used_info["base_total"] or 0,
                "gap_from": gap_from,
                "status": counter_status(c.next_number, used_info["used_max"] or 0, used_info["base_total"] or 0),
                }
            )
    # Paper counter update
    if request.method == "POST" and request.POST.get("target") == "paper":
        paper_type_id = request.POST.get("paper_type")
        year = request.POST.get("year") or year_now
        next_number = request.POST.get("next_number")
        try:
            year = int(year)
            next_number = int(next_number)
        except (TypeError, ValueError):
            messages.error(request, "Năm và số tiếp theo phải là số.")
            return redirect(counters_url("paper"))
        if not paper_type_id:
            messages.error(request, "Vui lòng chọn loại giấy.")
            return redirect(counters_url("paper"))
        exists_active = AdmPaperDocument.objects.filter(
            paper_type_id=paper_type_id,
            created_at__year=year,
            is_deleted=False,
            running_number=next_number,
        ).exists()
        if exists_active:
            messages.error(request, f"Số {next_number} đã được sử dụng cho giấy này.")
            return redirect(counters_url("paper"))
        counter, created = AdmPaperCounter.objects.get_or_create(
            paper_type_id=paper_type_id,
            year=year,
            defaults={"next_number": next_number},
        )
        if not created:
            counter.next_number = next_number
            counter.save(update_fields=["next_number", "updated_at"])
        pt_name = paper_types.filter(pk=paper_type_id).first()
        pt_name = pt_name.name if pt_name else paper_type_id
        messages.success(request, f"Cập nhật số tiếp theo cho giấy '{pt_name}' năm {year} -> {next_number}.")
        return redirect(counters_url("paper"))

    # Paper counters display
    paper_rows = []
    paper_years = (
        AdmPaperDocument.objects.values_list("created_at__year", flat=True).distinct()
    )
    years_set = set(y for y in paper_years if y)
    years_set.add(year_now)
    for pt in paper_types:
        for y in sorted(years_set):
            active_qs = AdmPaperDocument.objects.filter(
                paper_type_id=pt.id, created_at__year=y, is_deleted=False
            )
            void_qs = AdmPaperDocument.objects.filter(
                paper_type_id=pt.id, created_at__year=y, is_deleted=True
            )
            active_numbers = list(active_qs.values_list("running_number", flat=True))
            void_numbers = list(void_qs.values_list("running_number", flat=True))
            if not active_numbers and not void_numbers:
                continue
            active_numbers_sorted = sorted([n for n in active_numbers if n])
            candidate = 1
            for n in active_numbers_sorted:
                if n and n > candidate:
                    break
                candidate = (n or candidate) + 1
            counter = AdmPaperCounter.objects.filter(paper_type_id=pt.id, year=y).first()
            if counter and counter.next_number and counter.next_number > candidate:
                candidate = counter.next_number
            used_max = active_numbers_sorted[-1] if active_numbers_sorted else 0
            paper_rows.append(
                {
                    "paper_type": pt,
                    "paper_type_id": pt.id,
                    "year": y,
                    "used_max": used_max,
                    "next_number": candidate,
                    "gap_from": candidate,
                    "total": len(active_numbers_sorted),
                    "void_count": len([n for n in void_numbers if n]),
                    "status": counter_status(candidate, used_max, len(active_numbers_sorted)),
                    "counter": counter,
                }
            )

    paper_rows.sort(key=lambda row: (-row["year"], row["paper_type"].name))

    return render(
        request,
        "admindocuments/document_counters.html",
        {
            "doc_types": doc_types,
            "companies": companies,
            "paper_types": paper_types,
            "year_now": year_now,
            "counters": counter_rows,
            "paper_counters": paper_rows,
            "active_tab": active_tab,
            "counter_years": sorted({row["year"] for row in counter_rows}, reverse=True),
            "paper_counter_years": sorted({row["year"] for row in paper_rows}, reverse=True),
        },
    )


@login_required
@admin_staff_required
def document_counter_visual(request):
    doc_types = AdmDocumentType.objects.all().order_by("name")
    companies = AdmCompany.objects.all().order_by("name")
    year_now = timezone.now().year

    def _safe_int(val, default=None):
        try:
            return int(val)
        except (TypeError, ValueError):
            return default

    doc_type_id = _safe_int(request.GET.get("doc_type"))
    company_id = _safe_int(request.GET.get("company"))
    year = _safe_int(request.GET.get("year"), year_now)

    used_numbers = set()
    void_numbers = set()
    counter_next = None
    max_number = 0
    groups = {}

    if doc_type_id and company_id:
        docs = AdmAdministrativeDocument.objects.filter(
            doc_type_id=doc_type_id,
            issuing_company_id=company_id,
            number_year=year,
        ).values_list("running_number", "is_void", "id", "number_suffix")
        for rn, is_void, document_id, suffix in docs:
            if rn is None:
                continue
            group = groups.setdefault(rn, {"number": rn, "parent_id": None, "branches": 0})
            if not suffix:
                if group["parent_id"] is None or not is_void:
                    group["parent_id"] = document_id
            else:
                group["branches"] += 1
            if is_void:
                void_numbers.add(rn)
            else:
                used_numbers.add(rn)
        counter = AdmDocumentCounter.objects.filter(
            doc_type_id=doc_type_id, company_id=company_id, year=year
        ).first()
        if counter:
            counter_next = counter.next_number or 1
            max_number = max(max_number, (counter_next or 1) - 1)
        if used_numbers or void_numbers:
            max_number = max(max_number, max(used_numbers | void_numbers))
        if max_number == 0 and counter_next:
            max_number = (counter_next or 1) - 1

    years_set = {year_now}
    years_set.update(
        AdmDocumentCounter.objects.values_list("year", flat=True).distinct()
    )
    years_set.update(
        AdmAdministrativeDocument.objects.values_list("number_year", flat=True).distinct()
    )
    years = sorted({y for y in years_set if y}, reverse=True)
    context = {
        "doc_types": doc_types,
        "companies": companies,
        "year_now": year_now,
        "selected_doc_type": doc_type_id or "",
        "selected_company": company_id or "",
        "selected_year": year,
        "used_numbers": sorted(list(used_numbers)),
        "void_numbers": sorted(list(void_numbers)),
        "max_number": max_number,
        "counter_next": counter_next,
        "years": years,
        "numbers": list(range(1, max_number + 1)) if max_number else [],
        "number_cells": [dict(groups.get(n, {"number": n, "parent_id": None, "branches": 0}),
                              state="used" if n in used_numbers else "void" if n in void_numbers else "free")
                         for n in range(1, max_number + 1)],
    }
    if request.GET.get("partial") == "1":
        return render(request, "admindocuments/document_counter_visual_partial.html", context)
    return render(request, "admindocuments/document_counter_visual.html", context)


@login_required
@admin_staff_required
def paper_counter_visual(request):
    paper_type_id = request.GET.get("paper_type")
    year_now = timezone.now().year
    try:
        year = int(request.GET.get("year", year_now))
    except (TypeError, ValueError):
        year = year_now

    used_numbers = set()
    void_numbers = set()
    max_number = 0
    next_number = 1

    if paper_type_id:
        active_qs = AdmPaperDocument.objects.filter(
            paper_type_id=paper_type_id, created_at__year=year, is_deleted=False
        )
        void_qs = AdmPaperDocument.objects.filter(
            paper_type_id=paper_type_id, created_at__year=year, is_deleted=True
        )
        used_numbers.update([n for n in active_qs.values_list("running_number", flat=True) if n])
        void_numbers.update([n for n in void_qs.values_list("running_number", flat=True) if n])
        if used_numbers or void_numbers:
            max_number = max(used_numbers | void_numbers)
        if max_number == 0:
            max_number = len(used_numbers)

        # compute next_number (gap fill)
        existing = sorted(list(used_numbers))
        candidate = 1
        for n in existing:
            if n > candidate:
                break
            candidate = n + 1
        counter = AdmPaperCounter.objects.filter(paper_type_id=paper_type_id, year=year).first()
        if counter and counter.next_number and counter.next_number > candidate:
            candidate = counter.next_number
        next_number = candidate
        if max_number:
            max_number = max(max_number, next_number)

    numbers = list(range(1, max_number + 1)) if max_number else []
    return render(
        request,
        "admindocuments/paper_counter_visual_partial.html",
        {
            "used_numbers": used_numbers,
            "void_numbers": void_numbers,
            "max_number": max_number,
            "next_number": next_number,
            "numbers": numbers,
        },
    )


@login_required
@admin_staff_required
def master_data(request):
    if request.method == "POST" and request.POST.get("form_name") == "status_colors":
        if not request.user.has_perm("app_admindocuments.change_admdocumentstatus"):
            return HttpResponseForbidden("Bạn không có quyền chỉnh màu trạng thái văn bản.")
        try:
            status_id = int(request.POST.get("status_id", ""))
        except (TypeError, ValueError):
            return HttpResponse("Mã trạng thái không hợp lệ.", status=400)
        status = get_object_or_404(AdmDocumentStatus, pk=status_id)
        color_form = AdmDocumentStatusColorsForm(request.POST, instance=status)
        if color_form.is_valid():
            status = color_form.save(commit=False)
            status.updated_by = request.user
            status.save(update_fields=["badge_text_color", "badge_bg_color", "updated_by", "updated_at"])
            messages.success(request, "Đã cập nhật màu trạng thái văn bản.")
        else:
            messages.error(request, "Màu không hợp lệ. Vui lòng chọn màu chữ và màu nền theo định dạng #RRGGBB.")
        return redirect(f"{reverse('admindocuments:master_data')}#status")
    form_classes = {
        "doc_type": (AdmDocumentTypeForm, "Loại văn bản"),
        "content_type": (AdmContentTypeForm, "Loại nội dung"),
        "signer_role": (AdmSignerRoleForm, "Chức danh người ký"),
        "status": (AdmDocumentStatusForm, "Trạng thái văn bản"),
        "company": (AdmCompanyForm, "Công ty"),
        "department": (AdmDepartmentForm, "Phòng ban"),
        "receive_location": (AdmParcelReceiveLocationForm, "Nơi nhận"),
        "paper_type": (AdmPaperTypeForm, "Loại giấy (paper)"),
        "courier": (AdmCourierCompanyForm, "Đơn vị chuyển phát"),
    }
    perm_map = {
        "doc_type": "app_admindocuments.add_admdocumenttype",
        "content_type": "app_admindocuments.add_admcontenttype",
        "signer_role": "app_admindocuments.add_admsignerrole",
        "status": "app_admindocuments.add_admdocumentstatus",
        "company": "app_admindocuments.add_admcompany",
        "department": "app_admindocuments.add_admdepartment",
        "receive_location": "app_admindocuments.add_admparcelreceivelocation",
        "paper_type": "app_admindocuments.add_admpapertype",
        "courier": "app_admindocuments.add_admcouriercompany",
    }
    forms_map = {key: cls() for key, (cls, _) in form_classes.items()}

    if request.method == "POST":
        form_key = request.POST.get("form_name")
        form_entry = form_classes.get(form_key)
        if not form_entry:
            messages.error(request, "Form không hợp lệ.")
            return redirect("admindocuments:master_data")
        form_cls, label = form_entry
        required_perm = perm_map.get(form_key)
        if required_perm and not request.user.has_perm(required_perm):
            messages.error(request, "Bạn không có quyền thêm %s." % label.lower())
            return redirect(f"{reverse('admindocuments:master_data')}#{form_key}")
        bound_form = form_cls(request.POST)
        forms_map[form_key] = bound_form
        if bound_form.is_valid():
            obj = bound_form.save(commit=False)
            if hasattr(obj, "created_by"):
                obj.created_by = request.user
            if hasattr(obj, "updated_by"):
                obj.updated_by = request.user
            obj.save()
            messages.success(request, f"Đã thêm {label.lower()}.")
            return redirect(f"{reverse('admindocuments:master_data')}#{form_key}")
        messages.error(request, "Dữ liệu không hợp lệ, vui lòng kiểm tra lại.")

    # Ensure dropdowns ordered nicely
    if "department" in forms_map:
        forms_map["department"].fields["company"].queryset = AdmCompany.objects.order_by("name")

    context = {
        "doc_types": AdmDocumentType.objects.all().order_by("name"),
        "content_types": AdmContentType.objects.all().order_by("name"),
        "signer_roles": AdmSignerRole.objects.all().order_by("title"),
        "statuses": AdmDocumentStatus.objects.all().order_by("name"),
        "companies": AdmCompany.objects.all().order_by("name"),
        "departments": AdmDepartment.objects.select_related("company").all().order_by(
            "company__code", "name"
        ),
        "receive_locations": AdmParcelReceiveLocation.objects.all().order_by("name"),
        "paper_types": AdmPaperType.objects.all().order_by("name"),
        "couriers": AdmCourierCompany.objects.all().order_by("name"),
        "forms": forms_map,
        "has_admin_docs_access": _has_admin_docs_access(request.user),
        "form_permissions": {
            key: (not perm_map.get(key) or request.user.has_perm(perm_map[key]))
            for key in form_classes.keys()
        },
    }
    return render(request, "admindocuments/master_data.html", context)


@login_required
@admin_staff_required
def document_create(request):
    if request.method != "POST":
        return redirect("admindocuments:admindocuments_list")
    form = AdmAdministrativeDocumentForm(request.POST, request.FILES)
    if not form.is_valid():
        errors = form.errors.as_data()
        logger.warning(
            "Administrative document create validation failed user_id=%s field_errors=%s",
            request.user.pk,
            {name: [error.code for error in field_errors] for name, field_errors in errors.items()},
        )
        for name, field_errors in errors.items():
            label = form.fields[name].label if name in form.fields else "Thông tin văn bản"
            for error in field_errors:
                if error.code == "required":
                    detail = "Vui lòng nhập hoặc chọn trường này."
                elif name in {"issue_date", "effective_date", "expiry_date"} and error.code == "invalid":
                    detail = "Ngày không hợp lệ. Vui lòng nhập theo định dạng dd/mm/yyyy."
                else:
                    detail = " ".join(str(message) for message in error.messages)
                messages.error(request, f"{label}: {detail}")
        return redirect("admindocuments:admindocuments_list")

    doc = form.save(commit=False)
    doc.created_by = request.user

    if not getattr(doc, "status_id", None):
        default_status = AdmDocumentStatus.objects.filter(
            code=AdmDocumentStatus.CODE_DRAFT
        ).first() or AdmDocumentStatus.objects.order_by("id").first()
        if default_status:
            doc.status = default_status

    attachment_link = form.cleaned_data.get("attachment_link")
    issue_date = form.cleaned_data.get("issue_date")
    if not issue_date:
        issue_date = timezone.localdate() if settings.USE_TZ else date.today()
    doc.issue_date = issue_date

    # The upload is stored once, via _create_attachment_version below. Leaving it on
    # doc.attachment would make doc.save() consume the upload (and move its temp file).
    doc.attachment = None

    attempts = 0
    max_attempts = 30
    while attempts < max_attempts:
        attempts += 1
        try:
            with transaction.atomic():
                current_year = administrative_number_year()
                doc.number_year = current_year
                doc.running_number, void_conflicts = allocate_running_number(
                    doc.doc_type_id, doc.issuing_company_id, current_year
                )
                doc.save()
            break
        except IntegrityError:
            AdmDocumentCounter.objects.filter(
                doc_type_id=doc.doc_type_id,
                company_id=doc.issuing_company_id,
                year=current_year,
            ).update(next_number=F("next_number") + 1)
            if attempts >= max_attempts:
                messages.error(
                    request, "Could not allocate unique number. Please try again."
                )
                return redirect("admindocuments:admindocuments_list")
            continue

    try:
        _create_attachment_version(
            document=doc,
            user=request.user,
            uploaded_file=request.FILES.get("attachment"),
            link=attachment_link,
        )
    except Exception:
        logger.exception("Attachment upload failed for document %s", doc.pk)
        messages.warning(
            request,
            "Văn bản đã được tạo nhưng tải file đính kèm thất bại. "
            "Vui lòng tải lại file trong trang chi tiết.",
        )

    messages.success(
        request,
        f"Văn bản {doc.document_number_full} đã tạo thành công."
        + ("" if not void_conflicts else f" (Sử dụng lại số đã thu hồi từ {len(void_conflicts)} văn bản)"),
    )
    redirect_url = f"{reverse('admindocuments:admindocuments_list')}?new={doc.id}"
    return redirect(redirect_url)


@login_required
@admin_staff_required
def document_split(request, doc_id):
    is_popup = request.headers.get("x-requested-with") == "XMLHttpRequest"
    parent = get_object_or_404(AdmAdministrativeDocument.objects.select_related(
        "doc_type", "issuing_company"), pk=doc_id, split_parent__isnull=True)
    if request.method == "POST":
        try:
            key = uuid.UUID(request.POST.get("request_key", ""))
        except (ValueError, TypeError, AttributeError):
            key = None
        if key:
            existing = AdmAdministrativeDocument.objects.filter(
                split_request_key=key, split_parent=parent, created_by=request.user).first()
            if existing:
                if is_popup:
                    return JsonResponse({"redirect_url": reverse("admindocuments:admindocuments_detail", args=[existing.pk])})
                return redirect("admindocuments:admindocuments_detail", doc_id=existing.pk)
    try:
        suffix = next_split_suffix(parent)
    except ValueError as exc:
        suffix = None
        exhausted_message = str(exc)
    else:
        exhausted_message = ""
    counter = AdmDocumentCounter.objects.filter(
        doc_type_id=parent.doc_type_id, company_id=parent.issuing_company_id, year=parent.number_year).first()
    form = AdmAdministrativeDocumentSplitForm(
        request.POST if request.method == "POST" else None,
        request.FILES if request.method == "POST" else None,
        parent=parent, suffix=suffix or "",
    )
    if request.method == "POST" and suffix and form.is_valid():
        doc = form.save(commit=False)
        doc.attachment = None
        try:
            doc, created = create_split_document(
                parent=parent, document=doc, expected_suffix=form.cleaned_data["expected_suffix"],
                request_key=form.cleaned_data["request_key"], user=request.user)
        except ValueError as exc:
            form.add_error(None, str(exc))
        except IntegrityError:
            form.add_error(None, "Số dự kiến vừa được sử dụng. Vui lòng lưu lại.")
        else:
            if created:
                try:
                    _create_attachment_version(document=doc, user=request.user,
                        uploaded_file=request.FILES.get("attachment"), link=form.cleaned_data.get("attachment_link"))
                except Exception:
                    logger.exception("Attachment upload failed for split document %s", doc.pk)
                    messages.warning(request, "Văn bản đã tạo nhưng tải file thất bại. Vui lòng tải lại trong chi tiết.")
                messages.success(request, f"Đã tạo văn bản tách thửa {doc.document_number_full}.")
            if is_popup:
                return JsonResponse({"redirect_url": reverse("admindocuments:admindocuments_detail", args=[doc.pk])})
            return redirect("admindocuments:admindocuments_detail", doc_id=doc.pk)
        # Refresh the preview while keeping every entered field and the request key.
        try:
            suffix = next_split_suffix(parent)
        except ValueError as exc:
            suffix = None
            exhausted_message = str(exc)
        data = form.data.copy()
        data["expected_suffix"] = suffix or ""
        form.data = data
    column_names = [
        ["doc_type", "issuing_company", "issuing_department", "signer_role", "content_type", "title", "status", "ticket_code"],
        ["issue_date", "effective_date", "expiry_date", "reference_number", "reference_document", "is_reference_document", "attachment", "attachment_link", "note"],
    ]
    columns = [[form[name] for name in names if name in form.fields] for names in column_names]
    template = "admindocuments/includes/document_split_form.html" if is_popup else "admindocuments/document_split.html"
    return render(request, template, {
        "split_columns": columns,
        "parent": parent, "form": form, "suffix": suffix, "exhausted_message": exhausted_message,
        "counter_next": counter.next_number if counter else None,
        "group_documents": split_number_group(parent).select_related("issuing_department", "signer_role")
            .order_by("is_void", "number_suffix", "created_at"),
    })


@login_required
@admin_staff_required
def document_change_status(request, doc_id: int):
    if request.method != "POST":
        return redirect("admindocuments:admindocuments_list")

    status_code = request.POST.get("status_code")
    if not status_code:
        messages.error(request, "Missing status_code.")
        return redirect("admindocuments:admindocuments_list")

    try:
        doc = AdmAdministrativeDocument.objects.select_related("status").get(pk=doc_id)
    except AdmAdministrativeDocument.DoesNotExist:
        messages.error(request, "Document not found.")
        return redirect("admindocuments:admindocuments_list")

    new_status = AdmDocumentStatus.objects.filter(code=status_code).first()
    if not new_status:
        messages.error(request, "Invalid status.")
        return redirect("admindocuments:admindocuments_list")

    if doc.status_id == new_status.id:
        messages.info(request, "Status is already set to this value.")
        return redirect("admindocuments:admindocuments_list")

    doc.status = new_status
    if hasattr(doc, "updated_by"):
        doc.updated_by = request.user
    doc.save()
    messages.success(request, "Status updated successfully.")
    return redirect("admindocuments:admindocuments_list")


@login_required
@admin_staff_required
def document_detail(request, doc_id: int):
    try:
        doc = AdmAdministrativeDocument.objects.select_related(
            "doc_type",
            "content_type",
            "signer_role",
            "issuing_company",
            "issuing_department",
            "status",
        ).prefetch_related("attachments__created_by", "attachments__deleted_by").get(pk=doc_id)
    except AdmAdministrativeDocument.DoesNotExist:
        if request.GET.get("edit") == "1":
            return HttpResponse("Văn bản không tồn tại.", status=404)
        messages.error(request, "Document not found.")
        return redirect("admindocuments:admindocuments_list")

    as_drawer = request.GET.get("drawer") == "1"
    if request.GET.get("edit") == "1":
        return render(request, "admindocuments/includes/document_edit_form.html", _document_edit_context(doc))
    if not as_drawer:
        return redirect(f"{reverse('admindocuments:admindocuments_list')}?edit={doc.pk}")
    return render(request, "admindocuments/document_detail_drawer.html", {
        "doc": doc, "document_history": _document_history(doc),
        "STATUS_DRAFT": AdmDocumentStatus.CODE_DRAFT,
        "STATUS_ISSUED": AdmDocumentStatus.CODE_ISSUED,
        "STATUS_EXPIRED": AdmDocumentStatus.CODE_EXPIRED,
    })


@login_required
@admin_staff_required
def document_update(request, doc_id: int):
    ajax = request.headers.get("x-requested-with") == "XMLHttpRequest"
    try:
        doc = AdmAdministrativeDocument.objects.get(pk=doc_id)
    except AdmAdministrativeDocument.DoesNotExist:
        if ajax:
            return JsonResponse({"ok": False, "error": "Văn bản không tồn tại."}, status=404)
        messages.error(request, "Document not found.")
        return redirect("admindocuments:admindocuments_list")

    if request.method != "POST":
        return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)

    post_data = request.POST.copy()
    if post_data.get("reference_document") in {"None", "none", "null"}:
        post_data["reference_document"] = ""
    if post_data.get("reference_number") in {"None", "none", "null"}:
        post_data["reference_number"] = ""

    form_class = AdmAdministrativeDocumentEditPopupForm if ajax else AdmAdministrativeDocumentUpdateForm
    form = form_class(post_data, request.FILES, instance=doc)
    if not form.is_valid():
        if ajax:
            return render(request, "admindocuments/includes/document_edit_form.html",
                          _document_edit_context(doc, form), status=422)
        messages.error(request, "Invalid data. Please check required fields.")
        return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)

    doc = form.save(commit=False)
    attachment_link = form.cleaned_data.get("attachment_link")
    uploaded_file = request.FILES.get("attachment")
    issue_date = form.cleaned_data.get("issue_date") or doc.issue_date
    if not issue_date:
        issue_date = timezone.localdate() if settings.USE_TZ else date.today()
    doc.issue_date = issue_date
    doc.updated_by = request.user

    try:
        os.makedirs(settings.MEDIA_ROOT, exist_ok=True)
    except Exception:
        pass

    with transaction.atomic():
        doc.save()
        if uploaded_file or attachment_link:
            _create_attachment_version(
                document=doc,
                user=request.user,
                uploaded_file=uploaded_file,
                link=attachment_link,
            )

    if ajax:
        messages.success(request, "Đã cập nhật văn bản.")
        return JsonResponse({"ok": True, "document_id": doc.pk})
    if uploaded_file or attachment_link:
        messages.success(request, "Document updated successfully!")
    else:
        messages.success(request, "Document updated (không có file/link mới).")
    return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)


def _document_edit_context(doc, form=None):
    if form is None:
        form = AdmAdministrativeDocumentEditPopupForm(instance=doc)
    columns = [
        ["doc_type", "issuing_company", "issuing_department", "signer_role", "content_type", "title", "status", "ticket_code"],
        ["issue_date", "effective_date", "expiry_date", "reference_number", "reference_document", "is_reference_document", "attachment", "attachment_link", "note"],
    ]
    return {"doc": doc, "form": form, "edit_columns": [[form[name] for name in column] for column in columns]}


def _document_history(doc):
    labels = {**AdmAdministrativeDocumentForm.Meta.labels, "issue_date": "Ngày ban hành",
              "document_number_full": "Số hiệu", "is_void": "Vô hiệu", "number_suffix": "Hậu tố tách thửa",
              "running_number": "Số thứ tự", "number_year": "Năm số hiệu", "split_parent": "Văn bản gốc"}
    records = []
    for item in doc.history.select_related("changed_by").order_by("-changed_at", "-pk"):
        changes = item.changes if isinstance(item.changes, dict) else {}
        fields = []
        if item.change_type == AdmAdministrativeDocumentHistory.CHANGE_UPDATED:
            for name, values in changes.items():
                if isinstance(values, dict):
                    fields.append({"label": labels.get(name, name), "old": values.get("old"), "new": values.get("new")})
        records.append({"timestamp": item.changed_at, "user": item.changed_by,
                        "type": item.change_type, "fields": fields})
    return records


def _build_update_form_data(doc):
    def _fmt_date(value):
        return value.strftime("%d/%m/%Y") if value else ""

    return {
        "doc_type": doc.doc_type_id,
        "content_type": doc.content_type_id,
        "reference_number": doc.reference_number or "",
        "reference_document": doc.reference_document_id or "",
        "is_reference_document": "true" if doc.is_reference_document else "",
        "title": doc.title or "",
        "signer_role": doc.signer_role_id,
        "issuing_company": doc.issuing_company_id,
        "issuing_department": doc.issuing_department_id,
        "issue_date": _fmt_date(doc.issue_date),
        "effective_date": _fmt_date(doc.effective_date),
        "expiry_date": _fmt_date(doc.expiry_date),
        "status": doc.status_id,
        "ticket_code": doc.ticket_code or "",
        "note": doc.note or "",
    }


@login_required
@admin_staff_required
@transaction.atomic
def document_update_field(request, doc_id: int):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "Method not allowed."}, status=405)

    try:
        doc = AdmAdministrativeDocument.objects.select_for_update(of=("self",)).select_related(
            "content_type",
            "signer_role",
            "issuing_department",
            "status",
        ).get(pk=doc_id)
    except AdmAdministrativeDocument.DoesNotExist:
        return JsonResponse({"ok": False, "error": "Document not found."}, status=404)

    field = (request.POST.get("field") or "").strip()
    value = request.POST.get("value", "")
    allowed_fields = {
        "title",
        "content_type",
        "signer_role",
        "issuing_department",
        "issue_date",
        "effective_date",
        "expiry_date",
        "status",
        "reference_number",
        "note",
        "ticket_code",
        "reference_document",
        "is_reference_document",
    }
    if field not in allowed_fields:
        return JsonResponse({"ok": False, "error": "Field not allowed."}, status=400)

    data = _build_update_form_data(doc)
    data[field] = value

    if field == "reference_document" and value:
        data["reference_number"] = ""
        data["is_reference_document"] = "true"
    form = AdmAdministrativeDocumentEditPopupForm(data, None, instance=doc)
    if not form.is_valid():
        first_error = None
        for field_errors in form.errors.values():
            if field_errors:
                first_error = field_errors[0]
                break
        return JsonResponse(
            {"ok": False, "error": first_error or "Invalid data."}, status=400
        )

    # Validation may normalize other optional fields (None -> ""). A live edit
    # must persist only the requested value and its explicit reference fields.
    doc.refresh_from_db()
    changed_fields = [field]
    if field == "reference_document" and value:
        changed_fields.extend(["reference_number", "is_reference_document"])
    for name in changed_fields:
        setattr(doc, name, form.cleaned_data[name])
    updated_doc = doc
    updated_doc.updated_by = request.user
    updated_doc.save(update_fields=[*changed_fields, "updated_by", "updated_at"])

    def _fmt_date(value):
        return value.strftime("%d/%m/%Y") if value else "-"

    if field == "title":
        display = updated_doc.title
    elif field == "content_type":
        display = updated_doc.content_type.name if updated_doc.content_type_id else "-"
    elif field == "signer_role":
        display = updated_doc.signer_role.title if updated_doc.signer_role_id else "-"
    elif field == "issuing_department":
        display = (
            updated_doc.issuing_department.name
            if updated_doc.issuing_department_id
            else "-"
        )
    elif field == "issue_date":
        if updated_doc.issue_date:
            display = _fmt_date(updated_doc.issue_date)
        else:
            display = "—"
    elif field in {"effective_date", "expiry_date"}:
        display = _fmt_date(getattr(updated_doc, field))
    elif field == "status":
        display = updated_doc.status.name if updated_doc.status_id else "-"
    elif field == "reference_number":
        display = updated_doc.reference_number or "-"
    elif field == "note":
        display = updated_doc.note or "-"
    elif field == "reference_document":
        display = str(updated_doc.reference_document) if updated_doc.reference_document_id else "—"
    elif field == "is_reference_document":
        display = "Có" if updated_doc.is_reference_document else "Không"
    else:
        display = getattr(updated_doc, field) or "—"

    values = _build_update_form_data(updated_doc)
    payload = {"ok": True, "field": field, "display": display, "value": values[field], "document_id": updated_doc.pk}
    if field == "reference_document":
        payload["related_values"] = {"reference_number": updated_doc.reference_number or "",
                                     "is_reference_document": updated_doc.is_reference_document}
    if field == "status" and updated_doc.status_id:
        payload["status"] = {
            "code": updated_doc.status.code,
            "name": updated_doc.status.name,
            "text_color": updated_doc.status.badge_text_color,
            "bg_color": updated_doc.status.badge_bg_color,
        }
    return JsonResponse(payload)


@login_required
@admin_staff_required
def document_attachment_upload(request, doc_id: int):
    ajax = request.headers.get("x-requested-with") == "XMLHttpRequest"
    if request.method != "POST":
        return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)

    try:
        doc = AdmAdministrativeDocument.objects.get(pk=doc_id)
    except AdmAdministrativeDocument.DoesNotExist:
        if ajax:
            return JsonResponse({"ok": False, "error": "Văn bản không tồn tại."}, status=404)
        messages.error(request, "Document not found.")
        return redirect("admindocuments:admindocuments_list")

    uploaded_file = request.FILES.get("attachment")
    link = (request.POST.get("attachment_link") or "").strip() if ajax else ""
    if not uploaded_file and not link:
        if ajax:
            return JsonResponse({"ok": False, "error": "Vui lòng chọn tệp hoặc nhập link đính kèm."}, status=400)
        messages.error(request, "Vui lòng chọn file để tải lên.")
        return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)

    try:
        with transaction.atomic():
            locked_doc = AdmAdministrativeDocument.objects.select_for_update().get(pk=doc.pk)
            _create_attachment_version(document=locked_doc, user=request.user, uploaded_file=uploaded_file, link=link)
    except Exception:
        if ajax:
            return JsonResponse({"ok": False, "error": "Không thể thêm đính kèm. Vui lòng kiểm tra tệp hoặc đường dẫn và thử lại."}, status=400)
        messages.error(
            request,
            "Tải file thất bại. Vui lòng đổi tên file (không dùng ký tự lạ) và thử lại.",
        )
        return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)

    if ajax:
        return JsonResponse({"ok": True})
    messages.success(request, "Đã tải file lên thành công.")
    return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)


@login_required
@admin_staff_required
def document_delete(request, doc_id: int):
    if request.method != "POST":
        return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)

    snapshot = get_object_or_404(AdmAdministrativeDocument, pk=doc_id)
    with transaction.atomic():
        # Use the same lock order as ordinary allocation and splitting.
        counter = (AdmDocumentCounter.objects.select_for_update().filter(
            doc_type_id=snapshot.doc_type_id, company_id=snapshot.issuing_company_id,
            year=snapshot.number_year).first())
        doc = (
            AdmAdministrativeDocument.objects.select_for_update()
            .select_related("doc_type", "issuing_company")
            .get(pk=doc_id)
        )
        if doc.is_void:
            messages.info(request, "Văn bản đã ở trạng thái vô hiệu.")
            return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)

        if counter and not doc.number_suffix and doc.running_number == counter.next_number - 1:
            counter.next_number = doc.running_number
            counter.save(update_fields=["next_number", "updated_at"])

        original_number = doc.document_number_full
        doc.is_void = True
        doc.voided_at = timezone.now()
        doc.voided_by = request.user
        doc.updated_by = request.user
        doc.document_number_full = f"{original_number}-VOID"
        doc.save(
            update_fields=[
                "is_void",
                "voided_at",
                "voided_by",
                "updated_by",
                "document_number_full",
                "updated_at",
            ]
        )
    messages.success(request, "Đã vô hiệu văn bản và trả số hiệu về bộ đếm (nếu có thể).")
    return redirect("admindocuments:admindocuments_list")


@login_required
@admin_staff_required
def document_attachment_delete(request, doc_id: int, att_id: int):
    ajax = request.headers.get("x-requested-with") == "XMLHttpRequest"
    try:
        attachment = AdmDocumentAttachment.objects.select_related("document").get(
            pk=att_id, document_id=doc_id
        )
    except AdmDocumentAttachment.DoesNotExist:
        if ajax:
            return JsonResponse({"ok": False, "error": "Tệp không tồn tại."}, status=404)
        messages.error(request, "File không tồn tại.")
        return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)

    if request.method != "POST":
        return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)

    with transaction.atomic():
        attachment.is_deleted = True
        attachment.is_latest = False
        attachment.deleted_by = request.user
        attachment.deleted_at = timezone.now()
        attachment.save(update_fields=["is_deleted", "is_latest", "deleted_by", "deleted_at"])

        latest = (
            AdmDocumentAttachment.objects.filter(
                document_id=doc_id, is_deleted=False
            )
            .order_by("-version")
            .first()
        )
        if latest:
            latest.is_latest = True
            latest.save(update_fields=["is_latest"])

    if ajax:
        return JsonResponse({"ok": True})
    messages.success(request, "Đã đánh dấu xóa file.")
    return redirect("admindocuments:admindocuments_detail", doc_id=doc_id)
