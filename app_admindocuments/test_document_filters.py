from datetime import date, datetime
from io import BytesIO
from urllib.parse import parse_qs

from django.contrib.auth.models import User
from django.conf import settings
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import AdmAdministrativeDocument, AdmCompany, AdmContentType, AdmDepartment, AdmDocumentStatus, AdmDocumentType, AdmSignerRole


class DocumentDateFilterTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser('filter-admin', 'filter@example.com', 'test')
        company = AdmCompany.objects.create(code='FIL', name='Công ty thử')
        department = AdmDepartment.objects.create(code='FIL', name='Hành chính', company=company)
        dtype = AdmDocumentType.objects.create(code='FIL', name='Loại thử')
        content = AdmContentType.objects.create(code='FIL', name='Nội dung thử')
        status = AdmDocumentStatus.objects.create(code='FIL', name='Trạng thái thử')
        signer = AdmSignerRole.objects.create(code='FIL', title='Người ký thử')
        for number, title, created, issued in [
            (801, 'Trong khoảng', date(2026, 1, 1), date(2025, 6, 7)),
            (802, 'Ngoài khoảng', date(2026, 2, 1), date(2026, 5, 1)),
            (803, 'Năm cũ', date(2024, 2, 1), date(2025, 6, 7)),
        ]:
            doc = AdmAdministrativeDocument.objects.create(
                title=title, created_by=cls.user, issuing_company=company, issuing_department=department, doc_type=dtype,
                content_type=content, status=status, signer_role=signer, running_number=number,
                number_year=2026, issue_date=issued,
            )
            created_at = datetime.combine(created, datetime.min.time()).replace(hour=12)
            if settings.USE_TZ:
                created_at = timezone.make_aware(created_at)
            AdmAdministrativeDocument.objects.filter(pk=doc.pk).update(created_at=created_at)

    def setUp(self):
        self.client.force_login(self.user)
        self.url = reverse('admindocuments:admindocuments_list')
        self.export_url = reverse('admindocuments:admindocuments_list_export')

    def titles(self, response):
        return {doc.title for doc in response.context['documents']}

    def test_creation_range_is_inclusive_and_combines_with_issue_date(self):
        filters = {'start_date': '01/01/2026', 'end_date': '31/01/2026', 'issue_date': '07/06/2025'}
        response = self.client.get(self.url, filters)
        self.assertEqual(self.titles(response), {'Trong khoảng'})
        self.assertEqual(response.context['issue_date'], filters['issue_date'])
        query = parse_qs(response.context['list_query'])
        self.assertEqual(query['issue_date'], [filters['issue_date']])
        self.assertEqual(query['start_date'], [filters['start_date']])

    def test_removing_creation_chip_does_not_restore_default_year(self):
        response = self.client.get(self.url, {'created_all': '1'})
        self.assertEqual(self.titles(response), {'Trong khoảng', 'Ngoài khoảng', 'Năm cũ'})
        self.assertEqual(response.context['start_date'], '')
        self.assertEqual(response.context['created_all'], '1')
        self.assertEqual(parse_qs(response.context['list_query'])['created_all'], ['1'])

    def test_issue_date_can_filter_all_creation_years(self):
        response = self.client.get(self.url, {'created_all': '1', 'issue_date': '2025-06-07'})
        self.assertEqual(self.titles(response), {'Trong khoảng', 'Năm cũ'})
        self.assertEqual(response.context['issue_date'], '07/06/2025')

    def test_export_uses_same_creation_and_issue_date_filters_as_list(self):
        for filters in [
            {'start_date': '01/01/2026', 'end_date': '31/01/2026', 'issue_date': '07/06/2025'},
            {'created_all': '1', 'issue_date': '07/06/2025'},
        ]:
            with self.subTest(filters=filters):
                expected = self.titles(self.client.get(self.url, filters))
                response = self.client.get(self.export_url, filters)
                self.assertEqual(response.status_code, 200)
                workbook = load_workbook(BytesIO(response.content), read_only=True)
                self.assertEqual({row[2] for row in workbook.active.iter_rows(min_row=2, values_only=True)}, expected)
                workbook.close()
