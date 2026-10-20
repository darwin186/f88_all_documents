from datetime import date
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import AdmCompany, AdmIncomingDispatchStatus, AdmParcelReceipt


class ParcelCompletedScreenTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser('parcel-done-admin', 'p@example.com', 'test')
        cls.outsider = User.objects.create_user('parcel-done-outsider', password='test')
        cls.f88 = AdmCompany.objects.create(code='F88', name='F88')
        cls.nnx = AdmCompany.objects.create(code='NNX', name='Ngân hàng NNX')
        cls.done, _ = AdmIncomingDispatchStatus.objects.update_or_create(code='pkg_done', defaults={'name': 'Hoàn tất'})
        cls.other_status, _ = AdmIncomingDispatchStatus.objects.update_or_create(code='pkg_test_x', defaults={'name': 'Đã giao thử'})
        now = timezone.now()

        def make(number, company, parcel_type, status, code, tracking, sender='Công ty gửi', **extra):
            return AdmParcelReceipt.objects.create(
                document_number=number, received_by=cls.admin, recipient_name=f'Người {code}', recipient_employee_code=code,
                recipient_department='Phòng A', parcel_type=parcel_type, sender_unit=sender, tracking_code=tracking,
                receiving_company=company, status=status, created_by=cls.admin, completed_at=now, **extra)

        cls.p1 = make('PK-1', cls.f88, 'hoso', cls.done, 'NV001', 'VD001')
        cls.p2 = make('PK-2', cls.nnx, 'hanghoa', cls.done, 'NV002', 'VD002', actual_receiver_name='Người nhận hộ', actual_receiver_employee_code='NV777')
        cls.p3 = make('PK-3', cls.f88, 'khac', cls.other_status, 'NV003', 'VD003', sender='=HYPERLINK("http://x")')
        cls.url = reverse('admindocuments:parcel_receipt_completed_list')
        cls.export = reverse('admindocuments:parcel_receipt_completed_export')

    def setUp(self):
        self.client.force_login(self.admin)

    def _numbers(self, params):
        response = self.client.get(self.url, params)
        self.assertEqual(response.status_code, 200)
        return sorted(p.document_number for p in response.context['parcels'])

    def test_filter_by_company_type_and_status_uses_codes(self):
        self.assertEqual(self._numbers({}), ['PK-1', 'PK-2', 'PK-3'])
        self.assertEqual(self._numbers({'company': 'NNX'}), ['PK-2'])
        self.assertEqual(self._numbers({'parcel_type': 'khac'}), ['PK-3'])
        self.assertEqual(self._numbers({'status': 'pkg_test_x'}), ['PK-3'])
        self.assertEqual(self._numbers({'company': 'F88', 'parcel_type': 'hoso'}), ['PK-1'])

    def test_unknown_parcel_type_is_ignored_not_crashing(self):
        self.assertEqual(self._numbers({'parcel_type': 'bậy'}), ['PK-1', 'PK-2', 'PK-3'])

    def test_filter_by_employee_code_matches_recipient_and_actual_receiver(self):
        self.assertEqual(self._numbers({'employee_code': 'NV001'}), ['PK-1'])
        self.assertEqual(self._numbers({'employee_code': 'nv777'}), ['PK-2'])
        self.assertEqual(self._numbers({'employee_code': 'KHONGCO'}), [])

    def test_date_range_accepts_iso_and_vietnamese_formats(self):
        today = date.today()
        self.assertEqual(len(self._numbers({'start_date': today.strftime('%Y-%m-%d'), 'end_date': today.strftime('%d/%m/%Y')})), 3)
        self.assertEqual(self._numbers({'start_date': '2001-01-01', 'end_date': '2001-12-31'}), [])

    def test_filter_options_only_list_values_present(self):
        response = self.client.get(self.url)
        self.assertEqual(sorted(c.code for c in response.context['company_options']), ['F88', 'NNX'])
        self.assertEqual(sorted(s.code for s in response.context['status_options']), ['pkg_done', 'pkg_test_x'])
        self.assertContains(response, 'name="employee_code"')
        for name in ('company', 'parcel_type', 'status'):
            self.assertContains(response, f'<select id="pc-{ {"parcel_type": "type"}.get(name, name) }" name="{name}">', html=False)

    def test_page_uses_compact_controls_and_export_link_keeps_filters(self):
        html = self.client.get(self.url, {'company': 'F88', 'page': 1}).content.decode()
        self.assertIn('adm-pagination adm-pagination-compact', html)
        self.assertIn('ui-btn ui-btn-secondary ui-btn-sm', html)
        self.assertIn('data-export-excel', html)
        self.assertIn('export/?company=F88"', html)  # không mang theo tham số page

    def _rows(self, params=None):
        response = self.client.get(self.export, params or {})
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        self.assertIn('buu-kien-da-nhan-', response['Content-Disposition'])
        sheet = load_workbook(BytesIO(response.content)).active
        return [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]

    def test_export_contains_all_rows_by_default_with_header(self):
        rows = self._rows()
        self.assertEqual(rows[0][0], 'STT')
        self.assertIn('Mã NV người nhận', rows[0])
        self.assertEqual(len(rows) - 1, 3)

    def test_export_follows_the_applied_filters(self):
        rows = self._rows({'company': 'NNX'})
        self.assertEqual(len(rows) - 1, 1)
        record = dict(zip(rows[0], rows[1]))
        self.assertEqual(record['Mã vận đơn'], 'VD002')
        self.assertEqual(record['Mã NV người nhận'], 'NV002')
        self.assertEqual(record['Mã NV người nhận thực tế'], 'NV777')
        self.assertEqual(record['Công ty'], 'Ngân hàng NNX')
        self.assertEqual(record['Loại'], 'Hàng hóa')
        self.assertEqual(len(self._rows({'employee_code': 'NV001'})) - 1, 1)
        self.assertEqual(len(self._rows({'employee_code': 'KHONGCO'})) - 1, 0)

    def test_export_escapes_spreadsheet_formulas(self):
        rows = self._rows({'status': 'pkg_test_x'})
        sender = dict(zip(rows[0], rows[1]))['Đơn vị gửi']
        self.assertTrue(sender.startswith("'="), sender)

    def test_export_requires_access(self):
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(self.export).status_code, 403)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self.export).status_code, 302)
