from datetime import datetime
from io import BytesIO
from urllib.parse import parse_qs

from django.conf import settings
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from app_documents.models import Manager, Region, Shop
from .models import AdmCompany, AdmCourierCompany, AdmDepartment, AdmPaperDocument, AdmPaperType


class PaperDocumentRedesignTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser('paper-ui-admin', 'paper@example.com', 'test')
        cls.outsider = User.objects.create_user('paper-ui-outsider')
        cls.company = AdmCompany.objects.create(code='PUI', name='Công ty giấy')
        cls.department = AdmDepartment.objects.create(code='PUI', name='Hành chính', company=cls.company)
        cls.kind = AdmPaperType.objects.create(code='PUI', name='Giấy giới thiệu')
        cls.other_kind = AdmPaperType.objects.create(code='PU2', name='Giấy xác nhận')
        cls.courier = AdmCourierCompany.objects.create(name='CPN thử')
        cls.region = Region.objects.create(region_id=12, region_code='PUI', region_name='Miền Bắc')
        manager = Manager.objects.create(manager_code='PUI', qlkv_code='PUI', qlkv_name='Quản lý thử', qlv_code='PUI', qlv_name='Quản lý vùng thử')
        cls.shop = Shop.objects.bulk_create([Shop(shop_code=6123, shop_name='PGD thử', manager_id=manager, region_id=cls.region)])[0]
        cls.year = timezone.now().year
        cls.docs = []
        for index, (kind, status, year, month, day) in enumerate([
            (cls.kind, 'Đã gửi', cls.year, 1, 10),
            (cls.other_kind, 'Chờ gửi', cls.year, 2, 1),
            (cls.kind, 'Đã gửi', cls.year - 1, 1, 10),
        ], 1):
            doc = AdmPaperDocument.objects.create(paper_type=kind, running_number=index,
                document_number_full=f'{index}/{year}/PUI', region='Miền Bắc', responsible_person='Người phụ trách',
                requested_department=cls.shop if index != 2 else None, department=cls.department,
                summary=f'Giấy hành chính {index}', status=status, note='Ghi chú gốc',
                courier_company=cls.courier, courier_tracking_code=f'TRACK-{index}', created_by=cls.admin)
            created_at = datetime(year, month, day, 12)
            if settings.USE_TZ:
                created_at = timezone.make_aware(created_at)
            AdmPaperDocument.objects.filter(pk=doc.pk).update(created_at=created_at)
            cls.docs.append(doc)

    def setUp(self):
        self.client.force_login(self.admin)
        self.url = reverse('admindocuments:paper_document_list')
        self.detail = reverse('admindocuments:paper_document_detail', args=[self.docs[0].pk])
        self.field = reverse('admindocuments:paper_document_update_field', args=[self.docs[0].pk])
        self.export = reverse('admindocuments:paper_document_export')

    def test_combined_filters_preserve_query_and_export_matches(self):
        filters = {'paper_type': self.kind.pk, 'region': 'Miền Bắc', 'status': 'Đã gửi',
                   'requested_department': self.shop.pk, 'start_date': f'01/01/{self.year}', 'end_date': f'31/01/{self.year}'}
        response = self.client.get(self.url, filters)
        self.assertEqual([doc.pk for doc in response.context['documents']], [self.docs[0].pk])
        query = parse_qs(response.context['list_query'])
        self.assertEqual(query['region'], ['Miền Bắc'])
        self.assertEqual(query['requested_department'], [str(self.shop.pk)])
        formula_like_text = '=HYPERLINK("https://example.com", "Giấy")'
        AdmPaperDocument.objects.filter(pk=self.docs[0].pk).update(summary=formula_like_text)
        exported = self.client.get(self.export, filters)
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(BytesIO(exported.content), read_only=True)
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
        self.assertEqual([row[1] for row in rows], [self.docs[0].document_number_full])
        self.assertEqual(rows[0][5], self.shop.shop_name)
        self.assertEqual(rows[0][7], formula_like_text)
        self.assertEqual(workbook.active['H2'].data_type, 's')
        workbook.close()

    def test_clear_creation_range_includes_older_years(self):
        response = self.client.get(self.url, {'created_all': '1', 'status': 'Chờ gửi'})
        self.assertEqual(response.context['paginator'].count, 3)
        self.assertEqual(response.context['start_date'], '')
        self.assertNotIn('status', parse_qs(response.context['list_query']))
        self.assertNotContains(response, 'id="paper-filter-status"')
        default = self.client.get(self.url)
        self.assertEqual(default.context['paginator'].count, 2)

    def test_search_matches_ticket_and_existing_criteria_and_export(self):
        AdmPaperDocument.objects.filter(pk=self.docs[0].pk).update(ticket_code='F88-REQ-001')
        for query, expected in [('req-001', self.docs[0]), ('TRACK-2', self.docs[1]),
                                ('hành chính 3', self.docs[2]), (self.docs[0].document_number_full, self.docs[0])]:
            with self.subTest(query=query):
                response = self.client.get(self.url, {'created_all': '1', 'q': query})
                self.assertEqual([doc.pk for doc in response.context['documents']], [expected.pk])
        response = self.client.get(self.url, {'created_all': '1', 'q': 'req-001'})
        self.assertContains(response, 'Mã Ticket')
        self.assertContains(response, 'data-document-cell="ticket_code">F88-REQ-001</td>')
        self.assertNotContains(response, 'data-document-cell="status"')
        table_header = response.content.decode().split('adm-paper-table', 1)[1].split('</thead>', 1)[0]
        self.assertNotIn('Tình trạng', table_header)
        exported = self.client.get(self.export, {'created_all': '1', 'q': 'req-001'})
        workbook = load_workbook(BytesIO(exported.content), read_only=True)
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][1], self.docs[0].document_number_full)
        self.assertEqual(rows[0][11], 'F88-REQ-001')
        workbook.close()

    def paper_import_file(self, rows, headers=None):
        from .views import PAPER_IMPORT_HEADERS
        workbook = Workbook()
        sheet = workbook.active
        headers = headers if headers is not None else PAPER_IMPORT_HEADERS
        sheet.append(list(headers))
        defaults = {'Loại giấy': self.kind.name, 'Miền': self.region.region_name,
                    'Phòng ban nội bộ': self.department.name, 'Nội dung': 'Giấy nhập từ Excel',
                    'Ngày tạo (yyyy-mm-dd)': f'{self.year}-01-10'}
        for row in rows:
            values = {**defaults, **row}
            sheet.append([values.get(header) for header in headers])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return SimpleUploadedFile('paper-import.xlsx', output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def test_excel_template_has_ticket_instead_of_status(self):
        AdmPaperDocument.objects.filter(pk=self.docs[1].pk).update(ticket_code='TICKET-SAMPLE')
        response = self.client.get(reverse('admindocuments:paper_document_template'))
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.values)
        headers = list(rows[0])
        self.assertIn('Mã Ticket', headers)
        self.assertNotIn('Tình trạng', headers)
        self.assertEqual(rows[1][headers.index('Mã Ticket')], 'TICKET-SAMPLE')
        workbook.close()

    def test_excel_check_and_commit_preserves_ticket_for_each_paper(self):
        upload = self.paper_import_file([
            {'Nội dung': 'Giấy ticket thứ nhất', 'Mã Ticket': '  TICKET-001  '},
            {'Nội dung': 'Giấy ticket thứ hai', 'Mã Ticket': 'TICKET-002'},
        ])
        response = self.client.post(reverse('admindocuments:paper_document_import_url'), {'file': upload})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['rows'], 2)
        self.assertEqual(AdmPaperDocument.objects.count(), 3)
        committed = self.client.post(reverse('admindocuments:paper_document_import_commit'), {'token': response.json()['token']})
        self.assertEqual(committed.status_code, 200, committed.content)
        first = AdmPaperDocument.objects.get(summary='Giấy ticket thứ nhất')
        second = AdmPaperDocument.objects.get(summary='Giấy ticket thứ hai')
        self.assertEqual((first.ticket_code, second.ticket_code), ('TICKET-001', 'TICKET-002'))
        self.assertEqual(first.status, '')
        listing = self.client.get(self.url, {'created_all': '1', 'q': 'TICKET-001'})
        self.assertEqual([doc.pk for doc in listing.context['documents']], [first.pk])
        self.assertContains(listing, 'data-document-cell="ticket_code">TICKET-001</td>')

    def test_direct_excel_import_saves_ticket_without_status(self):
        upload = self.paper_import_file([{'Nội dung': 'Nhập trực tiếp', 'Mã Ticket': 'TICKET-DIRECT'}])
        response = self.client.post(reverse('admindocuments:paper_document_import'), {'file': upload})
        self.assertEqual(response.status_code, 302)
        doc = AdmPaperDocument.objects.get(summary='Nhập trực tiếp')
        self.assertEqual(doc.ticket_code, 'TICKET-DIRECT')
        self.assertEqual(doc.status, '')

    def test_import_accepts_legacy_and_optional_blank_columns(self):
        from .views import PAPER_IMPORT_HEADERS
        for label, headers, values, expected_ticket, expected_status in [
            ('legacy-blank', ['Tình trạng' if h == 'Mã Ticket' else h for h in PAPER_IMPORT_HEADERS], {}, '', ''),
            ('legacy-filled', ['Tình trạng' if h == 'Mã Ticket' else h for h in PAPER_IMPORT_HEADERS], {'Tình trạng': 'Đã gửi'}, '', 'Đã gửi'),
            ('both', [*PAPER_IMPORT_HEADERS, 'Tình trạng'], {'Mã Ticket': 'TICKET-BOTH'}, 'TICKET-BOTH', ''),
            ('blank-ticket', PAPER_IMPORT_HEADERS, {}, '', ''),
            ('neither', [h for h in PAPER_IMPORT_HEADERS if h != 'Mã Ticket'], {}, '', ''),
        ]:
            with self.subTest(template=label):
                summary = f'Excel {label}'
                upload = self.paper_import_file([{'Nội dung': summary, **values}], headers)
                response = self.client.post(reverse('admindocuments:paper_document_import_url'), {'file': upload})
                self.assertEqual(response.status_code, 200, response.content)
                committed = self.client.post(reverse('admindocuments:paper_document_import_commit'), {'token': response.json()['token']})
                self.assertEqual(committed.status_code, 200)
                doc = AdmPaperDocument.objects.get(summary=summary)
                self.assertEqual((doc.ticket_code, doc.status), (expected_ticket, expected_status))

    def test_import_rejects_overlong_ticket_before_creating_records(self):
        upload = self.paper_import_file([{'Mã Ticket': 'T' * 201}])
        response = self.client.post(reverse('admindocuments:paper_document_import_url'), {'file': upload})
        self.assertEqual(response.status_code, 400)
        self.assertIn('Mã Ticket', response.json()['details'][0])
        self.assertEqual(AdmPaperDocument.objects.count(), 3)

    def test_pagination_counts_continuously_and_keeps_escaped_filters(self):
        AdmPaperDocument.objects.bulk_create([AdmPaperDocument(paper_type=self.kind, running_number=100+i,
            document_number_full=f'PUI-{i}', responsible_person='Thử', summary='A&B + giấy', department=self.department) for i in range(26)])
        response = self.client.get(self.url, {'created_all': '1', 'q': 'A&B + giấy', 'page': '2'})
        self.assertEqual(response.context['documents'].start_index(), 26)
        self.assertEqual(parse_qs(response.context['list_query'])['q'], ['A&B + giấy'])
        self.assertContains(response, 'aria-label="Trang 2"')

    def test_invalid_create_keeps_entered_values_and_reopens_popup(self):
        response = self.client.post(self.url, {'paper_type': self.kind.pk, 'summary': 'Giữ nội dung vừa nhập'})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['form'].is_bound)
        self.assertEqual(response.context['form']['summary'].value(), 'Giữ nội dung vừa nhập')
        self.assertContains(response, 'data-paper-create-open')
        self.assertContains(response, 'Chọn Phòng giao dịch hoặc Phòng ban nội bộ')
        self.assertEqual(AdmPaperDocument.objects.count(), 3)

    def test_creation_uses_existing_number_allocator_and_records_responsible_person(self):
        response = self.client.post(self.url, {'paper_type': self.kind.pk, 'summary': 'Giấy mới', 'department': self.department.pk})
        self.assertEqual(response.status_code, 302)
        doc = AdmPaperDocument.objects.get(summary='Giấy mới')
        self.assertTrue(doc.document_number_full)
        self.assertEqual(doc.responsible_person, self.admin.username)
        self.assertEqual(doc.created_by, self.admin)

    def test_edit_fragment_and_old_links_open_popup(self):
        self.assertRedirects(self.client.get(self.detail), f'{self.url}?edit={self.docs[0].pk}', fetch_redirect_response=False)
        response = self.client.get(self.detail, {'edit': '1'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-live-field="summary"')
        self.assertContains(response, 'id_paper-edit-summary')
        self.assertNotContains(response, '<html')
        drawer = self.client.get(self.detail, {'drawer': '1'})
        self.assertContains(drawer, 'data-preview-tab="delivery"')
        self.assertContains(drawer, 'data-copy-document-number')

    def test_live_update_changes_only_requested_field_and_preserves_number(self):
        doc = self.docs[0]
        response = self.client.post(self.field, {'field': 'summary', 'value': 'Nội dung mới'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['display'], 'Nội dung mới')
        fresh = AdmPaperDocument.objects.get(pk=doc.pk)
        self.assertEqual(fresh.summary, 'Nội dung mới')
        self.assertEqual(fresh.note, 'Ghi chú gốc')
        self.assertEqual(fresh.document_number_full, doc.document_number_full)
        self.assertEqual(fresh.updated_by, self.admin)

    def test_live_update_rejects_empty_required_or_numbering_fields(self):
        for field, value in [('summary', ''), ('running_number', '999'), ('document_number_full', 'NEW')]:
            with self.subTest(field=field):
                self.assertEqual(self.client.post(self.field, {'field': field, 'value': value}).status_code, 400)
        doc = AdmPaperDocument.objects.get(pk=self.docs[0].pk)
        self.assertEqual(doc.summary, 'Giấy hành chính 1')
        self.assertEqual(doc.document_number_full, self.docs[0].document_number_full)

    def test_live_update_retains_at_least_one_receiving_department(self):
        self.assertEqual(self.client.post(self.field, {'field': 'requested_department', 'value': ''}).status_code, 200)
        response = self.client.post(self.field, {'field': 'department', 'value': ''})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(AdmPaperDocument.objects.get(pk=self.docs[0].pk).department_id, self.department.pk)

    def test_old_inactive_catalog_values_do_not_block_unrelated_live_updates(self):
        self.kind.is_active = False
        self.kind.save()
        self.courier.is_active = False
        self.courier.save()
        response = self.client.post(self.field, {'field': 'note', 'value': 'Ghi chú mới'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(AdmPaperDocument.objects.get(pk=self.docs[0].pk).note, 'Ghi chú mới')

    def test_deleted_papers_cannot_be_read_or_edited_or_exported(self):
        doc = self.docs[0]
        AdmPaperDocument.objects.filter(pk=doc.pk).update(is_deleted=True)
        self.assertEqual(self.client.get(self.detail, {'edit': '1'}).status_code, 404)
        self.assertEqual(self.client.post(self.field, {'field': 'note', 'value': 'x'}).status_code, 404)
        response = self.client.get(self.url, {'created_all': '1', 'q': doc.document_number_full})
        self.assertEqual(response.context['paginator'].count, 0)

    def test_permissions_and_csrf_protect_new_endpoints(self):
        self.assertEqual(self.client.get(self.field).status_code, 405)
        protected = Client(enforce_csrf_checks=True)
        protected.force_login(self.admin)
        self.assertEqual(protected.post(self.field, {'field': 'note', 'value': 'x'}).status_code, 403)
        self.client.force_login(self.outsider)
        for url in [self.url, self.export, self.detail]:
            self.assertEqual(self.client.get(url).status_code, 403)
        self.assertEqual(self.client.post(self.field, {'field': 'note', 'value': 'x'}).status_code, 403)
