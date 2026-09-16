from datetime import date
from io import BytesIO
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.contrib.admin.models import LogEntry
from django.test import TestCase, override_settings
from django.urls import reverse
from django.template.loader import render_to_string
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from app_documents.models import AreaManager, Manager, Region, Shop, UserProfile
from app_document_campaigns.models import (
    Campaign,
    CampaignType,
    CampaignError,
    CampaignImportJob,
    CampaignImportSource,
    CampaignStagingRow,
    CampaignVersion,
)
from app_document_campaigns.services.imports import (
    CampaignImportError,
    publish_excel_version,
    stage_import_rows,
)
from app_document_campaigns.services.excel_exports import build_cleaning_workbook
from app_document_campaigns.services.excel_imports import EXCEL_SOURCE_NAME, import_cleaning_workbook
from app_document_campaigns.services.sql_sources import (
    BUSINESS_TYPES,
    DOCUMENT_SOURCE_NAME,
    FOLDER_SOURCE_NAME,
    fetch_document_error_rows,
    stage_monthly_sql_sources,
)


class CampaignImportStagingTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="import-admin")
        UserProfile.objects.create(user=self.user, department="Vận hành")
        self.campaign_type = CampaignType.objects.get(code="hardcopy-document-error")
        region = Region.objects.create(region_code="IMPORT", region_name="Vùng import")
        area = AreaManager.objects.create(
            areaManager_code="IMPORT",
            areaManager_name="QLKV import",
            areaManager_email="area-import@example.com",
        )
        manager = Manager.objects.create(
            manager_code="IMPORT",
            areaManager=area,
            qlkv_code="IMP",
            qlkv_name="QLKV import",
            qlkv_email="area-import@example.com",
            qlv_code="IMP",
            qlv_name="QLV import",
            qlv_email="region-import@example.com",
        )
        self.shop = Shop.objects.create(
            shop_code=9101,
            shop_name="PGD Import",
            shop_email="shop-import@example.com",
            manager_id=manager,
            region_id=region,
        )
        self.campaign = Campaign.objects.create(
            code="CT-IMPORT-2026-09",
            name="Chiến dịch import tháng 09/2026",
            campaign_type=self.campaign_type,
            report_month=date(2026, 9, 1),
            status=Campaign.Status.DATA_REVIEW,
            created_by=self.user,
        )
        self.version = CampaignVersion.objects.create(
            campaign=self.campaign,
            version_number=1,
            source_type=CampaignVersion.SourceType.SQL,
            created_by=self.user,
        )
        self.source = CampaignImportSource.objects.create(
            version=self.version,
            name="missing-folders-sql",
            source_type=CampaignVersion.SourceType.SQL,
            created_by=self.user,
        )

    def _valid_row(self, source_key="folder:1001"):
        return {
            "source_key": source_key,
            "error_type": "folder",
            "source_object_id": "1001",
            "shop_code": "9101",
            "contract_code": "HD1001",
            "checking_issue": "Thiếu quyển chứng từ",
        }

    def test_reimport_replaces_staging_rows_without_duplicates(self):
        first = stage_import_rows(source=self.source, rows=[self._valid_row()])
        second = stage_import_rows(source=self.source, rows=[self._valid_row()])

        self.assertEqual(first["added"], 1)
        self.assertEqual(second["added"], 1)
        self.assertEqual(CampaignStagingRow.objects.filter(source=self.source).count(), 1)
        self.source.refresh_from_db()
        self.assertEqual(self.source.status, CampaignImportSource.Status.VALIDATED)
        self.assertEqual(self.source.invalid_count, 0)

    def test_invalid_and_duplicate_rows_are_kept_for_preview(self):
        summary = stage_import_rows(
            source=self.source,
            rows=[self._valid_row(), self._valid_row(), {"shop_code": "unknown"}],
        )

        self.assertEqual(summary["rows"], 3)
        self.assertEqual(summary["invalid"], 3)
        self.assertEqual(CampaignStagingRow.objects.filter(source=self.source).count(), 3)
        duplicate = CampaignStagingRow.objects.filter(source=self.source, source_key="folder:1001").first()
        self.assertIn(
            {"field": "source_key", "code": "duplicate_in_source"},
            duplicate.validation_errors,
        )

    def test_active_campaign_cannot_be_overwritten(self):
        self.campaign.status = Campaign.Status.ACTIVE
        self.campaign.save(update_fields=["status"])

        with self.assertRaisesMessage(CampaignImportError, "active"):
            stage_import_rows(source=self.source, rows=[self._valid_row()])

        self.assertFalse(CampaignStagingRow.objects.filter(source=self.source).exists())

    @patch("app_document_campaigns.services.sql_sources.fetch_document_error_rows")
    @patch("app_document_campaigns.services.sql_sources.fetch_folder_error_rows")
    def test_both_sql_sources_are_staged_independently(self, fetch_folders, fetch_documents):
        fetch_folders.return_value = [self._valid_row("folder:1001")]
        document_row = self._valid_row("document:1001")
        document_row["error_type"] = "document"
        fetch_documents.return_value = [document_row]

        result = stage_monthly_sql_sources(version=self.version, created_by=self.user)

        self.assertEqual(set(result), {FOLDER_SOURCE_NAME, DOCUMENT_SOURCE_NAME})
        self.assertEqual(
            set(CampaignImportSource.objects.filter(version=self.version).values_list("name", flat=True)),
            {self.source.name, FOLDER_SOURCE_NAME, DOCUMENT_SOURCE_NAME},
        )
        self.assertEqual(
            CampaignStagingRow.objects.filter(source__name__in=[FOLDER_SOURCE_NAME, DOCUMENT_SOURCE_NAME]).count(),
            2,
        )
        fetch_folders.assert_called_once_with(date(2026, 9, 1), using="default")
        fetch_documents.assert_called_once_with(date(2026, 9, 1), using="default")

    @patch("app_document_campaigns.services.sql_sources._fetch_rows", return_value=[])
    def test_sql_month_filter_uses_half_open_date_range(self, fetch_rows):
        fetch_document_error_rows(date(2026, 12, 1))

        self.assertEqual(
            fetch_rows.call_args.args[1],
            [date(2026, 12, 1), date(2027, 1, 1), *BUSINESS_TYPES],
        )

    def test_export_staging_workbook_contains_metadata_and_escapes_formulas(self):
        row = self._valid_row()
        row["customer_name"] = "=HYPERLINK(\"https://example.invalid\")"
        stage_import_rows(source=self.source, rows=[row])
        exported = build_cleaning_workbook(self.version)
        workbook = load_workbook(filename=exported, data_only=False)
        self.assertEqual(workbook["_Metadata"]["B1"].value, "dec-cleaning-v1")
        self.assertEqual(workbook["_Metadata"].sheet_state, "veryHidden")
        headers = [cell.value for cell in workbook["Lỗi chứng từ"][1]]
        customer_column = headers.index("Tên khách hàng") + 1
        self.assertTrue(workbook["Lỗi chứng từ"].cell(2, customer_column).value.startswith("'="))
        self.assertTrue(workbook["Lỗi chứng từ"].column_dimensions["A"].hidden)

    def test_export_requires_login(self):
        response = self.client.get(
            reverse("document_campaigns:export_staging_excel", kwargs={"version_id": self.version.pk})
        )

        self.assertEqual(response.status_code, 302)

    def test_reviewed_export_uses_uploaded_rows_instead_of_original_sql(self):
        stage_import_rows(source=self.source, rows=[self._valid_row(), self._valid_row("folder:1002")])
        reviewed_source = CampaignImportSource.objects.create(
            version=self.version, name="team-cleaning-excel", source_type="excel", created_by=self.user,
        )
        row = self._valid_row()
        row["checking_issue"] = "Nội dung đã được team sửa"
        stage_import_rows(source=reviewed_source, rows=[row])

        workbook = load_workbook(build_cleaning_workbook(self.version, reviewed=True))

        self.assertEqual(workbook["Lỗi chứng từ"].max_row, 2)
        self.assertEqual(workbook["Lỗi chứng từ"].cell(2, 21).value, "Nội dung đã được team sửa")

    def test_dataset_metrics_count_all_rows_and_distinct_contracts(self):
        from app_document_campaigns.services.dataset_metrics import summarize_dataset

        CampaignStagingRow.objects.bulk_create([
            CampaignStagingRow(
                source=self.source, row_number=i + 1, source_key=f"folder:{i}",
                normalized_payload={"contract_code": f"HD{i % 10}", "shop_code": 9101, "error_type": "folder"},
            ) for i in range(205)
        ])
        metrics = summarize_dataset(self.source.rows.all())
        self.assertEqual(len(metrics), 10)
        self.assertEqual(metrics[0]["value"], 205)
        self.assertEqual(metrics[1]["value"], 10)
        self.assertEqual(metrics[2]["value"], 1)

    def test_home_groups_cards_into_documents_and_services(self):
        self.client.force_login(self.user)
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertLess(content.index('id="documents-heading"'), content.index("Book lỗi chứng từ"))
        self.assertLess(content.index("Chứng từ bản mềm"), content.index('id="services-heading"'))
        self.assertLess(content.index('id="services-heading"'), content.index("Quản lý văn bản"))

    def test_round_trip_excel_is_staged_and_counts_removed_rows(self):
        stage_import_rows(
            source=self.source,
            rows=[self._valid_row("folder:1001"), self._valid_row("folder:1002")],
        )
        exported = build_cleaning_workbook(self.version)
        workbook = load_workbook(exported)
        sheet = workbook["Lỗi chứng từ"]
        headers = [cell.value for cell in sheet[1]]
        customer_column = headers.index("Tên khách hàng") + 1
        sheet.cell(2, customer_column).value = "Tên đã làm sạch"
        sheet.delete_rows(3)
        content = BytesIO()
        workbook.save(content)
        upload = SimpleUploadedFile(
            "cleaned.xlsx",
            content.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        source, summary = import_cleaning_workbook(
            version=self.version,
            created_by=self.user,
            uploaded_file=upload,
        )

        self.assertEqual(source.name, EXCEL_SOURCE_NAME)
        self.assertEqual(source.rows.count(), 1)
        self.assertEqual(source.rows.get().normalized_payload["customer_name"], "Tên đã làm sạch")
        self.assertEqual(summary["removed_from_sql"], 1)
        self.assertEqual(len(source.source_checksum), 64)

    def test_excel_for_another_campaign_is_rejected_before_staging(self):
        stage_import_rows(source=self.source, rows=[self._valid_row()])
        exported = build_cleaning_workbook(self.version)
        workbook = load_workbook(exported)
        workbook["_Metadata"]["B2"] = "OTHER-CAMPAIGN"
        content = BytesIO()
        workbook.save(content)
        upload = SimpleUploadedFile("wrong.xlsx", content.getvalue())

        with self.assertRaisesMessage(CampaignImportError, "campaign version"):
            import_cleaning_workbook(
                version=self.version,
                created_by=self.user,
                uploaded_file=upload,
            )

        self.assertFalse(
            CampaignImportSource.objects.filter(version=self.version, name=EXCEL_SOURCE_NAME).exists()
        )

    def test_confirmed_excel_is_published_and_missing_sql_row_is_soft_excluded(self):
        stage_import_rows(
            source=self.source,
            rows=[self._valid_row("folder:1001"), self._valid_row("folder:1002")],
        )
        removed_error = CampaignError.objects.create(
            campaign=self.campaign,
            version=self.version,
            source_key="folder:1002",
            source_object_id=1002,
            error_type=CampaignError.ErrorType.FOLDER,
            shop=self.shop,
            checking_issue="Dòng sẽ bị loại",
        )
        workbook = load_workbook(build_cleaning_workbook(self.version))
        workbook["Lỗi chứng từ"].delete_rows(3)
        content = BytesIO()
        workbook.save(content)
        excel_source, _ = import_cleaning_workbook(
            version=self.version,
            created_by=self.user,
            uploaded_file=SimpleUploadedFile("confirmed.xlsx", content.getvalue()),
        )

        summary = publish_excel_version(version=self.version, confirmed_by=self.user)

        self.version.refresh_from_db()
        self.campaign.refresh_from_db()
        excel_source.refresh_from_db()
        removed_error.refresh_from_db()
        self.assertEqual(self.version.status, CampaignVersion.Status.PUBLISHED)
        self.assertEqual(self.version.source_type, CampaignVersion.SourceType.EXCEL)
        self.assertEqual(self.campaign.current_version, self.version)
        self.assertEqual(self.campaign.status, Campaign.Status.READY)
        self.assertEqual(excel_source.status, CampaignImportSource.Status.CONFIRMED)
        self.assertEqual(removed_error.status, CampaignError.Status.EXCLUDED)
        self.assertEqual(CampaignError.objects.get(source_key="folder:1001").status, CampaignError.Status.READY)
        self.assertEqual(summary["applied_added"], 1)
        self.assertEqual(summary["applied_excluded"], 1)

    def test_campaign_data_preparation_ui_is_available(self):
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        stage_import_rows(source=self.source, rows=[self._valid_row()])
        self.client.force_login(self.user)

        listing = self.client.get(reverse("document_campaigns:campaign_list"))
        detail = self.client.get(
            reverse("document_campaigns:campaign_detail", kwargs={"campaign_id": self.campaign.pk})
        )

        self.assertEqual(listing.status_code, 200)
        self.assertContains(listing, "Tạo chiến dịch")
        self.assertContains(listing, "Khởi tạo kỳ book lỗi chứng từ.")
        self.assertContains(listing, "Tạo kỳ gửi lỗi để chuẩn bị dữ liệu lỗi theo từng tháng.")
        self.assertContains(listing, "Chiến dịch lỗi")
        self.assertContains(listing, "Danh sách và tiến độ chiến dịch")
        self.assertContains(listing, self.campaign.code)
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "Truy xuất lại dữ liệu SQL")
        self.assertContains(detail, "Tạo file Excel để team làm sạch")
        self.assertContains(detail, "Mã hợp đồng duy nhất")
        self.assertNotContains(detail, "Hiển thị tối đa 200 dòng")

    def test_non_admin_campaign_roles_are_read_only(self):
        self.client.force_login(self.user)
        campaign_endpoints = ["create_campaign_version", "update_campaign_deadline", "update_campaign_settings", "publish_campaign_to_shops"]
        version_endpoints = ["stage_sql_version", "queue_staging_excel_export", "import_staging_excel", "confirm_staging_excel"]
        for role in [None, "checker", "supervisor", "shop", "risk"]:
            self.user.groups.clear()
            if role:
                self.user.groups.add(Group.objects.get_or_create(name=role)[0])
            self.assertEqual(self.client.get(reverse("document_campaigns:campaign_create")).status_code, 403)
            self.assertEqual(self.client.post(reverse("document_campaigns:campaign_create"), {}).status_code, 403)
            for endpoint in campaign_endpoints:
                self.assertEqual(self.client.post(reverse(f"document_campaigns:{endpoint}", kwargs={"campaign_id": self.campaign.pk}), {}).status_code, 403)
            for endpoint in version_endpoints:
                self.assertEqual(self.client.post(reverse(f"document_campaigns:{endpoint}", kwargs={"version_id": self.version.pk}), {}).status_code, 403)
            listing = self.client.get(reverse("document_campaigns:campaign_list"))
            self.assertEqual(listing.status_code, 200)
            self.assertNotContains(listing, "Tạo chiến dịch mới")
            self.assertNotContains(listing, reverse("document_campaigns:campaign_create"))
            detail = self.client.get(reverse("document_campaigns:campaign_detail", kwargs={"campaign_id": self.campaign.pk}))
            self.assertEqual(detail.status_code, 200)
            self.assertNotContains(detail, "＋ Tạo version dữ liệu mới")
            self.assertNotContains(detail, 'enctype="multipart/form-data"')
        self.assertEqual(Campaign.objects.count(), 1)
        self.assertEqual(CampaignImportJob.objects.count(), 0)
        self.user.is_superuser = True
        self.user.save()
        self.assertEqual(self.client.get(reverse("document_campaigns:campaign_create")).status_code, 200)

    def test_gddb_navigation_is_separate_from_hardcopy(self):
        context = {
            "allowed_screens": ["gddb_registration", "receiving_v2", "checking_v2"],
            "is_admin": True,
            "gddb_nav": True,
        }
        navigation = render_to_string("app_documents/components/top_nav_v2.html", context)
        self.assertIn("Giao dịch bảo đảm", navigation)
        self.assertNotIn('<span class="text-gray-500">Dịch vụ</span>', navigation)
        self.assertIn("Cấu hình định danh", navigation)
        self.assertNotIn("Nhận chứng từ", navigation)
        self.assertNotIn("Chiến dịch lỗi", navigation)
        context["gddb_nav"] = False
        navigation = render_to_string("app_documents/components/top_nav_v2.html", context)
        self.assertIn("Nhận chứng từ", navigation)
        self.assertNotIn("Giao dịch bảo đảm", navigation)
        self.assertNotIn("Chiến dịch lỗi", navigation)
        self.assertNotIn("Tạo chiến dịch mới", navigation)
        self.assertNotIn("Danh sách và tiến độ chiến dịch", navigation)

    def test_master_data_requires_admin_and_hides_home_card(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("master_data")).status_code, 403)
        self.assertEqual(self.client.get(reverse("master_data_shops_api")).status_code, 403)
        self.assertNotContains(self.client.get(reverse("home")), reverse("master_data"))
        self.client.logout()
        self.assertEqual(self.client.get(reverse("master_data_shops_api")).status_code, 401)

    def test_master_data_admin_can_update_and_view_orgchart(self):
        import json
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        self.client.force_login(self.user)
        shop = self.shop
        url = reverse("master_data_shop_api", kwargs={"shop_id": shop.pk})
        self.assertContains(self.client.get(reverse("master_data")), "Danh mục phòng giao dịch")
        self.assertContains(self.client.get(reverse("master_data_api_docs")), "X-CSRFToken")
        self.assertContains(self.client.get(reverse("home")), reverse("master_data"))
        response = self.client.patch(url, data=json.dumps({"shop_email": "pgd@example.com", "is_shop_active": False}), content_type="application/json")
        self.assertEqual(response.status_code, 200)
        shop.refresh_from_db()
        self.assertEqual(shop.shop_email, "pgd@example.com")
        self.assertFalse(shop.is_shop_active)
        self.assertIsNotNone(shop.shop_closed_date)
        self.assertIn("region_manager", response.json()["orgchart"])
        self.assertEqual(LogEntry.objects.filter(user=self.user, object_id=str(shop.pk)).count(), 1)
        response = self.client.patch(url, data=json.dumps({"is_shop_active": True}), content_type="application/json")
        self.assertEqual(response.status_code, 200)
        shop.refresh_from_db()
        self.assertIsNone(shop.shop_closed_date)
        for payload in [{"shop_email": "not-an-email"}, {"is_shop_active": "false"}, {"manager_id": -1}, {"shop_name": "Unauthorized"}, []]:
            self.assertEqual(self.client.patch(url, data=json.dumps(payload), content_type="application/json").status_code, 400)
        self.assertEqual(self.client.get(reverse("master_data_shops_api"), {"q": "9101", "active": "true"}).json()["count"], 1)

    def test_master_data_patch_requires_csrf(self):
        from django.test import Client
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        shop = self.shop
        response = client.patch(reverse("master_data_shop_api", kwargs={"shop_id": shop.pk}), data='{"is_shop_active": false}', content_type="application/json")
        self.assertEqual(response.status_code, 403)

    def test_master_data_tokens_scopes_revocation_and_audit(self):
        import hashlib
        import json
        from django.test import Client
        from app_documents.models import ExternalDocumentIntakeToken, CollateralRegistrationApiToken
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        self.client.force_login(self.user)
        created = self.client.post(reverse("master_data_token_create"), {"name": "md-test", "scopes": ["master_data:shops:read", "master_data:shops:write"]})
        self.assertRedirects(created, reverse("master_data_tokens"))
        token = ExternalDocumentIntakeToken.objects.get(name="md-test")
        # assertRedirects consumed the single-display token page; create another.
        self.client.post(reverse("master_data_token_create"), {"name": "md-api", "scopes": ["master_data:shops:read", "master_data:shops:write"]})
        raw = self.client.session["master_data_new_token"]
        token = ExternalDocumentIntakeToken.objects.get(name="md-api")
        self.assertEqual(token.token_hash, hashlib.sha256(raw.encode()).hexdigest())
        first_page = self.client.get(reverse("master_data_tokens"))
        self.assertContains(first_page, raw)
        self.assertEqual(first_page["Cache-Control"], "no-store")
        self.assertNotContains(self.client.get(reverse("master_data_tokens")), raw)
        api_client = Client(enforce_csrf_checks=True)
        headers = {"HTTP_AUTHORIZATION": f"Bearer {raw}"}
        url = reverse("master_data_shop_api", kwargs={"shop_id": self.shop.pk})
        self.assertEqual(api_client.get(url, **headers).status_code, 200)
        self.assertEqual(api_client.patch(url, data=json.dumps({"shop_email": "token@example.com"}), content_type="application/json", **headers).status_code, 200)
        self.assertIn(f'"token_id": {token.pk}', LogEntry.objects.latest("action_time").change_message)
        token.refresh_from_db()
        self.assertIsNotNone(token.last_used_at)
        token.scopes = ["master_data:shops:read"]
        token.save()
        self.assertEqual(api_client.patch(url, data='{"is_shop_active": false}', content_type="application/json", **headers).status_code, 403)
        token.scopes = ["master_data:write"]
        token.save()
        self.assertEqual(api_client.get(url, **headers).status_code, 403)
        token.scopes = ["master_data:shops:read"]
        token.save()
        self.user.groups.clear()
        self.assertEqual(api_client.get(url, **headers).status_code, 403)
        self.user.groups.add(Group.objects.get(name="admin"))
        self.client.post(reverse("document_intake_token_revoke", kwargs={"token_id": token.pk}))
        self.assertEqual(api_client.get(url, **headers).status_code, 401)
        gddb_raw = "gddb_existing"
        existing = CollateralRegistrationApiToken.objects.create(name="Existing GDDB", token_prefix="gddb_exist", token_hash=hashlib.sha256(gddb_raw.encode()).hexdigest(), owner=self.user, created_by=self.user)
        self.assertContains(self.client.get(reverse("master_data_tokens")), "Existing GDDB")
        self.assertEqual(api_client.get(url, HTTP_AUTHORIZATION=f"Bearer {gddb_raw}").status_code, 401)
        existing.refresh_from_db()
        self.assertTrue(existing.is_active)
        from django.test import RequestFactory
        from app_documents.views import _gddb_intake_identity
        accepted, owner = _gddb_intake_identity(RequestFactory().post("/api/gddb/intake/", HTTP_AUTHORIZATION=f"Bearer {gddb_raw}"))
        self.assertTrue(accepted)
        self.assertEqual(owner, self.user)
        self.assertEqual(self.client.get(url, HTTP_AUTHORIZATION="Bearer invalid").status_code, 401)

    def test_master_data_token_ui_permissions_and_legacy_redirect(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("master_data_tokens")).status_code, 403)
        self.assertEqual(self.client.post(reverse("master_data_token_create"), {"name": "bad", "scopes": ["master_data:shops:write"]}).status_code, 403)
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        self.assertEqual(self.client.post(reverse("gddb_token_create"), {"name": "bad", "owner_id": self.user.pk}).status_code, 403)
        self.assertRedirects(self.client.post(reverse("document_intake_token_create"), {"name": "intake-test", "scopes": ["documents:write"]}), reverse("master_data_tokens"))
        self.assertNotContains(self.client.get(reverse("document_intake_management")), "Tạo token cho Prefect")
        self.user.is_superuser = True
        self.user.save()
        self.assertRedirects(self.client.get(reverse("gddb_registration_v2"), {"tab": "tokens"}), reverse("master_data_tokens"))
        self.client.post(reverse("gddb_token_create"), {"name": "gddb-new", "owner_id": self.user.pk})
        raw = self.client.session["gddb_new_api_token"]
        self.assertContains(self.client.get(reverse("master_data_tokens")), raw)
        self.assertNotContains(self.client.get(reverse("master_data_tokens")), raw)

    def test_master_data_screens_use_shared_navigation(self):
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        self.client.force_login(self.user)
        for name in ["master_data", "master_data_tokens", "master_data_api_docs"]:
            response = self.client.get(reverse(name))
            self.assertEqual(response.status_code, 200)
            self.assertTemplateUsed(response, "app_documents/base_app_documents_v2.html")
            self.assertTemplateUsed(response, "app_documents/components/top_nav_v2.html")
            self.assertContains(response, "data-user-menu")
            self.assertContains(response, self.user.username)
            self.assertContains(response, "Danh mục PGD")
            self.assertNotContains(response, '<span class="font-semibold text-f88green">Master Data</span>')
            self.assertContains(response, reverse("master_data_tokens"))
            self.assertContains(response, "logo-f88-primary.svg")
            self.assertNotContains(response, "Nhận chứng từ")
            if name == "master_data":
                self.assertContains(response, 'id="open-import"')
                self.assertContains(response, 'id="import-dialog" aria-labelledby="import-title"')
                self.assertNotContains(response, 'id="job-history"')
                self.assertContains(response, 'class="catalog-heading"')
                self.assertContains(response, "Kiểm tra &amp; cập nhật")

    def _catalog_excel_job(self, directory, edit=None):
        from django.core.files.base import ContentFile
        from app_documents.models import ShopCatalogJob
        from app_documents.shop_catalog_excel import process_job
        export = ShopCatalogJob.objects.create(kind="export", requested_by=self.user)
        process_job(export.pk)
        export.refresh_from_db()
        self.assertEqual(export.status, "succeeded", export.message)
        with export.output_file.open("rb") as source:
            book = load_workbook(source)
        if edit:
            edit(book["PGD"])
        stream = BytesIO()
        book.save(stream)
        imported = ShopCatalogJob.objects.create(kind="import", requested_by=self.user)
        imported.input_file.save("roundtrip.xlsx", ContentFile(stream.getvalue()))
        return imported

    def test_catalog_excel_roundtrip_celery_service_and_audit(self):
        from app_documents.shop_catalog_excel import process_job
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            def edit(sheet):
                sheet["D2"] = "excel@example.com"
                sheet["E2"] = False
            job = self._catalog_excel_job(directory, edit)
            process_job(job.pk)
            job.refresh_from_db()
            self.assertEqual(job.status, "succeeded", job.message)
            self.assertEqual(job.summary["updated"], 1)
            self.shop.refresh_from_db()
            self.assertEqual(self.shop.shop_email, "excel@example.com")
            self.assertFalse(self.shop.is_shop_active)
            self.assertIsNotNone(self.shop.shop_closed_date)
            self.assertIn("master_data_excel", LogEntry.objects.latest("action_time").change_message)
            process_job(job.pk)
            self.assertEqual(LogEntry.objects.count(), 1)

    def test_catalog_excel_invalid_file_and_concurrent_changes_are_atomic(self):
        from app_documents.shop_catalog_excel import process_job
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            Shop.objects.create(shop_code=9102, shop_name="Second PGD", shop_email="second@example.com", manager_id=self.shop.manager_id)
            def edit(sheet):
                sheet["D2"] = "valid-change@example.com"
                sheet["D3"] = "invalid-email"
            job = self._catalog_excel_job(directory, edit)
            process_job(job.pk)
            job.refresh_from_db()
            self.assertEqual(job.status, "failed")
            self.assertTrue(job.output_file)
            self.shop.refresh_from_db()
            self.assertEqual(self.shop.shop_email, "shop-import@example.com")
            job = self._catalog_excel_job(directory, lambda sheet: setattr(sheet["D2"], "value", "new@example.com"))
            self.shop.shop_email = "concurrent@example.com"
            self.shop.save()
            process_job(job.pk)
            job.refresh_from_db()
            self.assertEqual(job.status, "failed")
            self.shop.refresh_from_db()
            self.assertEqual(self.shop.shop_email, "concurrent@example.com")

    def test_catalog_excel_job_enqueue_failure_and_permissions(self):
        from app_documents.models import ShopCatalogJob
        self.client.force_login(self.user)
        self.assertEqual(self.client.post(reverse("master_data_job_create"), {"kind": "export"}).status_code, 403)
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        with patch("app_documents.tasks.process_shop_catalog_job.delay", side_effect=RuntimeError("broker offline")):
            response = self.client.post(reverse("master_data_job_create"), {"kind": "export"})
        self.assertEqual(response.status_code, 202)
        job = ShopCatalogJob.objects.get(pk=response.json()["id"])
        self.assertEqual(job.status, "failed")
        self.assertEqual(self.client.get(reverse("master_data_job_download", kwargs={"job_id": job.pk})).status_code, 409)
        self.assertEqual(self.client.post(reverse("master_data_job_create"), {"kind": "import"}).status_code, 400)

    def test_catalog_excel_unchanged_legacy_email_does_not_block_import(self):
        from app_documents.shop_catalog_excel import process_job
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        self.shop.shop_email = "legacy-invalid-email"
        self.shop.save()
        with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            job = self._catalog_excel_job(directory, lambda sheet: setattr(sheet["E2"], "value", False))
            process_job(job.pk)
            job.refresh_from_db()
            self.assertEqual(job.status, "succeeded", job.message)
            self.shop.refresh_from_db()
            self.assertFalse(self.shop.is_shop_active)
            self.assertEqual(self.shop.shop_email, "legacy-invalid-email")

    def test_catalog_excel_only_one_active_job_per_kind(self):
        from app_documents.models import ShopCatalogJob
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        self.client.force_login(self.user)
        url = reverse("master_data_job_create")
        with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory), patch("app_documents.tasks.process_shop_catalog_job.delay"):
            first = self.client.post(url, {"kind": "export"})
            self.assertEqual(first.status_code, 202)
            second = self.client.post(url, {"kind": "export"})
            self.assertEqual(second.status_code, 409)
            self.assertEqual(second.json()["id"], first.json()["id"])
            self.assertEqual(ShopCatalogJob.objects.filter(kind="export").count(), 1)
            ShopCatalogJob.objects.filter(pk=first.json()["id"]).update(status="running")
            self.assertEqual(self.client.post(url, {"kind": "export"}).status_code, 409)
            upload = SimpleUploadedFile("shops.xlsx", b"test", content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.assertEqual(self.client.post(url, {"kind": "import", "file": upload}).status_code, 202)
            upload = SimpleUploadedFile("shops.xlsx", b"test")
            self.assertEqual(self.client.post(url, {"kind": "import", "file": upload}).status_code, 409)
            ShopCatalogJob.objects.filter(kind="export").update(status="succeeded")
            self.assertEqual(self.client.post(url, {"kind": "export"}).status_code, 202)

    def test_home_page_links_to_document_error_campaigns(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("home"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Book lỗi chứng từ")
        self.assertContains(response, reverse("document_campaigns:campaign_list"))

    def test_create_campaign_ui_creates_data_review_campaign(self):
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("document_campaigns:campaign_create"),
            {
                "name": "Chiến dịch tháng 10/2026",
                "campaign_type": str(self.campaign_type.pk),
                "report_month": "2026-10",
                "response_deadline": "2026-10-20T17:00",
            },
        )

        campaign = Campaign.objects.get(code="DEC-HC-202610")
        self.assertRedirects(
            response,
            reverse("document_campaigns:campaign_detail", kwargs={"campaign_id": campaign.pk}),
        )
        self.assertEqual(campaign.report_month, date(2026, 10, 1))
        self.assertEqual(campaign.campaign_type, self.campaign_type)
        self.assertEqual(campaign.status, Campaign.Status.DATA_REVIEW)
        self.assertEqual((campaign.link_expires_at - campaign.response_deadline).days, 31)

    @patch("app_document_campaigns.views.run_campaign_sql_import.delay")
    def test_stage_sql_ui_queues_background_job(self, delay):
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        delay.return_value.id = "celery-task-id"
        self.client.force_login(self.user)

        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(
                reverse("document_campaigns:stage_sql_version", kwargs={"version_id": self.version.pk}),
                follow=True,
            )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Đã đưa yêu cầu vào hàng đợi")
        job = CampaignImportJob.objects.get(version=self.version)
        self.assertEqual(job.status, CampaignImportJob.Status.QUEUED)
        self.assertEqual(job.celery_task_id, "celery-task-id")
        delay.assert_called_once_with(job.pk)

    @patch("app_document_campaigns.views.run_campaign_excel_export.delay")
    def test_export_ui_queues_background_job(self, delay):
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        delay.return_value.id = "excel-task-id"
        stage_import_rows(source=self.source, rows=[self._valid_row()])
        self.client.force_login(self.user)

        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(
                reverse(
                    "document_campaigns:queue_staging_excel_export",
                    kwargs={"version_id": self.version.pk},
                ),
                follow=True,
            )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Đã đưa yêu cầu tạo file Excel vào hàng đợi")
        job = CampaignImportJob.objects.get(
            version=self.version,
            job_type=CampaignImportJob.JobType.EXCEL_EXPORT,
        )
        self.assertEqual(job.status, CampaignImportJob.Status.QUEUED)
        self.assertEqual(job.celery_task_id, "excel-task-id")
        delay.assert_called_once_with(job.pk)

    @patch("app_document_campaigns.views.run_campaign_excel_import.delay")
    def test_excel_upload_ui_queues_background_job(self, delay):
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        delay.return_value.id = "excel-import-task-id"
        stage_import_rows(source=self.source, rows=[self._valid_row()])
        workbook = build_cleaning_workbook(self.version)
        self.client.force_login(self.user)

        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            with self.captureOnCommitCallbacks(execute=True):
                response = self.client.post(
                    reverse(
                        "document_campaigns:import_staging_excel",
                        kwargs={"version_id": self.version.pk},
                    ),
                    {
                        "return_to": "campaign_detail",
                        "file": SimpleUploadedFile("cleaned.xlsx", workbook.getvalue()),
                    },
                    follow=True,
                )

            self.assertEqual(response.status_code, 200)
            self.assertContains(response, "Đã tải file lên và đưa vào hàng đợi đối chiếu")
            job = CampaignImportJob.objects.get(
                version=self.version,
                job_type=CampaignImportJob.JobType.EXCEL_IMPORT,
            )
            self.assertEqual(job.status, CampaignImportJob.Status.QUEUED)
            self.assertEqual(job.celery_task_id, "excel-import-task-id")
            self.assertTrue(job.input_file.storage.exists(job.input_file.name))
            delay.assert_called_once_with(job.pk)

    def test_background_excel_import_job_records_progress_and_result(self):
        from app_document_campaigns.tasks import run_campaign_excel_import

        stage_import_rows(source=self.source, rows=[self._valid_row()])
        workbook = build_cleaning_workbook(self.version)
        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            job = CampaignImportJob.objects.create(
                version=self.version,
                job_type=CampaignImportJob.JobType.EXCEL_IMPORT,
                total_steps=3,
                input_file=SimpleUploadedFile("cleaned.xlsx", workbook.getvalue()),
                input_filename="cleaned.xlsx",
                created_by=self.user,
            )

            result = run_campaign_excel_import.run(job.pk)

            job.refresh_from_db()
            self.assertEqual(result["status"], "succeeded")
            self.assertEqual(job.status, CampaignImportJob.Status.SUCCEEDED)
            self.assertEqual(job.current_step, 3)
            self.assertEqual(job.summary["rows"], 1)
            self.assertTrue(
                CampaignImportSource.objects.filter(
                    version=self.version,
                    name=EXCEL_SOURCE_NAME,
                ).exists()
            )

    def test_background_excel_import_job_exposes_safe_validation_error(self):
        from app_document_campaigns.tasks import run_campaign_excel_import

        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            job = CampaignImportJob.objects.create(
                version=self.version,
                job_type=CampaignImportJob.JobType.EXCEL_IMPORT,
                total_steps=3,
                input_file=SimpleUploadedFile("broken.xlsx", b"not an Excel workbook"),
                input_filename="broken.xlsx",
                created_by=self.user,
            )

            result = run_campaign_excel_import.run(job.pk)

            job.refresh_from_db()
            self.assertEqual(result["status"], "failed")
            self.assertEqual(job.status, CampaignImportJob.Status.FAILED)
            self.assertTrue(job.finished_at)
            self.assertEqual(job.error_message, "Không đọc được cấu trúc file Excel.")

    def test_background_excel_job_creates_downloadable_file(self):
        from app_document_campaigns.tasks import run_campaign_excel_export

        stage_import_rows(source=self.source, rows=[self._valid_row()])
        job = CampaignImportJob.objects.create(
            version=self.version,
            job_type=CampaignImportJob.JobType.EXCEL_EXPORT,
            total_steps=1,
            created_by=self.user,
        )

        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            result = run_campaign_excel_export.run(job.pk)
            job.refresh_from_db()
            self.assertEqual(result["status"], "succeeded")
            self.assertEqual(job.status, CampaignImportJob.Status.SUCCEEDED)
            self.assertTrue(job.output_file.storage.exists(job.output_file.name))

            self.client.force_login(self.user)
            response = self.client.get(
                reverse(
                    "document_campaigns:export_staging_excel",
                    kwargs={"version_id": self.version.pk},
                )
            )
            self.assertEqual(response.status_code, 200)
            self.assertEqual(
                response["Content-Type"],
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            downloaded = b"".join(response.streaming_content)
            self.assertEqual(load_workbook(BytesIO(downloaded))["_Metadata"]["B1"].value, "dec-cleaning-v1")

    @patch("app_document_campaigns.services.excel_exports.build_cleaning_workbook")
    def test_background_excel_job_records_failure_instead_of_hanging(self, build_workbook):
        from app_document_campaigns.tasks import run_campaign_excel_export

        build_workbook.side_effect = RuntimeError("broken workbook")
        job = CampaignImportJob.objects.create(
            version=self.version,
            job_type=CampaignImportJob.JobType.EXCEL_EXPORT,
            total_steps=1,
            created_by=self.user,
        )

        with self.assertRaises(RuntimeError):
            run_campaign_excel_export.run(job.pk)

        job.refresh_from_db()
        self.assertEqual(job.status, CampaignImportJob.Status.FAILED)
        self.assertTrue(job.finished_at)
        self.assertNotIn("broken workbook", job.error_message)

    @patch("app_document_campaigns.services.sql_sources.stage_monthly_sql_sources")
    def test_background_sql_job_records_success(self, stage_sources):
        from app_document_campaigns.tasks import run_campaign_sql_import

        expected = {
            FOLDER_SOURCE_NAME: {"rows": 2, "invalid": 0},
            DOCUMENT_SOURCE_NAME: {"rows": 3, "invalid": 1},
        }
        stage_sources.return_value = expected
        job = CampaignImportJob.objects.create(version=self.version, created_by=self.user)

        result = run_campaign_sql_import.run(job.pk)

        job.refresh_from_db()
        self.assertEqual(result["status"], "succeeded")
        self.assertEqual(job.status, CampaignImportJob.Status.SUCCEEDED)
        self.assertEqual(job.current_step, 2)
        self.assertEqual(job.summary, expected)

    @patch("app_document_campaigns.services.sql_sources.stage_monthly_sql_sources")
    def test_background_sql_job_records_failure_instead_of_hanging(self, stage_sources):
        from app_document_campaigns.tasks import run_campaign_sql_import

        stage_sources.side_effect = RuntimeError("database unavailable")
        job = CampaignImportJob.objects.create(version=self.version, created_by=self.user)

        with self.assertRaises(RuntimeError):
            run_campaign_sql_import.run(job.pk)

        job.refresh_from_db()
        self.assertEqual(job.status, CampaignImportJob.Status.FAILED)
        self.assertTrue(job.finished_at)
        self.assertNotIn("database unavailable", job.error_message)
