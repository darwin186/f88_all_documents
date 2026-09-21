import json
from datetime import date, timedelta
from unittest.mock import patch

from django.core import mail
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from app_documents.models import (
    AreaManager,
    Manager,
    Region,
    RegionManager,
    Shop,
    UserProfile,
)
from app_document_campaigns.models import (
    Campaign,
    CampaignError,
    CampaignType,
    CampaignVersion,
    ChecklistQuestion,
    ChecklistTemplate,
    ShopResponse,
    ShopSubmission,
    ShopResponseOption,
    CampaignResponseOption,
    ShopEmailDelivery,
    AreaManagerAccessLink,
)
from app_document_campaigns.services.access_links import issue_shop_access_link, resolve_shop_access_link
from app_document_campaigns.tasks import process_shop_submission


class ShopResponseFlowTests(TestCase):
    def test_review_excel_stalled_job_is_failed_and_can_be_recreated(self):
        from app_document_campaigns.models import TeamReviewExcelJob
        self._review_admin()
        self.client.force_login(self.user)
        job = TeamReviewExcelJob.objects.create(campaign=self.campaign, requested_by=self.user, kind="export", status="running")
        TeamReviewExcelJob.objects.filter(pk=job.pk).update(updated_at=timezone.now()-timedelta(minutes=13))
        url = reverse("document_campaigns:review_excel_status",kwargs={"campaign_id":self.campaign.pk,"job_id":job.pk})
        self.assertEqual(self.client.get(url).json()["status"], "failed")
        TeamReviewExcelJob.objects.create(campaign=self.campaign, requested_by=self.user, kind="export")

    def test_review_receipt_is_live_for_unlinked_folder_and_exported(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from app_documents.models import Folder, FolderStatus
        from app_document_campaigns.review_views import review_queryset
        from app_document_campaigns.services.folder_receipt import receipt_values
        from app_document_campaigns.services.team_review_excel import build_workbook
        status = FolderStatus.objects.create(folder_status_code="RXT", folder_status_name="Nhận test", created_by=self.user, is_received=False)
        folder = Folder.objects.create(folder_code="REVIEW-RECEIPT", shop_id=self.shop, folder_created_date=date(2026, 9, 1), folder_status_id=status)
        self.error_1.error_type = "folder"
        self.error_1.code = folder.folder_code
        self.error_1.save()
        self.assertEqual(receipt_values(review_queryset(self.campaign).get(pk=self.error_1.pk))[0], "Chưa nhận")
        status.is_received = True
        status.save()
        folder.lastest_received_date = timezone.now()
        folder.save()
        error = review_queryset(self.campaign).get(pk=self.error_1.pk)
        self.assertEqual(receipt_values(error)[0], "Đã nhận")
        self._review_admin()
        self.client.force_login(self.user)
        url = reverse("document_campaigns:refresh_folder_receipt",kwargs={"campaign_id":self.campaign.pk,"error_id":self.error_1.pk})
        self.assertEqual(self.client.get(url).json()["received"], True)
        content, _ = build_workbook(self.campaign)
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook["Team review"]["M1"].value, "Trạng thái nhận quyển")
        self.assertEqual(workbook["Team review"]["M2"].value, "Đã nhận")
        self.assertTrue(workbook["Team review"]["N2"].value)
        workbook.close()

    def test_review_response_filter_and_bulk_comments_are_atomic(self):
        from app_document_campaigns.models import TeamReview
        self._review_admin()
        self.client.force_login(self.user)
        original = ShopResponse.objects.create(error=self.error_1, shop=self.shop, answer_code="PGD hẹn bổ sung chứng từ", note="PGD gốc")
        page = self.client.get(reverse("document_campaigns:campaign_review", kwargs={"campaign_id":self.campaign.pk}), {"response":original.answer_code})
        self.assertEqual([error.pk for error in page.context["page"]], [self.error_1.pk])
        self.assertNotContains(page, "Chỉ review khi PGD đã gửi chính thức hoặc hết hạn phản hồi.")
        self.assertContains(page, "data-review-progress")
        self.assertNotContains(page, 'id="review-progress"')
        url = reverse("document_campaigns:bulk_team_review", kwargs={"campaign_id":self.campaign.pk})
        payload = {"rows":[{"id":self.error_1.pk,"expected_review_id":None},{"id":self.error_2.pk,"expected_review_id":None}],"decision":"approved","note":"Nhận xét chung","mode":"append"}
        response = self._post_json(url, payload)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["updated"], 2)
        self.assertEqual(TeamReview.objects.count(), 2)
        original.refresh_from_db()
        self.assertEqual(original.note, "PGD gốc")
        self.assertEqual(self._post_json(url, payload).status_code, 409)
        self.assertEqual(TeamReview.objects.count(), 2)
        latest = TeamReview.objects.get(error=self.error_1)
        payload["rows"] = [{"id":self.error_1.pk,"expected_review_id":latest.pk}]
        payload["decision"] = ""
        payload["note"] = "Thêm ghi chú"
        self.assertEqual(self._post_json(url, payload).status_code, 200)
        self.assertEqual(TeamReview.objects.filter(error=self.error_1).latest("pk").note,"Nhận xét chung\nThêm ghi chú")

    def test_bulk_review_is_admin_only_and_waits_for_submission(self):
        self.client.force_login(self.user)
        url = reverse("document_campaigns:bulk_team_review", kwargs={"campaign_id":self.campaign.pk})
        payload = {"rows":[{"id":self.error_1.pk,"expected_review_id":None}],"decision":"approved","note":"Test","mode":"replace"}
        self.assertEqual(self._post_json(url,payload).status_code, 403)
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        self.assertEqual(self._post_json(url,payload).status_code, 409)

    def _review_file(self, content, edits=(), delete_row=None):
        from io import BytesIO
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content))
        for row, col, value in edits:
            workbook["Team review"].cell(row, col, value)
        if delete_row:
            workbook["Team review"].delete_rows(delete_row)
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def _review_admin(self):
        self.user.groups.add(Group.objects.get_or_create(name="admin")[0])
        self.campaign.response_deadline = timezone.now() - timedelta(minutes=1)
        self.campaign.save(update_fields=["response_deadline"])

    def test_team_review_excel_roundtrip_dropdown_and_omitted_rows(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from app_document_campaigns.models import TeamReview
        from app_document_campaigns.services.team_review_excel import build_workbook, import_workbook
        self._review_admin()
        from datetime import datetime
        self.error_1.source_created_at = datetime(2026, 9, 17, 12, 30)
        self.error_1.save(update_fields=["source_created_at"])
        original = ShopResponse.objects.create(error=self.error_1, shop=self.shop, answer_code="PGD xác nhận lỗi", note="Phản hồi gốc", version_no=2)
        content, total = build_workbook(self.campaign)
        self.assertEqual(total, 2)
        self.assertEqual(import_workbook(self.campaign.pk, self.user, content)["updated"], 0)
        workbook = load_workbook(BytesIO(content))
        sheet = workbook["Team review"]
        validation = list(sheet.data_validations.dataValidation)[0]
        self.assertEqual(validation.formula1, '"Giữ lỗi,Loại lỗi"')
        self.assertEqual(str(validation.sqref), "O2:O3")
        self.assertTrue(validation.showErrorMessage)
        self.assertFalse(sheet["O2"].protection.locked)
        self.assertTrue(sheet["K2"].protection.locked)
        workbook.close()
        edited = self._review_file(content, [(2, 15, "Loại lỗi"), (2, 16, "Team đã kiểm tra")], delete_row=3)
        self.assertEqual(import_workbook(self.campaign.pk, self.user, edited)["updated"], 1)
        self.assertEqual(TeamReview.objects.get(error=self.error_1).decision, "excluded")
        self.assertFalse(TeamReview.objects.filter(error=self.error_2).exists())
        original.refresh_from_db()
        self.assertEqual(original.note, "Phản hồi gốc")
        self.assertEqual(original.version_no, 2)
        self.assertEqual(self.campaign.errors.count(), 2)

    def test_team_review_excel_formats_naive_aware_and_missing_source_dates(self):
        from datetime import datetime, timezone as datetime_timezone
        from app_document_campaigns.services.team_review_excel import reference_values
        with timezone.override("Asia/Ho_Chi_Minh"):
            self.error_1.source_created_at = datetime(2026, 9, 17, 0, 30)
            self.assertEqual(reference_values(self.error_1, {})[4], "17/09/2026")
            self.error_1.source_created_at = datetime(2026, 9, 16, 18, 30, tzinfo=datetime_timezone.utc)
            self.assertEqual(reference_values(self.error_1, {})[4], "17/09/2026")
            self.error_1.source_created_at = None
            self.assertEqual(reference_values(self.error_1, {})[4], "")

    def test_team_review_excel_tampering_invalid_dropdown_and_conflict_are_atomic(self):
        from app_document_campaigns.models import TeamReview
        from app_document_campaigns.services.team_review_excel import build_workbook, import_workbook
        self._review_admin()
        content, _ = build_workbook(self.campaign)
        for edits in [[(2, 15, "Giữ lỗi"), (3, 15, "Sai lựa chọn")], [(2, 15, "Giữ lỗi"), (2, 12, "Sửa phản hồi PGD")], [(2, 17, "Sai chữ ký")]]:
            with self.subTest(edits=edits), self.assertRaises(ValueError):
                import_workbook(self.campaign.pk, self.user, self._review_file(content, edits))
            self.assertEqual(TeamReview.objects.count(), 0)
        edited = self._review_file(content, [(2, 15, "Giữ lỗi")])
        with self.assertRaises(ValueError):
            import_workbook(self.campaign.pk + 1, self.user, edited)
        self.campaign.response_deadline = timezone.now() + timedelta(days=1)
        self.campaign.save()
        with self.assertRaisesMessage(ValueError, "PGD chưa gửi/hết hạn"):
            import_workbook(self.campaign.pk, self.user, edited)
        self._review_admin()
        TeamReview.objects.create(error=self.error_1, decision="excluded", reviewed_by=self.user)
        with self.assertRaisesMessage(ValueError, "dữ liệu đã thay đổi"):
            import_workbook(self.campaign.pk, self.user, edited)
        self.assertEqual(TeamReview.objects.count(), 1)

        TeamReview.objects.create(error=self.error_1, decision="approved", reviewed_by=self.user)
        _, total = build_workbook(self.campaign)
        self.assertEqual(total, 2)  # Multiple history records must not duplicate an error row.

    def test_team_review_excel_task_files_and_failure(self):
        from tempfile import TemporaryDirectory
        from django.test import override_settings
        from django.core.files.base import ContentFile
        from app_document_campaigns.models import TeamReviewExcelJob
        from app_document_campaigns.tasks import process_team_review_excel
        self._review_admin()
        with TemporaryDirectory() as media, override_settings(MEDIA_ROOT=media):
            job = TeamReviewExcelJob.objects.create(campaign=self.campaign, requested_by=self.user, kind="export")
            self.assertEqual(process_team_review_excel(job.pk)["status"], "succeeded")
            job.refresh_from_db()
            self.assertTrue(job.output_file.storage.exists(job.output_file.name))
            with job.output_file.open("rb") as file:
                content = file.read()
            imported = TeamReviewExcelJob.objects.create(campaign=self.campaign, requested_by=self.user, kind="import")
            imported.input_file.save("review.xlsx", ContentFile(self._review_file(content, [(2, 15, "Giữ lỗi")])))
            self.assertEqual(process_team_review_excel(imported.pk)["status"], "succeeded")
            imported.refresh_from_db()
            self.assertEqual(imported.summary["updated"], 1)
            bad = TeamReviewExcelJob.objects.create(campaign=self.campaign, requested_by=self.user, kind="import")
            bad.input_file.save("bad.xlsx", ContentFile(b"invalid"))
            self.assertEqual(process_team_review_excel(bad.pk)["status"], "failed")

    def test_team_review_excel_endpoints_permissions_single_job_and_missing_file(self):
        from app_document_campaigns.models import TeamReviewExcelJob
        self.client.force_login(self.user)
        url = reverse("document_campaigns:queue_review_excel", kwargs={"campaign_id": self.campaign.pk})
        self.assertEqual(self.client.post(url, {"kind": "export"}).status_code, 403)
        self._review_admin()
        with patch("app_document_campaigns.tasks.process_team_review_excel.delay") as delay:
            response = self.client.post(url, {"kind": "export"})
            self.assertEqual(response.status_code, 202)
            job = TeamReviewExcelJob.objects.get(pk=response.json()["id"])
            delay.assert_called_once_with(job.pk)
            self.assertEqual(self.client.post(url, {"kind": "export"}).status_code, 409)
        job.status = "succeeded"
        job.output_file = "document_campaigns/team_review/exports/missing-test.xlsx"
        job.save()
        url = reverse("document_campaigns:download_review_excel", kwargs={"campaign_id": self.campaign.pk, "job_id": job.pk})
        self.assertEqual(self.client.get(url).status_code, 404)

    def test_shop_rows_show_employee_business_and_sort_source_date(self):
        now = timezone.now()
        self.error_1.employee_name = "Đinh Thị Vui"
        self.error_1.business_type_name = "Đóng HĐCC chủ động"
        self.error_1.source_created_at = now
        self.error_1.save()
        self.error_2.source_created_at = now - timedelta(days=3)
        self.error_2.save(update_fields=["source_created_at"])
        url = reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token})
        ascending = self.client.get(url)
        self.assertContains(ascending, "Nhân viên: Đinh Thị Vui")
        self.assertContains(ascending, "Nghiệp vụ: Đóng HĐCC chủ động")
        self.assertContains(ascending, "Ngày phát sinh")
        self.assertContains(ascending, "data-date-sort")
        self.assertEqual([row["error"].pk for row in ascending.context["rows"]], [self.error_2.pk, self.error_1.pk])
        descending = self.client.get(url, {"date_order": "desc"})
        self.assertEqual([row["error"].pk for row in descending.context["rows"]], [self.error_1.pk, self.error_2.pk])
        self.error_2.source_created_at = None
        self.error_2.save(update_fields=["source_created_at"])
        missing_date = self.client.get(url, {"date_order": "invalid"})
        self.assertEqual(missing_date.context["date_order"], "asc")
        self.assertEqual([row["error"].pk for row in missing_date.context["rows"]], [self.error_1.pk, self.error_2.pk])
        self.assertEqual(ShopResponse.objects.filter(error__campaign=self.campaign).count(), 0)

    def test_area_contract_records_filter_is_scoped(self):
        foreign_shop = self._create_foreign_shop_error()
        area_user = self._scoped_user("records-area", "supervisor", email="area@example.com")
        self.client.force_login(area_user)
        url = reverse("document_campaigns:campaign_response_records", kwargs={"campaign_id": self.campaign.pk})
        page = self.client.get(url, {"shop": str(self.shop.pk), "date_order": "desc"})
        self.assertContains(page, "Tên phòng giao dịch")
        self.assertContains(page, "Tên nhân viên")
        self.assertContains(page, "Nghiệp vụ")
        self.assertContains(page, "PGD Test Autosave")
        self.assertNotContains(page, "PGD Ngoài phạm vi")
        self.assertEqual(page.context["page"].paginator.count, 2)
        forbidden = self.client.get(url, {"shop": str(foreign_shop.pk)})
        self.assertEqual(forbidden.context["page"].paginator.count, 0)
        self.assertNotContains(forbidden, "PGD Ngoài phạm vi")
        checker = self._scoped_user("records-checker", "checker", region=self.region)
        self.client.force_login(checker)
        self.assertEqual(self.client.get(url).context["page"].paginator.count, 0)

    def test_admin_contract_records_can_filter_one_shop_and_sort_dates(self):
        self._create_foreign_shop_error()
        self.error_1.source_created_at = timezone.now()
        self.error_1.save(update_fields=["source_created_at"])
        self.error_2.source_created_at = timezone.now() - timedelta(days=2)
        self.error_2.save(update_fields=["source_created_at"])
        self.client.force_login(self._scoped_user("records-admin", "admin"))
        url = reverse("document_campaigns:campaign_response_records", kwargs={"campaign_id": self.campaign.pk})
        response = self.client.get(url, {"shop": str(self.shop.pk), "date_order": "desc"})
        self.assertEqual([error.pk for error in response.context["page"]], [self.error_1.pk, self.error_2.pk])
        self.assertContains(response, "PGD Ngoài phạm vi")  # Available in the admin's filter options.
        detail = self.client.get(reverse("document_campaigns:shop_response_monitor_detail", kwargs={"campaign_id": self.campaign.pk, "shop_id": self.shop.pk}), {"date_order": "asc"})
        self.assertEqual([error.pk for error in detail.context["errors"]], [self.error_2.pk, self.error_1.pk])

    def setUp(self):
        self.user = get_user_model().objects.create_user(username="campaign-admin", password="test-password")
        self.campaign_type = CampaignType.objects.get(code="hardcopy-document-error")
        self.region = Region.objects.create(region_code="TEST", region_name="Vùng test")
        self.area = AreaManager.objects.create(
            areaManager_code="AREA01",
            areaManager_name="QLKV test",
            areaManager_email="area@example.com",
        )
        self.region_manager = RegionManager.objects.create(
            regionManager_code="QLV01",
            regionManager_name="QLV test",
            regionManager_email="region@example.com",
        )
        self.manager = Manager.objects.create(
            manager_code="MANAGER01",
            areaManager=self.area,
            regionManager=self.region_manager,
            qlkv_code="QLKV01",
            qlkv_name="QLKV test",
            qlkv_email="area@example.com",
            qlv_code="QLV01",
            qlv_name="QLV test",
            qlv_email="region@example.com",
        )
        self.shop = Shop.objects.create(
            shop_code=9001,
            shop_name="PGD Test Autosave",
            shop_email="shop@example.com",
            manager_id=self.manager,
            region_id=self.region,
        )
        now = timezone.now()
        self.campaign = Campaign.objects.create(
            code="CT-2026-09",
            name="Chiến dịch tháng 09/2026",
            campaign_type=self.campaign_type,
            report_month=date(2026, 9, 1),
            status=Campaign.Status.ACTIVE,
            response_opens_at=now - timedelta(hours=1),
            response_deadline=now + timedelta(days=7),
            link_expires_at=now + timedelta(days=8),
            created_by=self.user,
        )
        self.version = CampaignVersion.objects.create(
            campaign=self.campaign,
            version_number=1,
            source_type=CampaignVersion.SourceType.EXCEL,
            status=CampaignVersion.Status.PUBLISHED,
            created_by=self.user,
        )
        self.campaign.current_version = self.version
        self.campaign.save(update_fields=["current_version"])
        self.checklist = ChecklistTemplate.objects.create(
            code="SHOP_CONFIRM",
            name="PGD xác nhận lỗi",
            created_by=self.user,
        )
        ChecklistQuestion.objects.create(
            template=self.checklist,
            code="CONFIRM",
            label="PGD phản hồi",
            question_type=ChecklistQuestion.QuestionType.SINGLE,
            options=["PGD xác nhận lỗi", "PGD hẹn bổ sung chứng từ"],
            is_required=True,
        )
        self.error_1 = self._create_error("document:1001", 1001)
        self.error_2 = self._create_error("document:1002", 1002)
        self.link, self.raw_token = issue_shop_access_link(
            campaign=self.campaign,
            shop=self.shop,
            allowed_email=self.shop.shop_email,
            created_by=self.user,
        )

    def _create_error(self, source_key, source_id):
        return CampaignError.objects.create(
            campaign=self.campaign,
            version=self.version,
            source_key=source_key,
            source_object_id=source_id,
            error_type=CampaignError.ErrorType.DOCUMENT,
            shop=self.shop,
            area_manager=self.area,
            region=self.region,
            contract_code=f"HD{source_id}",
            checking_issue="Thiếu chữ ký khách hàng",
            checklist_template=self.checklist,
            status=CampaignError.Status.WAITING_SHOP,
        )

    def _post_json(self, url, payload):
        return self.client.post(url, data=json.dumps(payload), content_type="application/json")

    def test_response_page_and_health_are_available(self):
        health = self.client.get(reverse("document_campaigns:health"))
        page = self.client.get(
            reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token})
        )

        self.assertEqual(health.status_code, 200)
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "PGD Test Autosave")
        self.assertContains(page, "Thiếu chữ ký khách hàng", count=2)

    def test_response_dropdown_can_fall_back_to_campaign_type_checklist(self):
        self.campaign_type.shop_checklist_template = self.checklist
        self.campaign_type.save(update_fields=["shop_checklist_template"])
        CampaignError.objects.filter(pk__in=[self.error_1.pk, self.error_2.pk]).update(
            checklist_template=None
        )

        page = self.client.get(
            reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token})
        )

        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "PGD xác nhận lỗi")
        self.assertContains(page, "PGD hẹn bổ sung chứng từ")
        self.assertContains(page, '<option value="" selected>— Chọn phản hồi —</option>', count=2)
        self.assertNotContains(page, '<option value="PGD hẹn bổ sung chứng từ" selected>')

    def test_supplement_answer_is_valid_and_remains_selected_after_autosave(self):
        autosave = self._post_json(
            reverse("document_campaigns:autosave_batch", kwargs={"raw_token": self.raw_token}),
            {
                "changes": [{
                    "error_uid": str(self.error_1.error_uid),
                    "version_no": 0,
                    "answer_code": "PGD hẹn bổ sung chứng từ",
                    "note": "",
                }]
            },
        )

        page = self.client.get(
            reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token})
        )

        self.assertEqual(autosave.status_code, 200)
        self.assertContains(
            page,
            '<option value="PGD hẹn bổ sung chứng từ" selected>PGD hẹn bổ sung chứng từ</option>',
            count=1,
        )

    def test_response_table_header_does_not_overlap_first_invalid_row(self):
        page = self.client.get(
            reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token})
        )

        self.assertContains(page, '<thead class="bg-gray-50')
        self.assertNotContains(page, 'sticky top-[65px]')

    def _scoped_user(self, username, role, *, region=None, shop=None, email=""):
        user = get_user_model().objects.create_user(username=username, email=email)
        user.groups.add(Group.objects.get_or_create(name=role)[0])
        UserProfile.objects.create(
            user=user,
            department="Vận hành",
            region=region,
            shop=shop,
        )
        return user

    def _create_foreign_shop_error(self):
        region = Region.objects.create(region_code="OTHER", region_name="Vùng khác")
        area = AreaManager.objects.create(
            areaManager_code="AREA02",
            areaManager_name="QLKV khác",
            areaManager_email="other-area@example.com",
        )
        region_manager = RegionManager.objects.create(
            regionManager_code="QLV02",
            regionManager_name="QLV khác",
            regionManager_email="other-region@example.com",
        )
        manager = Manager.objects.create(
            manager_code="MANAGER02",
            areaManager=area,
            regionManager=region_manager,
            qlkv_code="QLKV02",
            qlkv_name="QLKV khác",
            qlkv_email="other-area@example.com",
            qlv_code="QLV02",
            qlv_name="QLV khác",
            qlv_email="other-region@example.com",
        )
        shop = Shop.objects.create(
            shop_code=9002,
            shop_name="PGD Ngoài phạm vi",
            shop_email="other-shop@example.com",
            manager_id=manager,
            region_id=region,
        )
        CampaignError.objects.create(
            campaign=self.campaign,
            version=self.version,
            source_key="document:foreign",
            source_object_id=2001,
            error_type=CampaignError.ErrorType.DOCUMENT,
            shop=shop,
            area_manager=area,
            region=region,
            contract_code="HD2001",
            checking_issue="Lỗi ngoài phạm vi",
            checklist_template=self.checklist,
            status=CampaignError.Status.WAITING_SHOP,
        )
        return shop

    def test_admin_monitor_lists_shops_and_opens_response_detail(self):
        admin = self._scoped_user("monitor-admin", "admin")
        self.client.force_login(admin)

        listing = self.client.get(
            reverse(
                "document_campaigns:campaign_response_monitor",
                kwargs={"campaign_id": self.campaign.pk},
            )
        )
        detail = self.client.get(
            reverse(
                "document_campaigns:shop_response_monitor_detail",
                kwargs={"campaign_id": self.campaign.pk, "shop_id": self.shop.pk},
            )
        )

        self.assertEqual(listing.status_code, 200)
        self.assertContains(listing, "PGD Test Autosave")
        self.assertContains(listing, "Chưa phản hồi")
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "Thiếu chữ ký khách hàng", count=2)

    def test_monitor_pagination_parent_filters_and_multi_status(self):
        admin = self._scoped_user("monitor-pages-admin", "admin")
        self.client.force_login(admin)
        for index in range(51):
            shop = Shop.objects.create(shop_code=9500 + index, shop_name=f"PGD phân trang {index}", shop_email=f"page{index}@example.com", manager_id=self.manager, region_id=self.region)
            CampaignError.objects.create(campaign=self.campaign, version=self.version, source_key=f"document:page-{index}", source_object_id=5000+index, error_type="document", shop=shop, area_manager=self.area, region=self.region, status="waiting_shop")
        ShopResponse.objects.create(error=self.error_1, shop=self.shop, answer_code="confirm_error")
        url = reverse("document_campaigns:campaign_response_monitor", kwargs={"campaign_id": self.campaign.pk})
        first = self.client.get(url)
        self.assertEqual(len(first.context["rows"]), 50)
        self.assertEqual(first.context["summary"]["shops"], 52)
        self.assertEqual(first.context["rows"][0]["progress"], 0)
        last = self.client.get(url, {"page": "2"})
        self.assertEqual(len(last.context["rows"]), 2)
        self.assertEqual(last.context["rows"][-1]["shop_id"], self.shop.pk)
        multiple = self.client.get(url, {"status": ["not_started", "in_progress"]})
        self.assertEqual(multiple.context["page"].paginator.count, 52)
        other = self._create_foreign_shop_error()
        filtered = self.client.get(url, {"region": self.region_manager.pk, "area": other.manager_id.areaManager_id, "shop": other.pk})
        self.assertEqual(filtered.context["selected_area"], "")
        self.assertEqual(filtered.context["selected_shop"], "")
        self.assertNotIn(other.pk, [row["shop_id"] for row in filtered.context["rows"]])
        self.assertContains(filtered, self.region_manager.regionManager_name)

    def test_individual_deadline_extension_changes_progress_and_blocks_review(self):
        admin = self._scoped_user("extend-admin", "admin")
        self.client.force_login(admin)
        self.campaign.response_deadline = timezone.now() - timedelta(hours=1)
        self.campaign.save(update_fields=["response_deadline"])
        monitor_url = reverse("document_campaigns:campaign_response_monitor", kwargs={"campaign_id": self.campaign.pk})
        expired_monitor = self.client.get(monitor_url)
        self.assertEqual(expired_monitor.context["shop_response_progress"], 100)
        self.assertEqual(expired_monitor.context["rows"][0]["status"], "expired")
        self.assertEqual(expired_monitor.context["rows"][0]["progress"], 0)
        deadline_url = reverse("document_campaigns:extend_shop_deadline", kwargs={"campaign_id": self.campaign.pk, "shop_id": self.shop.pk})
        response = self.client.post(deadline_url, {"deadline": (timezone.now()+timedelta(days=9)).strftime("%Y-%m-%dT%H:%M")})
        self.assertEqual(response.status_code, 200)
        self.link.refresh_from_db()
        self.assertTrue(self.link.has_deadline_extension)
        self.assertGreaterEqual(self.link.expires_at, self.link.response_deadline)
        self.assertEqual(self.client.get(monitor_url).context["shop_response_progress"], 0)
        review_url = reverse("document_campaigns:save_team_review", kwargs={"campaign_id": self.campaign.pk, "error_id": self.error_1.pk})
        self.assertEqual(self._post_json(review_url, {"decision":"approved", "note":"", "expected_review_id":None}).status_code, 409)
        old_deadline = self.link.response_deadline
        self.client.post(reverse("document_campaigns:issue_campaign_shop_link", kwargs={"campaign_id":self.campaign.pk,"shop_id":self.shop.pk}))
        self.link.refresh_from_db()
        self.assertEqual(self.link.response_deadline, old_deadline)
        self.client.post(reverse("document_campaigns:update_campaign_deadline", kwargs={"campaign_id":self.campaign.pk}), {"response_deadline": (timezone.now()+timedelta(days=2)).strftime("%Y-%m-%dT%H:%M"), "return_to":"response_monitor"})
        self.link.refresh_from_db()
        self.assertEqual(self.link.response_deadline, old_deadline)

    def test_shop_actions_are_admin_only_and_submitted_shops_stay_locked(self):
        deadline_url = reverse("document_campaigns:extend_shop_deadline", kwargs={"campaign_id":self.campaign.pk,"shop_id":self.shop.pk})
        email_url = reverse("document_campaigns:email_shop_link", kwargs={"campaign_id":self.campaign.pk,"shop_id":self.shop.pk})
        self.client.force_login(self.user)
        self.assertEqual(self.client.post(deadline_url).status_code, 403)
        self.assertEqual(self.client.post(email_url).status_code, 403)
        admin = self._scoped_user("submitted-actions-admin", "admin")
        self.client.force_login(admin)
        ShopSubmission.objects.create(campaign=self.campaign, shop=self.shop, idempotency_key="submitted-actions")
        self.assertEqual(self.client.post(deadline_url).status_code, 409)
        self.assertEqual(self.client.post(email_url).status_code, 409)
        monitor = self.client.get(reverse("document_campaigns:campaign_response_monitor", kwargs={"campaign_id":self.campaign.pk}))
        self.assertEqual(monitor.context["shop_response_progress"], 100)
        self.assertEqual(monitor.context["rows"][0]["progress"], 0)

    @patch("app_document_campaigns.tasks.send_campaign_shop_link.delay")
    def test_email_button_queues_new_link_and_mail_ccs_current_managers(self, delay):
        from app_document_campaigns.tasks import send_campaign_shop_link
        admin = self._scoped_user("email-admin", "admin")
        self.client.force_login(admin)
        response = self.client.post(reverse("document_campaigns:email_shop_link", kwargs={"campaign_id":self.campaign.pk,"shop_id":self.shop.pk}))
        self.assertEqual(response.status_code, 200)
        delivery = ShopEmailDelivery.objects.get()
        delay.assert_called_once_with(self.link.pk, response.json()["url"], delivery.pk)
        self.assertEqual(send_campaign_shop_link.run(self.link.pk, response.json()["url"], delivery.pk)["status"], "sent")
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, ShopEmailDelivery.Status.SENT)
        self.assertEqual(delivery.area_manager, self.area)
        self.assertEqual(mail.outbox[-1].to, [self.shop.shop_email])
        self.assertCountEqual(mail.outbox[-1].cc, [self.manager.qlkv_email, self.manager.qlv_email])
        self.assertIn("Hạn phản hồi:", mail.outbox[-1].body)
        status = self.client.get(response.json()["status_url"])
        self.assertEqual(status.json()["area_emailed"], 1)

    def test_admin_can_issue_and_copy_a_new_shop_link_from_response_monitor(self):
        admin = self._scoped_user("link-admin", "admin")
        self.client.force_login(admin)
        monitor = self.client.get(
            reverse(
                "document_campaigns:campaign_response_monitor",
                kwargs={"campaign_id": self.campaign.pk},
            )
        )
        issue_url = reverse(
            "document_campaigns:issue_campaign_shop_link",
            kwargs={"campaign_id": self.campaign.pk, "shop_id": self.shop.pk},
        )

        response = self.client.post(issue_url, HTTP_X_REQUESTED_WITH="XMLHttpRequest")

        self.assertContains(monitor, "PGD đã nhận phát hành")
        self.assertContains(monitor, "PGD Test Autosave")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])
        raw_token = response.json()["url"].rstrip("/").split("/")[-1]
        self.assertEqual(resolve_shop_access_link(raw_token).shop, self.shop)
        old_link = self.client.get(
            reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token})
        )
        self.assertEqual(old_link.status_code, 410)

    def test_admin_can_publish_all_shop_links_without_sending_email(self):
        admin = self._scoped_user("publish-admin", "admin")
        self.client.force_login(admin)

        response = self.client.post(
            reverse(
                "document_campaigns:publish_campaign_to_shops",
                kwargs={"campaign_id": self.campaign.pk},
            ),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["count"], 1)
        self.assertEqual(response.json()["links"][0]["shop_code"], self.shop.shop_code)
        self.assertIn("/respond/", response.json()["links"][0]["url"])
        self.assertEqual(len(mail.outbox), 0)
        self.campaign.refresh_from_db()
        self.assertEqual(self.campaign.status, Campaign.Status.ACTIVE)
        self.assertIsNotNone(self.campaign.response_opens_at)
        old_link = self.client.get(
            reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token})
        )
        self.assertEqual(old_link.status_code, 410)

    def test_non_admin_cannot_issue_shop_link(self):
        shop_user = self._scoped_user("link-shop", "shop", shop=self.shop)
        self.client.force_login(shop_user)

        response = self.client.post(
            reverse(
                "document_campaigns:issue_campaign_shop_link",
                kwargs={"campaign_id": self.campaign.pk, "shop_id": self.shop.pk},
            )
        )

        self.assertEqual(response.status_code, 403)

    def test_admin_can_configure_deadline_for_existing_campaign(self):
        admin = self._scoped_user("deadline-admin", "admin")
        self.client.force_login(admin)
        new_deadline = timezone.now() + timedelta(days=14)

        response = self.client.post(
            reverse(
                "document_campaigns:update_campaign_deadline",
                kwargs={"campaign_id": self.campaign.pk},
            ),
            {"response_deadline": new_deadline.strftime("%Y-%m-%dT%H:%M")},
        )

        self.assertRedirects(
            response,
            reverse("document_campaigns:campaign_detail", kwargs={"campaign_id": self.campaign.pk}),
        )
        self.campaign.refresh_from_db()
        self.assertEqual(self.campaign.link_expires_at - self.campaign.response_deadline, timedelta(days=31))

    def test_deadline_popup_returns_json_and_preserves_individual_extension(self):
        admin = self._scoped_user("deadline-popup-admin", "admin")
        self.client.force_login(admin)
        self.link.has_deadline_extension = True
        self.link.save(update_fields=["has_deadline_extension"])
        old_link_deadline = self.link.response_deadline
        new_deadline = timezone.now() + timedelta(days=14)

        response = self.client.post(
            reverse("document_campaigns:update_campaign_deadline", kwargs={"campaign_id": self.campaign.pk}),
            {"response_deadline": new_deadline.strftime("%Y-%m-%dT%H:%M"), "return_to": "response_monitor"},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
            HTTP_ACCEPT="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/json")
        self.assertTrue(response.json()["ok"])
        self.link.refresh_from_db()
        self.assertEqual(self.link.response_deadline, old_link_deadline)

    def test_admin_can_edit_campaign_name_and_deadline_from_settings_popup(self):
        admin = self._scoped_user("settings-admin", "admin")
        self.client.force_login(admin)
        new_deadline = timezone.now() + timedelta(days=20)

        response = self.client.post(
            reverse(
                "document_campaigns:update_campaign_settings",
                kwargs={"campaign_id": self.campaign.pk},
            ),
            {
                "name": "Book lỗi chứng từ tháng 09 - Đã cập nhật",
                "shop_instructions": "Chọn phản hồi cho từng hợp đồng.\nGhi chú nếu cần.",
                "area_manager_instructions": "QLKV theo dõi các PGD trong khu vực.",
                "response_options": list(ShopResponseOption.objects.filter(is_active=True).values_list("pk", flat=True)),
                "response_deadline": new_deadline.strftime("%Y-%m-%dT%H:%M"),
            },
        )

        self.assertRedirects(
            response,
            reverse("document_campaigns:campaign_detail", kwargs={"campaign_id": self.campaign.pk}),
        )
        self.campaign.refresh_from_db()
        self.link.refresh_from_db()
        self.assertEqual(self.campaign.name, "Book lỗi chứng từ tháng 09 - Đã cập nhật")
        self.assertEqual(self.campaign.shop_instructions, "Chọn phản hồi cho từng hợp đồng.\nGhi chú nếu cần.")
        self.assertEqual(self.campaign.area_manager_instructions, "QLKV theo dõi các PGD trong khu vực.")
        self.assertEqual(self.link.response_deadline, self.campaign.response_deadline)
        self.assertEqual(self.link.expires_at, self.campaign.link_expires_at)

    def test_shop_instructions_render_safely_and_edit_keeps_deadline(self):
        instructions = "Đọc hướng dẫn trước khi phản hồi.\n<script>alert(1)</script>"
        self.campaign.shop_instructions = instructions
        self.campaign.save()
        page = self.client.get(reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token}))
        self.assertContains(page, "Hướng Dẫn")
        self.assertContains(page, "&lt;script&gt;alert(1)&lt;/script&gt;")
        self.assertNotContains(page, "<script>alert(1)</script>")
        admin = self._scoped_user("instruction-admin", "admin")
        self.client.force_login(admin)
        deadline, expiry = self.campaign.response_deadline, self.campaign.link_expires_at
        response = self.client.post(reverse("document_campaigns:update_campaign_settings", kwargs={"campaign_id": self.campaign.pk}), {"name": self.campaign.name, "response_deadline": deadline.strftime("%Y-%m-%dT%H:%M"), "shop_instructions": "Hướng dẫn mới", "response_options": list(ShopResponseOption.objects.filter(is_active=True).values_list("pk", flat=True))})
        self.assertEqual(response.status_code, 302)
        self.campaign.refresh_from_db()
        self.assertEqual(self.campaign.shop_instructions, "Hướng dẫn mới")
        self.assertEqual(self.campaign.response_deadline, deadline)
        self.assertEqual(self.campaign.link_expires_at, expiry)

    def test_campaign_choices_allow_additions_but_not_removals_while_active(self):
        from app_document_campaigns.forms import CampaignSettingsForm
        first, second = list(ShopResponseOption.objects.order_by("pk")[:2])
        original = CampaignResponseOption.objects.create(campaign=self.campaign, option=first, value=first.code, label=first.label)
        self.campaign.status = Campaign.Status.ACTIVE
        self.campaign.save()
        data = {"name": self.campaign.name, "response_deadline": self.campaign.response_deadline.strftime("%Y-%m-%dT%H:%M"), "response_options": [first.pk, second.pk]}
        form = CampaignSettingsForm(data, instance=self.campaign)
        self.assertTrue(form.is_valid(), form.errors)
        form.save()
        self.assertEqual(self.campaign.response_options.count(), 2)
        first.label = "Đổi tên ở danh mục"
        first.is_active = False
        first.save()
        original.refresh_from_db()
        self.assertNotEqual(original.label, first.label)
        data["response_options"] = [second.pk]
        form = CampaignSettingsForm(data, instance=self.campaign)
        self.assertFalse(form.is_valid())
        self.assertIn("response_options", form.errors)
        self.assertEqual(self.campaign.response_options.count(), 2)
        page = self.client.get(reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token}))
        self.assertContains(page, original.label)
        self.assertNotContains(page, "Đổi tên ở danh mục")

    def test_campaign_choices_reject_unselected_response_and_master_data_is_admin_only(self):
        from app_document_campaigns.services.responses import save_response_batch, ResponseValidationError
        option = ShopResponseOption.objects.first()
        CampaignResponseOption.objects.create(campaign=self.campaign, option=option, value=option.code, label=option.label)
        error = self.campaign.errors.first()
        with self.assertRaises(ResponseValidationError):
            save_response_batch(link=self.link, changes=[{"error_uid": str(error.error_uid), "answer_code": "NOT_SELECTED", "expected_version": 0}])
        admin = self._scoped_user("choices-admin", "admin")
        self.client.force_login(admin)
        self.assertContains(self.client.get(reverse("master_data_response_options")), "Danh mục phản hồi PGD")
        response = self.client.post(reverse("master_data_response_options"), {"label": "Lựa chọn mới", "sort_order": 10, "is_active": "on", "description": "Hướng dẫn"})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(ShopResponseOption.objects.filter(label="Lựa chọn mới").exists())
        shop = self._scoped_user("choices-shop", "shop", shop=self.shop)
        self.client.force_login(shop)
        self.assertEqual(self.client.get(reverse("master_data_response_options")).status_code, 403)

    def test_campaign_settings_popup_is_admin_only(self):
        shop_user = self._scoped_user("settings-shop", "shop", shop=self.shop)
        self.client.force_login(shop_user)

        detail = self.client.get(
            reverse("document_campaigns:campaign_detail", kwargs={"campaign_id": self.campaign.pk})
        )
        update = self.client.post(
            reverse(
                "document_campaigns:update_campaign_settings",
                kwargs={"campaign_id": self.campaign.pk},
            ),
            {
                "name": "Không được phép",
                "response_deadline": (timezone.now() + timedelta(days=10)).strftime("%Y-%m-%dT%H:%M"),
            },
        )

        self.assertNotContains(detail, "Cài đặt chiến dịch")
        self.assertEqual(update.status_code, 403)

    def test_region_and_shop_monitor_cannot_view_outside_scope(self):
        foreign_shop = self._create_foreign_shop_error()
        region_user = self._scoped_user(
            "region-monitor", "supervisor", region=self.region, email="region@example.com"
        )
        self.client.force_login(region_user)
        listing = self.client.get(
            reverse(
                "document_campaigns:campaign_response_monitor",
                kwargs={"campaign_id": self.campaign.pk},
            )
        )
        self.assertContains(listing, "PGD Test Autosave")
        self.assertNotContains(listing, "PGD Ngoài phạm vi")

        profile_region_only = self._scoped_user(
            "profile-region-only", "checker", region=self.region
        )
        self.client.force_login(profile_region_only)
        no_manager_scope = self.client.get(
            reverse(
                "document_campaigns:campaign_response_monitor",
                kwargs={"campaign_id": self.campaign.pk},
            )
        )
        self.assertNotContains(no_manager_scope, "PGD Test Autosave")

        shop_user = self._scoped_user("shop-monitor", "shop", shop=self.shop)
        self.client.force_login(shop_user)
        forbidden_detail = self.client.get(
            reverse(
                "document_campaigns:shop_response_monitor_detail",
                kwargs={"campaign_id": self.campaign.pk, "shop_id": foreign_shop.pk},
            )
        )
        self.assertEqual(forbidden_detail.status_code, 404)

    def test_area_monitor_is_derived_from_manager_email(self):
        self._create_foreign_shop_error()
        area_user = self._scoped_user(
            "area-monitor",
            "supervisor",
            email="area@example.com",
        )
        self.client.force_login(area_user)

        listing = self.client.get(
            reverse(
                "document_campaigns:campaign_response_monitor",
                kwargs={"campaign_id": self.campaign.pk},
            )
        )

        self.assertNotContains(listing, "Phạm vi:")
        self.assertContains(listing, "PGD Test Autosave")
        self.assertNotContains(listing, "PGD Ngoài phạm vi")
        self.assertNotContains(listing, "data-open-common-deadline")
        self.assertNotContains(listing, "data-email-link")
        self.assertNotContains(listing, "data-extend-deadline")

        area_listing = self.client.get(reverse("document_campaigns:campaign_area_monitor", kwargs={"campaign_id": self.campaign.pk}))
        self.assertContains(area_listing, "QLKV test")
        self.assertNotContains(area_listing, "data-area-issue")
        self.assertNotContains(area_listing, "data-area-email action")
        self.assertNotContains(area_listing, "data-area-extend")
        forbidden_issue = self.client.post(reverse("document_campaigns:issue_area_manager_link", kwargs={"campaign_id": self.campaign.pk, "area_id": self.area.pk}))
        self.assertEqual(forbidden_issue.status_code, 403)

    def test_admin_issues_unique_readonly_area_link_and_tabs_replace_contract_link(self):
        admin = self._scoped_user("area-link-admin", "admin")
        self.client.force_login(admin)
        foreign_shop = self._create_foreign_shop_error()
        self.campaign.area_manager_instructions = "QLKV vui lòng theo dõi tiến độ PGD thuộc khu vực."
        self.campaign.save(update_fields=["area_manager_instructions"])
        shop_listing = self.client.get(reverse("document_campaigns:campaign_response_monitor", kwargs={"campaign_id": self.campaign.pk}))
        self.assertContains(shop_listing, "Theo Phòng giao dịch")
        self.assertContains(shop_listing, "Theo Quản lý khu vực")
        self.assertNotContains(shop_listing, "Xem lỗi theo hợp đồng")
        filtered_areas = self.client.get(
            reverse("document_campaigns:campaign_area_monitor", kwargs={"campaign_id": self.campaign.pk}),
            {"region": self.region_manager.pk, "area": self.area.pk, "shop": self.shop.pk, "status": ["not_started"]},
        )
        self.assertEqual([row["area_id"] for row in filtered_areas.context["rows"]], [self.area.pk])

        issue = self.client.post(reverse("document_campaigns:issue_area_manager_link", kwargs={"campaign_id": self.campaign.pk, "area_id": self.area.pk}))
        self.assertEqual(issue.status_code, 200)
        self.assertTrue(issue.json()["ok"])
        link = AreaManagerAccessLink.objects.get(campaign=self.campaign, area_manager=self.area)
        public = self.client.get(issue.json()["url"])
        self.assertEqual(public.status_code, 200)
        self.assertContains(public, "Bộ dữ liệu tổng hợp")
        self.assertContains(public, "QLKV vui lòng theo dõi tiến độ PGD thuộc khu vực.")
        self.assertContains(public, "logo-f88-primary.svg")
        self.assertContains(public, "Số lượng phòng giao dịch phát sinh lỗi trong kỳ")
        self.assertContains(public, "Số lượng hợp đồng phát sinh lỗi chứng từ")
        self.assertContains(public, self.error_1.contract_code)
        self.assertContains(public, self.error_2.contract_code)
        self.assertContains(public, "Chỉ xem")
        self.assertNotContains(public, "textarea")
        self.assertContains(public, 'name="q"')
        self.assertContains(public, 'name="shop"')
        self.assertContains(public, 'name="response"')
        self.assertNotContains(public, 'type="date"')
        for sort_key in ("shop", "contract", "date", "error_type", "employee", "issue", "response", "note"):
            self.assertContains(public, f"sort={sort_key}")
        sorted_page = self.client.get(issue.json()["url"], {"sort": "contract", "direction": "desc"})
        self.assertLess(sorted_page.content.index(self.error_2.contract_code.encode()), sorted_page.content.index(self.error_1.contract_code.encode()))
        filtered_page = self.client.get(issue.json()["url"], {"q": self.error_1.contract_code})
        self.assertContains(filtered_page, self.error_1.contract_code)
        self.assertNotContains(filtered_page, self.error_2.contract_code)
        link.refresh_from_db()
        self.assertIsNotNone(link.last_accessed_at)

    @patch("app_document_campaigns.tasks.send_area_manager_view_link.delay")
    def test_admin_can_email_and_extend_area_link(self, delay):
        from app_document_campaigns.tasks import send_area_manager_view_link
        admin = self._scoped_user("area-email-admin", "admin")
        self.client.force_login(admin)
        email = self.client.post(reverse("document_campaigns:email_area_manager_link", kwargs={"campaign_id": self.campaign.pk, "area_id": self.area.pk}))
        self.assertEqual(email.status_code, 200)
        link = AreaManagerAccessLink.objects.get(campaign=self.campaign, area_manager=self.area)
        delay.assert_called_once_with(link.pk, email.json()["url"])
        self.assertEqual(link.email_status, AreaManagerAccessLink.EmailStatus.QUEUED)
        self.assertEqual(send_area_manager_view_link.run(link.pk, email.json()["url"])["status"], "sent")
        self.assertEqual(mail.outbox[-1].to, [self.area.areaManager_email])
        new_expiration = timezone.now() + timedelta(days=30)
        extended = self.client.post(
            reverse("document_campaigns:extend_area_manager_link", kwargs={"campaign_id": self.campaign.pk, "area_id": self.area.pk}),
            {"expires_at": new_expiration.strftime("%Y-%m-%dT%H:%M")},
        )
        self.assertEqual(extended.status_code, 200)
        link.refresh_from_db()
        self.assertGreater(link.expires_at, timezone.now() + timedelta(days=29))

    def test_autosave_batches_rows_and_increments_versions(self):
        url = reverse("document_campaigns:autosave_batch", kwargs={"raw_token": self.raw_token})
        response = self._post_json(
            url,
            {
                "changes": [
                    {
                        "error_uid": str(self.error_1.error_uid),
                        "version_no": 0,
                        "answer_code": "PGD xác nhận lỗi",
                        "note": "Dòng một",
                    },
                    {
                        "error_uid": str(self.error_2.error_uid),
                        "version_no": 0,
                        "answer_code": "PGD hẹn bổ sung chứng từ",
                        "note": "Dòng hai",
                    },
                ]
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(ShopResponse.objects.count(), 2)
        self.assertEqual({item["version_no"] for item in response.json()["saved"]}, {1})

    def test_stale_autosave_returns_conflict_without_overwriting(self):
        url = reverse("document_campaigns:autosave_batch", kwargs={"raw_token": self.raw_token})
        first = {
            "changes": [{
                "error_uid": str(self.error_1.error_uid),
                "version_no": 0,
                "answer_code": "PGD xác nhận lỗi",
                "note": "Bản mới",
            }]
        }
        self.assertEqual(self._post_json(url, first).status_code, 200)
        stale = dict(first)
        stale["changes"] = [dict(first["changes"][0], note="Bản ghi đè")]
        response = self._post_json(url, stale)

        self.assertEqual(response.status_code, 409)
        self.assertEqual(ShopResponse.objects.get(error=self.error_1).note, "Bản mới")

    def test_submit_is_idempotent_and_locks_future_autosave(self):
        autosave_url = reverse("document_campaigns:autosave_batch", kwargs={"raw_token": self.raw_token})
        changes = [
            {"error_uid": str(error.error_uid), "version_no": 0, "answer_code": "PGD xác nhận lỗi"}
            for error in [self.error_1, self.error_2]
        ]
        self.assertEqual(self._post_json(autosave_url, {"changes": changes}).status_code, 200)
        submit_url = reverse("document_campaigns:submit_responses", kwargs={"raw_token": self.raw_token})

        first = self._post_json(submit_url, {"idempotency_key": "stable-submit-key"})
        second = self._post_json(submit_url, {"idempotency_key": "stable-submit-key"})
        locked = self._post_json(
            autosave_url,
            {"changes": [{"error_uid": str(self.error_1.error_uid), "version_no": 1, "answer_code": "Khác"}]},
        )

        self.assertEqual(first.status_code, 200)
        self.assertFalse(first.json()["idempotent"])
        self.assertEqual(second.status_code, 200)
        self.assertTrue(second.json()["idempotent"])
        self.assertEqual(ShopSubmission.objects.count(), 1)
        self.assertEqual(locked.status_code, 423)

    def test_submit_reports_exact_rows_missing_required_dropdown(self):
        autosave_url = reverse("document_campaigns:autosave_batch", kwargs={"raw_token": self.raw_token})
        self.assertEqual(
            self._post_json(
                autosave_url,
                {
                    "changes": [{
                        "error_uid": str(self.error_1.error_uid),
                        "version_no": 0,
                        "answer_code": "PGD xác nhận lỗi",
                        "note": "",
                    }]
                },
            ).status_code,
            200,
        )

        response = self._post_json(
            reverse("document_campaigns:submit_responses", kwargs={"raw_token": self.raw_token}),
            {"idempotency_key": "missing-dropdown-key"},
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["missing_uids"], [str(self.error_2.error_uid)])
        self.assertIn("1 dòng chưa chọn", response.json()["error"])

    @patch("app_document_campaigns.tasks.process_shop_submission.delay")
    def test_submit_enqueues_post_submit_task_only_after_commit(self, delay):
        autosave_url = reverse("document_campaigns:autosave_batch", kwargs={"raw_token": self.raw_token})
        changes = [
            {"error_uid": str(error.error_uid), "version_no": 0, "answer_code": "PGD xác nhận lỗi"}
            for error in [self.error_1, self.error_2]
        ]
        self.assertEqual(self._post_json(autosave_url, {"changes": changes}).status_code, 200)
        submit_url = reverse("document_campaigns:submit_responses", kwargs={"raw_token": self.raw_token})

        with self.captureOnCommitCallbacks(execute=False) as callbacks:
            response = self._post_json(submit_url, {"idempotency_key": "async-submit-key"})
            self.assertEqual(response.status_code, 200)
            delay.assert_not_called()

        self.assertEqual(len(callbacks), 1)
        callbacks[0]()
        submission = ShopSubmission.objects.get()
        delay.assert_called_once_with(submission.pk)

        with self.captureOnCommitCallbacks(execute=True) as duplicate_callbacks:
            duplicate = self._post_json(submit_url, {"idempotency_key": "async-submit-key"})
        self.assertTrue(duplicate.json()["idempotent"])
        self.assertEqual(duplicate_callbacks, [])
        delay.assert_called_once_with(submission.pk)

    @patch("app_document_campaigns.services.responses.logger.exception")
    @patch("app_document_campaigns.tasks.process_shop_submission.delay", side_effect=RuntimeError("broker down"))
    def test_broker_outage_does_not_break_committed_submission(self, delay, log_exception):
        autosave_url = reverse("document_campaigns:autosave_batch", kwargs={"raw_token": self.raw_token})
        changes = [
            {"error_uid": str(error.error_uid), "version_no": 0, "answer_code": "PGD xác nhận lỗi"}
            for error in [self.error_1, self.error_2]
        ]
        self.assertEqual(self._post_json(autosave_url, {"changes": changes}).status_code, 200)

        with self.captureOnCommitCallbacks(execute=True):
            response = self._post_json(
                reverse("document_campaigns:submit_responses", kwargs={"raw_token": self.raw_token}),
                {"idempotency_key": "broker-outage-key"},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(ShopSubmission.objects.count(), 1)
        delay.assert_called_once()
        log_exception.assert_called_once()

        read_only_page = self.client.get(
            reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token})
        )
        self.assertEqual(read_only_page.status_code, 200)
        self.assertNotContains(read_only_page, 'id="submit-responses"')
        self.assertEqual(read_only_page.content.count(b" disabled"), 4)

    def test_document_types_are_rendered_as_bullet_list(self):
        self.error_1.document_type_name = "Giấy đề nghị\nPhiếu chi giải ngân\nỦy quyền epay"
        self.error_1.save(update_fields=["document_type_name"])

        page = self.client.get(
            reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token})
        )

        self.assertContains(page, '<ul class="dec-document-types">')
        self.assertContains(page, "<li>Giấy đề nghị</li>")
        self.assertContains(page, "<li>Phiếu chi giải ngân</li>")
        self.assertContains(page, "<li>Ủy quyền epay</li>")

    def test_post_submit_task_sends_confirmation_without_access_token(self):
        autosave_url = reverse("document_campaigns:autosave_batch", kwargs={"raw_token": self.raw_token})
        changes = [
            {"error_uid": str(error.error_uid), "version_no": 0, "answer_code": "PGD xác nhận lỗi"}
            for error in [self.error_1, self.error_2]
        ]
        self.assertEqual(self._post_json(autosave_url, {"changes": changes}).status_code, 200)
        with self.captureOnCommitCallbacks(execute=False):
            self._post_json(
                reverse("document_campaigns:submit_responses", kwargs={"raw_token": self.raw_token}),
                {"idempotency_key": "email-submit-key"},
            )
        submission = ShopSubmission.objects.get()

        result = process_shop_submission.run(submission.pk)

        self.assertEqual(result["status"], "sent")
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, [self.shop.shop_email])
        self.assertNotIn(self.raw_token, mail.outbox[0].body)

    def test_expired_link_is_rejected(self):
        self.link.expires_at = timezone.now() - timedelta(seconds=1)
        self.link.save(update_fields=["expires_at"])
        response = self.client.get(
            reverse("document_campaigns:shop_response", kwargs={"raw_token": self.raw_token})
        )
        self.assertEqual(response.status_code, 410)
