"""Manual Excel round-trip for Step 6 QTRR error-code mapping."""
from io import BytesIO
from uuid import UUID
from zipfile import ZipFile

from django.core import signing
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app_document_campaigns.models import Campaign, CampaignError, CampaignErrorBooking
from documents import excel_snapshot


HEADERS = [
    "ID dòng lỗi", "Mã PGD", "Tên phòng giao dịch", "Mã hợp đồng", "Ngày phát sinh",
    "Loại lỗi", "Nội dung lỗi", "Tên nhân viên", "Nghiệp vụ", "QLKV xác nhận",
    "Mã lỗi QTRR", "Ghi chú mapping", "_snapshot",
]
SHEET_NAME = "Book loi QTRR"
CATALOG_SHEET = "Danh muc ma loi"
SALT = "campaign-qtrr-booking-excel-v1"
MAX_ROWS = 100000
MAX_FILE_SIZE = 50 * 1024 * 1024


def eligible_rows(campaign):
    return (
        CampaignError.objects.filter(campaign=campaign, status=CampaignError.Status.AREA_CONFIRMED)
        .select_related("shop", "risk_booking__risk_code")
        .order_by("shop__shop_code", "contract_code", "pk")
    )


def _local_date(value):
    if value and timezone.is_aware(value):
        value = timezone.localtime(value)
    return value.strftime("%d/%m/%Y") if value else ""


def _reference(error):
    return [
        str(error.error_uid), str(error.shop.shop_code), error.shop.shop_name,
        error.contract_code or error.code or "", _local_date(error.source_created_at),
        error.get_error_type_display(), error.checking_issue, error.employee_name,
        error.business_type_name, "Đã xác nhận",
    ]


def _normalized(values):
    return [str(value or "") for value in values]


def _code_label(code):
    return f"{code.code} · {code.name}"


def build_workbook(campaign):
    codes = list(campaign.risk_error_codes.all().order_by("sort_order", "code"))
    if not codes:
        raise ValueError("Kỳ chưa cấu hình mã lỗi QTRR.")
    rows = eligible_rows(campaign)
    total = rows.count()
    if not total:
        raise ValueError("Chưa có dòng QLKV xác nhận để book lỗi.")
    if total > MAX_ROWS:
        raise ValueError("Bộ dữ liệu vượt giới hạn 100.000 dòng.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    code_labels = {code.pk: _code_label(code) for code in codes}
    for index, error in enumerate(rows.iterator(chunk_size=1000), 2):
        booking = getattr(error, "risk_booking", None)
        selected = code_labels.get(booking.risk_code_id, "") if booking else ""
        note = booking.note if booking else ""
        reference = _reference(error)
        token = excel_snapshot.dumps(
            {
                "v": 1,
                "c": campaign.pk,
                "u": error.error_uid.hex,
                "r": excel_snapshot.digest(_normalized(reference)),
                "b": booking.pk if booking else None,
                "e": excel_snapshot.digest([selected, note]),
            },
            salt=SALT,
        )
        sheet.append(reference + [selected, note, token])
        for column in range(1, len(HEADERS) + 1):
            sheet.cell(index, column).data_type = "s"
        sheet.cell(index, 11).protection = Protection(locked=False)
        sheet.cell(index, 12).protection = Protection(locked=False)

    catalog = workbook.create_sheet(CATALOG_SHEET)
    catalog.append(["Mã lỗi được phép sử dụng"])
    for code in codes:
        catalog.append([_code_label(code)])
    catalog.sheet_state = "hidden"
    dropdown = DataValidation(
        type="list",
        formula1=f"'{CATALOG_SHEET}'!$A$2:$A${len(codes) + 1}",
        allow_blank=True,
    )
    dropdown.showErrorMessage = True
    dropdown.errorStyle = "stop"
    dropdown.errorTitle = "Mã lỗi không hợp lệ"
    dropdown.error = "Chọn mã lỗi trong danh sách của kỳ."
    sheet.add_data_validation(dropdown)
    dropdown.add(f"K2:K{max(total + 1, 2)}")
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = f"B1:L{total + 1}"
    sheet.column_dimensions["A"].hidden = True
    sheet.column_dimensions["M"].hidden = True
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    widths = {"B": 14, "C": 28, "D": 20, "E": 14, "F": 28, "G": 52, "H": 24, "I": 28, "J": 18, "K": 48, "L": 38}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="00844A")

    guide = workbook.create_sheet("Hướng dẫn")
    for text in [
        f"{campaign.code} · {campaign.name}",
        "Chỉ sửa cột Mã lỗi QTRR và Ghi chú mapping.",
        "Cột Mã lỗi QTRR có dropdown theo danh sách đã cấu hình cho kỳ.",
        "Dòng để trống mã lỗi sẽ không làm thay đổi dữ liệu hệ thống.",
        "Xóa dòng khỏi file không xóa dữ liệu trên hệ thống.",
        "Không sửa ID, dữ liệu nguồn hoặc cột _snapshot.",
    ]:
        guide.append([text])
    guide.column_dimensions["A"].width = 120

    output = BytesIO()
    workbook.save(output)
    content = output.getvalue()
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("File xuất vượt 50 MB.")
    return content, total


def import_workbook(campaign_id, actor, content):
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("File tối đa 50 MB.")
    with ZipFile(BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 512 * 1024 * 1024:
            raise ValueError("File Excel giải nén quá lớn.")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError("Sai template Step 6. Hãy tải lại file từ hệ thống.")
        sheet = workbook[SHEET_NAME]
        rows = sheet.iter_rows(values_only=True)
        if list(next(rows, ())) != HEADERS:
            raise ValueError("Sai cấu trúc template Step 6. Hãy tải lại file từ hệ thống.")
        changes, seen = [], set()
        for line, row in enumerate(rows, 2):
            if all(value is None for value in row):
                continue
            if len(row) != len(HEADERS):
                raise ValueError(f"Dòng {line}: sai số cột.")
            try:
                snapshot = excel_snapshot.loads(row[-1], salt=SALT)
                uid = str(UUID(str(row[0])))
                if snapshot.get("v") != 1 or snapshot["c"] != campaign_id or str(UUID(hex=snapshot["u"])) != uid or uid in seen:
                    raise ValueError("Sai kỳ, sai ID hoặc trùng dòng lỗi.")
                seen.add(uid)
                if excel_snapshot.digest(_normalized(row[:-3])) != snapshot["r"]:
                    raise ValueError("Dữ liệu nguồn bị sửa; chỉ sửa hai cột mapping.")
                selected, note = str(row[-3] or "").strip(), str(row[-2] or "").strip()
                if excel_snapshot.digest([selected, note]) == snapshot["e"]:
                    continue
                if not selected:
                    raise ValueError("Phải chọn Mã lỗi QTRR khi thay đổi dòng.")
                if len(note) > 2000 or note.startswith("="):
                    raise ValueError("Ghi chú tối đa 2.000 ký tự và không dùng công thức.")
                changes.append((line, uid, selected, note, snapshot))
            except (signing.BadSignature, ValueError, KeyError, TypeError) as exc:
                raise ValueError(f"Dòng {line}: {exc}") from exc
    finally:
        workbook.close()

    with transaction.atomic():
        campaign = Campaign.objects.select_for_update().get(pk=campaign_id)
        if not actor.is_active or not (actor.is_superuser or actor.groups.filter(name="admin").exists()):
            raise ValueError("Tài khoản không còn quyền Admin.")
        codes = {_code_label(code): code for code in campaign.risk_error_codes.all()}
        errors = {
            str(error.error_uid): error
            for error in eligible_rows(campaign).select_for_update(of=("self",)).filter(error_uid__in=[item[1] for item in changes])
        }
        updated = 0
        for line, uid, selected, note, snapshot in changes:
            error = errors.get(uid)
            if not error:
                raise ValueError(f"Dòng {line}: dữ liệu không còn đủ điều kiện book lỗi.")
            code = codes.get(selected)
            if not code:
                raise ValueError(f"Dòng {line}: mã lỗi không thuộc danh sách cấu hình của kỳ.")
            booking = getattr(error, "risk_booking", None)
            if (booking.pk if booking else None) != snapshot["b"]:
                raise ValueError(f"Dòng {line}: mapping đã thay đổi; hãy xuất lại file.")
            if excel_snapshot.digest(_normalized(_reference(error))) != snapshot["r"]:
                raise ValueError(f"Dòng {line}: dữ liệu nguồn đã thay đổi; hãy xuất lại file.")
            CampaignErrorBooking.objects.update_or_create(
                error=error,
                defaults={"risk_code": code, "note": note, "mapped_by": actor},
            )
            updated += 1
    return {"updated": updated, "rows": len(seen)}
