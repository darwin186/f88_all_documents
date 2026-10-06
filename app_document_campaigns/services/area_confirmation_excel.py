"""Round-trip Excel for Step 5 QLKV decisions."""
from io import BytesIO
from uuid import UUID
from zipfile import ZipFile

from django.core import signing
from django.db import transaction
from django.db.models import OuterRef, Subquery
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app_document_campaigns.models import AreaConfirmation, Campaign, CampaignError, TeamReview
from documents import excel_snapshot


HEADERS = [
    "ID dòng lỗi", "Mã QLKV", "Tên QLKV", "Mã PGD", "Tên phòng giao dịch",
    "Mã hợp đồng", "Ngày phát sinh", "Loại lỗi", "Nội dung lỗi",
    "Phòng vận hành xác nhận", "Nhận xét Phòng vận hành",
    "QLKV xác nhận", "Ghi chú QLKV", "_snapshot",
]
LEGACY_DECISIONS = {
    "Xác nhận lỗi": True,
    "Gỡ lỗi": False,
    "Đồng thuận": True,
    "Không đồng thuận": False,
}
SALT = "campaign-area-confirmation-excel-v1"
MAX_ROWS = 100000
MAX_FILE_SIZE = 50 * 1024 * 1024
STEP5_STATUSES = [
    CampaignError.Status.WAITING_AREA,
    CampaignError.Status.AREA_CONFIRMED,
    CampaignError.Status.AREA_RETURNED,
]


def step5_errors(campaign):
    latest_review = TeamReview.objects.filter(error_id=OuterRef("pk")).order_by("-created_at", "-pk")
    return (
        CampaignError.objects.filter(campaign=campaign, status__in=STEP5_STATUSES)
        .annotate(latest_team_decision=Subquery(latest_review.values("decision")[:1]))
        .filter(latest_team_decision=TeamReview.Decision.APPROVED)
    )


def decision_options(_campaign=None):
    return {"Đồng thuận": True, "Không đồng thuận": False}


def normalize_decision(value):
    if value is True or value in {"true", "confirmed"}:
        return True
    if value is False or value in {"false", "returned"}:
        return False
    raise ValueError("Kết luận QLKV trong file không hợp lệ.")


def _local_date(value):
    if value and timezone.is_aware(value):
        value = timezone.localtime(value)
    return value.strftime("%d/%m/%Y") if value else ""


def reference_values(error):
    area = error.shop.manager_id.areaManager if error.shop.manager_id_id else None
    review = error.latest_team_review
    return [
        str(error.error_uid),
        str(area.areaManager_code if area else ""),
        area.areaManager_name if area else "",
        str(error.shop.shop_code),
        error.shop.shop_name,
        error.contract_code or error.code or "",
        _local_date(error.source_created_at),
        error.get_error_type_display(),
        error.checking_issue,
        "Giữ lỗi",
        review.note if review else "",
    ]


def normalized(values):
    return [str(value or "") for value in values]


def unpack_snapshot(snapshot):
    if snapshot.get("v") != 3 or "c" not in snapshot:
        return snapshot
    return {
        "campaign": snapshot["c"],
        "uid": str(UUID(snapshot["u"])),
        "reference_hash": snapshot["r"],
        "review_id": snapshot["i"],
        "confirmation_id": snapshot["a"],
        "decision_options": snapshot["o"],
        "editable_hash": snapshot["e"],
    }


def _rows(campaign):
    latest_review = TeamReview.objects.filter(error_id=OuterRef("pk")).order_by("-created_at", "-pk")
    return (
        step5_errors(campaign)
        .annotate(latest_review_id=Subquery(latest_review.values("pk")[:1]))
        .select_related("shop", "shop__manager_id__areaManager")
        .order_by("shop__manager_id__areaManager__areaManager_name", "shop__shop_code", "id")
    )


def build_workbook(campaign, progress=None):
    rows = _rows(campaign)
    total = rows.count()
    if total > MAX_ROWS:
        raise ValueError("Bộ dữ liệu vượt giới hạn 100.000 dòng.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QLKV xac nhan"
    sheet.append(HEADERS)
    decisions = decision_options(campaign)
    reverse_decisions = {value: label for label, value in decisions.items()}
    for index, error in enumerate(rows.iterator(chunk_size=1000), 1):
        review = error.team_reviews.order_by("-created_at", "-pk").first()
        error.latest_team_review = review
        confirmation = error.area_confirmations.order_by("-created_at", "-pk").first()
        reference = reference_values(error)
        decision = reverse_decisions.get(confirmation.is_agreed, "") if confirmation else ""
        note = confirmation.note if confirmation else ""
        token = excel_snapshot.dumps(
            {
                "v": 3,
                "c": campaign.pk,
                "u": error.error_uid.hex,
                "r": excel_snapshot.digest(normalized(reference)),
                "i": review.pk if review else None,
                "a": confirmation.pk if confirmation else None,
                "o": decisions,
                "e": excel_snapshot.digest([decision, note]),
            },
            salt=SALT,
        )
        sheet.append(reference + [decision, note, token])
        row_number = index + 1
        for col in range(1, len(HEADERS) + 1):
            sheet.cell(row_number, col).data_type = "s"
        for col in (len(HEADERS) - 2, len(HEADERS) - 1):
            sheet.cell(row_number, col).protection = Protection(locked=False)
        if progress and index % 1000 == 0:
            progress(round(index * 80 / max(total, 1)))
    excel_choices = ",".join(decisions).replace('"', '""')
    dropdown = DataValidation(type="list", formula1=f'"{excel_choices}"', allow_blank=True)
    dropdown.showErrorMessage = True
    dropdown.errorStyle = "stop"
    dropdown.errorTitle = "Kết luận không hợp lệ"
    dropdown.error = "Chọn Đồng thuận hoặc Không đồng thuận trong danh sách."
    sheet.add_data_validation(dropdown)
    decision_column = get_column_letter(len(HEADERS) - 2)
    dropdown.add(f"{decision_column}2:{decision_column}{max(total + 1, 2)}")
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = f"B1:{get_column_letter(len(HEADERS) - 1)}{total + 1}"
    sheet.column_dimensions["A"].hidden = True
    sheet.column_dimensions[get_column_letter(len(HEADERS))].hidden = True
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    for col in range(2, len(HEADERS)):
        sheet.column_dimensions[get_column_letter(col)].width = 24
    for col in (5, 9, 11, 13):
        sheet.column_dimensions[get_column_letter(col)].width = 42
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="00844A")
    guide = workbook.create_sheet("Hướng dẫn")
    for text in [
        f"{campaign.code} · {campaign.name}",
        "Chỉ sửa QLKV xác nhận (droplist) và Ghi chú QLKV.",
        "Droplist gồm Đồng thuận / Không đồng thuận.",
        "Xóa dòng khỏi file không xóa dữ liệu trên hệ thống.",
        "Không sửa dữ liệu nguồn, ID hoặc cột _snapshot.",
    ]:
        guide.append([text])
    guide.column_dimensions["A"].width = 120
    stream = BytesIO()
    workbook.save(stream)
    content = stream.getvalue()
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("File xuất vượt 50 MB.")
    return content, total


def import_workbook(campaign_id, actor, content, progress=None):
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("File tối đa 50 MB.")
    with ZipFile(BytesIO(content)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 512 * 1024 * 1024:
            raise ValueError("File Excel giải nén quá lớn.")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = workbook["QLKV xac nhan"]
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        if headers != HEADERS:
            raise ValueError("Sai template Step 5. Hãy tải lại file từ hệ thống.")
        current_decisions = decision_options()
        changes, seen = [], set()
        for line, row in enumerate(rows, 2):
            if all(value is None for value in row):
                continue
            if len(row) != len(HEADERS):
                raise ValueError(f"Dòng {line}: sai số cột.")
            try:
                snapshot = unpack_snapshot(excel_snapshot.loads(row[-1], salt=SALT))
                uid = str(UUID(str(row[0])))
                if snapshot["campaign"] != campaign_id or snapshot["uid"] != uid or uid in seen:
                    raise ValueError("Sai kỳ, sai ID hoặc trùng dòng lỗi.")
                seen.add(uid)
                file_reference = normalized(row[:-3])
                reference_matches = (
                    excel_snapshot.digest(file_reference) == snapshot["reference_hash"]
                    if "reference_hash" in snapshot
                    else file_reference == snapshot["reference"]
                )
                if not reference_matches:
                    raise ValueError("Dữ liệu nguồn bị sửa; chỉ sửa hai cột QLKV.")
                decision, note = str(row[-3] or "").strip(), str(row[-2] or "").strip()
                editable_unchanged = (
                    excel_snapshot.digest([decision, note]) == snapshot["editable_hash"]
                    if "editable_hash" in snapshot
                    else decision == snapshot["decision"] and note == snapshot["note"]
                )
                if editable_unchanged:
                    continue
                allowed_decisions = snapshot.get("decision_options") or {**LEGACY_DECISIONS, **current_decisions}
                if decision not in allowed_decisions:
                    raise ValueError(
                        "Chọn Đồng thuận hoặc Không đồng thuận; không được xóa phản hồi đã có."
                    )
                if len(note) > 2000 or note.startswith("="):
                    raise ValueError("Ghi chú tối đa 2.000 ký tự và không dùng công thức.")
                changes.append(
                    (line, uid, normalize_decision(allowed_decisions[decision]), note, snapshot)
                )
            except (signing.BadSignature, ValueError, KeyError, TypeError) as exc:
                raise ValueError(f"Dòng {line}: {exc}") from exc
            if progress and line % 1000 == 0:
                progress(min(50, round(line * 50 / max(sheet.max_row, 1))))
    finally:
        workbook.close()

    with transaction.atomic():
        campaign = Campaign.objects.select_for_update().get(pk=campaign_id)
        if not actor.is_active or not (actor.is_superuser or actor.groups.filter(name="admin").exists()):
            raise ValueError("Tài khoản không còn quyền Admin.")
        uids = [item[1] for item in changes]
        errors = {
            str(error.error_uid): error
            for error in _rows(campaign).select_for_update(of=("self",)).filter(error_uid__in=uids)
        }
        updated = 0
        for line, uid, decision, note, snapshot in changes:
            error = errors.get(uid)
            if not error:
                raise ValueError(f"Dòng {line}: dữ liệu không còn ở Step 5.")
            review = error.team_reviews.order_by("-created_at", "-pk").first()
            confirmation = error.area_confirmations.order_by("-created_at", "-pk").first()
            if (review.pk if review else None) != snapshot["review_id"] or (confirmation.pk if confirmation else None) != snapshot["confirmation_id"]:
                raise ValueError(f"Dòng {line}: dữ liệu đã thay đổi; hãy xuất lại file.")
            if "reference_hash" in snapshot:
                error.latest_team_review = review
                if excel_snapshot.digest(normalized(reference_values(error))) != snapshot["reference_hash"]:
                    raise ValueError(f"Dòng {line}: dữ liệu đã thay đổi; hãy xuất lại file.")
            area = error.shop.manager_id.areaManager if error.shop.manager_id_id else None
            if not area:
                raise ValueError(f"Dòng {line}: PGD chưa có QLKV trong Master Data.")
            AreaConfirmation.objects.create(
                error=error,
                area_manager=area,
                is_agreed=decision,
                note=note,
                confirmed_by=actor,
            )
            error.status = (
                CampaignError.Status.AREA_CONFIRMED
                if decision
                else CampaignError.Status.AREA_RETURNED
            )
            error.save(update_fields=["status", "updated_at"])
            updated += 1
    return {"updated": updated, "rows": len(seen)}
