import hashlib
from io import BytesIO

from django.db import transaction
from openpyxl import load_workbook

from app_document_campaigns.models import CampaignImportSource, CampaignVersion
from app_document_campaigns.services.excel_exports import COLUMNS, WORKBOOK_SCHEMA_VERSION
from app_document_campaigns.services.imports import CampaignImportError, stage_import_rows


EXCEL_SOURCE_NAME = "team-cleaning-excel"
MAX_WORKBOOK_SIZE = 10 * 1024 * 1024
MAX_WORKBOOK_ROWS = 100_000


def _plain_value(cell):
    if cell.data_type == "f":
        raise CampaignImportError(f"Không chấp nhận công thức Excel tại ô {cell.coordinate}.")
    value = cell.value
    if isinstance(value, str) and len(value) > 1 and value[0] == "'" and value[1] in "=+-@":
        return value[1:]
    return value


def _read_metadata(workbook):
    if "_Metadata" not in workbook.sheetnames:
        raise CampaignImportError("File thiếu sheet _Metadata.")
    values = {}
    for key_cell, value_cell in workbook["_Metadata"].iter_rows(min_col=1, max_col=2):
        if key_cell.value:
            values[str(key_cell.value)] = value_cell.value
    return values


def _validate_metadata(metadata, version):
    expected = {
        "schema_version": WORKBOOK_SCHEMA_VERSION,
        "campaign_code": version.campaign.code,
        "campaign_version": version.version_number,
        "report_month": version.campaign.report_month.isoformat(),
    }
    for key, expected_value in expected.items():
        if str(metadata.get(key, "")) != str(expected_value):
            raise CampaignImportError(f"Metadata {key} không đúng với campaign version đang import.")


def _technical_key_errors(row, baseline):
    source_key = str(row.get("source_key") or "").strip()
    error_type = str(row.get("error_type") or "").strip()
    source_object_id = str(row.get("source_object_id") or "").strip()
    original = baseline.get(source_key)
    if original and (
        str(original.get("error_type")) != error_type
        or str(original.get("source_object_id")) != source_object_id
    ):
        return [{"field": "source_key", "code": "technical_key_changed"}]
    if original is None and source_key != f"{error_type}:{source_object_id}":
        return [{"field": "source_key", "code": "technical_key_changed"}]
    return []


def _parse_data_rows(workbook, version):
    if "Lỗi chứng từ" not in workbook.sheetnames:
        raise CampaignImportError("File thiếu sheet Lỗi chứng từ.")
    sheet = workbook["Lỗi chứng từ"]
    if sheet.max_row > MAX_WORKBOOK_ROWS + 1:
        raise CampaignImportError(f"File vượt quá giới hạn {MAX_WORKBOOK_ROWS} dòng.")

    label_to_field = {label: field for field, label, _, _ in COLUMNS}
    headers = [cell.value for cell in sheet[1]]
    if len(headers) != len(set(headers)):
        raise CampaignImportError("File có tên cột bị trùng.")
    missing = [label for label in label_to_field if label not in headers]
    if missing:
        raise CampaignImportError("File thiếu cột bắt buộc: " + ", ".join(missing))
    positions = {index: label_to_field[label] for index, label in enumerate(headers) if label in label_to_field}

    baseline = {
        row.source_key: row.normalized_payload
        for source in version.import_sources.filter(source_type=CampaignVersion.SourceType.SQL)
        for row in source.rows.filter(validation_errors=[])
    }
    parsed = []
    removed_formula_error = None
    for cells in sheet.iter_rows(min_row=2):
        if not any(cell.value not in (None, "") for cell in cells):
            continue
        try:
            row = {field: _plain_value(cells[index]) for index, field in positions.items()}
        except CampaignImportError as exc:
            removed_formula_error = exc
            break
        key_errors = _technical_key_errors(row, baseline)
        if key_errors:
            row["_validation_errors"] = key_errors
        parsed.append(row)
    if removed_formula_error:
        raise removed_formula_error
    return parsed, baseline


@transaction.atomic
def import_cleaning_workbook(*, version, created_by, uploaded_file, progress_callback=None):
    version = CampaignVersion.objects.select_for_update().select_related("campaign").get(pk=version.pk)
    filename = str(getattr(uploaded_file, "name", "upload.xlsx")).replace("\\", "/").split("/")[-1]
    if not filename.lower().endswith(".xlsx"):
        raise CampaignImportError("Chỉ chấp nhận file .xlsx không chứa macro.")
    content = uploaded_file.read()
    if not content or len(content) > MAX_WORKBOOK_SIZE:
        raise CampaignImportError("File Excel rỗng hoặc vượt quá giới hạn 10 MB.")
    try:
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=False)
    except Exception as exc:
        raise CampaignImportError("Không đọc được cấu trúc file Excel.") from exc

    metadata = _read_metadata(workbook)
    _validate_metadata(metadata, version)
    if progress_callback:
        progress_callback("Đã kiểm tra cấu trúc và thông tin chiến dịch trong file.", 1)
    rows, baseline = _parse_data_rows(workbook, version)
    if progress_callback:
        progress_callback(f"Đã đọc {len(rows)} dòng từ file Excel.", 2)
    source, _ = CampaignImportSource.objects.get_or_create(
        version=version,
        name=EXCEL_SOURCE_NAME,
        defaults={
            "source_type": CampaignVersion.SourceType.EXCEL,
            "created_by": created_by,
        },
    )
    if source.source_type != CampaignVersion.SourceType.EXCEL:
        raise CampaignImportError("Source team-cleaning-excel không phải nguồn Excel.")
    summary = stage_import_rows(source=source, rows=rows)
    excel_keys = {
        str(row.get("source_key") or "").strip()
        for row in rows
        if not row.get("_validation_errors")
    }
    summary["removed_from_sql"] = len(set(baseline) - excel_keys)
    source.source_filename = filename
    source.source_checksum = hashlib.sha256(content).hexdigest()
    source.preview_summary = summary
    source.save(
        update_fields=["source_filename", "source_checksum", "preview_summary", "updated_at"]
    )
    return source, summary
