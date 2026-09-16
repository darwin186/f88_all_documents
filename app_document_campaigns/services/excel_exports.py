from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app_document_campaigns.models import CampaignImportSource, CampaignStagingRow


WORKBOOK_SCHEMA_VERSION = "dec-cleaning-v1"
HEADER_FILL = PatternFill("solid", fgColor="00844A")
HEADER_FONT = Font(color="FFFFFF", bold=True)

COLUMNS = (
    ("source_key", "source_key", 28, True),
    ("error_type", "error_type", 14, True),
    ("source_object_id", "source_object_id", 18, True),
    ("code", "Mã quyển/chứng từ", 24, False),
    ("source_created_at", "Ngày phát sinh", 20, False),
    ("shop_code", "Mã PGD", 12, False),
    ("shop_name", "Tên phòng giao dịch", 28, False),
    ("shop_email", "Email PGD", 28, False),
    ("qlkv_name", "QLKV thời điểm phát sinh", 26, False),
    ("qlkv_email", "Email QLKV", 28, False),
    ("qlv_name", "QLV thời điểm phát sinh", 26, False),
    ("qlv_email", "Email QLV", 28, False),
    ("region_name", "Vùng", 18, False),
    ("contract_code", "Mã hợp đồng", 22, False),
    ("customer_code", "Mã khách hàng", 18, False),
    ("customer_name", "Tên khách hàng", 28, False),
    ("employee_code", "Mã nhân viên", 18, False),
    ("employee_name", "Tên nhân viên", 26, False),
    ("business_type_name", "Nghiệp vụ", 30, False),
    ("document_type_name", "Bộ chứng từ", 34, False),
    ("checking_issue", "Lỗi", 48, False),
)


class CampaignExcelExportError(Exception):
    pass


def _excel_safe(value):
    if value is None:
        return ""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _selected_sql_rows(version):
    sources = list(
        version.import_sources.filter(source_type="sql").order_by("id")
    )
    if not sources:
        raise CampaignExcelExportError("Version chưa có dữ liệu SQL trong staging.")
    invalid_count = sum(source.invalid_count for source in sources)
    if invalid_count:
        raise CampaignExcelExportError(
            f"Staging còn {invalid_count} dòng lỗi; cần xử lý trước khi xuất Excel."
        )
    rows = list(
        CampaignStagingRow.objects.filter(source__in=sources, validation_errors=[]).order_by(
            "source_key", "row_number"
        )
    )
    duplicate_keys = []
    seen = set()
    for row in rows:
        if row.source_key in seen:
            duplicate_keys.append(row.source_key)
        seen.add(row.source_key)
    if duplicate_keys:
        raise CampaignExcelExportError(
            "Có source_key trùng giữa các nguồn SQL: " + ", ".join(sorted(set(duplicate_keys))[:5])
        )
    return sources, rows


def build_cleaning_workbook(version, *, reviewed=False):
    if reviewed:
        source = version.import_sources.filter(source_type="excel", name="team-cleaning-excel").first()
        if not source or source.invalid_count or source.status not in ("validated", "confirmed"):
            raise CampaignExcelExportError("Chưa có file team review hợp lệ để xuất.")
        sources = [source]
        rows = list(source.rows.filter(validation_errors=[]).order_by("row_number"))
    else:
        sources, rows = _selected_sql_rows(version)
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Lỗi chứng từ"
    data_sheet.freeze_panes = "A2"
    data_sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    for column_number, (_, label, width, hidden) in enumerate(COLUMNS, start=1):
        cell = data_sheet.cell(row=1, column=column_number, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        column = data_sheet.column_dimensions[get_column_letter(column_number)]
        column.width = width
        column.hidden = hidden
    data_sheet.row_dimensions[1].height = 34

    for row_number, staging_row in enumerate(rows, start=2):
        payload = staging_row.normalized_payload
        manager = payload.get("manager_snapshot") or {}
        for column_number, (field, _, _, _) in enumerate(COLUMNS, start=1):
            value = payload.get(field, manager.get(field, ""))
            cell = data_sheet.cell(row=row_number, column=column_number, value=_excel_safe(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    guide = workbook.create_sheet("Hướng dẫn")
    guide.append(["HƯỚNG DẪN LÀM SẠCH DỮ LIỆU"])
    guide["A1"].fill = HEADER_FILL
    guide["A1"].font = HEADER_FONT
    guide.append(["1. Kiểm tra thông tin PGD, hợp đồng, khách hàng và nội dung lỗi."])
    guide.append(["2. Có thể xóa dòng không đưa vào chiến dịch hoặc sửa trường nghiệp vụ cần thiết."])
    guide.append(["3. Không bỏ ẩn hoặc thay đổi ba cột kỹ thuật đầu tiên."])
    guide.append(["4. Giữ nguyên tên sheet và định dạng .xlsx khi gửi lại hệ thống."])
    guide.column_dimensions["A"].width = 110

    metadata = workbook.create_sheet("_Metadata")
    metadata.append(["schema_version", WORKBOOK_SCHEMA_VERSION])
    metadata.append(["campaign_code", version.campaign.code])
    metadata.append(["campaign_version", version.version_number])
    metadata.append(["report_month", version.campaign.report_month.isoformat()])
    metadata.append(["row_count", len(rows)])
    metadata.append(["sources", ",".join(source.name for source in sources)])
    metadata.append(["export_kind", "reviewed" if reviewed else "sql"])
    metadata.sheet_state = "veryHidden"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
