from app_document_campaigns.services.response_deadlines import shop_deadlines, shop_deadline_passed
"""Round-trip review decisions only; immutable data is signed per row."""
from io import BytesIO
from uuid import UUID
from zipfile import ZipFile

from django.core import signing
from django.db import transaction
from django.db.models import Exists, OuterRef
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from app_document_campaigns.models import Campaign, CampaignError, ShopSubmission, TeamReview
from .folder_receipt import with_folder_receipt, receipt_values

HEADERS = ["ID dòng lỗi", "Mã PGD", "Tên phòng giao dịch", "Mã hợp đồng", "Ngày phát sinh", "Loại lỗi", "Tên nhân viên", "Lỗi", "Nghiệp vụ", "Chứng từ", "PGD phản hồi", "Ghi chú PGD", "Kết luận review", "Nhận xét team", "_snapshot"]
LEGACY_HEADERS = list(HEADERS)
HEADERS = HEADERS[:12] + ["Trạng thái nhận quyển", "Ngày nhận quyển"] + HEADERS[12:]
DECISIONS = {"Giữ lỗi": "approved", "Loại lỗi": "excluded"}
SALT = "campaign-team-review-excel-v1"
MAX_ROWS = 100000
MAX_FILE_SIZE = 50 * 1024 * 1024


def eligible_errors(campaign):
    history = TeamReview.objects.filter(error_id=OuterRef("pk"))
    return CampaignError.objects.filter(campaign=campaign).exclude(status="cancelled").annotate(has_review=Exists(history)).exclude(status="excluded", has_review=False)


def reference_values(error, labels):
    response = getattr(error, "shop_response", None)
    occurred_at = error.source_created_at
    # USE_TZ=False yields naive values already representing local wall time.
    if occurred_at and timezone.is_aware(occurred_at):
        occurred_at = timezone.localtime(occurred_at)
    occurred_date = occurred_at.strftime("%d/%m/%Y") if occurred_at else ""
    return [str(error.error_uid), str(error.shop.shop_code), error.shop.shop_name, error.contract_code or error.code or "", occurred_date, error.get_error_type_display(), error.employee_name, error.checking_issue, error.business_type_name, error.document_type_name, labels.get(response.answer_code, response.answer_code) if response else "", response.note if response else ""] + receipt_values(error)


def build_workbook(campaign, progress=None):
    from app_document_campaigns.review_views import review_queryset

    errors = review_queryset(campaign).order_by("shop__shop_code", "id")
    total = errors.count()
    if total > MAX_ROWS:
        raise ValueError("Bộ dữ liệu vượt giới hạn 100.000 dòng.")
    labels = dict(campaign.response_options.values_list("value", "label"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Team review"
    sheet.append(HEADERS)
    reverse_decisions = {value: label for label, value in DECISIONS.items()}
    for index, error in enumerate(errors.iterator(chunk_size=1000), 1):
        values = reference_values(error, labels)
        decision = reverse_decisions.get(error.review_decision, "")
        note = error.review_note or ""
        response = getattr(error, "shop_response", None)
        token = signing.dumps({"campaign": campaign.pk, "uid": str(error.error_uid), "reference": values, "review_id": error.review_id, "decision": decision, "note": note, "response_version": response.version_no if response else 0}, salt=SALT, compress=True)
        sheet.append(values + [decision, note, token])
        row_number = index + 1
        for col in range(1, len(HEADERS) + 1):
            cell = sheet.cell(row_number, col)
            cell.data_type = "s"  # Never execute user text as an Excel formula.
        for col in (15, 16):
            sheet.cell(row_number, col).protection = Protection(locked=False)
        if progress and index % 1000 == 0:
            progress(round(index * 80 / max(total, 1)))
    dropdown = DataValidation(type="list", formula1='"Giữ lỗi,Loại lỗi"', allow_blank=True)
    dropdown.showDropDown = False
    dropdown.showErrorMessage = True
    dropdown.errorStyle = "stop"
    dropdown.errorTitle = "Kết luận không hợp lệ"
    dropdown.error = "Chọn Giữ lỗi hoặc Loại lỗi trong danh sách."
    sheet.add_data_validation(dropdown)
    dropdown.add(f"O2:O{max(total + 1, 2)}")
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = f"B1:P{total + 1}"
    sheet.column_dimensions["A"].hidden = True
    sheet.column_dimensions["Q"].hidden = True
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    for col in ("B", "E", "F", "G", "I", "M", "N", "O"):
        sheet.column_dimensions[col].width = 24
    for col in ("C", "H", "J", "K", "L", "P"):
        sheet.column_dimensions[col].width = 45
    sheet.column_dimensions["D"].width = 22
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="00844A")
    guide = workbook.create_sheet("Hướng dẫn")
    for text in [f"{campaign.code} · {campaign.name}", "Chỉ sửa Kết luận review (droplist) và Nhận xét team (không bắt buộc).", "Dòng để trống kết luận và không thay đổi được giữ nguyên. Xóa bớt dòng không xóa lỗi.", "Không sửa ID, dữ liệu PGD hoặc cột _snapshot; file chỉ dùng cho đúng kỳ đã xuất.", "Import có lỗi sẽ không cập nhật dòng nào. Nếu dữ liệu đã thay đổi trên web, xuất lại file.", "Chỉ áp dụng khi PGD đã gửi chính thức hoặc hết hạn; không sửa phản hồi PGD."]:
        guide.append([text])
    guide.column_dimensions["A"].width = 120
    stream = BytesIO()
    workbook.save(stream)
    content = stream.getvalue()
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("File xuất vượt 50 MB; cần chia bộ dữ liệu trước khi xuất.")
    return content, total


def import_workbook(campaign_id, actor, content, progress=None):
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("File tối đa 50 MB.")
    with ZipFile(BytesIO(content)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 512 * 1024 * 1024:
            raise ValueError("File Excel giải nén quá lớn.")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = workbook["Team review"]
        if sheet.max_row > MAX_ROWS + 1:
            raise ValueError("File vượt giới hạn 100.000 dòng.")
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        if headers not in (HEADERS, LEGACY_HEADERS):
            raise ValueError("Sai template. Hãy tải file từ màn hình Team review.")
        reference_count = len(headers) - 3
        changes, seen = [], set()
        for line, row in enumerate(rows, 2):
            if all(value is None for value in row):
                continue
            if len(row) != len(headers):
                raise ValueError(f"Dòng {line}: sai số cột.")
            try:
                snapshot = signing.loads(str(row[-1] or ""), salt=SALT)
                uid = str(UUID(str(row[0])))
                if snapshot["campaign"] != campaign_id or snapshot["uid"] != uid or uid in seen:
                    raise ValueError("Sai kỳ, sai ID hoặc trùng dòng lỗi.")
                seen.add(uid)
                if [str(value or "") for value in row[:reference_count]] != snapshot["reference"]:
                    raise ValueError("Dữ liệu gốc/PGD bị sửa; chỉ sửa hai cột review.")
                decision, note = str(row[-3] or "").strip(), str(row[-2] or "").strip()
                if decision == snapshot["decision"] and note == snapshot["note"]:
                    continue
                if decision not in DECISIONS:
                    raise ValueError("Chọn Giữ lỗi hoặc Loại lỗi; không được xóa kết luận đã có.")
                if len(note) > 2000 or note.startswith("="):
                    raise ValueError("Nhận xét tối đa 2.000 ký tự và không dùng công thức.")
                changes.append((line, uid, DECISIONS[decision], note, snapshot))
            except (signing.BadSignature, ValueError, KeyError, TypeError) as exc:
                raise ValueError(f"Dòng {line}: {exc}") from exc
            if progress and line % 1000 == 0:
                progress(min(50, round(line * 50 / max(sheet.max_row, 1))))
    finally:
        workbook.close()
    with transaction.atomic():
        campaign = Campaign.objects.select_for_update().get(pk=campaign_id)
        if not actor.is_active or not (actor.is_superuser or actor.groups.filter(name="admin").exists()):
            raise ValueError("Tài khoản không còn quyền admin.")
        if campaign.status != Campaign.Status.ACTIVE:
            raise ValueError("Kỳ không ở giai đoạn team review.")
        deadlines = shop_deadlines(campaign)
        submitted = set(ShopSubmission.objects.filter(campaign=campaign).values_list("shop_id", flat=True))
        labels = dict(campaign.response_options.values_list("value", "label"))
        # Chunk IN queries so large workbooks also work on SQLite/local environments.
        errors = {}
        uids = [change[1] for change in changes]
        for start in range(0, len(uids), 500):
            errors.update({str(error.error_uid): error for error in with_folder_receipt(eligible_errors(campaign)).select_for_update(of=("self",)).select_related("shop", "shop_response").filter(error_uid__in=uids[start:start + 500])})
        latest_ids = {}
        for start in range(0, len(uids), 500):
            for review in TeamReview.objects.filter(error__error_uid__in=uids[start:start + 500], error__campaign=campaign).order_by("error_id", "-created_at", "-pk"):
                latest_ids.setdefault(review.error_id, review.pk)
        pending = []
        for line, uid, decision, note, snapshot in changes:
            error = errors.get(uid)
            if not error or (not shop_deadline_passed(campaign, error.shop_id, deadlines) and error.shop_id not in submitted):
                raise ValueError(f"Dòng {line}: lỗi không còn khả dụng hoặc PGD chưa gửi/hết hạn.")
            response = getattr(error, "shop_response", None)
            if latest_ids.get(error.pk) != snapshot["review_id"] or reference_values(error, labels)[:len(snapshot["reference"])] != snapshot["reference"] or (response.version_no if response else 0) != snapshot["response_version"]:
                raise ValueError(f"Dòng {line}: dữ liệu đã thay đổi; hãy xuất lại file.")
            pending.append(TeamReview(error=error, decision=decision, note=note, reviewed_by=actor))
        TeamReview.objects.bulk_create(pending, batch_size=500)
    return {"updated": len(pending), "rows": len(seen)}
