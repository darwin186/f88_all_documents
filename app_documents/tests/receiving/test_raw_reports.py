from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace as NS
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from app_documents.models import (
    BusinessType, ContractDetail, DocumentsDetail, DocumentType, Folder, FolderStatus, FolderType,
    LoanDetail, Manager, Package, PartnerPackage, Region, Shop, UserProfile,
)
from app_documents.raw_reports import (
    CHECKING_HEADERS, RECEIVING_HEADERS, EXCEL_MAX_DATA_ROWS, checking_row, date_range, excel_response,
    receiving_row,
)


class RawReportExcelTests(SimpleTestCase):
    def test_receiving_missing_and_received_rows(self):
        folder = NS(folder_code='2026083136151', shop_id=NS(shop_code=3615, shop_name='YBI23007.Dốc Đỏ Nghĩa Lộ'),
                    folder_created_date=date(2026, 8, 31), folder_type_id=NS(folder_type_name='Vận hành'),
                    is_original=True, is_issue=True, lastest_received_date=None,
                    folder_status_id=NS(folder_status_name='Chưa nhận'), is_late=None, is_on_time=None,
                    package_id=None, lastest_received_by=None)
        row = receiving_row(folder, 1)
        self.assertEqual(row[2], 'YBI23007.Dốc Đỏ Nghĩa Lộ')
        self.assertEqual(row[5:], ['Gốc PSCT', None, 'Chưa nhận', '—', ''])
        folder.lastest_received_date = datetime(2026, 9, 4, 12)
        folder.folder_status_id.folder_status_name = 'Đã nhận'
        folder.is_late = True
        folder.package_id = NS(package_code='VH-260903-201', partnerpackage=NS(partner_package_code='DT001'))
        folder.lastest_received_by = NS(username='hanntn2')
        row = receiving_row(folder, 2)
        self.assertEqual(row[7:], ['Đã nhận Trễ hạn', 'F88: VH-260903-201 | ĐT: DT001', 'hanntn2'])
        folder.is_late, folder.is_on_time = False, True
        folder.is_original, folder.is_issue = False, False
        row = receiving_row(folder, 3)
        self.assertEqual(row[5], 'Bổ sung')
        self.assertEqual(row[7], 'Đã nhận Đúng hạn')

    def test_checking_contract_first_customer_matches_ui_and_folder_dates(self):
        document = NS(shop_id=NS(shop_code='PGD01', shop_name='PGD01.Tên'),
                      folder_id=NS(folder_code='00123', folder_created_date=date(2026, 8, 1),
                                   lastest_received_date=datetime(2026, 8, 24), folder_type_id=None),
                      contract_id=NS(contract_code='0000456', customer_id=NS(customer_name='KH HĐCC')),
                      loan_id=NS(loan_code='0000789', customer_name='KH GNN'),
                      document_type_id=None, business_type_id=None, document_status_id=None,
                      status_id=None, lastest_checked_date=None, lastest_checked_by=None, package_id=None)
        row = checking_row(document, 1)
        self.assertEqual(row[3:5], [date(2026, 8, 1), datetime(2026, 8, 24)])
        self.assertEqual(row[6:8], ['0000456', 'KH GNN'])
        self.assertEqual(row[11:16], ['Chưa duyệt', None, '', '', ''])
        document.loan_id = None
        self.assertEqual(checking_row(document, 1)[7], 'KH HĐCC')
        document.contract_id = None
        self.assertEqual(checking_row(document, 1)[6:8], ['', ''])

    def test_excel_keeps_ids_dates_and_formula_like_strings(self):
        response = excel_response(RECEIVING_HEADERS, [[1, '00123456789012345', '=1+1', date(2026, 8, 31),
                                  'VH', 'Gốc PSCT', datetime(2026, 9, 4, 13, 53), 'Đã nhận', '—', 'user']], 'receiving')
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual([c.value for c in sheet[1]], RECEIVING_HEADERS)
        self.assertEqual(sheet['B2'].value, '00123456789012345')
        self.assertEqual(sheet['C2'].data_type, 's')
        self.assertEqual(sheet['D2'].value, datetime(2026, 8, 31))
        self.assertEqual(sheet['G2'].value, datetime(2026, 9, 4, 13, 53))
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:J2')
        self.assertEqual(response['Cache-Control'], 'no-store')

    def test_invalid_or_reversed_dates_are_rejected(self):
        for raw in ['2026-02-30', '2026-09-01 to 2026-08-01', 'bad', '2026-01-01 to 2026-02-01 to 2026-03-01']:
            with self.assertRaises(ValueError):
                date_range(raw, 'Ngày nhận')
        self.assertEqual(date_range('2026-08-01 to 2026-09-30', 'Ngày nhận'),
                         (date(2026, 8, 1), date(2026, 9, 30)))


class RawReportPortalTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(username='report_checker', password='test')
        cls.user.groups.add(Group.objects.create(name='admin'))
        cls.region = Region.objects.create(region_code='R1', region_name='Vùng 1')
        other_region = Region.objects.create(region_code='R2', region_name='Vùng 2')
        UserProfile.objects.create(user=cls.user, region=cls.region, department='Vận hành')
        cls.manager = Manager.objects.create(manager_code='MR', qlkv_code='A', qlkv_name='A', qlv_code='B', qlv_name='B')
        cls.shop = Shop.objects.create(shop_code=1001, shop_name='PGD 1', manager_id=cls.manager, region_id=cls.region)
        cls.other_shop = Shop.objects.create(shop_code=1002, shop_name='PGD 2', manager_id=cls.manager, region_id=other_region)
        cls.folder_type = FolderType.objects.create(folder_type_code='RPT', folder_type_name='Vận hành', created_by=cls.user)
        cls.status = FolderStatus.objects.create(folder_status_code='RP', folder_status_name='Đã nhận', is_received=True, created_by=cls.user)
        cls.package = Package.objects.create(package_code='VH-260824-208', created_by=cls.user)
        PartnerPackage.objects.create(package_id=cls.package, partner_package_code='DT001', created_date=datetime(2026, 8, 24))
        cls.business = BusinessType.objects.create(business_type_code='RP', business_type_name='Nghiệp vụ', allow_checking=True, created_by=cls.user)
        cls.contract = ContractDetail.objects.create(contract_code='0000123456')
        cls.loan = LoanDetail.objects.create(loan_code='0000987654', customer_name='KH GNN')
        cls.folders = []
        for i in range(27):
            cls.folders.append(Folder.objects.create(
                folder_code=f'00000{i}', shop_id=cls.shop, folder_type_id=cls.folder_type,
                folder_created_date=date(2026, 8, 15), lastest_received_date=datetime(2026, 8, 25),
                lastest_received_by=cls.user, folder_status_id=cls.status, package_id=cls.package, is_on_time=True,
            ))
        Folder.objects.create(folder_code='OUTSIDE', shop_id=cls.other_shop, folder_created_date=date(2026, 8, 15))
        cls.document = DocumentsDetail.objects.create(
            documents_code='RAW001', shop_id=cls.shop, manager_id=cls.manager,
            folder_id=cls.folders[0], documents_created_date=date(2026, 7, 1),
            business_type_id=cls.business, contract_id=cls.contract, loan_id=cls.loan, package_id=cls.package,
        )
        DocumentsDetail.objects.create(documents_code='OUTSIDE', shop_id=cls.other_shop, manager_id=cls.manager,
                                       folder_id=cls.folders[0], documents_created_date=date(2026, 7, 1), business_type_id=cls.business)

    def setUp(self):
        self.client.force_login(self.user)

    def test_admin_receiving_preview_and_export_include_all_regions(self):
        params = {'filter_folder_date': '2026-08-01 to 2026-08-31'}
        response = self.client.get(reverse('document_raw_export'), params)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['total'], 28)
        self.assertEqual(len(response.context['rows']), 25)
        self.assertContains(response, 'Xuất Excel')
        response = self.client.get(reverse('document_raw_export'), {**params, 'download': '1', 'page': '2'})
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 29)
        self.assertEqual(sheet['I2'].value, 'F88: VH-260824-208 | ĐT: DT001')
        self.assertIn('OUTSIDE', [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)])
        response = self.client.get(reverse('document_raw_export'), {'choice_shop': str(self.other_shop.pk), 'download': '1'})
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row, 2)

    def test_checking_export_dates_and_mapping(self):
        response = self.client.get(reverse('document_raw_export'), {'report_kind': 'checking', 'filter_folder_date': '2026-08-15', 'choice_shop': str(self.shop.pk), 'download': '1'})
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual([c.value for c in sheet[1]], CHECKING_HEADERS)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['D2'].value, datetime(2026, 8, 15))
        self.assertEqual(sheet['G2'].value, '0000123456')
        self.assertEqual(sheet['H2'].value, 'KH GNN')
        self.assertEqual(sheet['P2'].value, 'DT001')

    def test_download_checks_dates_empty_filters_and_limits(self):
        url = reverse('document_raw_export')
        self.assertEqual(self.client.get(url, {'download': '1'}).status_code, 400)
        self.assertEqual(self.client.get(url, {'filter_receive_date': 'oops', 'download': '1'}).status_code, 400)
        with override_settings(DOCUMENT_RAW_EXPORT_MAX_ROWS=5):
            response = self.client.get(url, {'filter_folder_date': '2026-08-15', 'download': '1'})
            self.assertContains(response, 'vượt giới hạn', status_code=400)

    def test_direct_access_requires_admin_and_login(self):
        url = reverse('document_raw_export')
        self.client.logout()
        self.assertEqual(self.client.get(url, {'download': '1'}).status_code, 302)
        self.client.force_login(self.user)
        self.user.groups.clear()
        self.user.groups.add(Group.objects.create(name='checker'))
        for kind in ('receiving', 'checking'):
            self.assertEqual(self.client.get(url, {'report_kind': kind}).status_code, 403)
            self.assertEqual(self.client.get(url, {'report_kind': kind, 'download': '1'}).status_code, 403)

    def test_multiple_pgds_and_selected_filters_are_retained(self):
        params = {'choice_shop': [str(self.shop.pk), str(self.other_shop.pk)]}
        response = self.client.get(reverse('document_raw_export'), params)
        self.assertEqual(response.context['total'], 28)
        self.assertEqual(response.context['selected_shop_count'], 2)
        for shop in (self.shop, self.other_shop):
            self.assertContains(response, f'value="{shop.pk}" checked')
        response = self.client.get(reverse('document_raw_export'), {**params, 'download': '1'})
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row, 29)

    def test_excel_capacity_blocks_preview_and_download_even_if_setting_is_larger(self):
        url = reverse('document_raw_export')
        with override_settings(DOCUMENT_RAW_EXPORT_MAX_ROWS=EXCEL_MAX_DATA_ROWS + 100):
            with patch('django.db.models.query.QuerySet.count', return_value=EXCEL_MAX_DATA_ROWS + 1):
                for params in ({'choice_shop': str(self.shop.pk)},
                               {'choice_shop': str(self.shop.pk), 'download': '1'}):
                    with patch('app_documents.raw_reports.excel_response') as export:
                        response = self.client.get(url, params)
                    self.assertEqual(response.status_code, 400)
                    self.assertFalse(response.context['can_download'])
                    self.assertContains(response, 'vượt giới hạn', status_code=400)
                    export.assert_not_called()

    def test_kpi_navigation_is_shared_and_old_exports_are_removed(self):
        response = self.client.get(reverse('document_raw_export'))
        self.assertEqual(response.context['active_tab'], 'export')
        self.assertContains(response, 'Cấu hình KPI')
        self.assertContains(response, 'Xuất dữ liệu')
        self.assertContains(response, 'table-xs')
        self.assertContains(response, 'filter-input')
        for path in ('/nhan-chung-tu-v2/bao-cao/', '/duyet-chung-tu-v2/bao-cao/'):
            self.assertEqual(self.client.get(path).status_code, 404)
        response = self.client.get(reverse('document_kpi_v2'), {'tab': 'config'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse('document_raw_export'))

    def test_no_filters_do_not_read_all_data_and_invalid_pgds_are_rejected(self):
        response = self.client.get(reverse('document_raw_export'))
        self.assertEqual(response.context['total'], 0)
        self.assertFalse(response.context['can_download'])
        response = self.client.get(reverse('document_raw_export'), {'choice_shop': 'bad'})
        self.assertEqual(response.status_code, 400)

    def test_only_requested_filters_and_primary_filters_remain(self):
        labels = ['Phòng giao dịch', 'Loại quyển chứng từ', 'Ngày phát sinh', 'Ngày nhận',
                  'Mã thùng', 'Mã thùng đối tác', 'Ngày duyệt', 'Người duyệt',
                  'Người nhận', 'Nghiệp vụ', 'Loại chứng từ']
        for kind in ('receiving', 'checking'):
            response = self.client.get(reverse('document_raw_export'), {'report_kind': kind})
            self.assertEqual([field['label'] for field in response.context['fields']], labels)
            self.assertEqual([field['label'] for field in response.context['fields'] if field['primary']], labels[:4])
            self.assertFalse(response.context['filters_expanded'])
            self.assertContains(response, 'role="tablist"')
            self.assertContains(response, 'role="tabpanel"')
        response = self.client.get(reverse('document_raw_export'), {'choice_lastest_receiver': self.user.username})
        self.assertTrue(response.context['filters_expanded'])

    def test_receiving_document_filters_match_same_document_without_duplicate_folders(self):
        type_a = DocumentType.objects.create(document_type_code='A', document_type_name='Loại A', created_by=self.user)
        type_b = DocumentType.objects.create(document_type_code='B', document_type_name='Loại B', created_by=self.user)
        self.document.document_type_id = type_a
        self.document.lastest_checked_date = datetime(2026, 9, 1, 10)
        self.document.lastest_checked_by = self.user
        self.document.save()
        for code, doc_type, checked_date in [('MATCH', type_a, datetime(2026, 9, 1, 11)),
                                              ('OTHER', type_b, datetime(2026, 9, 2, 10))]:
            DocumentsDetail.objects.create(documents_code=code, shop_id=self.shop, manager_id=self.manager,
                folder_id=self.folders[0], documents_created_date=date(2026, 8, 15),
                document_type_id=doc_type, business_type_id=self.business,
                lastest_checked_date=checked_date, lastest_checked_by=self.user)
        params = {'filtercheckdate': '2026-09-01', 'choice_user_duyet': self.user.username,
                  'choicebusinesstype': str(self.business.pk), 'choice_document_type': str(type_a.pk)}
        response = self.client.get(reverse('document_raw_export'), params)
        self.assertEqual(response.context['total'], 1)
        response = self.client.get(reverse('document_raw_export'), {**params, 'download': '1'})
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row, 2)
        response = self.client.get(reverse('document_raw_export'), {**params, 'choice_document_type': str(type_b.pk)})
        self.assertEqual(response.context['total'], 0)

    def test_switching_tab_keeps_multiple_pgds_and_receiver_filter(self):
        params = {'report_kind': ['receiving', 'checking'],
                  'choice_shop': [str(self.shop.pk), str(self.other_shop.pk)],
                  'choice_lastest_receiver': self.user.username}
        response = self.client.get(reverse('document_raw_export'), params)
        self.assertEqual(response.context['kind'], 'checking')
        self.assertEqual(response.context['selected_shop_count'], 2)
        self.assertEqual(response.context['total'], 2)
        response = self.client.get(reverse('document_raw_export'), {**params, 'choice_lastest_receiver': 'no_receiver'})
        self.assertEqual(response.context['total'], 0)
