"""Round-trip shop catalog workbooks; validate and commit imports atomically."""
import json
from io import BytesIO
from celery.exceptions import SoftTimeLimitExceeded

from django.contrib.admin.models import CHANGE, LogEntry
from django.contrib.contenttypes.models import ContentType
from django.core import signing
from django.core.files.base import ContentFile
from django.core.validators import validate_email
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import Manager, Shop, ShopCatalogJob
from .master_data import local_today, serialize, shop_queryset

HEADERS = ["shop_id", "shop_code", "shop_name", "shop_email", "is_shop_active", "manager_id", "shop_closed_date", "area_manager_name", "area_manager_email", "region_manager_name", "region_manager_email", "_snapshot"]
SALT = "shop-catalog-v1"


def update(job, **fields):
    ShopCatalogJob.objects.filter(pk=job.pk).update(updated_at=timezone.now(), **fields)


def snapshot(shop):
    return {"shop_id": shop.pk, "shop_code": shop.shop_code, "shop_name": shop.shop_name, "shop_email": shop.shop_email or "", "is_shop_active": shop.is_shop_active, "manager_id": shop.manager_id_id, "shop_closed_date": str(shop.shop_closed_date) if shop.shop_closed_date else None}


def write_row(sheet, values):
    # Explicit strings prevent catalog text from becoming Excel formulas.
    sheet.append(values)
    for cell, value in zip(sheet[sheet.max_row], values):
        if isinstance(value, str):
            cell.data_type = "s"


def save_workbook(job, workbook, filename):
    stream = BytesIO()
    workbook.save(stream)
    job.output_file.save(filename, ContentFile(stream.getvalue()), save=False)
    update(job, output_file=job.output_file.name)


def export_catalog(job):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PGD"
    write_row(sheet, HEADERS)
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = "A1:L1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="00854A")
    for index, shop in enumerate(shop_queryset().iterator(chunk_size=500), 1):
        data = serialize(shop)
        original = snapshot(shop)
        org = data["orgchart"]
        write_row(sheet, [shop.pk, shop.shop_code, shop.shop_name, shop.shop_email or "", shop.is_shop_active, shop.manager_id_id, data["shop_closed_date"], org["area_manager"]["name"], org["area_manager"]["email"], org["region_manager"]["name"], org["region_manager"]["email"], signing.dumps(original, salt=SALT, compress=True)])
        if index % 100 == 0:
            update(job, message=f"Đã xuất {index} PGD", progress=50)
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A":12,"B":14,"C":40,"D":38,"E":20,"F":18,"G":20,"H":30,"I":38,"J":30,"K":38}.items():
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["L"].hidden = True
    instructions = workbook.create_sheet("Huong dan")
    for text in ["Chỉ chỉnh shop_email, is_shop_active (TRUE/FALSE), manager_id.", "Các trường khác chỉ để đối chiếu. Không sửa/xóa cột _snapshot.", "Có thể chỉ giữ các dòng cần sửa; xóa dòng KHÔNG xóa PGD.", "Import không tạo PGD mới. Toàn bộ file được kiểm tra trước khi cập nhật.", "Nếu dữ liệu PGD đã thay đổi sau khi xuất, tải file mới và thực hiện lại.", "Tắt PGD tự ghi ngày đóng cửa; bật lại xóa ngày đóng cửa."]:
        instructions.append([text])
    instructions.column_dimensions["A"].width = 110
    managers = workbook.create_sheet("Manager")
    write_row(managers, ["manager_id", "manager_code", "area_manager_name", "area_manager_email", "region_manager_name", "region_manager_email", "is_current", "valid_from", "valid_to"])
    today = local_today()
    for manager in Manager.objects.select_related("areaManager", "regionManager").order_by("manager_code").iterator(chunk_size=500):
        area, region = manager.areaManager, manager.regionManager
        current = manager.is_valid and (not manager.valid_from or manager.valid_from <= today) and (not manager.valid_to or manager.valid_to >= today)
        write_row(managers, [manager.pk, manager.manager_code, area.areaManager_name if area else manager.qlkv_name, area.areaManager_email if area else manager.qlkv_email, region.regionManager_name if region else manager.qlv_name, region.regionManager_email if region else manager.qlv_email, current, str(manager.valid_from) if manager.valid_from else None, str(manager.valid_to) if manager.valid_to else None])
    managers.freeze_panes = "A2"
    for column in "ABCDEFGHI":
        managers.column_dimensions[column].width = 30
    instructions.append(["Sheet Manager dùng đối chiếu manager_id; chỉ chọn bộ có is_current=TRUE. Import không sửa sheet Manager."])
    save_workbook(job, workbook, f"danh-muc-pgd-{job.pk}.xlsx")
    update(job, status="succeeded", progress=100, message="Đã xuất danh mục PGD", summary={"rows": sheet.max_row - 1})


def parse_boolean(value):
    if type(value) is bool:
        return value
    if isinstance(value, str) and value.strip().upper() in {"TRUE", "FALSE"}:
        return value.strip().upper() == "TRUE"
    raise ValueError("is_shop_active phải là TRUE hoặc FALSE.")


def integer(value, field):
    if type(value) not in (int, float) or int(value) != value or value <= 0:
        raise ValueError(f"{field} phải là số nguyên dương.")
    return int(value)


def import_catalog(job):
    errors, changes, seen = [], [], set()
    with job.input_file.open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=False)
        try:
            if "PGD" not in workbook.sheetnames:
                raise ValueError("Thiếu sheet PGD. Hãy dùng file xuất từ Master Data.")
            rows = workbook["PGD"].iter_rows(values_only=True)
            if list(next(rows, [])) != HEADERS:
                raise ValueError("Cấu trúc cột không đúng file danh mục PGD đã xuất.")
            for line, values in enumerate(rows, 2):
                if line > 10001:
                    raise ValueError("File tối đa 10.000 dòng.")
                if all(value is None for value in values):
                    continue
                try:
                    data = dict(zip(HEADERS, values))
                    pk = integer(data["shop_id"], "shop_id")
                    if pk in seen:
                        raise ValueError("PGD trùng trong file.")
                    seen.add(pk)
                    original = signing.loads(data["_snapshot"], salt=SALT)
                    if original["shop_id"] != pk or original["shop_code"] != data["shop_code"] or original["shop_name"] != data["shop_name"]:
                        raise ValueError("Mã/tên PGD hoặc snapshot đã bị thay đổi.")
                    email = data["shop_email"] or ""
                    if not isinstance(email, str) or len(email.strip()) > 100:
                        raise ValueError("Email không hợp lệ hoặc quá 100 ký tự.")
                    email = email.strip()
                    if email and email != original["shop_email"]:
                        try:
                            validate_email(email)
                        except ValidationError:
                            raise ValueError("Email PGD mới không đúng định dạng. Vui lòng kiểm tra lại.") from None
                    active = parse_boolean(data["is_shop_active"])
                    manager_id = integer(data["manager_id"], "manager_id")
                    changes.append((line, original, {"shop_email": email, "is_shop_active": active, "manager_id": manager_id}))
                except SoftTimeLimitExceeded:
                    raise
                except Exception as exc:
                    errors.append({"row": line, "error": str(exc)})
                if line % 100 == 0:
                    update(job, progress=min(70, line // 150 + 10), message=f"Đang kiểm tra dòng {line}")
        finally:
            workbook.close()
    if not changes and not errors:
        raise ValueError("File không có dữ liệu PGD.")
    with transaction.atomic():
        locked = {shop.pk: shop for shop in Shop.objects.select_for_update().filter(pk__in=seen).order_by("shop_id")}
        managers = {m.pk: m for m in Manager.objects.filter(pk__in=[data["manager_id"] for _, _, data in changes])}
        today = local_today()
        for line, original, data in changes:
            shop = locked.get(original["shop_id"])
            if not shop or snapshot(shop) != original:
                errors.append({"row": line, "error": "PGD không tồn tại hoặc đã thay đổi sau khi xuất. Hãy xuất lại file mới."})
            if data["manager_id"] != original["manager_id"]:
                manager = managers.get(data["manager_id"])
                if not manager or not manager.is_valid or (manager.valid_from and manager.valid_from > today) or (manager.valid_to and manager.valid_to < today):
                    errors.append({"row": line, "error": "Bộ quản lý không tồn tại hoặc hết/chưa hiệu lực."})
        if not errors:
            updated = 0
            content_type = ContentType.objects.get_for_model(Shop)
            for line, original, data in changes:
                shop = locked[original["shop_id"]]
                if all(data[key] == original[key] for key in data):
                    continue
                if shop.is_shop_active != data["is_shop_active"]:
                    shop.shop_closed_date = None if data["is_shop_active"] else today
                shop.shop_email = data["shop_email"]
                shop.is_shop_active = data["is_shop_active"]
                shop.manager_id_id = data["manager_id"]
                shop.save(update_fields=["shop_email", "is_shop_active", "shop_closed_date", "manager_id"])
                LogEntry.objects.log_action(user_id=job.requested_by_id, content_type_id=content_type.pk, object_id=shop.pk, object_repr=str(shop), action_flag=CHANGE, change_message=json.dumps({"source": "master_data_excel", "job_id": job.pk, "before": original, "after": data}, ensure_ascii=False))
                updated += 1
            update(job, status="succeeded", progress=100, message=f"Đã cập nhật {updated}/{len(changes)} PGD", summary={"rows": len(changes), "updated": updated})
    if errors:
        report = Workbook()
        sheet = report.active
        sheet.title = "Loi import"
        write_row(sheet, ["Dòng Excel", "Lỗi"])
        for error in errors:
            write_row(sheet, [error["row"], error["error"]])
        sheet.column_dimensions["B"].width = 110
        save_workbook(job, report, f"loi-import-pgd-{job.pk}.xlsx")
        update(job, status="failed", progress=100, message=f"Có {len(errors)} lỗi; chưa cập nhật PGD nào. Tải báo cáo lỗi để sửa file.", summary={"errors": errors[:100], "error_count": len(errors), "updated": 0})


def process_job(job_id):
    job = ShopCatalogJob.objects.select_related("requested_by").get(pk=job_id)
    if not ShopCatalogJob.objects.filter(pk=job.pk, status="queued").update(status="running", progress=5, updated_at=timezone.now()):
        return
    try:
        user = job.requested_by
        if not user.is_active or not (user.is_superuser or user.groups.filter(name="admin").exists()):
            raise ValueError("Tài khoản tạo job không còn quyền admin.")
        update(job, message="Đang chuẩn bị file" if job.kind == "export" else "Đang kiểm tra file Excel")
        if job.kind == "export":
            export_catalog(job)
        else:
            import_catalog(job)
    except Exception as exc:
        update(job, status="failed", message=f"Không hoàn tất: {exc}")
