from datetime import date, datetime, timedelta
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .gddb import import_collateral_registrations
from .gddb_dashboard import calculate_sla_due_at
from .models import (
    CollateralRegistration,
    CollateralRegistrationHoliday,
    CollateralRegistrationImportBatch,
    CollateralRegistrationLog,
    CollateralRegistrationStatus,
    UserProfile,
)


class GddbDashboardTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        admin_group = Group.objects.create(name="admin")
        checker_group = Group.objects.create(name="checker")
        cls.admin = User.objects.create_user(username="dashboard_admin", password="test")
        cls.admin.groups.add(admin_group)
        UserProfile.objects.create(user=cls.admin, department="Vận hành")
        cls.checker = User.objects.create_user(username="dashboard_checker", password="test")
        cls.checker.groups.add(checker_group)
        UserProfile.objects.create(user=cls.checker, department="Vận hành")

    def setUp(self):
        now = timezone.now()
        previous_day = now - timedelta(days=1)

        CollateralRegistration.objects.create(
            contract_code="DASH-PENDING",
            gddb_status=CollateralRegistrationStatus.PENDING,
            disbursement_source="F88 Fund",
            asset_type="Ô tô",
        )
        CollateralRegistration.objects.create(
            contract_code="DASH-WAITING",
            gddb_status=CollateralRegistrationStatus.REGISTERED,
            postmini_updated="Chưa cập nhật",
            disbursement_source="CIMB",
            asset_type="Xe máy",
            registered_by=self.admin,
            registered_at=now,
        )
        CollateralRegistration.objects.create(
            contract_code="DASH-NOT-REGISTERED",
            gddb_status=CollateralRegistrationStatus.NOT_REGISTERED,
            disbursement_source=None,
            asset_type=None,
            processing_by=self.admin,
            archived_at=now,
        )
        completed = CollateralRegistration.objects.create(
            contract_code="DASH-COMPLETED",
            gddb_status=CollateralRegistrationStatus.REGISTERED,
            postmini_updated="Đã cập nhật",
            disbursement_source="F88",
            asset_type="Ô tô",
            registered_by=self.admin,
            registered_at=previous_day,
            archived_at=previous_day,
        )
        CollateralRegistration.objects.filter(pk=completed.pk).update(
            created_at=previous_day,
            archived_at=previous_day,
            registered_at=previous_day,
        )

    def test_dashboard_page_and_api_allow_checker_with_personal_scope(self):
        admin_client = Client()
        admin_client.force_login(self.admin)
        checker_client = Client()
        checker_client.force_login(self.checker)

        admin_page = admin_client.get(reverse("gddb_dashboard"))
        checker_page = checker_client.get(reverse("gddb_dashboard"))
        self.assertEqual(admin_page.status_code, 200)
        self.assertContains(admin_page, "Dashboard tổng quan")
        self.assertEqual(checker_page.status_code, 200)
        self.assertContains(checker_page, "Dashboard của tôi")
        response = checker_client.get(reverse("api_gddb_dashboard"))
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["role"], "checker")
        self.assertIn("personal", payload)
        self.assertNotIn("cards", payload)

    def test_dashboard_returns_realtime_cards_and_breakdowns(self):
        client = Client()
        client.force_login(self.admin)
        today = timezone.now().date()
        response = client.get(
            reverse("api_gddb_dashboard"),
            {"period": "month", "month": today.strftime("%Y-%m")},
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["success"])
        self.assertEqual(data["cards"]["month_total"], 4)
        self.assertEqual(data["cards"]["pending"], 1)
        self.assertEqual(data["cards"]["unassigned"], 1)
        self.assertEqual(data["cards"]["waiting_posmini"], 1)

        status = {row["label"]: row["value"] for row in data["status"]}
        self.assertEqual(status["Chưa đăng ký"], 1)
        self.assertEqual(status["Cập nhật PosMini"], 1)
        self.assertEqual(status["Không đăng ký"], 1)
        self.assertEqual(status["Hoàn tất đăng ký"], 1)

        source = {row["label"]: row["value"] for row in data["source"]}
        self.assertEqual(source["F88"], 2)
        self.assertEqual(source["CIMB"], 1)
        self.assertEqual(source["MB"], 1)

        asset = {row["label"]: row["value"] for row in data["asset"]}
        self.assertEqual(asset["Ô tô"], 2)
        self.assertEqual(asset["Xe máy"], 1)
        self.assertEqual(asset["Chưa xác định"], 1)

        checker = {row["label"]: row["value"] for row in data["checker"]}
        self.assertEqual(checker[self.admin.username], 2)

    def test_dashboard_rejects_ranges_over_one_year(self):
        client = Client()
        client.force_login(self.admin)
        response = client.get(
            reverse("api_gddb_dashboard"),
            {
                "period": "range",
                "date_from": "2025-01-01",
                "date_to": "2026-02-01",
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.json()["success"])

    def test_intake_persists_asset_type_for_dashboard(self):
        import_collateral_registrations(
            [
                {
                    "Mã hợp đồng": "DASH-ASSET-INTAKE",
                    "Loại tài sản": "Ô tô tải",
                }
            ],
            user=self.admin,
            source_type="api",
        )

        registration = CollateralRegistration.objects.get(
            contract_code="DASH-ASSET-INTAKE"
        )
        self.assertEqual(registration.asset_type, "Ô tô tải")

    def test_sla_adds_three_calendar_days_then_rolls_weekend_forward(self):
        effective_at = datetime(2026, 7, 1, 9, 0)
        self.assertEqual(
            calculate_sla_due_at(effective_at, set()),
            datetime(2026, 7, 6, 9, 0),
        )

    def test_dashboard_uses_batch_slot_hour_for_sla(self):
        batch = CollateralRegistrationImportBatch.objects.create(
            source_type="api",
            business_date=date(2026, 7, 1),
            slot_number=1,
        )
        assigned = CollateralRegistration.objects.get(contract_code="DASH-PENDING")
        assigned.import_batch = batch
        assigned.processing_by = self.checker
        assigned.disbursement_date = date(2026, 7, 1)
        assigned.save(
            update_fields=["import_batch", "processing_by", "disbursement_date"]
        )

        client = Client()
        client.force_login(self.checker)
        response = client.get(reverse("api_gddb_dashboard"))

        case = response.json()["personal"]["assigned_cases"][0]
        self.assertTrue(case["effective_at"].startswith("2026-07-01T09:00:00"))
        self.assertTrue(case["due_at"].startswith("2026-07-06T09:00:00"))

    def test_reconciliation_export_contains_sla_summary_and_case_detail(self):
        batch = CollateralRegistrationImportBatch.objects.create(
            source_type="api",
            business_date=date(2026, 7, 1),
            slot_number=1,
        )
        registration = CollateralRegistration.objects.get(
            contract_code="DASH-WAITING"
        )
        registration.import_batch = batch
        registration.disbursement_date = date(2026, 7, 1)
        registration.save(update_fields=["import_batch", "disbursement_date"])

        client = Client()
        client.force_login(self.admin)
        export_day = timezone.now().date()
        response = client.get(
            reverse("gddb_export"),
            {
                "period_type": "day",
                "export_day": export_day.isoformat(),
            },
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        self.assertIn("Tong quan SLA", workbook.sheetnames)
        self.assertIn("Doi soat GDDB", workbook.sheetnames)
        detail = workbook["Doi soat GDDB"]
        headers = [cell.value for cell in detail[1]]
        self.assertIn("Thời điểm hiệu lực SLA", headers)
        self.assertIn("Hạn SLA", headers)
        self.assertIn("Kết quả SLA", headers)
        row = next(
            values
            for values in detail.iter_rows(min_row=2, values_only=True)
            if values[0] == "DASH-WAITING"
        )
        self.assertEqual(row[15], datetime(2026, 7, 1, 9, 0))
        self.assertEqual(row[16], datetime(2026, 7, 6, 9, 0))

    def test_checker_dashboard_returns_assigned_cases_and_own_history(self):
        assigned = CollateralRegistration.objects.get(contract_code="DASH-PENDING")
        assigned.processing_by = self.checker
        assigned.disbursement_date = timezone.now().date() - timedelta(days=7)
        assigned.save(update_fields=["processing_by", "disbursement_date"])
        CollateralRegistrationLog.objects.create(
            registration=assigned,
            action="case_claim",
            created_by=self.checker,
        )

        client = Client()
        client.force_login(self.checker)
        response = client.get(reverse("api_gddb_dashboard"))

        self.assertEqual(response.status_code, 200)
        personal = response.json()["personal"]
        self.assertEqual(personal["cards"]["assigned"], 1)
        self.assertEqual(personal["cards"]["overdue"], 1)
        self.assertEqual(
            personal["assigned_cases"][0]["contract_code"],
            "DASH-PENDING",
        )
        self.assertEqual(personal["action_history"][0]["action_label"], "Nhận case")
