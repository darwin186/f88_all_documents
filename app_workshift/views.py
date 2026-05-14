import calendar
from io import BytesIO
from urllib.parse import urlencode
from datetime import datetime, timedelta
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.contrib.auth.models import User
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.db import models
from django.db.models import Q
from django.utils import timezone
from django.core.serializers.json import DjangoJSONEncoder
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app_documents.utils import get_user_context, require_ui_permission
from app_admindocuments.models import AdmParcelRecipientImportBatch

from .models import (
    WorkPolicy,
    WorkShiftTemplate,
    WorkWeek,
    WorkShift,
    TaskCatalog,
    TaskAssignment,
    WorkAuditLog,
)


DEFAULT_POLICY = {
    'working_days': [0, 1, 2, 3, 4],
    'open_time': '08:00',
    'close_time': '17:30',
    'lunch_start': '12:00',
    'lunch_end': '13:30',
    'max_shifts_per_day': 2,
    'max_hours_per_day': Decimal('8'),
    'max_hours_per_week': Decimal('40'),
    'max_hours_per_month': Decimal('160'),
    'min_hours_per_shift': Decimal('4'),
}


def _get_active_checker_users():
    checker_qs = (
        User.objects.filter(groups__name='checker', is_active=True)
        .select_related('userprofile')
        .distinct()
        .order_by('username')
    )
    current_batch = AdmParcelRecipientImportBatch.objects.filter(is_current=True).first()
    if not current_batch:
        return checker_qs

    active_recipients = current_batch.recipients.filter(is_active_member=True)
    employee_codes = list(
        active_recipients.exclude(employee_code='')
        .values_list('employee_code', flat=True)
        .distinct()
    )
    gapo_user_ids = list(
        active_recipients.exclude(gapo_user_id='')
        .values_list('gapo_user_id', flat=True)
        .distinct()
    )
    emails = list(
        active_recipients.exclude(email='')
        .values_list('email', flat=True)
        .distinct()
    )

    active_filter = Q()
    if employee_codes:
        active_filter |= Q(userprofile__employee_code__in=employee_codes)
    if gapo_user_ids:
        active_filter |= Q(userprofile__gapo_user_id__in=gapo_user_ids)
    if emails:
        active_filter |= Q(email__in=emails)
    if not active_filter:
        return checker_qs.none()
    return checker_qs.filter(active_filter)


def _parse_time(value):
    if not value:
        return None
    try:
        return datetime.strptime(value, '%H:%M').time()
    except ValueError:
        return None


def _parse_date_value(value):
    if not value:
        return None
    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except ValueError:
        return None


def _get_active_policy():
    return WorkPolicy.objects.filter(is_active=True).order_by('-effective_from', '-policy_id').first()


def _get_policy_values(policy):
    if not policy:
        return {
            'working_days': DEFAULT_POLICY['working_days'],
            'open_time': _parse_time(DEFAULT_POLICY['open_time']),
            'close_time': _parse_time(DEFAULT_POLICY['close_time']),
            'lunch_start': _parse_time(DEFAULT_POLICY['lunch_start']),
            'lunch_end': _parse_time(DEFAULT_POLICY['lunch_end']),
            'max_shifts_per_day': DEFAULT_POLICY['max_shifts_per_day'],
            'max_hours_per_day': DEFAULT_POLICY['max_hours_per_day'],
            'max_hours_per_week': DEFAULT_POLICY['max_hours_per_week'],
            'max_hours_per_month': DEFAULT_POLICY['max_hours_per_month'],
            'min_hours_per_shift': DEFAULT_POLICY['min_hours_per_shift'],
        }
    working_days = []
    if policy.working_days:
        for part in policy.working_days.split(','):
            try:
                working_days.append(int(part.strip()))
            except ValueError:
                continue
    if not working_days:
        working_days = DEFAULT_POLICY['working_days']
    return {
        'working_days': working_days,
        'open_time': policy.open_time,
        'close_time': policy.close_time,
        'lunch_start': policy.lunch_start,
        'lunch_end': policy.lunch_end,
        'max_shifts_per_day': policy.max_shifts_per_day,
        'max_hours_per_day': policy.max_hours_per_day,
        'max_hours_per_week': policy.max_hours_per_week,
        'max_hours_per_month': policy.max_hours_per_month,
        'min_hours_per_shift': policy.min_hours_per_shift,
    }


def _get_week_start(date_val):
    return date_val - timedelta(days=date_val.weekday())


def _get_week_days(week_start):
    return [week_start + timedelta(days=idx) for idx in range(5)]


def _get_month_range(date_val):
    if date_val.day >= 25:
        start = date_val.replace(day=25)
        if date_val.month == 12:
            end = date_val.replace(year=date_val.year + 1, month=1, day=24)
        else:
            end = date_val.replace(month=date_val.month + 1, day=24)
    else:
        if date_val.month == 1:
            start = date_val.replace(year=date_val.year - 1, month=12, day=25)
        else:
            start = date_val.replace(month=date_val.month - 1, day=25)
        end = date_val.replace(day=24)
    return start, end


def _calc_hours(day, start_time, end_time, lunch_start, lunch_end, include_lunch_break=True):
    if not start_time or not end_time:
        return None
    start_dt = datetime.combine(day, start_time)
    end_dt = datetime.combine(day, end_time)
    if end_dt <= start_dt:
        return None
    total = Decimal((end_dt - start_dt).total_seconds()) / Decimal(3600)
    if include_lunch_break and lunch_start and lunch_end:
        lunch_start_dt = datetime.combine(day, lunch_start)
        lunch_end_dt = datetime.combine(day, lunch_end)
        overlap = max(
            Decimal('0'),
            Decimal((min(end_dt, lunch_end_dt) - max(start_dt, lunch_start_dt)).total_seconds()) / Decimal(3600),
        )
        total = total - overlap
    if total <= 0:
        return None
    return total.quantize(Decimal('0.01'))


def _shift_template_payload(policy_values=None):
    templates = WorkShiftTemplate.objects.filter(is_active=True).order_by('sort_order', 'start_time', 'template_name')
    lunch_start = policy_values.get('lunch_start') if policy_values else _parse_time(DEFAULT_POLICY['lunch_start'])
    lunch_end = policy_values.get('lunch_end') if policy_values else _parse_time(DEFAULT_POLICY['lunch_end'])
    return [
        {
            'id': item.template_id,
            'name': item.template_name,
            'start': item.start_time.strftime('%H:%M'),
            'end': item.end_time.strftime('%H:%M'),
            'include_lunch_break': item.include_lunch_break,
            'effective_from': item.effective_from.isoformat() if item.effective_from else '',
            'effective_to': item.effective_to.isoformat() if item.effective_to else '',
            'hours': float(_calc_hours(datetime.today().date(), item.start_time, item.end_time, lunch_start, lunch_end, item.include_lunch_break) or Decimal('0')),
        }
        for item in templates
    ]


def _is_template_effective(template, day):
    if not template.is_active:
        return False
    if template.effective_from and day < template.effective_from:
        return False
    if template.effective_to and day > template.effective_to:
        return False
    return True


def _shift_payload_for_audit(item):
    shift_template = item.get('shift_template')
    return {
        'shift_date': item.get('shift_date'),
        'shift_index': item.get('shift_index'),
        'shift_template_id': shift_template.template_id if shift_template else item.get('shift_template_id'),
        'start_time': item.get('start_time'),
        'end_time': item.get('end_time'),
        'total_hours': item.get('total_hours'),
    }


def _build_time_slots(start='06:00', end='22:00', step_minutes=30):
    start_time = _parse_time(start)
    end_time = _parse_time(end)
    if not start_time or not end_time:
        return []
    cursor = datetime.combine(datetime.today().date(), start_time)
    end_dt = datetime.combine(datetime.today().date(), end_time)
    step = timedelta(minutes=step_minutes)
    slots = []
    while cursor <= end_dt:
        slots.append(cursor.strftime('%H:%M'))
        cursor += step
    return slots


def _build_period_shift_rows(user, period_days, policy_values):
    shift_indices = list(range(1, policy_values['max_shifts_per_day'] + 1))
    period_start = period_days[0]
    period_end = period_days[-1]
    shifts = list(
        WorkShift.objects.filter(
            work_week__user=user,
            shift_date__range=(period_start, period_end),
        ).select_related('shift_template')
    )
    shift_map = {}
    for shift in shifts:
        shift_map.setdefault(shift.shift_date.isoformat(), {})[shift.shift_index] = shift
    template_lookup = {
        (item.start_time, item.end_time): item.template_id
        for item in WorkShiftTemplate.objects.filter(is_active=True)
    }
    shift_rows = []
    total_hours = Decimal('0')
    for day in period_days:
        day_key = day.isoformat()
        daily_total = Decimal('0')
        row_shifts = {}
        for idx in shift_indices:
            shift = shift_map.get(day_key, {}).get(idx)
            if shift:
                daily_total += shift.total_hours
            row_shifts[str(idx)] = {
                'start_time': shift.start_time.strftime('%H:%M') if shift else '',
                'end_time': shift.end_time.strftime('%H:%M') if shift else '',
                'template_id': (
                    shift.shift_template_id
                    if shift and shift.shift_template_id
                    else template_lookup.get((shift.start_time, shift.end_time), '') if shift else ''
                ),
            }
        total_hours += daily_total
        shift_rows.append({
            'day': day,
            'day_key': day_key,
            'shifts': row_shifts,
            'total_hours': daily_total,
        })
    return shift_rows, total_hours


def _build_week_context(user, week_start, policy_values, selected_user=None):
    week_days = _get_week_days(week_start)
    week_end = week_days[-1]
    work_week = WorkWeek.objects.filter(user=user, week_start_date=week_start).first()
    shift_rows, total_hours = _build_period_shift_rows(user, week_days, policy_values)
    if work_week:
        total_hours = work_week.total_hours
    return {
        'week_days': week_days,
        'week_end': week_end,
        'work_week': work_week,
        'shift_rows': shift_rows,
        'total_hours': total_hours,
    }


def _get_month_days(month_start):
    _, last_day = calendar.monthrange(month_start.year, month_start.month)
    return [month_start.replace(day=day) for day in range(1, last_day + 1)]


def _get_calendar_month_range(month_start):
    month_days = _get_month_days(month_start)
    return month_days[0], month_days[-1]


def _build_month_calendar(month_start, month_days):
    month_day_set = set(month_days)
    first_week_start = _get_week_start(month_start)
    last_week_start = _get_week_start(month_days[-1])
    weeks = []
    cursor = first_week_start
    while cursor <= last_week_start:
        week_days = _get_week_days(cursor)
        weeks.append([
            {
                'day': day,
                'day_key': day.isoformat(),
                'in_month': day in month_day_set,
            }
            for day in week_days
        ])
        cursor += timedelta(days=7)
    return weeks


def _build_day_status_map(user, period_days):
    week_statuses = {
        item.week_start_date: item.status
        for item in WorkWeek.objects.filter(
            user=user,
            week_start_date__in={_get_week_start(day) for day in period_days},
        )
    }
    return {
        day.isoformat(): week_statuses.get(_get_week_start(day), '')
        for day in period_days
    }


def _build_day_template_map(shift_rows):
    day_template_map = {}
    for row in shift_rows:
        selected_ids = []
        for shift in row.get('shifts', {}).values():
            template_id = shift.get('template_id') if shift else ''
            if template_id:
                selected_ids.append(str(template_id))
        day_template_map[row['day_key']] = selected_ids
    return day_template_map


@login_required
@require_ui_permission('workshift_register')
def workshift_register_view(request):
    user_context = get_user_context(request.user)
    is_admin = user_context.get('is_admin') or user_context.get('is_super_admin')
    ctv_users = _get_active_checker_users() if is_admin else User.objects.none()

    policy = _get_active_policy()
    policy_values = _get_policy_values(policy)
    max_shift_hours = Decimal('0')
    if policy_values.get('max_shifts_per_day'):
        try:
            max_shift_hours = (policy_values['max_hours_per_day'] / Decimal(policy_values['max_shifts_per_day'])).quantize(Decimal('0.01'))
        except Exception:
            max_shift_hours = Decimal('0')
    min_shift_hours = policy_values.get('min_hours_per_shift') or Decimal('4')

    selected_user = request.user
    user_id = request.GET.get('user_id')
    if is_admin:
        if user_id:
            selected_user = ctv_users.filter(pk=user_id).first()
            if not selected_user:
                messages.error(request, 'Checker không còn trong danh sách đang làm việc.')
                selected_user = ctv_users.first() or request.user
        else:
            selected_user = ctv_users.first() or request.user

    view_mode = request.POST.get('mode') or request.GET.get('mode', 'week')
    if view_mode not in {'week', 'month'}:
        view_mode = 'week'

    month_param = request.POST.get('month') or request.GET.get('month')
    if month_param:
        try:
            selected_month = datetime.strptime(month_param, '%Y-%m').date().replace(day=1)
        except ValueError:
            selected_month = datetime.today().date().replace(day=1)
    else:
        selected_month = datetime.today().date().replace(day=1)

    week_param = request.GET.get('week_start')
    if week_param:
        try:
            week_start = datetime.strptime(week_param, '%Y-%m-%d').date()
        except ValueError:
            week_start = _get_week_start(datetime.today().date())
    else:
        week_start = _get_week_start(datetime.today().date())

    prev_week = week_start - timedelta(days=7)
    next_week = week_start + timedelta(days=7)
    prev_month = (selected_month - timedelta(days=1)).replace(day=1)
    next_month = (selected_month.replace(day=28) + timedelta(days=4)).replace(day=1)
    calendar_month_days = _get_month_days(selected_month)
    month_calendar = _build_month_calendar(selected_month, calendar_month_days)
    if view_mode == 'month':
        month_start, month_end = _get_calendar_month_range(selected_month)
    else:
        month_start, month_end = _get_month_range(week_start)
    month_days = calendar_month_days if view_mode == 'month' else _get_month_days(month_start)
    week_days = _get_week_days(week_start)
    week_end = week_days[-1]
    period_days = month_days if view_mode == 'month' else week_days
    period_day_set = set(period_days)
    day_status_map = _build_day_status_map(selected_user, period_days)
    shift_indices = list(range(1, policy_values['max_shifts_per_day'] + 1))
    work_week = WorkWeek.objects.filter(user=selected_user, week_start_date=week_start).first()
    month_hours = WorkShift.objects.filter(
        work_week__user=selected_user,
        shift_date__range=(month_start, month_end),
    ).aggregate(total=models.Sum('total_hours')).get('total') or Decimal('0')
    is_locked = work_week and work_week.status in [
        WorkWeek.Status.SUBMITTED,
        WorkWeek.Status.APPROVED,
        WorkWeek.Status.LOCKED,
    ]
    slot_times = []
    if policy_values.get('open_time') and policy_values.get('close_time'):
        slot_dt = datetime.combine(datetime.today().date(), policy_values['open_time'])
        end_dt = datetime.combine(datetime.today().date(), policy_values['close_time'])
        step = timedelta(minutes=30)
        while slot_dt < end_dt:
            slot_times.append(slot_dt.strftime('%H:%M'))
            slot_dt += step
    shift_templates = _shift_template_payload(policy_values)
    active_shift_templates = {
        str(item.template_id): item
        for item in WorkShiftTemplate.objects.filter(is_active=True)
    }

    assignment_map = {}
    assignments = TaskAssignment.objects.filter(
        work_shift__work_week__user=selected_user,
        work_shift__shift_date__range=(period_days[0], period_days[-1]),
    ).select_related('task', 'work_shift')
    for assignment in assignments:
        day_key = assignment.work_shift.shift_date.isoformat()
        shift_key = str(assignment.work_shift.shift_index)
        assignment_map.setdefault(day_key, {}).setdefault(shift_key, []).append({
            'task': assignment.task.task_name,
            'hours': float(assignment.planned_hours),
        })

    errors = []
    if request.method == 'POST':
        action = request.POST.get('action', 'draft')
        locked_weeks = WorkWeek.objects.filter(
            user=selected_user,
            week_start_date__in={_get_week_start(day) for day in period_days},
            status__in=[WorkWeek.Status.SUBMITTED, WorkWeek.Status.APPROVED, WorkWeek.Status.LOCKED],
        )
        if locked_weeks.exists() and not is_admin:
            messages.error(request, 'Có tuần làm việc đã submit. Vui lòng liên hệ admin để chỉnh sửa.')
            redirect_url = f"{reverse('workshift_register')}?mode={view_mode}&week_start={week_start}&month={selected_month:%Y-%m}"
            if is_admin:
                redirect_url += f"&user_id={selected_user.id}"
            return redirect(redirect_url)

        shifts_payload = []
        daily_totals = {day: Decimal('0') for day in period_days}
        weekly_totals = {}
        form_values = {}
        period_total = Decimal('0')

        for day in period_days:
            day_allowed = day.weekday() in policy_values['working_days']
            daily_shifts = []
            for idx in range(1, policy_values['max_shifts_per_day'] + 1):
                start_key = f'shift_{day.isoformat()}_{idx}_start'
                end_key = f'shift_{day.isoformat()}_{idx}_end'
                template_key = f'shift_{day.isoformat()}_{idx}_template'
                start_val = request.POST.get(start_key, '').strip()
                end_val = request.POST.get(end_key, '').strip()
                template_id = request.POST.get(template_key, '').strip()
                form_values[f'{day.isoformat()}_{idx}_start'] = start_val
                form_values[f'{day.isoformat()}_{idx}_end'] = end_val
                form_values[f'{day.isoformat()}_{idx}_template'] = template_id
                if not start_val and not end_val:
                    continue
                shift_template = active_shift_templates.get(template_id)
                if not shift_template:
                    errors.append(f'{day.strftime("%d/%m")} ca {idx}: ca làm việc không hợp lệ hoặc đã tạm dừng.')
                    continue
                if not _is_template_effective(shift_template, day):
                    errors.append(f'{day.strftime("%d/%m")} ca {idx}: ca làm việc chưa có hiệu lực hoặc đã hết hiệu lực.')
                    continue
                if not day_allowed:
                    errors.append(f'{day.strftime("%d/%m")} không thuộc ngày làm việc.')
                    continue
                start_time = _parse_time(start_val)
                end_time = _parse_time(end_val)
                if not start_time or not end_time:
                    errors.append(f'{day.strftime("%d/%m")} ca {idx}: giờ không hợp lệ.')
                    continue
                if start_time != shift_template.start_time or end_time != shift_template.end_time:
                    errors.append(f'{day.strftime("%d/%m")} ca {idx}: giờ không khớp danh mục ca.')
                    continue
                if start_time < policy_values['open_time'] or end_time > policy_values['close_time']:
                    errors.append(f'{day.strftime("%d/%m")} ca {idx}: ngoài giờ mở/kết kho.')
                    continue
                hours = _calc_hours(
                    day,
                    start_time,
                    end_time,
                    policy_values['lunch_start'],
                    policy_values['lunch_end'],
                    shift_template.include_lunch_break,
                )
                if hours is None:
                    errors.append(f'{day.strftime("%d/%m")} ca {idx}: ca làm việc không hợp lệ.')
                    continue
                daily_shifts.append((start_time, end_time))
                daily_totals[day] += hours
                period_total += hours
                week_key = _get_week_start(day)
                weekly_totals[week_key] = weekly_totals.get(week_key, Decimal('0')) + hours
                shifts_payload.append({
                    'shift_date': day,
                    'shift_index': idx,
                    'start_time': start_time,
                    'end_time': end_time,
                    'total_hours': hours,
                    'shift_template': shift_template,
                })
            if len(daily_shifts) > policy_values['max_shifts_per_day']:
                errors.append(f'{day.strftime("%d/%m")} vượt quá số ca tối đa.')
            if len(daily_shifts) > 1:
                daily_shifts = sorted(daily_shifts, key=lambda item: item[0])
                for current, next_shift in zip(daily_shifts, daily_shifts[1:]):
                    if current[1] > next_shift[0]:
                        errors.append(f'{day.strftime("%d/%m")} ca làm việc bị chồng giờ.')
                        break
            if daily_totals[day] > policy_values['max_hours_per_day']:
                errors.append(f'{day.strftime("%d/%m")} vượt quá {policy_values["max_hours_per_day"]} giờ/ngày.')

        touched_week_keys = {_get_week_start(day) for day in period_days}
        outside_week_totals = WorkShift.objects.filter(
            work_week__user=selected_user,
            work_week__week_start_date__in=touched_week_keys,
        ).exclude(
            shift_date__in=period_day_set,
        ).values('work_week__week_start_date').annotate(total=models.Sum('total_hours'))
        for item in outside_week_totals:
            week_key = item['work_week__week_start_date']
            weekly_totals[week_key] = weekly_totals.get(week_key, Decimal('0')) + (item['total'] or Decimal('0'))

        for week_key, week_total in sorted(weekly_totals.items()):
            if week_total > policy_values['max_hours_per_week']:
                errors.append(f'Tuần {week_key.strftime("%d/%m")} vượt quá {policy_values["max_hours_per_week"]} giờ/tuần.')
        saved_period_hours = month_hours if view_mode == 'month' else (work_week.total_hours if work_week else Decimal('0'))
        projected_month_hours = month_hours - saved_period_hours + period_total
        if projected_month_hours > policy_values['max_hours_per_month']:
            errors.append(f'Vượt quá {policy_values["max_hours_per_month"]} giờ/tháng.')

        if errors:
            shift_rows = []
            for day in period_days:
                day_key = day.isoformat()
                shifts = {}
                for idx in range(1, policy_values['max_shifts_per_day'] + 1):
                    shifts[str(idx)] = {
                        'start_time': form_values.get(f'{day_key}_{idx}_start', ''),
                        'end_time': form_values.get(f'{day_key}_{idx}_end', ''),
                        'template_id': form_values.get(f'{day_key}_{idx}_template', ''),
                    }
                shift_rows.append({
                    'day': day,
                    'day_key': day_key,
                    'shifts': shifts,
                    'total_hours': daily_totals.get(day, Decimal('0')),
                })
            day_template_map = _build_day_template_map(shift_rows)
            context = {
                **user_context,
                'policy': policy,
                'policy_values': policy_values,
                'week_start': week_start,
                'week_end': week_end,
                'week_days': week_days,
                'view_mode': view_mode,
                'selected_month': selected_month,
                'prev_month': prev_month,
                'next_month': next_month,
                'month_calendar': month_calendar,
                'shift_indices': shift_indices,
                'prev_week': prev_week,
                'next_week': next_week,
                'period_days': period_days,
                'shift_rows': shift_rows,
                'total_hours': period_total,
                'month_hours': month_hours,
                'month_start': month_start,
                'month_end': month_end,
                'slot_times': slot_times,
                'shift_templates': shift_templates,
                'shift_templates_json': json.dumps(shift_templates, cls=DjangoJSONEncoder),
                'assignment_map_json': json.dumps(assignment_map, cls=DjangoJSONEncoder),
                'day_status_json': json.dumps(day_status_map, cls=DjangoJSONEncoder),
                'day_status_map': day_status_map,
                'day_template_map': day_template_map,
                'errors': errors,
                'selected_user': selected_user,
                'ctv_users': ctv_users,
                'is_admin': is_admin,
                'is_locked': is_locked,
            }
            return render(request, 'app_workshift/workshift_register.html', context)

        payload_by_week = {}
        for item in shifts_payload:
            payload_by_week.setdefault(_get_week_start(item['shift_date']), []).append(item)

        for week_key in sorted(touched_week_keys):
            week_period_days = {day for day in period_days if _get_week_start(day) == week_key}
            week_payload = payload_by_week.get(week_key, [])
            current_week_days = _get_week_days(week_key)
            current_week_end = current_week_days[-1]
            current_week = WorkWeek.objects.filter(user=selected_user, week_start_date=week_key).first()
            if not current_week and not week_payload:
                continue
            if not current_week:
                current_week = WorkWeek.objects.create(
                    user=selected_user,
                    week_start_date=week_key,
                    week_end_date=current_week_end,
                    status=WorkWeek.Status.DRAFT,
                    total_hours=Decimal('0'),
                    created_by=request.user,
                    updated_by=request.user,
                )
            else:
                current_week.updated_by = request.user
                current_week.updated_at = timezone.now()
                current_week.save(update_fields=['updated_by', 'updated_at'])

            existing_shifts = WorkShift.objects.filter(work_week=current_week)
            before_data = list(existing_shifts.values('shift_date', 'shift_index', 'shift_template_id', 'start_time', 'end_time', 'total_hours'))
            editable_existing = existing_shifts.filter(shift_date__in=week_period_days)
            existing_map = {(s.shift_date, s.shift_index): s for s in editable_existing}
            payload_keys = set()
            new_shifts = []
            for item in week_payload:
                key = (item['shift_date'], item['shift_index'])
                payload_keys.add(key)
                if key in existing_map:
                    shift = existing_map[key]
                    shift.start_time = item['start_time']
                    shift.end_time = item['end_time']
                    shift.total_hours = item['total_hours']
                    shift.shift_template = item['shift_template']
                    shift.updated_by = request.user
                    shift.save(update_fields=['shift_template', 'start_time', 'end_time', 'total_hours', 'updated_by'])
                else:
                    new_shifts.append(WorkShift(
                        work_week=current_week,
                        shift_template=item['shift_template'],
                        shift_date=item['shift_date'],
                        shift_index=item['shift_index'],
                        start_time=item['start_time'],
                        end_time=item['end_time'],
                        total_hours=item['total_hours'],
                        created_by=request.user,
                        updated_by=request.user,
                    ))
            if new_shifts:
                WorkShift.objects.bulk_create(new_shifts)
            for key, shift in existing_map.items():
                if key not in payload_keys:
                    shift.delete()

            week_total = WorkShift.objects.filter(work_week=current_week).aggregate(total=models.Sum('total_hours')).get('total') or Decimal('0')
            current_week.total_hours = week_total
            update_fields = ['total_hours', 'updated_by', 'updated_at']
            if action == 'submit':
                current_week.status = WorkWeek.Status.SUBMITTED
                current_week.approved_by = None
                current_week.approved_at = None
                current_week.rejected_reason = None
                update_fields.extend(['status', 'approved_by', 'approved_at', 'rejected_reason'])
            elif current_week.status == WorkWeek.Status.REJECTED:
                current_week.status = WorkWeek.Status.DRAFT
                current_week.rejected_reason = None
                update_fields.extend(['status', 'rejected_reason'])
            current_week.updated_by = request.user
            current_week.save(update_fields=update_fields)

            WorkAuditLog.objects.create(
                entity_type='WorkWeek',
                entity_id=current_week.week_id,
                action='submit' if action == 'submit' else 'save',
                before_data=json.loads(json.dumps({'shifts': before_data}, cls=DjangoJSONEncoder)),
                after_data=json.loads(json.dumps({
                    'shifts': [_shift_payload_for_audit(item) for item in week_payload],
                }, cls=DjangoJSONEncoder)),
                actor=request.user,
            )

        if action == 'submit':
            messages.success(request, 'Đăng ký lịch làm việc thành công.')
        else:
            messages.success(request, 'Đã lưu bản nháp.')

        redirect_url = f"{reverse('workshift_register')}?mode={view_mode}&week_start={week_start}&month={selected_month:%Y-%m}"
        if is_admin:
            redirect_url += f"&user_id={selected_user.id}"
        return redirect(redirect_url)

    week_context = _build_week_context(selected_user, week_start, policy_values, selected_user=selected_user)
    if view_mode == 'month':
        shift_rows, period_total = _build_period_shift_rows(selected_user, month_days, policy_values)
    else:
        shift_rows = week_context['shift_rows']
        period_total = week_context['total_hours']
    day_template_map = _build_day_template_map(shift_rows)

    context = {
        **user_context,
        'policy': policy,
        'policy_values': policy_values,
        'week_start': week_start,
        'week_end': week_context['week_end'],
        'week_days': week_context['week_days'],
        'view_mode': view_mode,
        'selected_month': selected_month,
        'prev_month': prev_month,
        'next_month': next_month,
        'month_calendar': month_calendar,
        'period_days': period_days,
        'shift_indices': shift_indices,
        'shift_rows': shift_rows,
        'total_hours': period_total,
        'month_hours': month_hours,
        'month_start': month_start,
        'month_end': month_end,
        'slot_times': slot_times,
        'shift_templates': shift_templates,
        'shift_templates_json': json.dumps(shift_templates, cls=DjangoJSONEncoder),
        'assignment_map_json': json.dumps(assignment_map, cls=DjangoJSONEncoder),
        'day_status_json': json.dumps(day_status_map, cls=DjangoJSONEncoder),
        'day_status_map': day_status_map,
        'day_template_map': day_template_map,
        'work_week': week_context['work_week'],
        'selected_user': selected_user,
        'ctv_users': ctv_users,
        'is_admin': is_admin,
        'is_locked': is_locked,
        'prev_week': prev_week,
        'next_week': next_week,
    }
    return render(request, 'app_workshift/workshift_register.html', context)


@login_required
@require_ui_permission('workshift_tasks')
def workshift_tasks_view(request):
    user_context = get_user_context(request.user)
    is_admin = user_context.get('is_admin') or user_context.get('is_super_admin')
    task_tab = request.GET.get('tab', 'assign') if is_admin else 'assign'
    if task_tab not in {'assign', 'approve', 'catalog'}:
        task_tab = 'assign'

    week_param = request.GET.get('week_start')
    if week_param:
        try:
            week_start = datetime.strptime(week_param, '%Y-%m-%d').date()
        except ValueError:
            week_start = _get_week_start(datetime.today().date())
    else:
        week_start = _get_week_start(datetime.today().date())
    prev_week = week_start - timedelta(days=7)
    next_week = week_start + timedelta(days=7)

    ctv_users = _get_active_checker_users() if is_admin else User.objects.none()
    users_to_show = ctv_users if is_admin else [request.user]
    status_labels = {
        WorkWeek.Status.DRAFT: 'Nháp',
        WorkWeek.Status.SUBMITTED: 'Chờ duyệt',
        WorkWeek.Status.APPROVED: 'Đã duyệt',
        WorkWeek.Status.REJECTED: 'Trả lại',
        WorkWeek.Status.LOCKED: 'Đã khóa',
    }
    review_status = request.GET.get('status', WorkWeek.Status.SUBMITTED) if is_admin else WorkWeek.Status.SUBMITTED
    if review_status not in status_labels and review_status != 'all':
        review_status = WorkWeek.Status.SUBMITTED
    review_checker_id = request.GET.get('checker_id', '')
    review_from_date_raw = request.GET.get('from_date', '')
    review_to_date_raw = request.GET.get('to_date', '')
    review_from_date = _parse_date_value(review_from_date_raw)
    review_to_date = _parse_date_value(review_to_date_raw)
    review_filters = {
        'checker_id': review_checker_id,
        'from_date': review_from_date_raw,
        'to_date': review_to_date_raw,
        'status': review_status,
    }
    review_query = urlencode({
        'tab': 'approve',
        **{key: value for key, value in review_filters.items() if value},
    })

    work_weeks = WorkWeek.objects.filter(user__in=users_to_show, week_start_date=week_start).select_related('user', 'approved_by')
    review_queryset = WorkWeek.objects.filter(user__in=users_to_show).select_related('user', 'approved_by')
    if review_status != 'all':
        review_queryset = review_queryset.filter(status=review_status)
    if review_checker_id:
        review_queryset = review_queryset.filter(user_id=review_checker_id)
    if review_from_date:
        review_queryset = review_queryset.filter(week_end_date__gte=review_from_date)
    if review_to_date:
        review_queryset = review_queryset.filter(week_start_date__lte=review_to_date)
    approved_work_weeks = work_weeks.filter(status=WorkWeek.Status.APPROVED)
    shifts = WorkShift.objects.filter(work_week__in=approved_work_weeks).select_related('work_week', 'work_week__user') if approved_work_weeks else []
    policy = _get_active_policy()
    policy_values = _get_policy_values(policy)
    week_days = _get_week_days(week_start)
    shift_indices = list(range(1, policy_values['max_shifts_per_day'] + 1))
    shift_grid = {
        user.id: {day.isoformat(): {idx: None for idx in shift_indices} for day in week_days}
        for user in users_to_show
    }
    for shift in shifts:
        day_key = shift.shift_date.isoformat()
        user_id = shift.work_week.user_id
        if user_id in shift_grid and day_key in shift_grid[user_id]:
            shift_grid[user_id][day_key][shift.shift_index] = shift
    tasks = TaskCatalog.objects.filter(is_active=True).order_by('task_name')
    tasks_catalog = TaskCatalog.objects.all().order_by('task_name') if is_admin else []

    if request.method == 'POST' and is_admin:
        action = request.POST.get('action', 'assign')
        if action == 'create_task':
            task_code = request.POST.get('task_code', '').strip()
            task_name = request.POST.get('task_name', '').strip()
            is_active = request.POST.get('is_active') == '1'
            if not task_code or not task_name:
                messages.error(request, 'Vui lòng nhập mã task và tên task.')
            elif TaskCatalog.objects.filter(task_code=task_code).exists():
                messages.error(request, 'Mã task đã tồn tại.')
            else:
                TaskCatalog.objects.create(
                    task_code=task_code,
                    task_name=task_name,
                    is_active=is_active,
                    created_by=request.user,
                    updated_by=request.user,
                )
                messages.success(request, 'Đã tạo task mới.')
            return redirect(f"{reverse('workshift_tasks')}?tab=catalog&week_start={week_start}")

        if action in {'approve_week', 'reject_week'}:
            work_week = get_object_or_404(WorkWeek.objects.filter(user__in=users_to_show), pk=request.POST.get('week_id'))
            next_query = request.POST.get('next_query') or 'tab=approve'
            before_data = {
                'status': work_week.status,
                'approved_by_id': work_week.approved_by_id,
                'approved_at': work_week.approved_at,
                'rejected_reason': work_week.rejected_reason,
            }
            if action == 'approve_week':
                if work_week.status not in [WorkWeek.Status.SUBMITTED, WorkWeek.Status.REJECTED]:
                    messages.error(request, 'Chỉ duyệt lịch đang chờ duyệt hoặc đã bị trả lại.')
                else:
                    work_week.status = WorkWeek.Status.APPROVED
                    work_week.approved_by = request.user
                    work_week.approved_at = timezone.now()
                    work_week.rejected_reason = None
                    work_week.updated_by = request.user
                    work_week.save(update_fields=['status', 'approved_by', 'approved_at', 'rejected_reason', 'updated_by', 'updated_at'])
                    WorkAuditLog.objects.create(
                        entity_type='WorkWeek',
                        entity_id=work_week.week_id,
                        action='approve',
                        before_data=json.loads(json.dumps(before_data, cls=DjangoJSONEncoder)),
                        after_data=json.loads(json.dumps({
                            'status': work_week.status,
                            'approved_by_id': work_week.approved_by_id,
                            'approved_at': work_week.approved_at,
                            'rejected_reason': work_week.rejected_reason,
                        }, cls=DjangoJSONEncoder)),
                        actor=request.user,
                    )
                    messages.success(request, 'Đã duyệt lịch làm việc.')
            else:
                has_assignments = TaskAssignment.objects.filter(work_shift__work_week=work_week).exists()
                if work_week.status not in [WorkWeek.Status.SUBMITTED, WorkWeek.Status.APPROVED]:
                    messages.error(request, 'Chỉ trả lại lịch đang chờ duyệt hoặc đã duyệt.')
                elif has_assignments:
                    messages.error(request, 'Lịch đã có phân công công việc, không thể trả lại.')
                else:
                    work_week.status = WorkWeek.Status.REJECTED
                    work_week.approved_by = None
                    work_week.approved_at = None
                    work_week.rejected_reason = request.POST.get('rejected_reason', '').strip() or None
                    work_week.updated_by = request.user
                    work_week.save(update_fields=['status', 'approved_by', 'approved_at', 'rejected_reason', 'updated_by', 'updated_at'])
                    WorkAuditLog.objects.create(
                        entity_type='WorkWeek',
                        entity_id=work_week.week_id,
                        action='reject',
                        before_data=json.loads(json.dumps(before_data, cls=DjangoJSONEncoder)),
                        after_data=json.loads(json.dumps({
                            'status': work_week.status,
                            'approved_by_id': work_week.approved_by_id,
                            'approved_at': work_week.approved_at,
                            'rejected_reason': work_week.rejected_reason,
                        }, cls=DjangoJSONEncoder)),
                        actor=request.user,
                    )
                    messages.success(request, 'Đã trả lại lịch làm việc.')
            return redirect(f"{reverse('workshift_tasks')}?{next_query}")

        shift_id = request.POST.get('shift_id')
        task_id = request.POST.get('task_id')
        planned_hours = request.POST.get('planned_hours')
        note = request.POST.get('note', '').strip()
        try:
            planned_hours_val = Decimal(planned_hours)
        except (TypeError, ValueError, ArithmeticError):
            planned_hours_val = Decimal('0')
        if not shift_id or not task_id:
            messages.error(request, 'Vui lòng chọn ca và task.')
        else:
            shift = get_object_or_404(WorkShift, pk=shift_id)
            task = get_object_or_404(TaskCatalog, pk=task_id)
            current_total = shift.task_assignments.aggregate(total=models.Sum('planned_hours')).get('total') or Decimal('0')
            if shift.work_week.status != WorkWeek.Status.APPROVED:
                messages.error(request, 'Chỉ được phân công công việc cho lịch đã duyệt.')
            elif planned_hours_val <= 0:
                messages.error(request, 'Giờ task không hợp lệ.')
            elif current_total + planned_hours_val > shift.total_hours:
                messages.error(request, 'Tổng giờ task vượt quá giờ ca.')
            else:
                TaskAssignment.objects.create(
                    work_shift=shift,
                    task=task,
                    planned_hours=planned_hours_val,
                    note=note or None,
                    assigned_by=request.user,
                )
                messages.success(request, 'Đã phân công task.')
        return redirect(f"{reverse('workshift_tasks')}?tab=assign&week_start={week_start}")

    assignments = TaskAssignment.objects.filter(work_shift__in=shifts).select_related('task', 'work_shift')
    assignments_by_shift = {}
    assignment_totals_by_shift = {}
    for assignment in assignments:
        assignments_by_shift.setdefault(assignment.work_shift_id, []).append(assignment)
        assignment_totals_by_shift[assignment.work_shift_id] = (
            assignment_totals_by_shift.get(assignment.work_shift_id, Decimal('0')) + assignment.planned_hours
        )
    review_weeks = list(
        review_queryset.prefetch_related('shifts__shift_template').order_by('-week_start_date', 'user__username')
    ) if is_admin else []

    context = {
        **user_context,
        'week_start': week_start,
        'week_end': week_start + timedelta(days=4),
        'prev_week': prev_week,
        'next_week': next_week,
        'week_days': week_days,
        'shift_indices': shift_indices,
        'shift_grid': shift_grid,
        'users_to_show': users_to_show,
        'shifts': shifts,
        'assignments_by_shift': assignments_by_shift,
        'assignment_totals_by_shift': assignment_totals_by_shift,
        'tasks': tasks,
        'tasks_catalog': tasks_catalog,
        'review_weeks': review_weeks,
        'review_filters': review_filters,
        'review_query': review_query,
        'status_labels': status_labels,
        'ctv_users': ctv_users,
        'is_admin': is_admin,
        'task_tab': task_tab,
    }
    return render(request, 'app_workshift/workshift_tasks.html', context)


@login_required
@require_ui_permission('workshift_tasks')
def workshift_report_view(request):
    user_context = get_user_context(request.user)
    is_admin = user_context.get('is_admin') or user_context.get('is_super_admin')
    ctv_users = _get_active_checker_users() if is_admin else User.objects.none()
    users_to_show = ctv_users if is_admin else User.objects.filter(pk=request.user.pk)

    month_param = request.GET.get('month')
    if month_param:
        try:
            selected_month = datetime.strptime(month_param, '%Y-%m').date().replace(day=1)
        except ValueError:
            selected_month = datetime.today().date().replace(day=1)
    else:
        selected_month = datetime.today().date().replace(day=1)
    prev_month = (selected_month - timedelta(days=1)).replace(day=1)
    next_month = (selected_month.replace(day=28) + timedelta(days=4)).replace(day=1)
    month_days = _get_month_days(selected_month)
    month_calendar = _build_month_calendar(selected_month, month_days)
    month_start, month_end = _get_calendar_month_range(selected_month)

    checker_id = request.GET.get('checker_id', '')
    status_filter = request.GET.get('status', WorkWeek.Status.APPROVED)
    status_labels = {
        WorkWeek.Status.DRAFT: 'Nháp',
        WorkWeek.Status.SUBMITTED: 'Chờ duyệt',
        WorkWeek.Status.APPROVED: 'Đã duyệt',
        WorkWeek.Status.REJECTED: 'Trả lại',
        WorkWeek.Status.LOCKED: 'Đã khóa',
    }
    if status_filter not in status_labels and status_filter != 'all':
        status_filter = WorkWeek.Status.APPROVED

    shifts = WorkShift.objects.filter(
        work_week__user__in=users_to_show,
        shift_date__range=(month_start, month_end),
    ).select_related(
        'work_week',
        'work_week__user',
        'shift_template',
    ).prefetch_related(
        'task_assignments__task',
    ).order_by('shift_date', 'start_time', 'work_week__user__username')
    if checker_id and is_admin:
        shifts = shifts.filter(work_week__user_id=checker_id)
    if status_filter != 'all':
        shifts = shifts.filter(work_week__status=status_filter)

    report_by_day = {day.isoformat(): [] for day in month_days}
    user_totals = {}
    task_totals = {}
    total_hours = Decimal('0')
    assigned_hours = Decimal('0')
    for shift in shifts:
        assignments = list(shift.task_assignments.all())
        shift_assigned_hours = sum((item.planned_hours for item in assignments), Decimal('0'))
        total_hours += shift.total_hours
        assigned_hours += shift_assigned_hours
        user_key = shift.work_week.user_id
        user_totals.setdefault(user_key, {
            'user': shift.work_week.user,
            'hours': Decimal('0'),
            'assigned_hours': Decimal('0'),
            'shifts': 0,
        })
        user_totals[user_key]['hours'] += shift.total_hours
        user_totals[user_key]['assigned_hours'] += shift_assigned_hours
        user_totals[user_key]['shifts'] += 1
        for assignment in assignments:
            task_key = assignment.task_id
            task_totals.setdefault(task_key, {
                'task': assignment.task,
                'hours': Decimal('0'),
            })
            task_totals[task_key]['hours'] += assignment.planned_hours
        report_by_day.setdefault(shift.shift_date.isoformat(), []).append({
            'shift': shift,
            'assignments': assignments,
            'assigned_hours': shift_assigned_hours,
        })

    report_weeks = []
    for week in month_calendar:
        report_weeks.append([
            {
                **cell,
                'items': report_by_day.get(cell['day_key'], []),
            }
            for cell in week
        ])

    context = {
        **user_context,
        'is_admin': is_admin,
        'ctv_users': ctv_users,
        'selected_month': selected_month,
        'prev_month': prev_month,
        'next_month': next_month,
        'report_weeks': report_weeks,
        'status_filter': status_filter,
        'status_labels': status_labels,
        'checker_id': checker_id,
        'total_hours': total_hours,
        'assigned_hours': assigned_hours,
        'total_shifts': shifts.count(),
        'user_totals': sorted(user_totals.values(), key=lambda item: item['hours'], reverse=True),
        'task_totals': sorted(task_totals.values(), key=lambda item: item['hours'], reverse=True),
    }
    return render(request, 'app_workshift/workshift_report.html', context)


@login_required
@require_ui_permission('workshift_tasks')
def workshift_report_export_view(request):
    user_context = get_user_context(request.user)
    is_admin = user_context.get('is_admin') or user_context.get('is_super_admin')
    ctv_users = _get_active_checker_users() if is_admin else User.objects.none()
    users_to_show = ctv_users if is_admin else User.objects.filter(pk=request.user.pk)

    month_param = request.GET.get('month')
    if month_param:
        try:
            selected_month = datetime.strptime(month_param, '%Y-%m').date().replace(day=1)
        except ValueError:
            selected_month = datetime.today().date().replace(day=1)
    else:
        selected_month = datetime.today().date().replace(day=1)
    month_start, month_end = _get_calendar_month_range(selected_month)

    checker_id = request.GET.get('checker_id', '')
    status_filter = request.GET.get('status', WorkWeek.Status.APPROVED)
    valid_statuses = {choice[0] for choice in WorkWeek.Status.choices}
    if status_filter not in valid_statuses and status_filter != 'all':
        status_filter = WorkWeek.Status.APPROVED

    shifts = WorkShift.objects.filter(
        work_week__user__in=users_to_show,
        shift_date__range=(month_start, month_end),
    ).select_related(
        'work_week',
        'work_week__user',
        'shift_template',
    ).prefetch_related(
        'task_assignments__task',
    ).order_by('work_week__user__username', 'shift_date', 'start_time')
    if checker_id and is_admin:
        shifts = shifts.filter(work_week__user_id=checker_id)
    if status_filter != 'all':
        shifts = shifts.filter(work_week__status=status_filter)

    daily_rows = {}
    task_rows = {}
    for shift in shifts:
        full_name = shift.work_week.user.get_full_name() or shift.work_week.user.username
        daily_key = (full_name, shift.shift_date)
        daily_rows.setdefault(daily_key, {
            'name': full_name,
            'date': shift.shift_date,
            'hours': Decimal('0'),
            'notes': [],
        })
        daily_rows[daily_key]['hours'] += shift.total_hours
        shift_label = shift.shift_template.template_name if shift.shift_template else f'Ca {shift.shift_index}'
        assignments = list(shift.task_assignments.all())
        if assignments:
            task_names = ', '.join(item.task.task_name for item in assignments)
            daily_rows[daily_key]['notes'].append(f'{shift_label}: {task_names}')
        else:
            daily_rows[daily_key]['notes'].append(f'{shift_label}: Chưa phân công')

        for assignment in assignments:
            task_key = (full_name, assignment.task.task_name)
            task_rows.setdefault(task_key, {
                'name': full_name,
                'task': assignment.task.task_name,
                'hours': Decimal('0'),
            })
            task_rows[task_key]['hours'] += assignment.planned_hours

    wb = Workbook()
    ws_daily = wb.active
    ws_daily.title = 'Lich lam viec theo ngay'
    ws_tasks = wb.create_sheet('Tom tat cong viec')

    header_fill = PatternFill('solid', fgColor='E5F4EC')
    thin = Side(style='thin', color='D9E2E7')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_sheet(ws, widths):
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical='center', wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

    ws_daily.append(['Họ tên', 'Ngày-tháng-năm', 'Tổng giờ làm việc', 'Tính phụ cấp', 'Ghi chú'])
    for item in sorted(daily_rows.values(), key=lambda row: (row['name'], row['date'])):
        ws_daily.append([
            item['name'],
            item['date'].strftime('%d/%m/%Y'),
            float(item['hours']),
            'YES' if item['hours'] >= Decimal('8') else 'NO',
            '; '.join(item['notes']),
        ])
    style_sheet(ws_daily, [28, 18, 18, 16, 60])

    ws_tasks.append(['Họ tên', 'Công việc', 'Tổng giờ làm (sum cả tháng)'])
    for item in sorted(task_rows.values(), key=lambda row: (row['name'], row['task'])):
        ws_tasks.append([
            item['name'],
            item['task'],
            float(item['hours']),
        ])
    style_sheet(ws_tasks, [28, 40, 24])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'workshift_report_{selected_month:%Y_%m}.xlsx'
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
@require_ui_permission('workshift_policy')
def workshift_policy_view(request):
    user_context = get_user_context(request.user)
    is_admin = user_context.get('is_admin') or user_context.get('is_super_admin')
    if not is_admin:
        messages.error(request, 'Bạn không có quyền truy cập cấu hình.')
        return redirect('workshift_register')

    policies = WorkPolicy.objects.all().order_by('-effective_from', '-policy_id')
    selected_policy = None
    policy_id = request.GET.get('policy_id')
    is_new = request.GET.get('new') == '1'
    if is_new:
        selected_policy = None
    elif policy_id:
        selected_policy = get_object_or_404(WorkPolicy, pk=policy_id)
    else:
        selected_policy = policies.filter(is_active=True).first() or policies.first()

    errors = []
    form_data = {}
    selected_days = []
    if request.method == 'POST':
        action = request.POST.get('action', 'save_policy')
        if action == 'create_shift_template':
            template_name = request.POST.get('template_name', '').strip()
            template_start = _parse_time(request.POST.get('template_start'))
            template_end = _parse_time(request.POST.get('template_end'))
            include_lunch_break = request.POST.get('template_include_lunch_break') == '1'
            template_effective_from_raw = request.POST.get('template_effective_from', '').strip()
            template_effective_to_raw = request.POST.get('template_effective_to', '').strip()
            template_effective_from = None
            template_effective_to = None
            if template_effective_from_raw:
                try:
                    template_effective_from = datetime.strptime(template_effective_from_raw, '%Y-%m-%d').date()
                except ValueError:
                    messages.error(request, 'Ngày bắt đầu hiệu lực của ca không hợp lệ.')
                    return redirect('workshift_policy')
            if template_effective_to_raw:
                try:
                    template_effective_to = datetime.strptime(template_effective_to_raw, '%Y-%m-%d').date()
                except ValueError:
                    messages.error(request, 'Ngày kết thúc hiệu lực của ca không hợp lệ.')
                    return redirect('workshift_policy')
            try:
                sort_order = int(request.POST.get('template_sort_order') or 0)
            except ValueError:
                sort_order = 0
            template_active = request.POST.get('template_is_active') == '1'
            if not template_name or not template_start or not template_end:
                messages.error(request, 'Vui lòng nhập tên ca và khung giờ.')
            elif template_start >= template_end:
                messages.error(request, 'Giờ bắt đầu ca phải trước giờ kết thúc.')
            elif template_effective_from and template_effective_to and template_effective_from > template_effective_to:
                messages.error(request, 'Ngày bắt đầu hiệu lực phải trước hoặc bằng ngày kết thúc.')
            else:
                WorkShiftTemplate.objects.create(
                    template_name=template_name,
                    start_time=template_start,
                    end_time=template_end,
                    include_lunch_break=include_lunch_break,
                    effective_from=template_effective_from,
                    effective_to=template_effective_to,
                    sort_order=sort_order,
                    is_active=template_active,
                    created_by=request.user,
                    updated_by=request.user,
                )
                messages.success(request, 'Đã tạo ca làm việc.')
            return redirect('workshift_policy')
        if action == 'toggle_shift_template':
            template_id = request.POST.get('template_id')
            template = get_object_or_404(WorkShiftTemplate, pk=template_id)
            template.is_active = not template.is_active
            template.updated_by = request.user
            template.save(update_fields=['is_active', 'updated_by', 'updated_at'])
            messages.success(request, 'Đã cập nhật trạng thái ca làm việc.')
            return redirect('workshift_policy')
        if action == 'update_shift_template':
            template = get_object_or_404(WorkShiftTemplate, pk=request.POST.get('template_id'))
            template_name = request.POST.get('template_name', '').strip()
            effective_from_raw = request.POST.get('template_effective_from', '').strip()
            effective_to_raw = request.POST.get('template_effective_to', '').strip()
            effective_from = _parse_date_value(effective_from_raw)
            effective_to = _parse_date_value(effective_to_raw)
            if effective_from_raw and effective_from is None:
                messages.error(request, 'Ngày bắt đầu hiệu lực của ca không hợp lệ.')
                return redirect('workshift_policy')
            if effective_to_raw and effective_to is None:
                messages.error(request, 'Ngày kết thúc hiệu lực của ca không hợp lệ.')
                return redirect('workshift_policy')
            try:
                sort_order = int(request.POST.get('template_sort_order') or 0)
            except ValueError:
                sort_order = 0
            is_active = request.POST.get('template_is_active') == '1'
            is_used = WorkShift.objects.filter(shift_template=template).exists()

            if not template_name:
                messages.error(request, 'Vui lòng nhập tên ca.')
            elif effective_from and effective_to and effective_from > effective_to:
                messages.error(request, 'Ngày bắt đầu hiệu lực phải trước hoặc bằng ngày kết thúc.')
            else:
                template.template_name = template_name
                template.effective_from = effective_from
                template.effective_to = effective_to
                template.sort_order = sort_order
                template.is_active = is_active
                update_fields = [
                    'template_name',
                    'effective_from',
                    'effective_to',
                    'sort_order',
                    'is_active',
                    'updated_by',
                    'updated_at',
                ]
                if not is_used:
                    template_start = _parse_time(request.POST.get('template_start'))
                    template_end = _parse_time(request.POST.get('template_end'))
                    include_lunch_break = request.POST.get('template_include_lunch_break') == '1'
                    if not template_start or not template_end:
                        messages.error(request, 'Vui lòng nhập khung giờ ca.')
                        return redirect('workshift_policy')
                    if template_start >= template_end:
                        messages.error(request, 'Giờ bắt đầu ca phải trước giờ kết thúc.')
                        return redirect('workshift_policy')
                    template.start_time = template_start
                    template.end_time = template_end
                    template.include_lunch_break = include_lunch_break
                    update_fields.extend(['start_time', 'end_time', 'include_lunch_break'])
                template.updated_by = request.user
                template.save(update_fields=update_fields)
                if is_used:
                    messages.success(request, 'Đã cập nhật ca. Ca đã được đăng ký nên khung giờ được giữ nguyên.')
                else:
                    messages.success(request, 'Đã cập nhật ca làm việc.')
            return redirect('workshift_policy')

        policy_id = request.POST.get('policy_id') or ''
        policy_name = request.POST.get('policy_name', '').strip() or 'Default'
        working_days = request.POST.getlist('working_days')
        open_time = _parse_time(request.POST.get('open_time'))
        close_time = _parse_time(request.POST.get('close_time'))
        lunch_start = _parse_time(request.POST.get('lunch_start'))
        lunch_end = _parse_time(request.POST.get('lunch_end'))
        effective_from_raw = request.POST.get('effective_from')
        is_active = request.POST.get('is_active') == '1'
        try:
            max_shifts_per_day = int(request.POST.get('max_shifts_per_day') or 0)
        except ValueError:
            max_shifts_per_day = 0
        try:
            max_hours_per_day = Decimal(request.POST.get('max_hours_per_day') or '0')
        except (ValueError, ArithmeticError):
            max_hours_per_day = Decimal('0')
        try:
            max_hours_per_week = Decimal(request.POST.get('max_hours_per_week') or '0')
        except (ValueError, ArithmeticError):
            max_hours_per_week = Decimal('0')
        try:
            max_hours_per_month = Decimal(request.POST.get('max_hours_per_month') or '0')
        except (ValueError, ArithmeticError):
            max_hours_per_month = Decimal('0')
        try:
            min_hours_per_shift = Decimal(request.POST.get('min_hours_per_shift') or '0')
        except (ValueError, ArithmeticError):
            min_hours_per_shift = Decimal('0')

        effective_from = None
        if effective_from_raw:
            try:
                effective_from = datetime.strptime(effective_from_raw, '%Y-%m-%d').date()
            except ValueError:
                errors.append('Ngày hiệu lực không hợp lệ.')

        if not working_days:
            errors.append('Vui lòng chọn ngày làm việc.')
        if not open_time or not close_time:
            errors.append('Vui lòng nhập giờ mở/kết kho.')
        if open_time and close_time and open_time >= close_time:
            errors.append('Giờ mở kho phải trước giờ kết kho.')
        if lunch_start and lunch_end and lunch_start >= lunch_end:
            errors.append('Giờ nghỉ trưa không hợp lệ.')
        if max_shifts_per_day <= 0:
            errors.append('Số ca/ngày không hợp lệ.')
        if max_hours_per_day <= 0 or max_hours_per_week <= 0 or max_hours_per_month <= 0:
            errors.append('Hạn mức giờ không hợp lệ.')
        if min_hours_per_shift <= 0:
            errors.append('Giờ tối thiểu/ca không hợp lệ.')

        form_data = {
            'policy_id': policy_id,
            'policy_name': policy_name,
            'working_days': [int(day) for day in working_days if day.isdigit()],
            'open_time': request.POST.get('open_time') or '',
            'close_time': request.POST.get('close_time') or '',
            'lunch_start': request.POST.get('lunch_start') or '',
            'lunch_end': request.POST.get('lunch_end') or '',
            'max_shifts_per_day': max_shifts_per_day,
            'max_hours_per_day': max_hours_per_day,
            'max_hours_per_week': max_hours_per_week,
            'max_hours_per_month': request.POST.get('max_hours_per_month') or '',
            'min_hours_per_shift': request.POST.get('min_hours_per_shift') or '',
            'effective_from': effective_from_raw or '',
            'is_active': is_active,
        }
        selected_days = form_data['working_days']

        if not errors:
            if policy_id:
                selected_policy = get_object_or_404(WorkPolicy, pk=policy_id)
                selected_policy.policy_name = policy_name
                selected_policy.working_days = ','.join(sorted(working_days))
                selected_policy.open_time = open_time
                selected_policy.close_time = close_time
                selected_policy.lunch_start = lunch_start
                selected_policy.lunch_end = lunch_end
                selected_policy.max_shifts_per_day = max_shifts_per_day
                selected_policy.max_hours_per_day = max_hours_per_day
                selected_policy.max_hours_per_week = max_hours_per_week
                selected_policy.max_hours_per_month = max_hours_per_month
                selected_policy.min_hours_per_shift = min_hours_per_shift
                selected_policy.effective_from = effective_from
                selected_policy.updated_by = request.user
                selected_policy.is_active = is_active
                selected_policy.save()
            else:
                selected_policy = WorkPolicy.objects.create(
                    policy_name=policy_name,
                    working_days=','.join(sorted(working_days)),
                    open_time=open_time,
                    close_time=close_time,
                    lunch_start=lunch_start,
                    lunch_end=lunch_end,
                    max_shifts_per_day=max_shifts_per_day,
                    max_hours_per_day=max_hours_per_day,
                    max_hours_per_week=max_hours_per_week,
                    max_hours_per_month=max_hours_per_month,
                    min_hours_per_shift=min_hours_per_shift,
                    effective_from=effective_from,
                    is_active=is_active,
                    created_by=request.user,
                    updated_by=request.user,
                )
            if is_active:
                WorkPolicy.objects.exclude(pk=selected_policy.pk).update(is_active=False)
            messages.success(request, 'Đã lưu cấu hình work policy.')
            return redirect(f"{reverse('workshift_policy')}?policy_id={selected_policy.policy_id}")

    if not form_data and selected_policy:
        selected_days = [int(x) for x in (selected_policy.working_days or '').split(',') if x.strip().isdigit()]
        form_data = {
            'policy_id': selected_policy.policy_id,
            'policy_name': selected_policy.policy_name,
            'working_days': selected_days,
            'open_time': selected_policy.open_time.strftime('%H:%M') if selected_policy.open_time else '',
            'close_time': selected_policy.close_time.strftime('%H:%M') if selected_policy.close_time else '',
            'lunch_start': selected_policy.lunch_start.strftime('%H:%M') if selected_policy.lunch_start else '',
            'lunch_end': selected_policy.lunch_end.strftime('%H:%M') if selected_policy.lunch_end else '',
            'max_shifts_per_day': selected_policy.max_shifts_per_day,
            'max_hours_per_day': selected_policy.max_hours_per_day,
            'max_hours_per_week': selected_policy.max_hours_per_week,
            'max_hours_per_month': selected_policy.max_hours_per_month,
            'min_hours_per_shift': selected_policy.min_hours_per_shift,
            'effective_from': selected_policy.effective_from.strftime('%Y-%m-%d') if selected_policy.effective_from else '',
            'is_active': selected_policy.is_active,
        }

    day_labels = [
        (0, 'Thứ 2'),
        (1, 'Thứ 3'),
        (2, 'Thứ 4'),
        (3, 'Thứ 5'),
        (4, 'Thứ 6'),
    ]
    time_slots_24h = _build_time_slots()

    context = {
        **user_context,
        'policies': policies,
        'selected_policy': selected_policy,
        'form_data': form_data,
        'selected_days': selected_days,
        'errors': errors,
        'day_labels': day_labels,
        'time_slots_24h': time_slots_24h,
        'shift_templates': WorkShiftTemplate.objects.annotate(
            usage_count=models.Count('work_shifts')
        ).order_by('sort_order', 'start_time', 'template_name'),
        'is_admin': is_admin,
    }
    return render(request, 'app_workshift/workshift_policy.html', context)
